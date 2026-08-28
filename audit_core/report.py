from __future__ import annotations

from decimal import Decimal
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from .common import AuditError


NAVY = "17365D"
PALE_BLUE = "EAF3F8"
TOTAL_BLUE = "D9EAF7"
PASS_GREEN = "E2F0D9"
REVIEW_YELLOW = "FFF2CC"
BORDER_BLUE = "B7C9D6"
INPUT_BLUE = "0000FF"
FORMULA_ERRORS = {"#REF!", "#DIV/0!", "#VALUE!", "#N/A", "#NAME?"}


def _font(*, bold: bool = False, color: str = "000000", size: float = 10) -> Font:
    return Font(name="Microsoft YaHei", size=size, bold=bold, color=color)


def _border() -> Border:
    side = Side(style="thin", color=BORDER_BLUE)
    return Border(left=side, right=side, top=side, bottom=side)


def _number(value: Any) -> str:
    if value is None:
        return "未识别"
    if isinstance(value, bool):
        return "是" if value else "否"
    try:
        decimal = Decimal(str(value))
    except Exception:
        return str(value)
    if decimal == decimal.to_integral_value():
        return str(int(decimal))
    return format(decimal.normalize(), "f")


CONFIDENCE_PRESENTATION = {
    "exact": "高",
    "fuzzy": "中",
    "unmatched": "低",
}


def _match_confidence(level: str) -> str:
    return f"置信度：{CONFIDENCE_PRESENTATION[level]}"


def _style_title(ws: Any, title: str, note: str) -> None:
    ws.merge_cells("A1:F1")
    ws["A1"] = title
    ws["A1"].font = _font(bold=True, color="FFFFFF", size=16)
    ws["A1"].fill = PatternFill("solid", fgColor=NAVY)
    ws["A1"].alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 28

    ws.merge_cells("A2:F2")
    ws["A2"] = note
    ws["A2"].font = _font(color="666666", size=9)
    ws["A2"].fill = PatternFill("solid", fgColor=PALE_BLUE)
    ws["A2"].alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[2].height = 42


def _write_header(ws: Any, values: list[str], *, height: float) -> None:
    for column, value in enumerate(values, 1):
        cell = ws.cell(3, column, value)
        cell.font = _font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.border = _border()
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[3].height = height


def _status_fill(value: Any) -> PatternFill | None:
    text = str(value or "")
    if text in {
        "通过", "全部通过", "匹配", "三方金额匹配", "身份与日期已验证",
        "完全对应", "模糊对应", "精确匹配",
    }:
        return PatternFill("solid", fgColor=PASS_GREEN)
    if any(
        token in text
        for token in (
            "需", "未验证", "不一致", "部分", "人工", "未匹配",
            "模糊匹配", "完全不匹配", "暂不能核销", "不满足",
        )
    ):
        return PatternFill("solid", fgColor=REVIEW_YELLOW)
    return None


def _write_row(
    ws: Any,
    row: int,
    values: list[Any],
    *,
    height: float,
    total: bool = False,
    font_size: float = 10,
) -> None:
    for column, value in enumerate(values, 1):
        cell = ws.cell(row, column, value)
        cell.border = _border()
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        cell.font = _font(
            bold=total,
            color="000000" if total else INPUT_BLUE,
            size=font_size,
        )
        if total:
            cell.fill = PatternFill("solid", fgColor=TOTAL_BLUE)
        elif column == 6:
            fill = _status_fill(value)
            if fill:
                cell.fill = fill
    ws.row_dimensions[row].height = height


def _sheet_base(ws: Any, *, zoom: int, tab_color: str) -> None:
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = zoom
    ws.sheet_properties.tabColor = tab_color
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToHeight = 0
    ws.print_title_rows = "1:3"


PASS_PERSONNEL_KNOWLEDGE_MATCHES = {"matched", "fuzzy_matched"}


def _personnel_knowledge_match_text(item: dict[str, Any]) -> str:
    status = str(item.get("knowledge_status") or "unmatched")
    mapping_resolved = item.get("mapping_status") == "matched"
    barcode = {
        "exact": "69码精确匹配",
        "invalid": "69码完全不匹配（格式无效）",
        "missing": "69码完全不匹配（缺失）",
        "not_found": (
            "69码未核验（结算行尚未映射到Excel）"
            if not mapping_resolved
            else "69码完全不匹配（知识库未登记）"
        ),
    }.get(str(item.get("knowledge_barcode_match") or "not_found"), "69码完全不匹配")
    name = {
        "exact": "产品名精确匹配",
        "fuzzy": "产品名模糊匹配",
        "unmatched": (
            "产品名未核验"
            if not mapping_resolved
            else "产品名未形成唯一模糊对应"
        ),
    }.get(
        str(item.get("knowledge_name_match_type") or "unmatched"),
        "产品名未核验" if not mapping_resolved else "产品名未形成唯一模糊对应",
    )
    if name == "产品名模糊匹配":
        name += "（严格高于0.5）"
    level = "exact" if status == "matched" else "fuzzy" if status == "fuzzy_matched" else "unmatched"
    return f"{barcode}；{name}\n知识库结论：{_match_confidence(level)}"


def _personnel_sales_knowledge_text(item: dict[str, Any]) -> str:
    knowledge_identity = (
        f"{item.get('knowledge_product_code')} / {item.get('knowledge_product_name')} / "
        f"{item.get('knowledge_barcode_69')}"
        if item.get("knowledge_product_name")
        else "未确认"
    )
    return (
        f"Excel商品：{item.get('excel_product_name') or '未匹配'}\n"
        f"条码：{item.get('mapped_barcode') or '未匹配'}；数量：{_number(item.get('excel_quantity'))}\n"
        f"知识库商品：{knowledge_identity}\n"
        f"{_personnel_knowledge_match_text(item)}\n"
        f"计算奖励：{_number(item['calculated_settlement_reward'])}元"
    )


def _personnel_row_product_name(item: dict[str, Any]) -> str:
    """Use the selected catalog name as the human-facing settlement-row title."""

    knowledge_name = str(item.get("knowledge_product_name") or "").strip()
    if knowledge_name:
        return knowledge_name
    return str(item.get("settlement_product_name") or "未识别商品").strip()


def _personnel_error_attribution(
    item: dict[str, Any],
    *,
    sales_file: str,
    settlement_file: str,
) -> str | None:
    if item.get("mapping_status") != "matched":
        return (
            f"问题文件：{sales_file}、{settlement_file}\n"
            f"错误原因：{settlement_file}识别到商品“"
            f"{item.get('settlement_product_name') or '未识别'}”、数量"
            f"{_number(item.get('settlement_quantity'))}，但未能在{sales_file}中唯一对应到"
            "一条已通过知识库的销售商品记录。现有材料不能判断是结算行识别信息不完整，"
            "还是销售Excel缺少/写错对应商品，因此商品、数量和奖励不能继续核销。"
        )
    if item.get("knowledge_status") not in PASS_PERSONNEL_KNOWLEDGE_MATCHES:
        return (
            f"问题文件：{sales_file}\n"
            "对照基准：商品知识库\n"
            f"错误原因：{sales_file}中的商品“{item.get('excel_product_name') or '未识别'}”、"
            f"69码“{item.get('mapped_barcode') or '未提供'}”未能唯一通过知识库身份核验；"
            "该销售商品没有可用于核销的权威商品身份。"
        )
    if Decimal(str(item.get("quantity_difference") or 0)) != 0:
        return (
            f"问题文件：{sales_file}、{settlement_file}\n"
            f"错误原因：{sales_file}读取数量为{_number(item.get('excel_quantity'))}，"
            f"{settlement_file}识别数量为{_number(item.get('settlement_quantity'))}，"
            "两份材料数量不一致，现有证据不能判断应以哪一份数量为准。"
        )
    if Decimal(str(item.get("settlement_line_amount_difference") or 0)) != 0:
        return (
            f"问题文件：{settlement_file}\n"
            f"对照文件：{sales_file}\n"
            f"错误原因：按{sales_file}对应销售数量和结算奖励规则计算为"
            f"{_number(item.get('calculated_settlement_reward'))}元，"
            f"但{settlement_file}该商品行识别奖励为"
            f"{_number(item.get('settlement_reward_amount'))}元，奖励金额不一致。"
        )
    if item.get("mapping_confidence") == "low":
        return (
            f"问题文件：{settlement_file}\n"
            f"对照文件：{sales_file}\n"
            f"错误原因：{settlement_file}该商品行的名称或条码清晰度不足，不能与"
            f"{sales_file}中的销售商品形成唯一对应。"
        )
    return None


def _personnel_comparison_text(
    item: dict[str, Any],
    *,
    sales_file: str,
    settlement_file: str,
) -> str:
    knowledge_result = _personnel_knowledge_match_text(item).splitlines()[0]
    knowledge_result = (
        knowledge_result.replace("69码精确匹配", "69码一致")
        .replace("69码完全不匹配", "69码不一致")
        .replace("产品名精确匹配", "商品名称一致")
        .replace("产品名模糊匹配", "商品名称对应")
        .replace("产品名未形成唯一模糊对应", "商品名称未形成唯一模糊对应")
    )

    mapping_confidence = str(item.get("mapping_confidence") or "low")
    mapping_matched = item.get("mapping_status") == "matched" and mapping_confidence != "low"
    if not mapping_matched:
        mapping_result = "商品无法确认（置信度：低）"
    elif mapping_confidence == "high":
        mapping_result = "商品对应（置信度：高）"
    else:
        mapping_result = "商品对应（置信度：中）"

    quantity_difference = item.get("quantity_difference")
    quantity_matched = (
        quantity_difference is not None and Decimal(str(quantity_difference)) == 0
    )
    if quantity_matched:
        quantity_result = "数量一致"
    elif item.get("excel_quantity") is None:
        quantity_result = (
            f"数量未核验：结算单 {_number(item.get('settlement_quantity'))}"
        )
    else:
        quantity_result = (
            f"数量不一致：Excel {_number(item.get('excel_quantity'))}；"
            f"结算单 {_number(item.get('settlement_quantity'))}"
        )

    amount_difference = Decimal(str(item.get("settlement_line_amount_difference") or 0))
    amount_matched = amount_difference == 0
    if amount_matched:
        amount_result = "奖励金额一致"
    else:
        amount_result = (
            f"奖励金额不一致：计算奖励 {_number(item.get('calculated_settlement_reward'))}元；"
            f"结算单 {_number(item.get('settlement_reward_amount'))}元"
        )

    settlement_result = (
        f"商品、数量、奖励金额全部对应（置信度："
        f"{'高' if mapping_confidence == 'high' else '中'}）"
        if mapping_matched and quantity_matched and amount_matched
        else f"{mapping_result}；{quantity_result}；{amount_result}"
    )
    comparison = (
        f"Excel商品与知识库：{knowledge_result}\n"
        f"结算单与销售记录：{settlement_result}"
    )
    attribution = _personnel_error_attribution(
        item,
        sales_file=sales_file,
        settlement_file=settlement_file,
    )
    return comparison if attribution is None else f"{comparison}\n{attribution}"


def _personnel_conclusion(
    item: dict[str, Any],
    *,
    sales_file: str,
    settlement_file: str,
) -> str:
    if item.get("mapping_status") != "matched":
        return (
            f"{_match_confidence('unmatched')}\n"
            f"要重新提交：核实正确商品后，更正{sales_file}中的缺失/错误商品，或重新拍摄"
            f"{settlement_file}中的该商品行并完整显示商品名称、条码和数量；重新提交后两份材料"
            "必须能够唯一对应"
        )
    if item.get("knowledge_status") not in PASS_PERSONNEL_KNOWLEDGE_MATCHES:
        return (
            f"{_match_confidence('unmatched')}\n"
            f"要重新提交：更正{sales_file}中的该商品名称或69码后重新提交；"
            f"{item.get('knowledge_resubmission') or '如销售文件无误，则补齐商品知识库后重新核销'}"
        )
    if Decimal(str(item.get("quantity_difference") or 0)) != 0:
        return (
            f"{_match_confidence('unmatched')}\n"
            f"要重新提交：核实正确数量后更正{sales_file}或{settlement_file}中的错误值，"
            "重新提交后两份材料的该商品数量必须一致"
        )
    if Decimal(str(item.get("settlement_line_amount_difference") or 0)) != 0:
        return (
            f"{_match_confidence('unmatched')}\n"
            f"要重新提交：更正{settlement_file}中的该商品奖励金额，使其与"
            f"{sales_file}对应销售数量按奖励规则计算的金额一致"
        )
    if item.get("mapping_confidence") == "low":
        return (
            f"{_match_confidence('unmatched')}\n"
            f"要重新提交：补拍{settlement_file}中的该商品行，让商品名称或条码清晰可见，"
            f"能够与{sales_file}唯一对应"
        )
    if (
        item.get("mapping_confidence") == "medium"
        or item.get("knowledge_status") == "fuzzy_matched"
    ):
        return (
            f"{_match_confidence('fuzzy')}\n"
            "69码精确、商品名称模糊匹配严格高于0.5，数量和金额一致，已通过；"
            "无需重新提交"
        )
    return f"{_match_confidence('exact')}\n无需重新提交"


def _add_personnel_sheet(wb: Workbook, result: dict[str, Any]) -> Any:
    ws = wb.create_sheet("人员激励核销")
    summary = result["summary"]
    sales = result["sales"]
    settlement = result["settlement"]
    transfer_files = sorted(
        {
            Path(str(name)).name
            for item in result.get("transfer_evidence") or []
            for name in item.get("source_files") or []
        }
    )
    sales_file = Path(str(sales["source_file"])).name
    settlement_file = Path(str(settlement["source_file"])).name
    transfer_file_text = "、".join(transfer_files) or "未识别到转账凭证文件"
    _style_title(
        ws,
        "人员激励核销（简明对比）",
        "销售Excel由代码读取并先核验商品知识库：69码必须精确匹配，产品名模糊相似度须严格高于0.5；结算单和转账截图由视觉AI识别。",
    )
    _write_header(
        ws,
        [
            "核验对象",
            f"销售Excel + 商品知识库（代码核验）\n{sales_file}",
            f"结算单（视觉AI识别）\n{settlement_file}",
            "转账凭证（视觉AI识别）\n" + transfer_file_text,
            "具体对比结果",
            "结论 / 要重新提交什么",
        ],
        height=30,
    )

    row = 4
    for item in result["sku_reconciliation"]:
        _write_row(
            ws,
            row,
            [
                _personnel_row_product_name(item),
                _personnel_sales_knowledge_text(item),
                f"视觉识别商品：{item['settlement_product_name']}\n数量：{_number(item['settlement_quantity'])}\n行奖励：{_number(item['settlement_reward_amount'])}元",
                "转账仅核对总额",
                _personnel_comparison_text(
                    item,
                    sales_file=sales_file,
                    settlement_file=settlement_file,
                ),
                _personnel_conclusion(
                    item,
                    sales_file=sales_file,
                    settlement_file=settlement_file,
                ),
            ],
            height=(
                174
                if _personnel_error_attribution(
                    item,
                    sales_file=sales_file,
                    settlement_file=settlement_file,
                )
                else 126
            ),
            font_size=9,
        )
        row += 1

    values = [
        Decimal(str(summary["calculated_line_reward"])),
        Decimal(str(summary["settlement_line_reward"])),
        Decimal(str(summary["transfer_total"])),
    ]
    three_way_difference = max(values) - min(values)
    knowledge_problem_count = int(summary.get("sales_knowledge_problem_count") or 0)
    if three_way_difference != 0:
        total_comparison = (
            f"问题文件：{sales_file}、{settlement_file}、{transfer_file_text}\n"
            f"错误原因：{sales_file}与{settlement_file}逐商品核算奖励为"
            f"{_number(summary['calculated_line_reward'])}元，{settlement_file}各商品行奖励合计为"
            f"{_number(summary['settlement_line_reward'])}元，{transfer_file_text}去重后的转账合计为"
            f"{_number(summary['transfer_total'])}元，三者最大相差{_number(three_way_difference)}元；"
            "现有材料不能确认完整、唯一的应核销金额。"
        )
    elif knowledge_problem_count:
        total_comparison = (
            f"问题文件：{sales_file}\n"
            "对照基准：商品知识库\n"
            f"错误原因：三方金额均为{_number(summary['calculated_line_reward'])}元，"
            f"但{sales_file}中仍有{knowledge_problem_count}个商品未通过知识库身份核验；"
            "金额相同不能替代商品身份核验。"
        )
    else:
        total_comparison = "三方金额一致，所有销售商品均已通过知识库身份核验。"
    _write_row(
        ws,
        row,
        [
            "合计金额",
            (
                f"数量：{_number(summary['excel_mapped_quantity'])}\n"
                f"计算奖励：{_number(summary['calculated_line_reward'])}元\n"
                f"知识库：{summary.get('sales_knowledge_matched_count', 0)}个产品名精确，"
                f"{summary.get('sales_knowledge_fuzzy_count', 0)}个产品名模糊，"
                f"{summary.get('sales_knowledge_problem_count', 0)}个未通过"
            ),
            f"视觉识别数量：{_number(summary['settlement_line_quantity'])}\n行奖励合计：{_number(summary['settlement_line_reward'])}元",
            f"视觉识别并去重：{_number(summary['transfer_total'])}元",
            (
                total_comparison
            ),
            (
                f"{_match_confidence('exact')}\n三方金额一致，无需重新提交"
                if three_way_difference == 0
                and knowledge_problem_count == 0
                else (
                    f"{_match_confidence('unmatched')}\n"
                    + (
                        f"要重新提交：按对应商品错误行更正并重新提交{sales_file}；"
                        "如销售文件无误，则补齐商品知识库后重新核销"
                        if three_way_difference == 0
                        else (
                            f"要重新提交：逐项核实正确金额，更正{sales_file}、{settlement_file}或"
                            f"{transfer_file_text}中的错误值；重新提交后三方金额必须一致"
                        )
                    )
                )
            ),
        ],
        height=132 if three_way_difference != 0 or knowledge_problem_count else 90,
        total=True,
    )
    row += 1
    claim_difference = Decimal(str(summary["claimed_amount"])) - Decimal(str(summary["transfer_total"]))
    if claim_difference > 0:
        claim_difference_text = f"申请金额比转账金额多{_number(abs(claim_difference))}元"
    elif claim_difference < 0:
        claim_difference_text = f"申请金额比转账金额少{_number(abs(claim_difference))}元"
    else:
        claim_difference_text = "申请金额与转账金额一致"
    claim_comparison = (
        f"{settlement_file}中的申请金额与{transfer_file_text}中的转账金额一致。"
        if claim_difference == 0
        else (
            f"问题文件：{settlement_file}、{transfer_file_text}\n"
            f"错误原因：{settlement_file}识别申请金额为{_number(summary['claimed_amount'])}元，"
            f"{transfer_file_text}识别转账金额为{_number(summary['transfer_total'])}元，"
            f"{claim_difference_text}。现有材料不能判断应以哪一份金额为准，金额证据链未闭合。"
        )
    )
    _write_row(
        ws,
        row,
        [
            "实际申请金额",
            "销售Excel不参与本项金额核对",
            f"文件：{settlement_file}\n视觉识别申请：{_number(summary['claimed_amount'])}元",
            f"文件：{transfer_file_text}\n视觉识别转账：{_number(summary['transfer_total'])}元",
            claim_comparison,
            (
                f"{_match_confidence('exact')}\n申请金额与转账金额一致，无需重新提交"
                if claim_difference == 0
                else (
                    f"{_match_confidence('unmatched')}\n"
                    f"要重新提交：更正{settlement_file}中的申请金额，使其与{transfer_file_text}一致；"
                    "若差异有业务原因，补交一份能同时关联上述文件的金额差异说明"
                )
            ),
        ],
        height=96,
    )
    row += 1

    transfers = result.get("transfer_evidence") or []
    recipient_visible = bool(transfers) and all(
        item.get("identity_visible") and item.get("recipient_name") for item in transfers
    )
    store_visible = bool(transfers) and all(item.get("store_name") for item in transfers)
    dates_visible = bool(transfers) and all(
        item.get("event_date_visible") and item.get("transfer_date") for item in transfers
    )
    identity_verified = recipient_visible and store_visible and dates_visible
    if identity_verified:
        transfer_note = "收款人、门店对应关系和完整交易日期可识别"
        comparison = (
            f"{sales_file}中的门店与{transfer_file_text}中的收款人、对应门店和完整交易日期"
            "均可逐笔确认。"
        )
    else:
        missing: list[str] = []
        if not recipient_visible:
            missing.append("完整收款人")
        if not store_visible:
            missing.append("门店对应关系")
        if not dates_visible:
            missing.append("完整交易日期")
        transfer_note = "截图金额可识别；" + "、".join(missing) + "未显示"
        comparison = (
            f"问题文件：{transfer_file_text}\n"
            f"对照文件：{sales_file}\n"
            f"错误原因：{sales_file}读取到{summary['store_count']}家门店；"
            f"{transfer_file_text}没有完整显示每笔收款人、对应门店和完整交易日期。"
            f"因此无法确认Excel中的{summary['store_count']}家门店分别由谁收款、"
            "对应哪一笔转账以及具体转账日期。"
        )
    identity_level = "exact" if identity_verified else (
        "fuzzy" if recipient_visible or store_visible or dates_visible else "unmatched"
    )
    amount_note = "金额一致" if three_way_difference == 0 else "金额不一致"
    _write_row(
        ws,
        row,
        [
            "收款人与日期",
            f"文件：{sales_file}\nExcel读取到{summary['store_count']}家门店",
            f"文件：{settlement_file}\n视觉识别活动期：{settlement.get('activity_start')} 至 {settlement.get('activity_end')}",
            f"文件：{transfer_file_text}\n{transfer_note}",
            comparison,
            (
                f"{_match_confidence(identity_level)}\n{amount_note}；"
                + (
                    "身份与日期均可确认，无需重新提交"
                    if identity_verified
                    else (
                        f"要重新提交：重新拍摄或导出{transfer_file_text}对应的完整转账凭证；"
                        f"每笔同时显示收款人、对应门店和完整交易日期，能够与{sales_file}中的"
                        f"{summary['store_count']}家门店逐一核对"
                    )
                )
            ),
        ],
        height=126,
    )

    # Freeze only the two title rows and the header. Freezing product rows makes
    # the lower reconciliation rows effectively unreachable on shorter screens.
    ws.freeze_panes = "A4"
    ws.auto_filter.ref = f"A3:F{row}"
    widths = {"A": 27, "B": 48, "C": 28, "D": 28, "E": 38, "F": 34}
    for column, width in widths.items():
        ws.column_dimensions[column].width = width
    _sheet_base(ws, zoom=85, tab_color="70AD47")
    return ws


KNOWLEDGE_FIELD_LABELS = {
    "product_code": "商品编码",
    "product_name": "商品名称",
    "barcode_69": "69码",
}

def _failed_display_controls(
    item: dict[str, Any],
    contract: dict[str, Any],
    contract_knowledge: dict[str, Any],
    contract_core: dict[str, Any],
    contract_sales: dict[str, Any],
) -> list[str]:
    failed: list[str] = []
    if contract_core.get("status") != "pass":
        failed.append("合同核心六项")
    if contract_sales.get("status") != "pass":
        failed.append("销售Excel对合同")
    if item.get("contract_attachment_sales_status") != "pass":
        failed.append("销售Excel对合同附件")
    if item.get("sales_internal_status") != "pass":
        failed.append("销售Excel文件内部")
    if not item.get("photo_files"):
        failed.append("现场照片")
    if item.get("period_match") != "match":
        failed.append("活动日期")
    if item.get("store_match") == "mismatch":
        failed.append("门店水印错误")
    elif item.get("store_match") not in {"exact", "compatible"}:
        failed.append("门店水印缺失或无法核对")
    if item.get("display_match") != "pass":
        failed.append("陈列标准")
    if item.get("duplicate_check") != "none":
        failed.append("照片复用")
    if contract_knowledge.get("status") == "fail":
        failed.append("合同商品知识库")
    if item.get("photo_knowledge_match") != "exact":
        failed.append("现场商品知识库")
    if item.get("photo_contract_product_status") not in {
        "not_applicable",
        "exact",
    }:
        failed.append("现场商品与合同")
    if contract.get("requires_promotion") and not item.get("promotion_present"):
        failed.append("合同促销要求")
    if item.get("amount_rule_status") != "pass":
        failed.append("金额计费口径")
    return list(dict.fromkeys(failed))


def _wrapped_line_count(value: Any, characters_per_line: int) -> int:
    lines = str(value or "").splitlines() or [""]
    return sum(max(1, (len(line) + characters_per_line - 1) // characters_per_line) for line in lines)


def _display_row_height(values: list[Any]) -> float:
    characters = [34, 46, 44, 44, 24, 34]
    longest = max(
        _wrapped_line_count(value, width)
        for value, width in zip(values, characters, strict=True)
    )
    return min(409.0, max(132.0, 13.0 * longest + 18.0))


def _clip(value: Any, limit: int = 180) -> str:
    text = str(value or "").strip()
    if len(text) <= limit:
        return text
    return text[: max(0, limit - 1)].rstrip() + "…"


def _short_list(values: list[Any], *, limit: int = 5, empty: str = "未识别") -> str:
    clean = list(dict.fromkeys(str(value).strip() for value in values if str(value).strip()))
    if not clean:
        return empty
    visible = clean[:limit]
    suffix = f"（另{len(clean) - limit}项）" if len(clean) > limit else ""
    return "、".join(visible) + suffix



def _contract_sales_text(contract_sales: dict[str, Any]) -> str:
    customer_labels = {
        "exact": "主体一致",
        "fuzzy": "主体模糊一致",
        "mismatch": "主体不一致",
        "unverifiable": "主体无法确认",
    }
    period_labels = {
        "covered": "日期在合同期内",
        "mismatch": "日期超出合同期",
        "unverifiable": "日期无法确认",
    }
    watermark_labels = {
        "present": "水印可见（仅记录，不参与核销）",
        "absent": "未见水印（仅记录，不参与核销）",
        "unverifiable": "水印无法确认（仅记录，不参与核销）",
    }
    seal_labels = {
        "present": "盖章可见",
        "absent": "未见盖章",
        "unverifiable": "盖章无法确认",
    }
    status = str(contract_sales.get("status") or "unverifiable")
    if status == "pass":
        level = (
            "exact"
            if contract_sales.get("customer_status") == "exact"
            else "fuzzy"
        )
    elif status == "unverifiable":
        level = "fuzzy"
    else:
        level = "unmatched"
    details = "；".join(
        [
            customer_labels.get(
                str(contract_sales.get("customer_status")), "主体未核验"
            ),
            period_labels.get(str(contract_sales.get("period_status")), "日期未核验"),
            watermark_labels.get(
                str(contract_sales.get("watermark_status")),
                "水印未核验（仅记录，不参与核销）",
            ),
            seal_labels.get(str(contract_sales.get("seal_status")), "盖章未核验"),
        ]
    )
    return f"{_match_confidence(level)}（{details}）"



def _photo_product_text(item: dict[str, Any]) -> str:
    hits = list(item.get("product_reference_hits") or [])
    if not hits:
        return f"未对应知识库商品（{_match_confidence('unmatched')}）"
    values = [
        f"{hit.get('product_name') or '未命名'}（{hit.get('product_code') or '缺编码'} / "
        f"{hit.get('barcode_69') or '缺69码'}，"
        f"{_match_confidence('exact' if hit.get('confidence') == 'exact' else 'fuzzy')}）"
        for hit in hits
    ]
    return "、".join(values)



def _contract_photo_text(
    item: dict[str, Any],
    contract: dict[str, Any],
    contract_knowledge: dict[str, Any],
    *,
    store: str,
    period: str,
    display: str,
    promotion_result: str,
    contract_product_result: str,
) -> str:
    hard_failure = (
        not item.get("photo_files")
        or item.get("period_match") == "mismatch"
        or item.get("store_match") == "mismatch"
        or item.get("display_match") == "fail"
        or (
            contract.get("requires_promotion")
            and not item.get("promotion_present")
        )
        or contract_knowledge.get("status") == "fail"
    )
    uncertain = (
        item.get("period_match") != "match"
        or item.get("store_match") != "exact"
        or item.get("display_match") != "pass"
    )
    level = "unmatched" if hard_failure else ("fuzzy" if uncertain else "exact")
    parts = [f"门店{store}", f"日期{period}", f"陈列{display}"]
    if contract.get("requires_promotion"):
        parts.append(f"促销{promotion_result}")
    if contract.get("requires_specific_products"):
        parts.append(f"合同商品{contract_product_result}")
    return f"{_match_confidence(level)}（{'；'.join(parts)}）"


def _plain_advice(value: Any) -> str:
    text = str(value)
    replacements = {
        "商品视觉RAG": "商品知识库",
        "视觉RAG": "商品知识库",
        "RAG": "商品知识库",
        "候选命中": "模糊匹配",
        "候选": "模糊匹配",
        "商品ID": "商品身份",
        "SKU": "具体商品",
        "唯一收敛": "唯一确认",
        "customer_name": "客户名称",
        "business_date": "业务日期",
        "product_code": "商品编码",
        "product_name": "商品名称",
        "barcode_69": "69码",
        "retail_price": "零售价",
        "total_amount": "合计金额",
        "quantity": "数量",
        "不一致或无法核对字段：": "需要核对：",
    }
    for source, target in replacements.items():
        text = text.replace(source, target)
    return text


def _row_ranges(values: list[int]) -> str:
    numbers = sorted(set(int(value) for value in values))
    if not numbers:
        return ""
    groups: list[tuple[int, int]] = []
    start = previous = numbers[0]
    for value in numbers[1:]:
        if value == previous + 1:
            previous = value
            continue
        groups.append((start, previous))
        start = previous = value
    groups.append((start, previous))
    return "、".join(
        str(start) if start == end else f"{start}-{end}"
        for start, end in groups
    )


def _display_resubmission_items(
    item: dict[str, Any],
    contract: dict[str, Any],
    contract_knowledge: dict[str, Any],
    contract_core: dict[str, Any],
    contract_sales: dict[str, Any],
) -> list[str]:
    """Build at most one plain-language request per source file type."""

    sales_details: list[str] = []
    if contract_sales.get("customer_status") in {"mismatch", "unverifiable"}:
        sales_details.append("客户名称要能和合同签订方对应")
    if contract_sales.get("period_status") in {"mismatch", "unverifiable"}:
        sales_details.append("业务日期要落在合同活动期内")
    if item.get("contract_attachment_sales_status") != "pass":
        sales_details.append("核验字段要逐行对齐合同附件")
    if item.get("sales_internal_status") != "pass":
        sales_details.append("更正数量×零售价与行合计，并保留清晰打印总计")

    contract_details: list[str] = []
    if contract_core.get("status") != "pass":
        labels = [
            str(check.get("label") or check.get("field"))
            for check in contract_core.get("field_checks") or []
            if check.get("status") != "pass"
        ]
        contract_details.append("补清合同核心字段：" + "、".join(labels))
    if contract_sales.get("integrity_status") != "pass":
        contract_details.append("盖章要清楚")
    if contract_knowledge.get("status") == "fail":
        contract_details.append("合同写到的具体商品要能对应知识库商品")
    if item.get("amount_rule_status") != "pass":
        contract_details.append("写清按店或按堆头的单价和数量")

    photo_details: list[str] = []
    if not item.get("photo_files"):
        photo_details.append(f"补交{item['contract_store_name']}的现场原图")
    if item.get("photo_knowledge_match") == "candidate":
        photo_details.append("把商品短码、规格、口味或正面包装再拍清楚一点")
    elif item.get("photo_knowledge_match") == "unmatched":
        photo_details.append("重新拍商品正面，至少让商品短码或明显名称看清")
    if item.get("photo_contract_product_status") not in {
        "not_applicable",
        "exact",
    }:
        photo_details.append("补拍合同已确认商品的清晰包装，并保留完整现场环境")
    if item.get("period_match") != "match":
        photo_details.append("让完整拍摄日期看得见")
    if item.get("store_match") == "mismatch":
        photo_details.append(
            f"更正照片水印地点：当前识别为“{item.get('visible_location') or '未识别'}”，"
            f"应显示可与合同门店“{item.get('contract_store_name')}”唯一对应的门店名称或地址；"
            "这是水印地点错误，不是清晰度问题"
        )
    elif item.get("store_match") not in {"exact", "compatible"}:
        photo_details.append(
            f"让照片水印显示可与合同门店“{item.get('contract_store_name')}”唯一对应的门店名称或地址"
        )
    if item.get("display_match") != "pass":
        photo_details.append("补一张完整堆头全景，能看清1平方米或数清4列")
    if item.get("duplicate_check") != "none":
        photo_details.append("提交这家门店自己拍的、没有在别店用过的原图")
    if contract.get("requires_promotion") and not item.get("promotion_present"):
        photo_details.append("把合同要求的促销牌或促销文字拍清楚")

    result: list[str] = []
    if contract_details:
        result.append("重新提交合同PDF或补充说明：" + "；".join(dict.fromkeys(contract_details)))
    if photo_details:
        result.append("重新提交现场照片：" + "；".join(dict.fromkeys(photo_details)))
    if sales_details:
        result.append("重新提交销售Excel：" + "；".join(dict.fromkeys(sales_details)))
    return result



def _contract_attachment_sales_text(value: dict[str, Any]) -> str:
    status = str(value.get("status") or "not_applicable")
    if status == "not_applicable":
        return ""
    confidence = "高" if status == "pass" else "低"
    return f"置信度：{confidence}（{_clip(value.get('basis') or '未完成逐行核对', 180)}）"


def _contract_attachment_knowledge_text(value: dict[str, Any]) -> str:
    status = str(value.get("status") or "not_applicable")
    if status == "not_applicable":
        return "不适用"
    confidence = "高" if status == "pass" else "低"
    records = list(value.get("records") or [])
    passed = (
        int(value.get("matched_count") or 0)
        + int(value.get("fuzzy_count") or 0)
        if "matched_count" in value or "fuzzy_count" in value
        else sum(
            record.get("knowledge_status") in {"matched", "fuzzy_matched"}
            for record in records
        )
    )
    problems = (
        int(value.get("problem_count") or 0)
        if "problem_count" in value
        else sum(
            record.get("knowledge_status") not in {"matched", "fuzzy_matched"}
            for record in records
        )
    )
    return (
        f"置信度：{confidence}（通过{passed}项；问题{problems}项）"
    )


def _contract_requirement_lines(contract: dict[str, Any]) -> list[str]:
    """Show only contract requirements that actually exist in the source."""

    lines = [f"陈列：{contract.get('display_standard') or '未识别'}"]
    if contract.get("requires_specific_products"):
        product_requirement = _short_list(
            list(contract.get("required_products") or []),
            limit=3,
            empty="合同限定了具体商品，但未识别到商品名称",
        )
        lines.append(f"商品：{product_requirement}")
    if contract.get("requires_promotion"):
        promotion_requirement = str(
            contract.get("required_promotion")
            or "合同要求促销，但未识别到具体形式"
        )
        lines.append(f"促销：{promotion_requirement}")
    return lines


def _control_status_text(value: Any) -> str:
    return {
        "pass": "通过",
        "fail": "不通过",
        "unverifiable": "无法核对",
        "exact": "一致",
        "fuzzy": "模糊匹配",
        "mismatch": "不一致",
    }.get(str(value), "无法核对")


def _contract_core_text(contract_core: dict[str, Any]) -> str:
    checks = list(contract_core.get("field_checks") or [])
    if not checks:
        return "合同核心六项：无法核对"
    return "\n".join(
        f"{check.get('label') or check.get('field')}：{check.get('contract_value') or '未识别'}"
        f"｜核销：{_control_status_text(check.get('status'))}"
        for check in checks
    )


def _photo_contract_product_text(item: dict[str, Any]) -> str:
    status = str(item.get("photo_contract_product_status") or "unmatched")
    if status == "not_applicable":
        return "不适用（合同未形成具体商品范围）"
    confidence = {
        "exact": "高",
        "candidate": "中",
        "unmatched": "低",
    }.get(status, "低")
    checks = list(item.get("photo_contract_product_checks") or [])
    sources = list(
        dict.fromkeys(
            str(source)
            for check in checks
            for source in check.get("contract_sources") or []
            if str(source)
        )
    )
    source_text = "、".join(sources) or "未找到合同商品来源"
    return (
        f"置信度：{confidence}（{_clip(item.get('photo_contract_product_basis') or '未完成核对', 140)}；"
        f"合同来源：{source_text}）"
    )


ATTACHMENT_FIELD_LABELS = {
    "customer_name": "客户名称",
    "business_date": "业务日期",
    "product_code": "商品编码",
    "product_name": "商品名称",
    "barcode_69": "条形码",
    "unit": "单位",
    "quantity": "数量",
    "retail_price": "零售价",
    "total_amount": "合计金额",
}


def _attachment_source_lines(
    source: dict[str, Any],
    *,
    side: str,
) -> list[str]:
    if side == "contract":
        fields = (
            ("customer_name", "客户名称"),
            ("business_date", "业务日期"),
            ("product_code", "商品编码"),
            ("product_name", "商品名称"),
            ("barcode_69", "条形码"),
            ("unit", "单位"),
            ("quantity", "数量"),
            ("retail_price", "零售价"),
            ("total_amount", "合计金额"),
        )
    else:
        fields = (
            ("customer_name", "客户名称"),
            ("period_text", "业务日期"),
            ("product_code", "商品编码"),
            ("product_name", "商品名称"),
            ("barcode", "条形码"),
            ("unit", "单位"),
            ("quantity", "数量"),
            ("retail_price", "零售价"),
            ("total_amount", "合计金额"),
        )
    lines: list[str] = []
    for field, label in fields:
        value = source.get(field)
        suffix = "元" if field in {"retail_price", "total_amount"} and value is not None else ""
        rendered = _number(value) if isinstance(value, (int, float, Decimal)) else value or "未识别"
        lines.append(f"{label}：{rendered}{suffix}")
    return lines


def _attachment_comparison_lines(record: dict[str, Any]) -> list[str]:
    lines = ["合同附件 → 销售Excel"]
    for comparison in record.get("field_comparisons") or []:
        field = str(comparison.get("field") or "")
        status = str(comparison.get("status") or "unverifiable")
        if field == "product_name" and status == "unverifiable":
            status_text = "模糊辅助（未识别，不单独判错）"
        elif field == "product_name" and status in {"exact", "fuzzy", "mismatch"}:
            status_text = "模糊匹配（辅助项）"
        elif field in {"product_code", "barcode_69"} and status == "exact":
            status_text = "精确匹配"
        elif field == "customer_name" and status == "exact":
            status_text = "精确匹配"
        elif field == "customer_name" and status == "fuzzy":
            status_text = "模糊匹配"
        else:
            status_text = _control_status_text(status)
        lines.append(
            f"{ATTACHMENT_FIELD_LABELS.get(field, field)}："
            f"{status_text}"
        )
    lines.append(
        "合同附件行内金额："
        + _control_status_text(record.get("contract_line_amount_status"))
    )
    if record.get("basis") and record.get("status") != "pass":
        lines.append(f"配对说明：{_plain_advice(record.get('basis'))}")
    if record.get("status") == "pass":
        lines.append("本项结果：全部对应")
    return lines


def _attachment_knowledge_lines(record: dict[str, Any] | None) -> list[str]:
    lines = ["合同商品 → 商品知识库"]
    if not record:
        return [*lines, "本行没有知识库核验结果", "本项结论：置信度：低"]

    knowledge_code = record.get("knowledge_product_code")
    knowledge_name = record.get("knowledge_product_name")
    knowledge_barcode = record.get("knowledge_barcode_69")
    if any((knowledge_code, knowledge_name, knowledge_barcode)):
        lines.extend(
            [
                f"知识库商品：{knowledge_name or '未找到'}",
                f"知识库商品编码：{knowledge_code or '未找到'}",
                f"知识库69码：{knowledge_barcode or '未找到'}",
            ]
        )
    else:
        lines.append("知识库商品：未确认")

    comparisons = {
        str(item.get("field") or ""): str(item.get("comparison") or "")
        for item in record.get("field_comparisons") or []
    }
    status = str(record.get("knowledge_status") or "unmatched")
    passed = status in {"matched", "fuzzy_matched"}
    if not comparisons:
        comparisons = {
            "product_code": "source_only",
            "product_name": "matched" if passed else "conflict",
            "barcode_69": (
                "matched"
                if record.get("source_barcode_69")
                and record.get("source_barcode_69") == knowledge_barcode
                else "conflict"
            ),
        }
    for field in ("product_code", "product_name", "barcode_69"):
        comparison = comparisons.get(field, "missing")
        if field == "product_code":
            result_text = (
                "精确匹配（知识库主编码或编码别名）"
                if comparison == "matched"
                else "无法核验"
                if comparison in {"missing", "not_applicable"}
                else "不匹配（知识库未登记合同业务编码）"
            )
        elif field == "product_name" and comparison in {"matched", "fuzzy"}:
            result_text = "模糊匹配（辅助项）"
        elif comparison == "matched":
            result_text = "精确匹配"
        elif (
            field == "product_name"
            and knowledge_name
            and record.get("source_barcode_69") == knowledge_barcode
        ):
            result_text = "模糊匹配（辅助项）"
        elif field == "product_name" and comparison in {"missing", "not_applicable"}:
            result_text = (
                "合同PDF漏识别（需重新识别以区分同69码商品）"
                if not passed and comparisons.get("barcode_69") == "matched"
                else "模糊辅助（未识别，不单独判错）"
            )
        elif field == "product_name":
            result_text = "模糊辅助（不单独判错）"
        elif comparison in {"missing", "not_applicable"}:
            result_text = "无法核验"
        else:
            result_text = "不匹配"
        lines.append(f"{ATTACHMENT_FIELD_LABELS[field]}：{result_text}")

    confidence = "高" if status == "matched" else "中" if status == "fuzzy_matched" else "低"
    lines.append(
        f"本项结果：{'商品存在已确认' if passed else '商品存在条件未满足'}；置信度：{confidence}"
    )
    return lines


def _add_display_sheet(wb: Workbook, result: dict[str, Any]) -> Any:
    ws = wb.create_sheet("堆头核销")
    summary = result["summary"]
    contract = result["contract"]
    sales = result["sales"]
    contract_sales = result.get("contract_sales_reconciliation") or {}
    contract_core = result.get("contract_core_reconciliation") or {}
    contract_knowledge = result.get("contract_product_knowledge") or {}
    attachment_knowledge = result.get("contract_attachment_product_knowledge") or {}
    attachment_sales = result.get("contract_attachment_sales_reconciliation") or {}
    _style_title(
        ws,
        "堆头陈列核销（简明对比）",
        "合同PDF是主核销文件：现场照片文字先独立对照完整商品知识库并做包装视觉比对，确认商品后再单独对照合同范围；销售Excel逐行对照合同附件。",
    )
    _write_header(
        ws,
        [
            f"合同PDF（视觉AI识别）\n{Path(str(contract['source_file'])).name}",
            "现场照片（视觉AI识别）\n逐张显示文件名和识别内容",
            f"销售Excel（代码读取）\n{Path(str(sales['source_file'])).name}",
            "与合同具体对比结果",
            "核销金额",
            "结论 / 要重新提交什么",
        ],
        height=60,
    )

    period_labels = {"match": "一致", "mismatch": "不一致", "unverifiable": "无法确认"}
    store_labels = {
        "exact": "一致",
        "compatible": "基本一致",
        "mismatch": "不一致",
        "filename_only": "无法确认（仅见文件名）",
    }
    display_labels = {"pass": "符合", "fail": "不符合", "uncertain": "无法确认"}
    display_basis_labels = {
        "stack_1sqm": "达到1平米堆头",
        "four_vertical": "达到4纵陈列",
        "both": "同时达到1平米堆头和4纵陈列",
        "none": "未达到1平米堆头或4纵陈列",
        "unclear": "照片不足以判断1平米堆头或4纵陈列",
    }
    duplicate_labels = {
        "none": "未发现跨门店照片复用",
        "exact": "发现跨门店照片复用",
        "possible": "疑似跨门店照片复用",
        "unverifiable": "无法判断是否存在跨门店照片复用",
    }
    photo_knowledge_labels = {
        "exact": _match_confidence("exact"),
        "candidate": _match_confidence("fuzzy"),
        "unmatched": _match_confidence("unmatched"),
    }
    contract_stores = {
        int(store["line_no"]): store for store in contract.get("stores") or []
    }
    contract_requirement_text = "\n".join(_contract_requirement_lines(contract))
    all_failed_controls: list[str] = []
    sales_internal_status = str(
        summary.get("sales_internal_status")
        or sales.get("internal_status")
        or "pass"
    )
    sales_internal_problem_rows = list(
        summary.get("sales_internal_problem_rows")
        or sales.get("internal_problem_rows")
        or []
    )
    sales_internal_text = (
        "通过（逐行金额、打印总数量和打印合计金额均一致）"
        if sales_internal_status == "pass"
        else (
            "未通过"
            if sales_internal_status == "fail"
            else "无法完整核对"
        )
    )
    if sales_internal_problem_rows:
        sales_internal_text += (
            f"（问题行："
            f"{_row_ranges([int(value) for value in sales_internal_problem_rows])}）"
        )

    row = 4
    activity_requests: list[str] = []
    core_problem_labels = [
        str(check.get("label") or check.get("field"))
        for check in contract_core.get("field_checks") or []
        if check.get("status") != "pass"
    ]
    if core_problem_labels:
        activity_requests.append(
            "重新提交合同PDF：补清" + "、".join(core_problem_labels)
        )
    if contract_sales.get("customer_status") in {"mismatch", "unverifiable"}:
        activity_requests.append("重新提交销售Excel：客户名称要与合同签订方对应")
    if contract_sales.get("period_status") in {"mismatch", "unverifiable"}:
        activity_requests.append("重新提交销售Excel：业务日期要落在合同执行周期内")
    if sales_internal_status != "pass":
        row_suffix = (
            f"（第{_row_ranges([int(value) for value in sales_internal_problem_rows])}行）"
            if sales_internal_problem_rows
            else ""
        )
        activity_requests.append(
            "重新导出销售Excel：更正数量×零售价与行合计，"
            f"并保留唯一清晰的合计/总计行{row_suffix}"
        )
    core_product_knowledge_line = (
        "合同核心商品 → 知识库："
        f"{_contract_attachment_knowledge_text(contract_knowledge)}\n"
        if contract_knowledge.get("status") != "not_applicable"
        else ""
    )
    activity_values = [
        (
            "活动概况｜合同PDF主核销文件\n"
            f"来源：{Path(str(contract.get('source_file') or '')).name}\n"
            f"{_contract_core_text(contract_core)}\n"
            f"商家/门店：共{len(contract.get('stores') or [])}家\n"
            f"堆头数量：{_number(contract.get('contract_stack_count'))}\n"
            f"{contract_requirement_text}\n"
            f"{core_product_knowledge_line}"
            f"合同附件商品 → 知识库：{_contract_attachment_knowledge_text(attachment_knowledge)}"
        ),
        (
            "现场照片执行证据\n"
            f"共{summary.get('photo_count') or 0}张；逐店对照合同门店、日期、活动、陈列和促销要求\n"
            "商品仅与合同已确认商品的知识库参考图片做视觉比对"
        ),
        (
            "销售概况\n"
            f"客户：{_short_list(list(sales.get('customers') or []), limit=4)}\n"
            f"业务日期：{_short_list(list(sales.get('period_values') or []), limit=6)}\n"
            f"商品明细：{len(sales.get('records') or [])}行\n"
            f"数量：{_number(sales.get('total_quantity'))}\n"
            f"合计金额：{_number(sales.get('retail_amount'))}元\n"
            f"Excel内部核对：{sales_internal_text}"
        ),
        (
            f"合同PDF核心六项：{_contract_attachment_sales_text(contract_core)}\n"
            f"销售Excel → 合同签订方/执行周期/盖章：{_contract_sales_text(contract_sales)}\n"
            f"销售Excel文件内部：{sales_internal_text}"
        ),
        "",
        (
            "置信度：高\n无需重新提交"
            if not activity_requests
            else "置信度：低\n要重新提交：" + "；".join(activity_requests)
        ),
    ]
    _write_row(
        ws,
        row,
        activity_values,
        height=_display_row_height(activity_values),
        total=True,
        font_size=9,
    )
    row += 1

    attachment = contract.get("sales_attachment") or {}
    attachment_requests: list[str] = []
    if not attachment.get("present"):
        attachment_requests.append("重新提交包含附加销售明细的完整合同PDF")
    if attachment_knowledge.get("status") == "fail":
        attachment_requests.append("更正合同附件中未通过知识库的商品编码、商品名称或条形码")
    if attachment_sales.get("status") != "pass":
        attachment_requests.append("让销售Excel逐行对齐合同附件")
    source_pages = "、".join(str(value) for value in attachment.get("source_pages") or [])
    attachment_values = [
        (
            "合同销售附件｜合同主基准\n"
            f"来源页：第{source_pages or '未提供'}页\n"
            f"商品明细：{len(attachment.get('records') or [])}行\n"
            f"打印总数量：{_number(attachment.get('total_quantity'))}\n"
            f"打印合计金额：{_number(attachment.get('total_amount'))}元\n"
            f"合同附件商品 → 知识库：{_contract_attachment_knowledge_text(attachment_knowledge)}"
        ),
        "现场照片不参与销售行核对；现场核验在门店记录中单独完成",
        (
            "销售Excel（只向合同附件对齐）\n"
            f"商品明细：{len(sales.get('records') or [])}行\n"
            f"数量：{_number(sales.get('total_quantity'))}\n"
            f"合计金额：{_number(sales.get('retail_amount'))}元\n"
            f"文件内部核对：{sales_internal_text}"
        ),
        (
            f"合同附件 → 销售Excel：{_contract_attachment_sales_text(attachment_sales)}\n"
            "核销范围：客户名称、业务日期、商品编码、商品名称、条形码、数量、零售价、合计金额\n"
            "单位：仅展示，不参与对应或结论"
        ),
        "",
        (
            "置信度：高\n无需重新提交"
            if not attachment_requests
            else "置信度：低\n要重新提交：" + "；".join(attachment_requests)
        ),
    ]
    _write_row(
        ws,
        row,
        attachment_values,
        height=_display_row_height(attachment_values),
        total=True,
        font_size=9,
    )
    row += 1

    attachment_sources = {
        int(item.get("line_no") or 0): item
        for item in attachment.get("records") or []
    }
    sales_sources = {
        int(item.get("excel_row") or 0): item
        for item in sales.get("records") or []
    }
    attachment_knowledge_rows = {
        int(item.get("contract_line_no") or item.get("excel_row") or 0): item
        for item in attachment_knowledge.get("records") or []
    }
    for record in attachment_sales.get("records") or []:
        contract_line_no = int(record.get("contract_line_no") or 0)
        source_page = int(record.get("contract_source_page") or 0)
        sales_row = record.get("sales_excel_row")
        contract_source = attachment_sources.get(contract_line_no, {})
        sales_source = sales_sources.get(int(sales_row or 0), {})
        knowledge_record = attachment_knowledge_rows.get(contract_line_no)
        knowledge_lines = _attachment_knowledge_lines(knowledge_record)
        comparison_lines = [
            *knowledge_lines,
            "",
            *_attachment_comparison_lines(record),
        ]
        problem_labels = [
            ATTACHMENT_FIELD_LABELS.get(str(value.get("field")), str(value.get("field")))
            for value in record.get("field_comparisons") or []
            if value.get("field") != "product_name"
            and value.get("status") not in {"exact", "fuzzy"}
        ]
        knowledge_problem_labels = [
            ATTACHMENT_FIELD_LABELS.get(str(value.get("field")), str(value.get("field")))
            for value in (knowledge_record or {}).get("field_comparisons") or []
            if value.get("field") != "product_name"
            and value.get("comparison") not in {
                "matched", "fuzzy", "source_only", "not_applicable"
            }
        ]
        if (
            knowledge_record
            and knowledge_record.get("knowledge_status") not in {"matched", "fuzzy_matched"}
            and any(
                value.get("field") == "product_name"
                and value.get("comparison") in {"missing", "not_applicable"}
                for value in knowledge_record.get("field_comparisons") or []
            )
        ):
            knowledge_problem_labels.append("合同PDF商品名称漏识别")
        if knowledge_record and not knowledge_problem_labels and knowledge_record.get(
            "knowledge_status"
        ) not in {"matched", "fuzzy_matched"}:
            knowledge_problem_labels.append("商品对应关系")
        if record.get("contract_line_amount_status") != "exact":
            problem_labels.append("合同附件行内金额")
        knowledge_passed = bool(knowledge_record) and knowledge_record.get(
            "knowledge_status"
        ) in {"matched", "fuzzy_matched"}
        passed = record.get("status") == "pass" and knowledge_passed
        confidence = (
            "高"
            if passed
            and record.get("confidence") == "high"
            and knowledge_record.get("knowledge_status") == "matched"
            else "中"
            if passed
            else "低"
        )
        problem_groups: list[str] = []
        if knowledge_problem_labels:
            problem_groups.append(
                "合同商品与知识库的"
                + "、".join(dict.fromkeys(knowledge_problem_labels))
            )
        if problem_labels:
            problem_groups.append(
                "合同附件与销售Excel的"
                + "、".join(dict.fromkeys(problem_labels))
            )
        attachment_row_values = [
            "\n".join(
                [
                    f"合同销售附件第{contract_line_no}行｜PDF第{source_page}页",
                    *_attachment_source_lines(contract_source, side="contract"),
                ]
            ),
            "现场照片不参与本行销售明细核对",
            "\n".join(
                [
                    f"销售Excel第{sales_row}行" if sales_row is not None else "销售Excel：未找到唯一对应行",
                    *_attachment_source_lines(sales_source, side="sales"),
                ]
            ),
            "\n".join(comparison_lines),
            "",
            (
                f"置信度：{confidence}\n无需重新提交"
                if passed
                else (
                    "置信度：低\n要重新提交："
                    f"核对PDF第{source_page}页附件第{contract_line_no}行"
                    + (f"与Excel第{sales_row}行" if sales_row is not None else "并补齐对应Excel行")
                    + "的"
                    + "；".join(problem_groups or ["商品对应关系"])
                )
            ),
        ]
        _write_row(
            ws,
            row,
            attachment_row_values,
            height=_display_row_height(attachment_row_values),
            total=True,
            font_size=9,
        )
        row += 1

    for unmatched_sales in attachment_sales.get("unmatched_sales_records") or []:
        excel_row = int(unmatched_sales.get("excel_row") or 0)
        sales_lines = [
            f"客户名称：{unmatched_sales.get('customer_name') or '未识别'}",
            f"业务日期：{unmatched_sales.get('period_text') or '未识别'}",
            f"商品编码：{unmatched_sales.get('product_code') or '未识别'}",
            f"商品名称：{unmatched_sales.get('product_name') or '未识别'}",
            f"条形码：{unmatched_sales.get('barcode') or '未识别'}",
            f"单位：{unmatched_sales.get('unit') or '未识别'}",
            f"数量：{_number(unmatched_sales.get('quantity'))}",
            f"零售价：{_number(unmatched_sales.get('retail_price'))}元",
            f"合计金额：{_number(unmatched_sales.get('total_amount'))}元",
        ]
        unmatched_values = [
            f"合同销售附件未找到Excel第{excel_row}行的对应基准",
            "现场照片不参与本行销售明细核对",
            "\n".join([f"销售Excel第{excel_row}行", *sales_lines]),
            "合同附件 → 销售Excel：无法完成；该Excel行没有唯一合同附件行",
            "",
            f"置信度：低\n要重新提交：核实Excel第{excel_row}行对应的合同附件行，不能用知识库或照片替代合同",
        ]
        _write_row(
            ws,
            row,
            unmatched_values,
            height=_display_row_height(unmatched_values),
            total=True,
            font_size=9,
        )
        row += 1

    for item in result["store_reconciliation"]:
        photo_names = "、".join(Path(str(value)).name for value in item["photo_files"]) or "未提交照片"
        period = period_labels.get(item["period_match"], str(item["period_match"]))
        store = store_labels.get(item["store_match"], str(item["store_match"]))
        display = display_labels.get(item["display_match"], str(item["display_match"]))
        display_basis = display_basis_labels.get(
            item.get("display_standard_basis"),
            str(item.get("display_standard_basis") or "未提供陈列标准依据"),
        )
        display_description = _clip(
            item.get("display_description") or "未提供可见陈列依据",
            130,
        )
        vertical_count = item.get("display_vertical_facing_count")
        vertical_count_text = "无法可靠计数" if vertical_count is None else str(vertical_count)
        duplicate = duplicate_labels.get(item["duplicate_check"], str(item["duplicate_check"]))
        photo_knowledge = photo_knowledge_labels.get(
            item.get("photo_knowledge_match"),
            str(item.get("photo_knowledge_match") or "未对应知识库"),
        )
        store_record = contract_stores.get(int(item["store_line_no"]), {})
        contract_address = str(store_record.get("address") or "合同未列地址")
        failed_controls = _failed_display_controls(
            item,
            contract,
            contract_knowledge,
            contract_core,
            contract_sales,
        )
        all_failed_controls.extend(failed_controls)
        failed_text = "、".join(failed_controls) if failed_controls else "无"
        resubmission_items = _display_resubmission_items(
            item,
            contract,
            contract_knowledge,
            contract_core,
            contract_sales,
        )
        resubmission_text = "\n".join(
            f"{index}. {value}"
            for index, value in enumerate(resubmission_items, 1)
        )
        promotion_result = (
            "符合"
            if not contract.get("requires_promotion") or item.get("promotion_present")
            else "不符合"
        )
        contract_text = (
            f"门店：{item['contract_store_name']}\n"
            f"地址：{contract_address}\n"
            f"本店堆头：{_number(store_record.get('stack_count'))}\n"
            "本行合同基准：门店、周期、活动内容、陈列/促销要求及合同已确认商品"
        )
        visible_text = _short_list(
            list(item.get("visible_text") or []),
            limit=6,
            empty="未识别到有效文字",
        )
        visual_basis = _clip(display_description, 120)
        photo_column_text = (
            f"文件：{photo_names}\n"
            f"可见文字：{visible_text}\n"
            f"识别日期：{item.get('visible_date') or '未识别'}\n"
            f"识别地点：{item.get('visible_location') or '未识别'}\n"
            f"陈列标准核验：{display}（{display_basis}）\n"
            f"视觉依据：{visual_basis}\n"
            f"可见纵列数：{vertical_count_text}\n"
            f"照片复用检查：{duplicate}\n"
            f"现场商品：{_photo_product_text(item)}\n"
            f"知识库结论：{photo_knowledge}\n"
            f"促销信息：{item.get('promotion_summary') or '未识别'}"
        )
        sales_column_text = (
            "销售Excel不作为本店现场证据，也不与现场照片互相核对。\n"
            "逐行核销统一见上方“合同销售附件”区：销售Excel只向合同附件对齐。"
        )
        contract_product_result = {
            "not_applicable": "不适用",
            "pass": "具体商品已对上知识库",
            "fail": "具体商品未对上知识库",
        }.get(str(contract_knowledge.get("status")), "未核验")
        comparison_column_text = (
            "合同 → 现场："
            + _contract_photo_text(
                item,
                contract,
                contract_knowledge,
                store=store,
                period=period,
                display=display,
                promotion_result=promotion_result,
                contract_product_result=contract_product_result,
            )
            + "\n"
            f"现场照片文字及包装关键特征 → 商品知识库：{photo_knowledge}\n"
            f"已确认现场商品 → 合同商品范围：{_photo_contract_product_text(item)}\n"
            "核销边界：现场照片不参与销售明细核销；销售Excel仅核对合同附件"
        )
        units_text = _number(item.get("claim_units"))
        calculation_text = (
            f"{units_text} × {_number(summary['fee_per_store'])}元"
            if item.get("amount_rule_status") == "pass"
            else "不自动分摊，转人工确认"
        )
        amount_column_text = (
            f"口径：{_clip(item.get('amount_rule_basis') or '未识别', 90)}\n"
            f"计算：{calculation_text}\n"
            f"本店支持：{_number(item['supported_amount'])}元"
        )
        conclusion_column_text = (
            "结论：通过\n无需重新提交"
            if item["status"] == "pass"
            else (
                f"结论：暂不能核销\n主要问题：{failed_text}\n"
                f"要重新提交：\n{resubmission_text or '1. 补交能看清问题项的原始材料'}"
            )
        )
        row_values = [
            contract_text,
            photo_column_text,
            sales_column_text,
            comparison_column_text,
            amount_column_text,
            conclusion_column_text,
        ]
        _write_row(
            ws,
            row,
            row_values,
            height=_display_row_height(row_values),
            font_size=9,
        )
        row += 1

    contract_core_issue = contract_core.get("status") != "pass"
    contract_sales_issue = contract_sales.get("status") != "pass"
    attachment_issue = (
        attachment_knowledge.get("status") not in {"not_applicable", "pass"}
        or attachment_sales.get("status") != "pass"
    )
    sales_internal_issue = sales_internal_status != "pass"
    if (
        summary["supplement_store_count"] == 0
        and not contract_core_issue
        and not contract_sales_issue
        and not attachment_issue
        and not sales_internal_issue
    ):
        conclusion = "通过"
    elif (
        summary["passed_store_count"]
        and summary["supplement_store_count"]
        and not contract_core_issue
        and not contract_sales_issue
        and not attachment_issue
        and not sales_internal_issue
    ):
        conclusion = "部分通过"
    else:
        conclusion = "暂不能核销"
    retail = _number(summary.get("sales_retail_amount"))
    failed_control_summary = "、".join(dict.fromkeys(all_failed_controls)) or "无"
    if contract_core_issue:
        failed_control_summary = (
            failed_control_summary + "、" if failed_control_summary != "无" else ""
        ) + "合同核心六项"
    if contract_sales_issue:
        failed_control_summary = (
            failed_control_summary + "、" if failed_control_summary != "无" else ""
        ) + "销售Excel对合同签订方/执行周期/盖章"
    if attachment_issue:
        failed_control_summary = (
            failed_control_summary + "、" if failed_control_summary != "无" else ""
        ) + "合同销售附件"
    if sales_internal_issue:
        failed_control_summary = (
            failed_control_summary + "、" if failed_control_summary != "无" else ""
        ) + "销售Excel文件内部"
    _write_row(
        ws,
        row,
        [
            f"合计金额\n合同{summary['contract_store_count']}家；预算{_number(summary.get('activity_budget'))}元；申报{_number(summary['claimed_amount'])}元",
            f"共识别{summary['photo_count']}张照片",
            f"Excel读取{summary['sales_sku_count']}行商品\n数量{_number(summary['sales_quantity'])}；零售额{retail}元\nExcel内部核对：{sales_internal_text}\n销售明细只对合同附件核销",
            f"合同核心六项：{_control_status_text(contract_core.get('status'))}\n合同附件 → 销售Excel：{_control_status_text(attachment_sales.get('status'))}\n{summary['passed_store_count']}家通过；{summary['supplement_store_count']}家需补证\n未通过字段汇总：{failed_control_summary}",
            f"建议核销{_number(summary['suggested_approved_amount'])}元\n暂缓{_number(summary['temporarily_held_amount'])}元",
            conclusion,
        ],
        height=108,
        total=True,
        font_size=9,
    )
    ws.freeze_panes = "A4"
    ws.auto_filter.ref = f"A3:F{row}"
    widths = {"A": 40, "B": 54, "C": 52, "D": 48, "E": 28, "F": 42}
    for column, width in widths.items():
        ws.column_dimensions[column].width = width
    _sheet_base(ws, zoom=75, tab_color="ED7D31")
    return ws


def _poster_row_height(values: list[str]) -> float:
    longest = max(
        (
            len(line)
            for value in values
            for line in str(value or "").splitlines()
        ),
        default=0,
    )
    explicit_lines = max(
        (len(str(value or "").splitlines()) for value in values),
        default=1,
    )
    wrapped_lines = max(explicit_lines, (longest // 30) + explicit_lines)
    return min(300.0, max(66.0, 15.0 * wrapped_lines))


def _add_poster_material_sheet(wb: Workbook, result: dict[str, Any]) -> Any:
    ws = wb.create_sheet("海报物料核销")
    summary = result["summary"]
    audit = result["poster_material_audit"]
    issues = list(audit.get("issues") or [])
    note = (
        "本页只列影响核销的错误，不重复展示已满足项。"
        f"申报金额{_number(summary.get('claimed_amount'))}元；"
        f"建议暂缓{_number(summary.get('temporarily_held_amount'))}元；"
        f"现场照片{_number(summary.get('photo_count'))}张。"
    )
    _style_title(ws, "海报/物料制作核销｜只显示错误", note)
    _write_header(
        ws,
        [
            "错误来源",
            "已识别内容",
            "合同/规则要求",
            "核验结果",
            "核销影响",
            "需要补交",
        ],
        height=36,
    )

    row = 4
    if issues:
        confidence_labels = {"high": "高", "medium": "中", "low": "低"}
        for issue in issues:
            sources = "、".join(str(value) for value in issue.get("source_files") or [])
            values = [
                f"错误项：{issue['title']}\n来源：{sources or '未识别'}",
                f"识别内容：{issue['observed']}",
                f"合同/规则要求：{issue['expected']}",
                "核验结果：不满足\n"
                f"置信度：{confidence_labels.get(str(issue.get('confidence')), '中')}",
                f"核销影响：{issue['impact']}",
                "暂不能核销\n"
                f"要重新提交什么：{issue['resubmission']}",
            ]
            _write_row(
                ws,
                row,
                values,
                height=_poster_row_height(values),
                font_size=9,
            )
            row += 1
    else:
        values = [
            "未发现影响核销的错误",
            "本批材料未形成阻断项",
            "合同、票据、结算单和现场照片要求均已满足",
            "核验结果：满足\n置信度：高",
            "可按已确认金额进入后续流程",
            "无需补交",
        ]
        _write_row(ws, row, values, height=66, font_size=9)

    ws.freeze_panes = "A4"
    ws.auto_filter.ref = f"A3:F{ws.max_row}"
    widths = {"A": 42, "B": 66, "C": 54, "D": 24, "E": 44, "F": 58}
    for column, width in widths.items():
        ws.column_dimensions[column].width = width
    _sheet_base(ws, zoom=78, tab_color="C65911")
    return ws


def _add_other_expense_sheet(wb: Workbook, result: dict[str, Any]) -> Any:
    ws = wb.create_sheet("其他费用核销")
    summary = result["summary"]
    audit = result["other_expense_audit"]
    issues = list(audit.get("issues") or [])
    decision = str(summary.get("decision_label") or "")
    if decision == "类型归类错误":
        decision_action = "应回归已有费用类型，按相应规则拆分或重新提交"
    elif decision == "需特殊审批":
        decision_action = "须先取得批准新增费用类型的独立特殊审批"
    elif decision == "资料需补正":
        decision_action = "须按问题文件补正后再进入人工特殊审批"
    else:
        decision_action = "须由有权审批人核定新增类型及最终金额"
    note = (
        "其他费用不是兜底分类：先排除现有费用类型，再检查特殊审批和基础资料。"
        f"当前结论：{decision}；"
        f"申报金额{_number(summary.get('claimed_amount'))}元；"
        f"系统自动建议金额0元；{decision_action}。"
    )
    _style_title(ws, "其他费用核销｜特殊审批通道", note)
    _write_header(
        ws,
        [
            "问题文件",
            "已识别内容",
            "分类/审批要求",
            "审核结论",
            "核销影响",
            "处理方式",
        ],
        height=36,
    )

    confidence_labels = {"high": "高", "medium": "中", "low": "低"}
    row = 4
    for issue in issues:
        sources = "、".join(str(value) for value in issue.get("source_files") or [])
        values = [
            f"问题：{issue['title']}\n文件：{sources or '未识别'}",
            f"识别结果：{issue['observed']}",
            f"分类/审批要求：{issue['expected']}",
            f"审核结论：{summary.get('decision_label')}\n"
            f"置信度：{confidence_labels.get(str(issue.get('confidence')), '中')}",
            f"核销影响：{issue['impact']}\n自动建议核销：0元",
            f"处理方式：{issue['resubmission']}",
        ]
        _write_row(
            ws,
            row,
            values,
            height=_poster_row_height(values),
            font_size=9,
        )
        row += 1

    ws.freeze_panes = "A4"
    ws.auto_filter.ref = f"A3:F{ws.max_row}"
    widths = {"A": 48, "B": 70, "C": 58, "D": 28, "E": 46, "F": 62}
    for column, width in widths.items():
        ws.column_dimensions[column].width = width
    _sheet_base(ws, zoom=78, tab_color="8064A2")
    return ws


def _add_maintenance_fee_sheet(wb: Workbook, result: dict[str, Any]) -> Any:
    ws = wb.create_sheet("维护费用核销")
    summary = result["summary"]
    audit = result["maintenance_fee_audit"]
    issues = list(audit.get("issues") or [])
    note = (
        "本页只列影响维护费用核销的问题；其他费用仍按独立特殊审批规则审核。"
        f"当前结论：{summary.get('decision_label')}；"
        f"申报金额{_number(summary.get('claimed_amount'))}元；"
        f"建议核销{_number(summary.get('suggested_approved_amount'))}元；"
        f"暂缓{_number(summary.get('temporarily_held_amount'))}元。"
    )
    _style_title(ws, "维护费用核销｜只显示错误", note)
    _write_header(
        ws,
        [
            "问题来源",
            "已识别内容",
            "维护费用规则",
            "审核结论",
            "核销影响",
            "需要补交",
        ],
        height=36,
    )

    confidence_labels = {"high": "高", "medium": "中", "low": "低"}
    row = 4
    if issues:
        for issue in issues:
            sources = "、".join(str(value) for value in issue.get("source_files") or [])
            values = [
                f"问题：{issue['title']}\n文件：{sources or '本包缺失资料'}",
                f"识别结果：{issue['observed']}",
                f"规则要求：{issue['expected']}",
                "审核结论：资料需补正\n"
                f"置信度：{confidence_labels.get(str(issue.get('confidence')), '中')}",
                f"核销影响：{issue['impact']}\n暂不能核销",
                f"要重新提交什么：{issue['resubmission']}",
            ]
            _write_row(
                ws,
                row,
                values,
                height=_poster_row_height(values),
                font_size=9,
            )
            row += 1
    else:
        values = [
            "未发现影响维护费用核销的问题",
            "盖章POS、电子表、合同、结算和专项资料已形成闭环",
            "合同规则与POS电子表已完成确定性复算",
            "审核结论：可核销\n置信度：高",
            f"建议核销：{_number(summary.get('suggested_approved_amount'))}元",
            "无需补交",
        ]
        _write_row(ws, row, values, height=66, font_size=9)

    ws.freeze_panes = "A4"
    ws.auto_filter.ref = f"A3:F{ws.max_row}"
    widths = {"A": 48, "B": 70, "C": 60, "D": 28, "E": 48, "F": 62}
    for column, width in widths.items():
        ws.column_dimensions[column].width = width
    _sheet_base(ws, zoom=78, tab_color="2F75B5")
    return ws


def create_combined_report(results: list[dict[str, Any]], output_path: str | Path) -> Path:
    if not results:
        raise AuditError("没有可生成工作簿的核销结果")
    by_scenario: dict[str, dict[str, Any]] = {}
    for result in results:
        scenario = str(result.get("scenario") or "")
        if scenario in by_scenario:
            raise AuditError(f"同一工作簿不能包含两份同类结果：{scenario}")
        by_scenario[scenario] = result
    unknown = set(by_scenario) - {
        "personnel_incentive",
        "promotional_display",
        "poster_material",
        "other_expense",
        "maintenance_fee",
    }
    if unknown:
        raise AuditError("不支持的核销结果类型：" + "、".join(sorted(unknown)))

    workbook = Workbook()
    workbook.remove(workbook.active)
    if "personnel_incentive" in by_scenario:
        _add_personnel_sheet(workbook, by_scenario["personnel_incentive"])
    if "promotional_display" in by_scenario:
        _add_display_sheet(workbook, by_scenario["promotional_display"])
    if "poster_material" in by_scenario:
        _add_poster_material_sheet(workbook, by_scenario["poster_material"])
    if "other_expense" in by_scenario:
        _add_other_expense_sheet(workbook, by_scenario["other_expense"])
    if "maintenance_fee" in by_scenario:
        _add_maintenance_fee_sheet(workbook, by_scenario["maintenance_fee"])
    target = Path(output_path)
    target.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(target)
    workbook.close()
    return target


def verify_workbook(path: str | Path, scenarios: list[str]) -> dict[str, Any]:
    source = Path(path)
    workbook = load_workbook(source, data_only=False)
    expected = [
        name
        for scenario, name in (
            ("personnel_incentive", "人员激励核销"),
            ("promotional_display", "堆头核销"),
            ("poster_material", "海报物料核销"),
            ("other_expense", "其他费用核销"),
            ("maintenance_fee", "维护费用核销"),
        )
        if scenario in scenarios
    ]
    if workbook.sheetnames != expected:
        raise AuditError(f"工作表名称或顺序错误：{workbook.sheetnames}，期望{expected}")
    formulas: list[str] = []
    errors: list[str] = []
    for ws in workbook.worksheets:
        if ws.max_column != 6:
            raise AuditError(f"{ws.title} 不是6列结构")
        merges = {str(value) for value in ws.merged_cells.ranges}
        if not {"A1:F1", "A2:F2"}.issubset(merges):
            raise AuditError(f"{ws.title} 缺少两行合并标题")
        expected_freeze = "A4"
        if str(ws.freeze_panes) != expected_freeze:
            raise AuditError(f"{ws.title} 冻结窗格错误：{ws.freeze_panes}")
        if ws.auto_filter.ref != f"A3:F{ws.max_row}":
            raise AuditError(f"{ws.title} 筛选范围错误：{ws.auto_filter.ref}")
        for row in ws.iter_rows():
            for cell in row:
                if cell.data_type == "f" or (isinstance(cell.value, str) and cell.value.startswith("=")):
                    formulas.append(f"{ws.title}!{cell.coordinate}")
                if isinstance(cell.value, str) and cell.value in FORMULA_ERRORS:
                    errors.append(f"{ws.title}!{cell.coordinate}")
    workbook.close()
    if formulas:
        raise AuditError("输出工作簿不允许公式：" + "、".join(formulas))
    if errors:
        raise AuditError("输出工作簿含公式错误值：" + "、".join(errors))
    return {
        "path": str(source.resolve()),
        "sheet_names": expected,
        "formula_count": 0,
        "formula_error_count": 0,
    }
