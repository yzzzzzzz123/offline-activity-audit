from __future__ import annotations

import re
from decimal import Decimal
from math import cos, pi
from collections import defaultdict
from pathlib import Path
from statistics import median
from typing import Any

from openpyxl import load_workbook
from PIL import ExifTags, Image
from pypdf import PdfReader

from .common import AuditError, barcode_text, normalize_text, sha256_file


PERSONNEL_HEADERS = {
    "period": {"日期", "活动日期", "期间"},
    "store": {"门店名称", "门店", "店名"},
    "barcode": {"条码", "条形码", "商品条码"},
    "product": {"名称", "商品名称", "产品名称"},
    "quantity": {"数量", "销售数量", "销量"},
    "unit_price": {"单价", "零售价", "销售单价"},
    "amount": {"金额", "销售金额", "合计金额"},
}

DISPLAY_HEADERS = {
    "customer": {"客户名称", "客户"},
    "period": {"业务日期", "日期", "期间"},
    "product_code": {"商品编码", "产品编码"},
    "product": {"商品名称", "名称", "产品名称"},
    "barcode": {"条形码", "条码"},
    "unit": {"单位", "销售单位"},
    "quantity": {"数量", "销售数量", "销量"},
    "retail_price": {"零售价", "销售单价"},
    "total": {"合计金额", "零售金额", "金额"},
}

MAINTENANCE_POS_HEADERS = {
    "period": {"日期", "活动日期", "业务日期", "期间", "销售日期"},
    "store": {"门店名称", "门店", "店名", "客户名称", "客户"},
    "product_code": {"商品编码", "产品编码", "货号"},
    "barcode": {"条形码", "条码", "商品条码"},
    "product": {"货品名称", "商品名称", "产品名称", "名称"},
    "quantity": {"货品数量", "销售数量", "数量", "销量"},
    "unit_price": {"含税单价", "销售单价", "零售价", "单价"},
    "amount": {"含税售额", "销售金额", "合计金额", "零售金额", "金额"},
}


def _header_map(
    ws: Any,
    candidates: dict[str, set[str]],
    *,
    required: set[str] | None = None,
) -> tuple[int, dict[str, int]]:
    normalized = {
        key: {normalize_text(value) for value in values}
        for key, values in candidates.items()
    }
    for row in range(1, min(ws.max_row, 30) + 1):
        mapping: dict[str, int] = {}
        for col in range(1, ws.max_column + 1):
            value = normalize_text(ws.cell(row, col).value)
            for key, options in normalized.items():
                if key not in mapping and value in options:
                    mapping[key] = col
        required_fields = set(required or {"product", "quantity"})
        if required is None:
            if candidates is PERSONNEL_HEADERS:
                required_fields |= {"store", "barcode"}
            else:
                required_fields |= {"customer", "product_code", "barcode"}
        if required_fields.issubset(mapping):
            return row, mapping
    raise AuditError(f"Could not find required Excel headers in sheet {ws.title!r}")


def _number(value: Any) -> float | None:
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace(",", "")
    if not text or text.startswith("="):
        return None
    try:
        return float(text)
    except ValueError:
        return None


def _json_quantity(value: float) -> int | float:
    return int(value) if value.is_integer() else value


def _money_pair_status(left: Any, right: Any) -> str:
    if left is None or right is None:
        return "unverifiable"
    try:
        difference = abs(Decimal(str(left)) - Decimal(str(right)))
    except Exception:
        return "unverifiable"
    return "exact" if difference <= Decimal("0.01") else "mismatch"


def read_personnel_sales(path: str | Path) -> dict[str, Any]:
    source = Path(path).resolve()
    workbook = load_workbook(source, read_only=True, data_only=False)
    worksheet = workbook[workbook.sheetnames[0]]
    header_row, columns = _header_map(worksheet, PERSONNEL_HEADERS)
    records: list[dict[str, Any]] = []
    for row in range(header_row + 1, worksheet.max_row + 1):
        store = worksheet.cell(row, columns["store"]).value
        barcode = barcode_text(worksheet.cell(row, columns["barcode"]).value)
        product = worksheet.cell(row, columns["product"]).value
        quantity = _number(worksheet.cell(row, columns["quantity"]).value)
        if not store or not barcode or not product or quantity is None:
            continue
        if normalize_text(store) in {"合计", "总计"}:
            continue
        record = {
            "excel_row": row,
            "period_text": str(worksheet.cell(row, columns.get("period", 0)).value or ""),
            "store_name": str(store).strip(),
            "barcode": barcode,
            "product_name": str(product).strip(),
            "quantity": _json_quantity(quantity),
            "unit_price": (
                _number(worksheet.cell(row, columns["unit_price"]).value)
                if "unit_price" in columns
                else None
            ),
            "amount_cell": (
                worksheet.cell(row, columns["amount"]).value
                if "amount" in columns
                else None
            ),
        }
        records.append(record)
    if not records:
        raise AuditError(f"No valid sales detail rows found in {source}")

    sku_acc: dict[str, dict[str, Any]] = {}
    store_acc: dict[str, dict[str, Any]] = {}
    for record in records:
        barcode = record["barcode"]
        sku = sku_acc.setdefault(
            barcode,
            {
                "barcode": barcode,
                "product_name": record["product_name"],
                "quantity": 0.0,
                "excel_rows": [],
                "store_quantities": defaultdict(float),
            },
        )
        if normalize_text(sku["product_name"]) != normalize_text(record["product_name"]):
            raise AuditError(f"Barcode {barcode} has conflicting product names")
        sku["quantity"] += float(record["quantity"])
        sku["excel_rows"].append(record["excel_row"])
        sku["store_quantities"][record["store_name"]] += float(record["quantity"])

        store = store_acc.setdefault(
            record["store_name"],
            {"store_name": record["store_name"], "quantity": 0.0, "excel_rows": []},
        )
        store["quantity"] += float(record["quantity"])
        store["excel_rows"].append(record["excel_row"])

    skus: list[dict[str, Any]] = []
    for sku in sku_acc.values():
        sku["quantity"] = _json_quantity(sku["quantity"])
        sku["store_quantities"] = [
            {"store_name": name, "quantity": _json_quantity(quantity)}
            for name, quantity in sorted(sku["store_quantities"].items())
        ]
        skus.append(sku)
    stores: list[dict[str, Any]] = []
    for store in store_acc.values():
        store["quantity"] = _json_quantity(store["quantity"])
        stores.append(store)
    skus.sort(key=lambda item: item["excel_rows"][0])
    stores.sort(key=lambda item: item["excel_rows"][0])
    total = sum(float(item["quantity"]) for item in records)
    result = {
        "source_file": str(source),
        "sheet": worksheet.title,
        "header_row": header_row,
        "detail_row_count": len(records),
        "sku_count": len(skus),
        "store_count": len(stores),
        "total_quantity": _json_quantity(total),
        "period_values": sorted({item["period_text"] for item in records if item["period_text"]}),
        "records": records,
        "skus": skus,
        "stores": stores,
    }
    workbook.close()
    return result


def read_display_sales(path: str | Path) -> dict[str, Any]:
    source = Path(path).resolve()
    formula_book = load_workbook(source, read_only=True, data_only=False)
    value_book = load_workbook(source, read_only=True, data_only=True)
    formula_ws = formula_book[formula_book.sheetnames[0]]
    value_ws = value_book[value_book.sheetnames[0]]
    header_row, columns = _header_map(formula_ws, DISPLAY_HEADERS)
    records: list[dict[str, Any]] = []
    external_formula_cells: list[str] = []
    printed_totals: list[dict[str, Any]] = []
    for row in range(header_row + 1, formula_ws.max_row + 1):
        formula_values = [
            formula_ws.cell(row, col).value
            for col in range(1, formula_ws.max_column + 1)
        ]
        for col, value in enumerate(formula_values, 1):
            if isinstance(value, str) and value.startswith("=") and "[" in value:
                external_formula_cells.append(formula_ws.cell(row, col).coordinate)
        normalized_cells = {
            normalize_text(value)
            for value in formula_values
            if value is not None
        }
        if normalized_cells & {"合计", "总计"}:
            printed_totals.append(
                {
                    "excel_row": row,
                    "quantity": _number(value_ws.cell(row, columns["quantity"]).value),
                    "total_amount": (
                        _number(value_ws.cell(row, columns["total"]).value)
                        if "total" in columns
                        else None
                    ),
                }
            )
            continue
        customer = formula_ws.cell(row, columns["customer"]).value
        product = formula_ws.cell(row, columns["product"]).value
        quantity = _number(value_ws.cell(row, columns["quantity"]).value)
        if not customer or not product or quantity is None:
            continue
        total_value = (
            _number(value_ws.cell(row, columns["total"]).value)
            if "total" in columns
            else None
        )
        retail_price = (
            _number(value_ws.cell(row, columns["retail_price"]).value)
            if "retail_price" in columns
            else None
        )
        calculated_total = (
            quantity * retail_price
            if retail_price is not None
            else None
        )
        line_amount_status = _money_pair_status(calculated_total, total_value)
        unit_column = columns.get("unit", 0)
        unit_value = value_ws.cell(row, unit_column).value if unit_column else None
        if unit_value is None and unit_column:
            formula_unit_value = formula_ws.cell(row, unit_column).value
            if not (
                isinstance(formula_unit_value, str)
                and formula_unit_value.startswith("=")
            ):
                unit_value = formula_unit_value
        records.append(
            {
                "excel_row": row,
                "customer_name": str(customer).strip(),
                "period_text": str(formula_ws.cell(row, columns.get("period", 0)).value or ""),
                "product_code": str(formula_ws.cell(row, columns.get("product_code", 0)).value or ""),
                "barcode": barcode_text(formula_ws.cell(row, columns.get("barcode", 0)).value),
                "product_name": str(product).strip(),
                "unit": str(unit_value or "").strip(),
                "quantity": _json_quantity(quantity),
                "retail_price": retail_price,
                "total_amount": total_value,
                "calculated_total_amount": (
                    _json_quantity(calculated_total)
                    if calculated_total is not None
                    else None
                ),
                "line_amount_status": line_amount_status,
                "line_amount_difference": (
                    _json_quantity(total_value - calculated_total)
                    if total_value is not None and calculated_total is not None
                    else None
                ),
            }
        )
    if not records:
        raise AuditError(f"No valid sales detail rows found in {source}")
    total_quantity = sum(float(item["quantity"]) for item in records)
    known_total_amounts = [item["total_amount"] for item in records if item["total_amount"] is not None]
    calculated_total_amount = (
        sum(known_total_amounts)
        if len(known_total_amounts) == len(records)
        else None
    )
    printed_total = printed_totals[0] if len(printed_totals) == 1 else None
    printed_quantity_status = (
        _money_pair_status(printed_total.get("quantity"), total_quantity)
        if printed_total is not None
        else "unverifiable"
    )
    printed_amount_status = (
        _money_pair_status(printed_total.get("total_amount"), calculated_total_amount)
        if printed_total is not None
        else "unverifiable"
    )
    line_problem_rows = [
        int(item["excel_row"])
        for item in records
        if item["line_amount_status"] != "exact"
    ]
    internal_problem_rows = list(line_problem_rows)
    if len(printed_totals) != 1:
        internal_problem_rows.extend(int(item["excel_row"]) for item in printed_totals)
    elif printed_quantity_status != "exact" or printed_amount_status != "exact":
        internal_problem_rows.append(int(printed_total["excel_row"]))
    internal_problem_rows = sorted(set(internal_problem_rows))
    internal_statuses = {
        *(str(item["line_amount_status"]) for item in records),
        printed_quantity_status,
        printed_amount_status,
    }
    if "mismatch" in internal_statuses:
        internal_status = "fail"
    elif "unverifiable" in internal_statuses or len(printed_totals) != 1:
        internal_status = "unverifiable"
    else:
        internal_status = "pass"
    result = {
        "source_file": str(source),
        "sheet": formula_ws.title,
        "header_row": header_row,
        "sku_count": len(records),
        "total_quantity": _json_quantity(total_quantity),
        "retail_amount": (
            _json_quantity(calculated_total_amount)
            if calculated_total_amount is not None
            else None
        ),
        "printed_total_rows": printed_totals,
        "printed_total_row": (
            int(printed_total["excel_row"]) if printed_total is not None else None
        ),
        "printed_total_quantity": (
            _json_quantity(float(printed_total["quantity"]))
            if printed_total is not None and printed_total.get("quantity") is not None
            else None
        ),
        "printed_total_amount": (
            _json_quantity(float(printed_total["total_amount"]))
            if printed_total is not None and printed_total.get("total_amount") is not None
            else None
        ),
        "calculated_total_quantity": _json_quantity(total_quantity),
        "calculated_total_amount": (
            _json_quantity(calculated_total_amount)
            if calculated_total_amount is not None
            else None
        ),
        "line_amount_status": (
            "fail"
            if any(item["line_amount_status"] == "mismatch" for item in records)
            else "unverifiable"
            if line_problem_rows
            else "pass"
        ),
        "line_amount_problem_rows": line_problem_rows,
        "printed_quantity_status": printed_quantity_status,
        "printed_amount_status": printed_amount_status,
        "internal_status": internal_status,
        "internal_problem_rows": internal_problem_rows,
        "internal_basis": (
            f"销售Excel明细{len(records)}行逐行核对数量×零售价=合计金额；"
            f"行金额{'全部一致' if not line_problem_rows else '问题行' + str(line_problem_rows)}；"
            f"打印总数量{'一致' if printed_quantity_status == 'exact' else '不一致或无法核对'}，"
            f"打印合计金额{'一致' if printed_amount_status == 'exact' else '不一致或无法核对'}。"
        ),
        "customers": sorted({item["customer_name"] for item in records}),
        "period_values": sorted({item["period_text"] for item in records if item["period_text"]}),
        "external_formula_cells": sorted(set(external_formula_cells)),
        "records": records,
        "store_level_available": False,
    }
    formula_book.close()
    value_book.close()
    return result


def read_maintenance_pos(path: str | Path) -> dict[str, Any]:
    """Read one POS workbook without exposing it to the visual extraction stage."""

    source = Path(path).resolve()
    formula_book = load_workbook(source, read_only=True, data_only=False)
    value_book = load_workbook(source, read_only=True, data_only=True)
    try:
        formula_ws = formula_book[formula_book.sheetnames[0]]
        value_ws = value_book[value_book.sheetnames[0]]
        header_row, columns = _header_map(
            formula_ws,
            MAINTENANCE_POS_HEADERS,
            required={"product", "quantity", "amount"},
        )
        records: list[dict[str, Any]] = []
        printed_totals: list[dict[str, Any]] = []
        external_formula_cells: list[str] = []
        formula_cells: list[str] = []
        for row in range(header_row + 1, formula_ws.max_row + 1):
            formula_values = [
                formula_ws.cell(row, col).value
                for col in range(1, formula_ws.max_column + 1)
            ]
            for col, value in enumerate(formula_values, 1):
                if isinstance(value, str) and value.startswith("="):
                    coordinate = formula_ws.cell(row, col).coordinate
                    formula_cells.append(coordinate)
                    if "[" in value:
                        external_formula_cells.append(coordinate)
            normalized_cells = {
                normalize_text(value)
                for value in formula_values
                if value is not None
            }
            if normalized_cells & {"合计", "总计"}:
                printed_totals.append(
                    {
                        "excel_row": row,
                        "quantity": _number(value_ws.cell(row, columns["quantity"]).value),
                        "sales_amount": _number(value_ws.cell(row, columns["amount"]).value),
                    }
                )
                continue

            product = formula_ws.cell(row, columns["product"]).value
            quantity = _number(value_ws.cell(row, columns["quantity"]).value)
            amount = _number(value_ws.cell(row, columns["amount"]).value)
            if not product or quantity is None or amount is None:
                continue
            unit_price = (
                _number(value_ws.cell(row, columns["unit_price"]).value)
                if "unit_price" in columns
                else None
            )
            calculated_amount = quantity * unit_price if unit_price is not None else None
            period_value = (
                formula_ws.cell(row, columns["period"]).value
                if "period" in columns
                else None
            )
            store_value = (
                formula_ws.cell(row, columns["store"]).value
                if "store" in columns
                else None
            )
            product_code_value = (
                formula_ws.cell(row, columns["product_code"]).value
                if "product_code" in columns
                else None
            )
            barcode_value = (
                formula_ws.cell(row, columns["barcode"]).value
                if "barcode" in columns
                else None
            )
            records.append(
                {
                    "excel_row": row,
                    "period_text": str(period_value or "").strip(),
                    "store_name": str(store_value or "").strip(),
                    "product_code": str(product_code_value or "").strip(),
                    "barcode": barcode_text(barcode_value),
                    "product_name": str(product).strip(),
                    "quantity": _json_quantity(quantity),
                    "unit_price": unit_price,
                    "sales_amount": amount,
                    "calculated_sales_amount": calculated_amount,
                    "line_amount_status": _money_pair_status(calculated_amount, amount),
                }
            )
        if not records:
            raise AuditError(f"POS电子表没有可读取的商品、数量和销售金额明细：{source.name}")

        total_quantity = sum(float(item["quantity"]) for item in records)
        total_sales_amount = sum(float(item["sales_amount"]) for item in records)
        printed_total = printed_totals[0] if len(printed_totals) == 1 else None
        printed_quantity_status = (
            _money_pair_status(printed_total.get("quantity"), total_quantity)
            if printed_total is not None
            else "unverifiable"
        )
        printed_amount_status = (
            _money_pair_status(printed_total.get("sales_amount"), total_sales_amount)
            if printed_total is not None
            else "unverifiable"
        )
        return {
            "source_file": source.name,
            "sheet": formula_ws.title,
            "header_row": header_row,
            "detail_row_count": len(records),
            "total_quantity": _json_quantity(total_quantity),
            "total_sales_amount": _json_quantity(total_sales_amount),
            "printed_total_rows": printed_totals,
            "printed_quantity_status": printed_quantity_status,
            "printed_amount_status": printed_amount_status,
            "period_values": sorted(
                {item["period_text"] for item in records if item["period_text"]}
            ),
            "store_values": sorted(
                {item["store_name"] for item in records if item["store_name"]}
            ),
            "formula_cells": sorted(set(formula_cells)),
            "external_formula_cells": sorted(set(external_formula_cells)),
            "records": records,
        }
    finally:
        formula_book.close()
        value_book.close()


def read_pos_target_summary(path: str | Path) -> dict[str, Any]:
    """Read the unique store-period-sales summary used by POS target incentives."""

    source = Path(path).resolve()
    formula_book = load_workbook(source, read_only=True, data_only=False)
    value_book = load_workbook(source, read_only=True, data_only=True)
    try:
        selected: tuple[Any, Any, int, dict[str, int]] | None = None
        aliases = {
            "period": {"日期", "期间", "活动时间", "销售期间"},
            "store": {"系统门店名称", "门店名称", "门店", "客户名称"},
            "amount": {"销售金额", "销售额", "销售合计", "金额"},
        }
        for sheet_name in formula_book.sheetnames:
            formula_ws = formula_book[sheet_name]
            value_ws = value_book[sheet_name]
            scan_rows = min(int(formula_ws.max_row), 50)
            scan_columns = min(int(formula_ws.max_column), 30)
            for row in range(1, scan_rows + 1):
                columns: dict[str, int] = {}
                for column in range(1, scan_columns + 1):
                    text = normalize_text(formula_ws.cell(row, column).value)
                    for field, candidates in aliases.items():
                        if text in {normalize_text(value) for value in candidates}:
                            columns.setdefault(field, column)
                if set(columns) == {"period", "store", "amount"}:
                    if selected is not None:
                        raise AuditError(
                            f"POS电子表存在多个门店销售汇总候选：{selected[0].title}、{sheet_name}"
                        )
                    selected = (formula_ws, value_ws, row, columns)
                    break
        if selected is None:
            raise AuditError("POS电子表未找到日期、系统门店名称、销售金额三列表头")
        formula_ws, value_ws, header_row, columns = selected
        records: list[dict[str, Any]] = []
        printed_totals: list[dict[str, Any]] = []
        formula_cells: list[str] = []
        external_formula_cells: list[str] = []
        last_row = min(int(formula_ws.max_row), header_row + 10_000)
        for row in range(header_row + 1, last_row + 1):
            raw_values = [
                formula_ws.cell(row, column).value
                for column in range(1, int(formula_ws.max_column) + 1)
            ]
            for column, raw in enumerate(raw_values, 1):
                if isinstance(raw, str) and raw.startswith("="):
                    coordinate = formula_ws.cell(row, column).coordinate
                    formula_cells.append(coordinate)
                    if "[" in raw:
                        external_formula_cells.append(coordinate)
            normalized = {normalize_text(value) for value in raw_values if value is not None}
            amount = _number(value_ws.cell(row, columns["amount"]).value)
            if normalized & {"合计", "总计"}:
                printed_totals.append(
                    {"excel_row": row, "sales_amount": _json_quantity(amount) if amount is not None else None}
                )
                continue
            period_value = formula_ws.cell(row, columns["period"]).value
            store_value = formula_ws.cell(row, columns["store"]).value
            if period_value is None and store_value is None and amount is None:
                continue
            if not store_value or amount is None:
                continue
            records.append(
                {
                    "excel_row": row,
                    "period_text": str(period_value or "").strip(),
                    "store_name": str(store_value).strip(),
                    "sales_amount": _json_quantity(amount),
                }
            )
        if not records:
            raise AuditError(f"POS电子表没有可读取的门店销售金额明细：{source.name}")
        stores = [normalize_text(item["store_name"]) for item in records]
        if len(stores) != len(set(stores)):
            raise AuditError("POS电子表门店汇总存在重复门店，无法唯一对应盖章POS")
        total = sum(float(item["sales_amount"]) for item in records)
        printed_total = printed_totals[0] if len(printed_totals) == 1 else None
        printed_status = (
            _money_pair_status(printed_total.get("sales_amount"), total)
            if printed_total is not None
            else "unverifiable"
        )
        return {
            "source_file": source.name,
            "sheet": formula_ws.title,
            "header_row": header_row,
            "detail_row_count": len(records),
            "store_count": len(records),
            "total_sales_amount": _json_quantity(total),
            "printed_total_rows": printed_totals,
            "printed_amount_status": printed_status,
            "period_values": sorted({item["period_text"] for item in records if item["period_text"]}),
            "formula_cells": sorted(set(formula_cells)),
            "external_formula_cells": sorted(set(external_formula_cells)),
            "records": records,
        }
    finally:
        formula_book.close()
        value_book.close()


def read_self_procured_pos_summary(path: str | Path) -> dict[str, Any]:
    """Read one unique store-period-quantity-amount POS summary for gift-material audit."""

    source = Path(path).resolve()
    formula_book = load_workbook(source, read_only=True, data_only=False)
    value_book = load_workbook(source, read_only=True, data_only=True)
    try:
        selected: tuple[Any, Any, int, dict[str, int]] | None = None
        aliases = {
            "period": {"日期", "期间", "活动时间", "活动周期", "销售期间"},
            "store": {"系统门店名称", "门店名称", "门店", "客户名称"},
            "quantity": {"销售数量", "销量", "数量"},
            "amount": {"销售金额", "销售收入", "销售额", "销售合计", "金额"},
        }
        normalized_aliases = {
            field: {normalize_text(value) for value in values}
            for field, values in aliases.items()
        }
        for sheet_name in formula_book.sheetnames:
            formula_ws = formula_book[sheet_name]
            value_ws = value_book[sheet_name]
            for row in range(1, min(int(formula_ws.max_row), 50) + 1):
                columns: dict[str, int] = {}
                for column in range(1, min(int(formula_ws.max_column), 30) + 1):
                    cell_text = normalize_text(formula_ws.cell(row, column).value)
                    for field, candidates in normalized_aliases.items():
                        if cell_text in candidates:
                            columns.setdefault(field, column)
                if set(columns) == {"period", "store", "quantity", "amount"}:
                    if selected is not None:
                        raise AuditError(
                            f"POS电子表存在多个门店数量金额汇总候选：{selected[0].title}、{sheet_name}"
                        )
                    selected = (formula_ws, value_ws, row, columns)
                    break
        if selected is None:
            raise AuditError("POS电子表未找到日期、门店名称、销售数量、销售金额四列表头")
        formula_ws, value_ws, header_row, columns = selected
        records: list[dict[str, Any]] = []
        printed_totals: list[dict[str, Any]] = []
        formula_cells: list[str] = []
        external_formula_cells: list[str] = []
        for row in range(header_row + 1, min(int(formula_ws.max_row), header_row + 10_000) + 1):
            raw_values = [
                formula_ws.cell(row, column).value
                for column in range(1, int(formula_ws.max_column) + 1)
            ]
            for column, raw in enumerate(raw_values, 1):
                if isinstance(raw, str) and raw.startswith("="):
                    coordinate = formula_ws.cell(row, column).coordinate
                    formula_cells.append(coordinate)
                    if "[" in raw:
                        external_formula_cells.append(coordinate)
            normalized = {normalize_text(value) for value in raw_values if value is not None}
            quantity = _number(value_ws.cell(row, columns["quantity"]).value)
            amount = _number(value_ws.cell(row, columns["amount"]).value)
            if normalized & {"合计", "总计"}:
                printed_totals.append(
                    {
                        "excel_row": row,
                        "sales_quantity": _json_quantity(quantity) if quantity is not None else None,
                        "sales_amount": _json_quantity(amount) if amount is not None else None,
                    }
                )
                continue
            period_value = formula_ws.cell(row, columns["period"]).value
            store_value = formula_ws.cell(row, columns["store"]).value
            if period_value is None and store_value is None and quantity is None and amount is None:
                continue
            if not store_value or quantity is None or amount is None:
                continue
            records.append(
                {
                    "excel_row": row,
                    "period_text": str(period_value or "").strip(),
                    "store_name": str(store_value).strip(),
                    "sales_quantity": _json_quantity(quantity),
                    "sales_amount": _json_quantity(amount),
                }
            )
        if not records:
            raise AuditError(f"POS电子表没有可读取的门店数量金额明细：{source.name}")
        normalized_stores = [normalize_text(item["store_name"]) for item in records]
        if len(normalized_stores) != len(set(normalized_stores)):
            raise AuditError("POS电子表门店汇总存在重复门店，无法唯一对应盖章POS")
        total_quantity = sum(float(item["sales_quantity"]) for item in records)
        total_amount = sum(float(item["sales_amount"]) for item in records)
        printed_total = printed_totals[0] if len(printed_totals) == 1 else None
        return {
            "source_file": source.name,
            "sheet": formula_ws.title,
            "header_row": header_row,
            "detail_row_count": len(records),
            "store_count": len(records),
            "total_sales_quantity": _json_quantity(total_quantity),
            "total_sales_amount": _json_quantity(total_amount),
            "printed_total_rows": printed_totals,
            "printed_quantity_status": (
                _money_pair_status(printed_total.get("sales_quantity"), total_quantity)
                if printed_total is not None
                else "unverifiable"
            ),
            "printed_amount_status": (
                _money_pair_status(printed_total.get("sales_amount"), total_amount)
                if printed_total is not None
                else "unverifiable"
            ),
            "period_values": sorted({item["period_text"] for item in records if item["period_text"]}),
            "formula_cells": sorted(set(formula_cells)),
            "external_formula_cells": sorted(set(external_formula_cells)),
            "records": records,
        }
    finally:
        formula_book.close()
        value_book.close()


def _difference_hash(path: Path) -> str:
    with Image.open(path) as image:
        gray = image.convert("L").resize((9, 8))
        pixels = list(gray.get_flattened_data())
    bits = []
    for row in range(8):
        offset = row * 9
        bits.extend(pixels[offset + col] > pixels[offset + col + 1] for col in range(8))
    value = sum((1 << index) for index, bit in enumerate(bits) if bit)
    return f"{value:016x}"


def _perceptual_hash(path: Path) -> str:
    size = 32
    low = 8
    with Image.open(path) as image:
        gray = image.convert("L").resize((size, size), Image.Resampling.LANCZOS)
        pixels = list(gray.get_flattened_data())
    cosine = [
        [cos((2 * position + 1) * frequency * pi / (2 * size)) for position in range(size)]
        for frequency in range(low)
    ]
    row_transform = [[0.0 for _ in range(low)] for _ in range(size)]
    for row in range(size):
        offset = row * size
        for horizontal_frequency in range(low):
            row_transform[row][horizontal_frequency] = sum(
                pixels[offset + column] * cosine[horizontal_frequency][column]
                for column in range(size)
            )
    coefficients: list[float] = []
    for vertical_frequency in range(low):
        for horizontal_frequency in range(low):
            coefficients.append(
                sum(
                    row_transform[row][horizontal_frequency]
                    * cosine[vertical_frequency][row]
                    for row in range(size)
                )
            )
    threshold = median(coefficients[1:])
    value = 0
    for coefficient in coefficients:
        value = (value << 1) | int(coefficient > threshold)
    return f"{value:016x}"


def _hamming_distance(left: str, right: str) -> int:
    return (int(left, 16) ^ int(right, 16)).bit_count()


def image_file_inventory(paths: list[str | Path], *, directory: str | Path | None = None) -> dict[str, Any]:
    resolved_paths = sorted(
        {Path(item).resolve() for item in paths},
        key=lambda item: item.name.casefold(),
    )
    files: list[dict[str, Any]] = []
    sha_groups: dict[str, list[str]] = defaultdict(list)
    dhash_groups: dict[str, list[str]] = defaultdict(list)
    phash_groups: dict[str, list[str]] = defaultdict(list)
    for path in resolved_paths:
        if not path.is_file():
            raise AuditError(f"Image file does not exist: {path}")
        sha = sha256_file(path)
        dhash = _difference_hash(path)
        phash = _perceptual_hash(path)
        exif_datetime = None
        gps_present = False
        try:
            with Image.open(path) as image:
                exif = image.getexif()
                for tag, value in exif.items():
                    name = ExifTags.TAGS.get(tag, str(tag))
                    if name in {"DateTimeOriginal", "DateTimeDigitized", "DateTime"} and value:
                        exif_datetime = str(value)
                    if name == "GPSInfo" and value:
                        gps_present = True
                width, height = image.size
        except Exception as exc:
            raise AuditError(f"Could not inspect image {path}: {exc}") from exc
        files.append(
            {
                "file_name": path.name,
                "absolute_path": str(path),
                "size": path.stat().st_size,
                "width": width,
                "height": height,
                "sha256": sha,
                "dhash": dhash,
                "phash": phash,
                "exif_datetime": exif_datetime,
                "gps_present": gps_present,
            }
        )
        sha_groups[sha].append(path.name)
        dhash_groups[dhash].append(path.name)
        phash_groups[phash].append(path.name)
    phash_candidate_pairs: list[dict[str, Any]] = []
    for index, left in enumerate(files):
        for right in files[index + 1 :]:
            distance = _hamming_distance(str(left["phash"]), str(right["phash"]))
            if distance <= 8:
                phash_candidate_pairs.append(
                    {
                        "files": [left["file_name"], right["file_name"]],
                        "distance": distance,
                    }
                )
    return {
        "directory": str(Path(directory).resolve()) if directory is not None else None,
        "file_count": len(files),
        "files": files,
        "exact_duplicate_groups": [group for group in sha_groups.values() if len(group) > 1],
        "same_dhash_groups": [group for group in dhash_groups.values() if len(group) > 1],
        "same_phash_groups": [group for group in phash_groups.values() if len(group) > 1],
        "phash_candidate_pairs": phash_candidate_pairs,
        "cross_activity_reuse_checked": False,
        "files_with_exif_datetime": sum(bool(item["exif_datetime"]) for item in files),
        "files_with_gps": sum(bool(item["gps_present"]) for item in files),
    }


def image_inventory(directory: str | Path) -> dict[str, Any]:
    root = Path(directory).resolve()
    if not root.is_dir():
        raise AuditError(f"Photo directory does not exist: {root}")
    paths = [
        item
        for item in root.iterdir()
        if item.suffix.lower() in {".jpg", ".jpeg", ".png", ".webp"}
    ]
    return image_file_inventory(paths, directory=root)


def pdf_inventory(path: str | Path) -> dict[str, Any]:
    source = Path(path).resolve()
    reader = PdfReader(source)
    metadata = reader.metadata or {}
    return {
        "source_file": str(source),
        "page_count": len(reader.pages),
        "sha256": sha256_file(source),
        "title": str(metadata.get("/Title") or ""),
    }
