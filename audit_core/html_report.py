from __future__ import annotations

import json
import re
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from .common import AuditError


SHEET_SCENARIOS = {
    "人员激励核销": "personnel_incentive",
    "堆头核销": "promotional_display",
}
SCENARIO_SHEETS = {value: key for key, value in SHEET_SCENARIOS.items()}
SUMMARY_PREFIXES = (
    "合计",
    "总计",
    "实际申报",
    "收款人与日期",
    "收款人/日期",
)
FORBIDDEN_VISIBLE_TERMS = (
    "详见审计JSON",
    "见内部结果",
    "候选命中",
    "商品ID",
    "视觉RAG",
    "唯一收敛",
)
DATA_OPEN = '<script id="audit-data" type="application/json">'
DATA_CLOSE = "</script>"


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).replace("\r\n", "\n").replace("\r", "\n").strip()


def _row_kind(values: list[str]) -> str:
    first = values[0].lstrip() if values else ""
    if first.startswith(SUMMARY_PREFIXES):
        return "summary"
    return "record"


def _row_status(values: list[str], kind: str) -> str:
    conclusion = values[5] if len(values) > 5 else ""
    no_resubmission = (
        "无需重新提交" in conclusion
        or "要重新提交：不用" in conclusion
        or "要重新提交什么：不用" in conclusion
    )
    if kind == "summary":
        if "要重新提交：" in conclusion and not no_resubmission:
            return "issue"
        return "summary"
    if no_resubmission and not any(
        token in conclusion for token in ("暂不能核销", "完全不匹配", "需补")
    ):
        return "pass"
    return "issue"


def _row_confidence(values: list[str]) -> str:
    text = "\n".join(values)
    if "置信度：低" in text:
        return "low"
    if "置信度：中" in text:
        return "medium"
    if "置信度：高" in text:
        return "high"
    return "not_applicable"


def _display_store_heading(values: list[str]) -> str | None:
    contract_text = values[0] if values else ""
    numbered_store = re.search(r"第\s*\d+\s*家[：:]\s*([^\n；]+)", contract_text)
    if numbered_store:
        return numbered_store.group(1).strip()
    current_store = re.search(
        r"本店(?!堆头)(?:[：:]\s*)?([^\n；]+)",
        contract_text,
    )
    if current_store:
        return current_store.group(1).strip()
    photo_text = values[1] if len(values) > 1 else ""
    visible_location = re.search(r"识别地点[：:]\s*([^\n]+)", photo_text)
    if visible_location:
        return visible_location.group(1).strip()
    return None


def _row_heading(
    values: list[str],
    row_number: int,
    kind: str,
    scenario: str,
) -> str:
    if kind == "summary":
        first = values[0] if values else ""
        for line in first.splitlines():
            label = line.strip()
            if label:
                return label[:90]
        return "汇总"
    if scenario == "promotional_display":
        store = _display_store_heading(values)
        if store:
            return store[:90]
    first = values[0] if values else ""
    if scenario == "personnel_incentive":
        title_lines = [line.strip() for line in first.splitlines() if line.strip()]
        if len(title_lines) >= 2 and title_lines[0].startswith("结算第"):
            return title_lines[1][:90]
    for line in first.splitlines():
        label = line.strip()
        if label:
            return label[:90]
    conclusion = values[5] if len(values) > 5 else ""
    for line in conclusion.splitlines():
        label = line.strip()
        if label:
            return label[:90]
    return f"核销记录 {row_number - 3}"


def _workbook_payload(workbook_path: str | Path) -> dict[str, Any]:
    source = Path(workbook_path)
    if not source.is_file():
        raise AuditError(f"找不到待生成HTML的工作簿：{source}")

    workbook = load_workbook(source, read_only=True, data_only=False)
    try:
        unknown = set(workbook.sheetnames) - set(SHEET_SCENARIOS)
        if unknown:
            raise AuditError("HTML不支持的工作表：" + "、".join(sorted(unknown)))
        if not workbook.sheetnames:
            raise AuditError("工作簿没有可展示的工作表")

        sheets: list[dict[str, Any]] = []
        for worksheet in workbook.worksheets:
            headers = [_cell_text(worksheet.cell(3, column).value) for column in range(1, 7)]
            if len(headers) != 6 or any(not header for header in headers):
                raise AuditError(f"{worksheet.title} 缺少完整六列表头，不能生成HTML")

            rows: list[dict[str, Any]] = []
            for row_number in range(4, worksheet.max_row + 1):
                values = [
                    _cell_text(worksheet.cell(row_number, column).value)
                    for column in range(1, 7)
                ]
                if not any(values):
                    continue
                kind = _row_kind(values)
                rows.append(
                    {
                        "excel_row": row_number,
                        "kind": kind,
                        "status": _row_status(values, kind),
                        "confidence": _row_confidence(values),
                        "heading": _row_heading(
                            values,
                            row_number,
                            kind,
                            SHEET_SCENARIOS[worksheet.title],
                        ),
                        "values": values,
                    }
                )

            sheets.append(
                {
                    "name": worksheet.title,
                    "scenario": SHEET_SCENARIOS[worksheet.title],
                    "title": _cell_text(worksheet["A1"].value) or worksheet.title,
                    "note": _cell_text(worksheet["A2"].value),
                    "headers": headers,
                    "rows": rows,
                }
            )
    finally:
        workbook.close()

    return {
        "schema_version": "1.0",
        "title": "线下活动核销结果",
        "sheets": sheets,
    }


def _json_for_script(payload: dict[str, Any]) -> str:
    return (
        json.dumps(payload, ensure_ascii=False, separators=(",", ":"))
        .replace("<", "\\u003c")
        .replace(">", "\\u003e")
        .replace("&", "\\u0026")
        .replace("\u2028", "\\u2028")
        .replace("\u2029", "\\u2029")
    )


HTML_TEMPLATE = r'''<!doctype html>
<html lang="zh-CN">
<head>
  <meta charset="utf-8">
  <meta name="viewport" content="width=device-width, initial-scale=1">
  <meta name="color-scheme" content="light">
  <title>线下活动核销结果 - 本地查看页</title>
  <style>
    :root {
      color-scheme: light;
      --page: #f1f3f1;
      --surface: #ffffff;
      --surface-muted: #f7f8f7;
      --ink: #1f2723;
      --ink-soft: #5e6a64;
      --ink-faint: #7b8781;
      --line: #d9dfdb;
      --line-strong: #bcc6c0;
      --accent: #245c49;
      --accent-strong: #174535;
      --accent-soft: #e8f1ec;
      --danger: #9b4034;
      --danger-soft: #f8ece9;
      --summary: #755b24;
      --summary-soft: #f5f0e3;
      --panel-radius: 12px;
      --control-radius: 8px;
      --label-radius: 6px;
      --font-sans: "MiSans", "HarmonyOS Sans SC", "Microsoft YaHei UI", "PingFang SC", sans-serif;
      --font-number: "Bahnschrift", "Segoe UI", var(--font-sans);
    }

    * { box-sizing: border-box; }
    html { scroll-behavior: smooth; }
    body {
      margin: 0;
      min-width: 320px;
      color: var(--ink);
      background: var(--page);
      font-family: var(--font-sans);
      font-size: 14px;
      line-height: 1.65;
      text-rendering: optimizeLegibility;
    }

    button, input, a { font: inherit; }
    button, a { -webkit-tap-highlight-color: transparent; }
    button:focus-visible, input:focus-visible, a:focus-visible, summary:focus-visible {
      outline: 3px solid rgba(36, 92, 73, .24);
      outline-offset: 2px;
    }
    button:active, a:active { transform: translateY(1px); }

    .masthead {
      color: var(--ink);
      background: var(--surface);
      border-bottom: 1px solid var(--line);
    }
    .masthead-inner {
      display: grid;
      grid-template-columns: minmax(240px, 1fr) auto auto;
      gap: 28px;
      align-items: center;
      width: min(1560px, calc(100% - 48px));
      min-height: 82px;
      margin: 0 auto;
      padding: 14px 0;
    }
    h1 {
      margin: 0;
      font-size: clamp(22px, 2.4vw, 30px);
      font-weight: 760;
      line-height: 1.25;
      letter-spacing: -.02em;
    }
    .masthead-meta {
      display: grid;
      gap: 2px;
      min-width: 0;
      color: var(--ink-soft);
      font-size: 12px;
      text-align: right;
    }
    .masthead-meta #fileName {
      overflow: hidden;
      max-width: 340px;
      color: var(--ink);
      font-family: var(--font-number);
      font-size: 13px;
      font-weight: 650;
      text-overflow: ellipsis;
      white-space: nowrap;
    }
    .masthead-actions { display: flex; align-items: center; gap: 8px; }
    .excel-link {
      display: inline-flex;
      align-items: center;
      justify-content: center;
      min-height: 42px;
      padding: 9px 15px;
      border-radius: var(--control-radius);
      cursor: pointer;
      font-size: 13px;
      font-weight: 750;
      text-decoration: none;
      white-space: nowrap;
      transition: border-color .16s ease, background .16s ease, color .16s ease;
    }
    .excel-link { color: #f7fbf8; background: var(--accent); border: 1px solid var(--accent); }
    .excel-link:hover { background: var(--accent-strong); border-color: var(--accent-strong); }

    .shell {
      display: grid;
      grid-template-columns: 300px minmax(0, 1fr);
      gap: 24px;
      width: min(1560px, calc(100% - 48px));
      margin: 24px auto 56px;
    }
    .side-panel {
      position: sticky;
      top: 16px;
      align-self: start;
      overflow: hidden;
      background: var(--surface);
      border: 1px solid var(--line);
      border-radius: var(--panel-radius);
    }
    .workspace { min-width: 0; }
    .scenario-tabs {
      display: grid;
      grid-template-columns: 1fr;
      gap: 6px;
      padding: 12px;
      border-bottom: 1px solid var(--line);
    }
    .tab-button {
      display: flex;
      align-items: center;
      justify-content: space-between;
      gap: 12px;
      width: 100%;
      min-height: 44px;
      padding: 10px 12px;
      color: var(--ink-soft);
      text-align: left;
      background: transparent;
      border: 1px solid transparent;
      border-radius: var(--control-radius);
      cursor: pointer;
      font-weight: 700;
      transition: border-color .16s ease, background .16s ease, color .16s ease;
    }
    .tab-button:hover { color: var(--ink); background: var(--surface-muted); }
    .tab-button[aria-selected="true"] {
      color: var(--accent-strong);
      background: var(--accent-soft);
      border-color: #bfd2c8;
      box-shadow: inset 3px 0 0 var(--accent);
    }
    .tab-button small { color: var(--ink-faint); font-family: var(--font-number); font-size: 11px; font-weight: 600; }

    .sheet-intro { padding: 20px 18px 18px; border-bottom: 1px solid var(--line); }
    .sheet-intro h2 {
      margin: 0 0 8px;
      font-size: 21px;
      font-weight: 760;
      line-height: 1.35;
      letter-spacing: -.015em;
    }
    .sheet-note { margin: 0; color: var(--ink-soft); font-size: 12px; line-height: 1.7; }

    .summary-grid {
      display: grid;
      grid-template-columns: repeat(2, minmax(0, 1fr));
      margin: 0;
      background: var(--surface);
      border-bottom: 1px solid var(--line);
    }
    .metric {
      min-height: 104px;
      padding: 15px 16px;
      background: var(--surface);
      border-right: 1px solid var(--line);
      border-bottom: 1px solid var(--line);
    }
    .metric:nth-child(2n) { border-right: 0; }
    .metric:nth-last-child(-n + 2) { border-bottom: 0; }
    .metric-label { display: block; color: var(--ink-soft); font-size: 12px; font-weight: 650; }
    .metric-value {
      display: block;
      margin-top: 6px;
      color: var(--ink);
      font-family: var(--font-number);
      font-size: 30px;
      font-weight: 700;
      line-height: 1;
    }
    .metric-hint { display: block; margin-top: 8px; color: var(--ink-faint); font-size: 11px; line-height: 1.45; }
    .metric.pass .metric-value { color: var(--accent); }
    .metric.issue .metric-value { color: var(--danger); }

    .side-actions {
      display: grid;
      grid-template-columns: repeat(2, minmax(0, 1fr));
      gap: 8px;
      padding: 12px;
    }
    .toolbar {
      position: sticky;
      top: 12px;
      z-index: 20;
      display: grid;
      grid-template-columns: minmax(300px, 1fr) auto;
      gap: 16px;
      align-items: end;
      padding: 14px;
      margin-bottom: 12px;
      background: rgba(255, 255, 255, .96);
      border: 1px solid var(--line);
      border-radius: var(--panel-radius);
      box-shadow: 0 8px 24px rgba(40, 55, 47, .07);
      backdrop-filter: blur(10px);
    }
    .control-group { min-width: 0; }
    .control-label {
      display: block;
      margin: 0 0 6px;
      color: var(--ink-soft);
      font-size: 11px;
      font-weight: 700;
    }
    .search-wrap { position: relative; min-width: 0; }
    .search-input {
      width: 100%;
      height: 42px;
      padding: 0 66px 0 12px;
      color: var(--ink);
      background: var(--surface);
      border: 1px solid var(--line-strong);
      border-radius: var(--control-radius);
      outline: none;
    }
    .search-input::placeholder { color: #88928d; }
    .search-input:focus { border-color: var(--accent); box-shadow: 0 0 0 3px rgba(36, 92, 73, .10); }
    .clear-search {
      position: absolute;
      right: 5px;
      top: 5px;
      min-width: 52px;
      height: 32px;
      padding: 0 8px;
      color: var(--ink-soft);
      background: var(--surface-muted);
      border: 0;
      border-radius: 5px;
      cursor: pointer;
      font-size: 12px;
      font-weight: 650;
    }
    .clear-search[hidden] { display: none; }
    .segmented { display: flex; flex-wrap: wrap; align-items: center; gap: 6px; }
    .filter-button, .tool-button {
      display: inline-flex;
      align-items: center;
      justify-content: center;
      gap: 7px;
      min-height: 38px;
      padding: 7px 11px;
      color: var(--ink-soft);
      background: var(--surface);
      border: 1px solid var(--line-strong);
      border-radius: var(--control-radius);
      cursor: pointer;
      font-size: 12px;
      font-weight: 700;
      white-space: nowrap;
      transition: border-color .16s ease, color .16s ease, background .16s ease;
    }
    .filter-button:hover, .tool-button:hover { color: var(--ink); border-color: var(--accent); }
    .filter-button[aria-pressed="true"] { color: #f7fbf8; background: var(--accent); border-color: var(--accent); }
    .filter-count {
      min-width: 20px;
      padding: 1px 5px;
      color: inherit;
      background: rgba(31, 39, 35, .07);
      border-radius: 4px;
      font-family: var(--font-number);
      font-size: 10px;
      text-align: center;
    }
    .filter-button[aria-pressed="true"] .filter-count { background: rgba(255, 255, 255, .18); }
    .side-actions .tool-button { width: 100%; }

    .results-line {
      display: flex;
      justify-content: space-between;
      align-items: center;
      gap: 16px;
      min-height: 32px;
      margin: 0 2px 8px;
      color: var(--ink-soft);
      font-size: 12px;
    }
    .results-line strong { color: var(--ink); font-family: var(--font-number); }

    .records { display: grid; gap: 8px; }
    .record-card {
      overflow: clip;
      background: var(--surface);
      border: 1px solid var(--line);
      border-left: 3px solid var(--line-strong);
      border-radius: var(--panel-radius);
      transition: border-color .16s ease, box-shadow .16s ease;
    }
    .record-card:hover { border-color: var(--line-strong); box-shadow: 0 6px 18px rgba(40, 55, 47, .055); }
    .record-card.pass { border-left-color: var(--accent); }
    .record-card.issue { border-left-color: var(--danger); }
    .record-card.summary { border-left-color: var(--summary); background: #fdfcf8; }
    .record-card details > summary {
      display: grid;
      grid-template-columns: minmax(0, 1fr) auto 20px;
      gap: 14px;
      align-items: center;
      min-height: 86px;
      padding: 16px 18px;
      cursor: pointer;
      list-style: none;
    }
    .record-card summary::-webkit-details-marker { display: none; }
    .record-title { min-width: 0; }
    .record-title h3 { margin: 0 0 4px; font-size: 16px; font-weight: 740; line-height: 1.45; }
    .record-title p {
      display: -webkit-box;
      margin: 0;
      overflow: hidden;
      color: var(--ink-soft);
      font-size: 12px;
      line-height: 1.55;
      -webkit-line-clamp: 2;
      -webkit-box-orient: vertical;
    }
    .status-pill {
      display: inline-flex;
      align-items: center;
      justify-content: center;
      min-width: 68px;
      padding: 5px 9px;
      border: 1px solid transparent;
      border-radius: var(--label-radius);
      font-size: 11px;
      font-weight: 750;
    }
    .status-pill.pass { color: var(--accent-strong); background: var(--accent-soft); border-color: #c9dbd1; }
    .status-pill.issue { color: #813127; background: var(--danger-soft); border-color: #eccac4; }
    .status-pill.summary { color: #674d18; background: var(--summary-soft); border-color: #e5d8b8; }
    .chevron { display: grid; place-items: center; color: var(--ink-faint); font-family: var(--font-number); font-size: 17px; }
    .chevron::before { content: "+"; }
    details[open] .chevron::before { content: "-"; }

    .record-body { padding: 18px; background: #fbfcfb; border-top: 1px solid var(--line); }
    .evidence-grid {
      display: grid;
      grid-template-columns: repeat(2, minmax(0, 1fr));
      gap: 10px;
    }
    .evidence-cell {
      min-width: 0;
      padding: 14px;
      background: var(--surface);
      border: 1px solid var(--line);
      border-radius: 10px;
    }
    .record-card.pass .conclusion-cell { background: #f4f8f5; border-color: #cadbd2; }
    .record-card.issue .conclusion-cell { background: #fdf7f5; border-color: #ebccc6; }
    .record-card.summary .conclusion-cell { background: #faf7ee; border-color: #e5d8b8; }
    .evidence-label {
      display: flex;
      align-items: baseline;
      justify-content: space-between;
      gap: 10px;
      margin-bottom: 10px;
      padding-bottom: 8px;
      border-bottom: 1px solid var(--line);
    }
    .evidence-label strong { color: var(--ink); font-size: 12px; font-weight: 750; }
    .evidence-label small {
      max-width: 58%;
      color: var(--ink-faint);
      font-size: 10px;
      overflow-wrap: anywhere;
      text-align: right;
    }
    .cell-body { color: #36413b; font-size: 12px; word-break: break-word; }
    .cell-line { margin: 0 0 6px; white-space: pre-wrap; }
    .cell-line:last-child { margin-bottom: 0; }
    .cell-line b { color: var(--ink); font-weight: 720; }
    .identity-card {
      margin: 9px 0;
      overflow: hidden;
      background: var(--surface);
      border: 1px solid var(--line-strong);
      border-radius: var(--control-radius);
    }
    .identity-title {
      padding: 8px 10px;
      color: var(--accent-strong);
      background: var(--accent-soft);
      border-bottom: 1px solid #c9d8d0;
      font-size: 12px;
      font-weight: 750;
    }
    .identity-table { width: 100%; border-collapse: collapse; table-layout: fixed; }
    .identity-table th,
    .identity-table td { padding: 7px 9px; vertical-align: top; text-align: left; }
    .identity-table tr + tr th,
    .identity-table tr + tr td { border-top: 1px solid var(--line); }
    .identity-table th {
      width: 78px;
      color: var(--ink-soft);
      background: var(--surface-muted);
      font-weight: 650;
      white-space: nowrap;
    }
    .identity-table td { color: var(--ink); font-weight: 650; }
    .identity-spacer { height: 3px; }
    .empty-value { color: var(--ink-faint); }
    .record-actions { display: flex; justify-content: flex-end; gap: 8px; margin-top: 12px; }

    .empty-state {
      padding: 56px 24px;
      text-align: center;
      background: var(--surface);
      border: 1px dashed var(--line-strong);
      border-radius: var(--panel-radius);
    }
    .empty-state strong { display: block; margin-bottom: 5px; font-size: 17px; }
    .empty-state span { color: var(--ink-soft); font-size: 12px; }

    .compact .record-card details > summary { min-height: 66px; padding-block: 10px; }
    .compact .record-title p { -webkit-line-clamp: 1; }
    .compact .evidence-grid { grid-template-columns: repeat(3, minmax(0, 1fr)); }
    .compact .evidence-cell { padding: 11px; }

    .back-top {
      position: fixed;
      right: 20px;
      bottom: 20px;
      min-width: 76px;
      height: 40px;
      padding: 0 12px;
      color: #f7fbf8;
      background: var(--accent);
      border: 1px solid var(--accent);
      border-radius: var(--control-radius);
      cursor: pointer;
      opacity: 0;
      pointer-events: none;
      transform: translateY(6px);
      transition: opacity .16s ease, transform .16s ease, background .16s ease;
    }
    .back-top:hover { background: var(--accent-strong); }
    .back-top.visible { opacity: 1; pointer-events: auto; transform: translateY(0); }
    .toast {
      position: fixed;
      left: 50%;
      bottom: 24px;
      z-index: 50;
      max-width: min(420px, calc(100% - 32px));
      padding: 10px 14px;
      color: #f7fbf8;
      background: #27312c;
      border-radius: var(--control-radius);
      box-shadow: 0 10px 24px rgba(31, 39, 35, .18);
      opacity: 0;
      pointer-events: none;
      transform: translate(-50%, 10px);
      transition: opacity .16s ease, transform .16s ease;
    }
    .toast.show { opacity: 1; transform: translate(-50%, 0); }
    .page-footer {
      width: min(1560px, calc(100% - 48px));
      margin: 0 auto 30px;
      color: var(--ink-faint);
      font-size: 11px;
      text-align: right;
    }

    @media (max-width: 1180px) {
      .masthead-inner, .shell, .page-footer { width: min(100% - 32px, 1560px); }
      .masthead-inner { grid-template-columns: minmax(220px, 1fr) auto; }
      .masthead-meta { grid-column: 1 / -1; grid-row: 2; text-align: left; }
      .masthead-meta #fileName { max-width: none; }
      .shell { grid-template-columns: 1fr; }
      .side-panel { position: relative; top: 0; }
      .scenario-tabs { grid-template-columns: repeat(2, minmax(0, 1fr)); }
      .summary-grid { grid-template-columns: repeat(4, minmax(0, 1fr)); }
      .metric { border-right: 1px solid var(--line); border-bottom: 0; }
      .metric:nth-child(2n) { border-right: 1px solid var(--line); }
      .metric:last-child { border-right: 0; }
      .side-actions { display: flex; }
      .side-actions .tool-button { width: auto; }
    }
    @media (max-width: 760px) {
      .masthead-inner, .shell, .page-footer { width: min(100% - 24px, 1560px); }
      .masthead-inner { gap: 14px; min-height: 0; padding: 14px 0; }
      .masthead-actions { justify-self: end; }
      .shell { gap: 14px; margin-top: 14px; }
      .toolbar { grid-template-columns: 1fr; position: relative; top: 0; gap: 12px; }
      .evidence-grid, .compact .evidence-grid { grid-template-columns: 1fr; }
      .record-body { padding: 14px; }
      .page-footer { text-align: left; }
    }
    @media (max-width: 520px) {
      body { line-height: 1.55; }
      .masthead-inner { grid-template-columns: 1fr; gap: 8px; padding: 10px 0; }
      h1 { font-size: 20px; }
      .masthead-meta, .masthead-actions { grid-column: 1; grid-row: auto; justify-self: stretch; text-align: left; }
      .masthead-meta { display: block; line-height: 1.4; }
      .masthead-meta span:last-child { display: none; }
      .masthead-actions > * { flex: 1; }
      .excel-link { min-height: 36px; padding: 6px 10px; font-size: 12px; }
      .shell { gap: 10px; margin-top: 10px; }
      .scenario-tabs { grid-template-columns: repeat(2, minmax(0, 1fr)); gap: 4px; padding: 6px; }
      .tab-button { min-height: 38px; padding: 7px 8px; font-size: 12px; }
      .sheet-intro { padding: 9px 11px; }
      .sheet-intro h2 { margin-bottom: 2px; font-size: 16px; }
      .sheet-note { font-size: 10px; line-height: 1.45; }
      .summary-grid { grid-template-columns: repeat(4, minmax(0, 1fr)); }
      .metric { min-height: 76px; padding: 8px 6px; border-right: 1px solid var(--line); border-bottom: 0; }
      .metric:nth-child(2n) { border-right: 1px solid var(--line); }
      .metric:last-child { border-right: 0; }
      .metric-label { font-size: 10px; }
      .metric-value { margin-top: 3px; font-size: 21px; }
      .metric-hint { display: block; margin-top: 4px; font-size: 9px; line-height: 1.25; }
      .side-actions { display: grid; gap: 6px; padding: 6px; }
      .side-actions .tool-button { width: 100%; }
      .toolbar { gap: 8px; padding: 10px; margin-bottom: 6px; }
      .control-label { margin-bottom: 4px; }
      .search-input { height: 38px; }
      .segmented { display: grid; grid-template-columns: repeat(4, minmax(0, 1fr)); gap: 4px; }
      .filter-button { width: 100%; min-height: 34px; gap: 3px; padding: 4px 3px; font-size: 10px; }
      .filter-count { min-width: 16px; padding-inline: 3px; }
      .results-line { min-height: 28px; margin-bottom: 6px; }
      .record-card details > summary { grid-template-columns: minmax(0, 1fr) 18px; gap: 10px; padding: 14px; }
      .status-pill { grid-column: 1; justify-self: start; }
      .chevron { grid-column: 2; grid-row: 1 / span 2; }
      .evidence-label { align-items: flex-start; flex-direction: column; }
      .evidence-label small { max-width: 100%; text-align: left; }
      .results-line { align-items: center; flex-direction: row; gap: 8px; font-size: 11px; }
      .back-top { right: 12px; bottom: 12px; }
    }

    @media (prefers-reduced-motion: reduce) {
      html { scroll-behavior: auto; }
      *, *::before, *::after { transition-duration: .01ms !important; }
    }

    @media print {
      @page { margin: 12mm; }
      body { background: #fff; font-size: 11px; }
      .masthead { border-bottom: 1px solid #777; }
      .masthead-inner, .shell, .page-footer { width: 100%; margin: 0; }
      .masthead-inner { display: block; min-height: 0; padding: 0 0 10px; }
      h1 { font-size: 22px; }
      .masthead-meta { margin-top: 4px; text-align: left; }
      .masthead-actions, .scenario-tabs, .toolbar, .side-actions, .record-actions, .back-top, .toast { display: none !important; }
      .shell { display: block; }
      .side-panel { position: static; overflow: visible; border: 0; }
      .sheet-intro { padding: 10px 0; border-bottom: 1px solid #999; }
      .summary-grid { grid-template-columns: repeat(4, 1fr); border-bottom: 1px solid #999; }
      .metric { min-height: 0; padding: 8px; box-shadow: none; break-inside: avoid; }
      .metric-value { font-size: 22px; }
      .workspace { margin-top: 10px; }
      .records { display: block; }
      .record-card { margin-bottom: 8px; box-shadow: none; break-inside: avoid; }
      .record-card details > summary { min-height: 0; padding: 9px; }
      .record-card details[open] .record-body { display: block; }
      .record-body { padding: 9px; }
      .evidence-grid, .compact .evidence-grid { grid-template-columns: 1fr 1fr; }
      .evidence-cell { padding: 9px; }
      .page-footer { margin-top: 12px; }
    }
  </style>
</head>
<body>
  <div id="topSentinel" aria-hidden="true"></div>
  <header class="masthead">
    <div class="masthead-inner">
      <div class="brand-lockup"><h1>线下活动核销结果</h1></div>
      <div class="masthead-meta">
        <span id="fileName">本地查看页</span>
        <span>页面可离线打开，内容与同名 Excel 保持一致</span>
      </div>
      <div class="masthead-actions">
        <a class="excel-link" id="excelLink" href="#" title="打开同名 Excel 文件">打开同版 Excel</a>
      </div>
    </div>
  </header>

  <main class="shell" id="top">
    <aside class="side-panel" aria-label="场景概览">
      <nav class="scenario-tabs" id="scenarioTabs" role="tablist" aria-label="核销场景"></nav>

      <section class="sheet-intro" aria-labelledby="sheetTitle">
        <h2 id="sheetTitle"></h2>
        <p class="sheet-note" id="sheetNote"></p>
      </section>

      <section class="summary-grid" id="summaryGrid" aria-label="当前场景概览"></section>

      <div class="side-actions" aria-label="记录显示方式">
        <button class="tool-button" id="densityButton" type="button" aria-pressed="false">紧凑显示</button>
        <button class="tool-button" id="expandButton" type="button" aria-pressed="false">展开全部</button>
      </div>
    </aside>

    <section class="workspace" aria-label="核销记录">
      <section class="toolbar" aria-label="查找与筛选">
        <div class="control-group">
          <label class="control-label" for="searchInput">搜索当前场景</label>
          <div class="search-wrap">
            <input class="search-input" id="searchInput" type="search" placeholder="输入门店、商品、69码、文件名或待补问题" autocomplete="off">
            <button class="clear-search" id="clearSearch" type="button" aria-label="清空搜索" hidden>清空</button>
          </div>
        </div>
        <div class="control-group">
          <span class="control-label">查看范围</span>
          <div class="segmented" id="statusFilters" aria-label="结论筛选">
            <button class="filter-button" type="button" data-status="all" data-label="全部" aria-pressed="true">全部</button>
            <button class="filter-button" type="button" data-status="pass" data-label="只看通过" aria-pressed="false">只看通过</button>
            <button class="filter-button" type="button" data-status="issue" data-label="只看待补" aria-pressed="false">只看待补</button>
            <button class="filter-button" type="button" data-status="summary" data-label="只看汇总" aria-pressed="false">只看汇总</button>
          </div>
        </div>
      </section>

      <div class="results-line">
        <span id="resultsCount" aria-live="polite"></span>
        <span>点击记录，查看六列完整内容</span>
      </div>

      <section class="records" id="records"></section>
    </section>
  </main>

  <p class="page-footer">本页用于查看核销结果，完整内容与同名 Excel 一致。</p>
  <button class="back-top" id="backTop" type="button">返回顶部</button>
  <div class="toast" id="toast" role="status" aria-live="polite"></div>
  __AUDIT_DATA__
  <script>
    (() => {
      'use strict';
      const data = JSON.parse(document.getElementById('audit-data').textContent);
      const state = { sheetIndex: 0, status: 'all', query: '', compact: false, expanded: false };
      const tabs = document.getElementById('scenarioTabs');
      const records = document.getElementById('records');
      const searchInput = document.getElementById('searchInput');
      const clearSearch = document.getElementById('clearSearch');
      const densityButton = document.getElementById('densityButton');
      const expandButton = document.getElementById('expandButton');
      const backTop = document.getElementById('backTop');
      const toast = document.getElementById('toast');
      let toastTimer = null;

      const escapeHtml = (value) => String(value ?? '')
        .replaceAll('&', '&amp;').replaceAll('<', '&lt;').replaceAll('>', '&gt;')
        .replaceAll('"', '&quot;').replaceAll("'", '&#039;');

      const splitHeader = (header) => {
        const lines = String(header || '').split('\n').filter(Boolean);
        return { label: lines[0] || '未命名字段', source: lines.slice(1).join(' / ') };
      };

      const renderLine = (line) => {
        const raw = String(line || '');
        const colon = raw.indexOf('：');
        if (colon > 0 && colon <= 18) {
          return `<p class="cell-line"><b>${escapeHtml(raw.slice(0, colon + 1))}</b>${escapeHtml(raw.slice(colon + 1))}</p>`;
        }
        return `<p class="cell-line">${escapeHtml(raw) || '&nbsp;'}</p>`;
      };

      const renderCellBody = (value) => {
        const lines = String(value || '').split('\n');
        const identityHeading = /^(现场商品\d+（知识库）|对应销售Excel第\d+行)$/;
        const identityField = /^(商品编码|商品名称|69码|数量)：(.*)$/;
        const blocks = [];
        let index = 0;
        while (index < lines.length) {
          const heading = lines[index].match(identityHeading);
          if (heading) {
            const rows = [];
            index += 1;
            while (index < lines.length) {
              const field = lines[index].match(identityField);
              if (!field) break;
              rows.push(
                `<tr><th scope="row">${escapeHtml(field[1])}</th><td>${escapeHtml(field[2]) || '未提供'}</td></tr>`
              );
              index += 1;
            }
            if (rows.length) {
              blocks.push(
                `<section class="identity-card"><div class="identity-title">${escapeHtml(heading[1])}</div>`
                + `<table class="identity-table"><tbody>${rows.join('')}</tbody></table></section>`
              );
              continue;
            }
          }
          if (!lines[index]) {
            blocks.push('<div class="identity-spacer" aria-hidden="true"></div>');
          } else {
            blocks.push(renderLine(lines[index]));
          }
          index += 1;
        }
        return blocks.join('');
      };

      const renderCell = (header, value, cellIndex) => {
        const meta = splitHeader(header);
        const body = value
          ? renderCellBody(value)
          : '<span class="empty-value">没有可核验内容</span>';
        const emphasis = cellIndex === 5 ? ' conclusion-cell' : '';
        return `<section class="evidence-cell${emphasis}">
          <div class="evidence-label"><strong>${escapeHtml(meta.label)}</strong>${meta.source ? `<small title="${escapeHtml(meta.source)}">${escapeHtml(meta.source)}</small>` : ''}</div>
          <div class="cell-body">${body}</div>
        </section>`;
      };

      const statusText = (status) => ({ pass: '通过', issue: '待补材料', summary: '汇总' }[status] || '查看');

      const conclusionExcerpt = (row) => {
        const conclusion = row.values[5] || row.values.find(Boolean) || '';
        const lines = conclusion.split('\n').map((line) => line.trim()).filter(Boolean);
        const priority = lines.filter((line) => line.startsWith('结论：') || line.startsWith('主要问题：') || line.startsWith('置信度：'));
        if (priority.length >= 2) return priority.slice(0, 2).join('；');
        if (priority.length === 1) {
          const next = lines.find((line) => line !== priority[0] && line !== '要重新提交：');
          return [priority[0], next].filter(Boolean).join('；');
        }
        return lines.slice(0, 2).join('；');
      };

      const activeSheet = () => data.sheets[state.sheetIndex];

      const currentRows = () => {
        const query = state.query.trim().toLocaleLowerCase('zh-CN');
        return activeSheet().rows.filter((row) => {
          if (state.status === 'summary' && row.kind !== 'summary') return false;
          if ((state.status === 'pass' || state.status === 'issue') && row.status !== state.status) return false;
          if (!query) return true;
          return [row.heading, ...row.values].join('\n').toLocaleLowerCase('zh-CN').includes(query);
        });
      };

      const showToast = (message) => {
        window.clearTimeout(toastTimer);
        toast.textContent = message;
        toast.classList.add('show');
        toastTimer = window.setTimeout(() => toast.classList.remove('show'), 1800);
      };

      const legacyCopy = (text) => {
        const area = document.createElement('textarea');
        area.value = text;
        area.setAttribute('readonly', '');
        area.style.position = 'fixed';
        area.style.left = '-9999px';
        document.body.appendChild(area);
        area.select();
        area.setSelectionRange(0, area.value.length);
        let copied = false;
        try { copied = document.execCommand('copy'); } catch (_) { copied = false; }
        area.remove();
        return copied;
      };

      const copyText = async (text) => {
        let copied = false;
        if (navigator.clipboard && window.isSecureContext) {
          try {
            await navigator.clipboard.writeText(text);
            copied = true;
          } catch (_) {
            copied = legacyCopy(text);
          }
        } else {
          copied = legacyCopy(text);
        }
        showToast(copied ? '本条结论已复制' : '浏览器未允许复制，请手动选择文字');
      };

      const setExcelLink = () => {
        const pathParts = decodeURIComponent(window.location.pathname).split('/');
        const pageName = pathParts[pathParts.length - 1] || '核销结果.html';
        const excelName = pageName.replace(/\.html$/i, '.xlsx');
        document.getElementById('fileName').textContent = pageName;
        const excelUrl = new URL(window.location.href);
        excelUrl.hash = '';
        excelUrl.search = '';
        excelUrl.pathname = excelUrl.pathname.replace(/[^/]*$/, encodeURIComponent(excelName));
        document.getElementById('excelLink').href = excelUrl.href;
      };

      const renderTabs = () => {
        tabs.innerHTML = data.sheets.map((sheet, index) => `
          <button class="tab-button" id="scenarioTab${index}" type="button" role="tab" data-sheet="${index}" aria-controls="records" aria-selected="${index === state.sheetIndex}" tabindex="${index === state.sheetIndex ? '0' : '-1'}">
            <span>${escapeHtml(sheet.name)}</span><small>${sheet.rows.length}项</small>
          </button>`).join('');
      };

      const renderIntro = () => {
        const sheet = activeSheet();
        document.getElementById('sheetTitle').textContent = sheet.title;
        document.getElementById('sheetNote').textContent = sheet.note;
      };

      const renderMetrics = () => {
        const rows = activeSheet().rows;
        const passed = rows.filter((row) => row.status === 'pass').length;
        const issues = rows.filter((row) => row.status === 'issue').length;
        const summaries = rows.filter((row) => row.kind === 'summary').length;
        const metrics = [
          ['全部条目', rows.length, '当前场景完整记录', ''],
          ['通过', passed, '无需重新提交', 'pass'],
          ['待补材料', issues, '点开查看补交要求', 'issue'],
          ['汇总项', summaries, '金额与身份等汇总', ''],
        ];
        document.getElementById('summaryGrid').innerHTML = metrics.map(([label, value, hint, tone]) => `
          <article class="metric ${tone}"><span class="metric-label">${label}</span><strong class="metric-value">${value}</strong><span class="metric-hint">${hint}</span></article>`).join('');
      };

      const renderFilters = () => {
        const rows = activeSheet().rows;
        const counts = {
          all: rows.length,
          pass: rows.filter((row) => row.status === 'pass').length,
          issue: rows.filter((row) => row.status === 'issue').length,
          summary: rows.filter((row) => row.kind === 'summary').length,
        };
        document.querySelectorAll('[data-status]').forEach((button) => {
          const status = button.dataset.status;
          button.setAttribute('aria-pressed', String(status === state.status));
          button.innerHTML = `${escapeHtml(button.dataset.label)}<span class="filter-count">${counts[status]}</span>`;
        });
      };

      const renderRecords = () => {
        const sheet = activeSheet();
        const rows = currentRows();
        document.getElementById('resultsCount').innerHTML = `当前显示 <strong>${rows.length}</strong> / ${sheet.rows.length} 条`;
        if (!rows.length) {
          records.innerHTML = '<div class="empty-state"><strong>没有找到符合条件的记录</strong><span>可以清空搜索或切换上方结论筛选。</span></div>';
          return;
        }
        records.innerHTML = rows.map((row) => {
          const kindClass = row.kind === 'summary' ? ' summary-row' : '';
          return `<article class="record-card ${row.status}${kindClass}">
            <details ${state.expanded ? 'open' : ''} data-row="${row.excel_row}">
              <summary>
                <div class="record-title"><h3>${escapeHtml(row.heading)}</h3><p>${escapeHtml(conclusionExcerpt(row))}</p></div>
                <span class="status-pill ${row.status}">${statusText(row.status)}</span>
                <span class="chevron" aria-hidden="true"></span>
              </summary>
              <div class="record-body">
                <div class="evidence-grid">${row.values.map((value, cellIndex) => renderCell(sheet.headers[cellIndex], value, cellIndex)).join('')}</div>
                <div class="record-actions"><button class="tool-button copy-conclusion" type="button" data-copy-row="${row.excel_row}">复制本条结论</button></div>
              </div>
            </details>
          </article>`;
        }).join('');
      };

      const renderAll = () => {
        renderTabs();
        renderIntro();
        renderMetrics();
        renderFilters();
        renderRecords();
      };

      tabs.addEventListener('click', (event) => {
        const button = event.target.closest('[data-sheet]');
        if (!button) return;
        state.sheetIndex = Number(button.dataset.sheet);
        state.status = 'all';
        state.query = '';
        state.expanded = false;
        searchInput.value = '';
        clearSearch.hidden = true;
        renderAll();
      });

      tabs.addEventListener('keydown', (event) => {
        if (event.key !== 'ArrowLeft' && event.key !== 'ArrowRight') return;
        const direction = event.key === 'ArrowRight' ? 1 : -1;
        const nextIndex = (state.sheetIndex + direction + data.sheets.length) % data.sheets.length;
        const nextButton = tabs.querySelector(`[data-sheet="${nextIndex}"]`);
        if (nextButton) {
          event.preventDefault();
          nextButton.click();
          tabs.querySelector(`[data-sheet="${nextIndex}"]`)?.focus();
        }
      });

      document.getElementById('statusFilters').addEventListener('click', (event) => {
        const button = event.target.closest('[data-status]');
        if (!button) return;
        state.status = button.dataset.status;
        renderFilters();
        renderRecords();
      });

      searchInput.addEventListener('input', () => {
        state.query = searchInput.value;
        clearSearch.hidden = !state.query;
        renderRecords();
      });
      clearSearch.addEventListener('click', () => {
        state.query = '';
        searchInput.value = '';
        clearSearch.hidden = true;
        searchInput.focus();
        renderRecords();
      });

      densityButton.addEventListener('click', () => {
        state.compact = !state.compact;
        document.body.classList.toggle('compact', state.compact);
        densityButton.setAttribute('aria-pressed', String(state.compact));
        densityButton.textContent = state.compact ? '舒展显示' : '紧凑显示';
      });

      expandButton.addEventListener('click', () => {
        state.expanded = !state.expanded;
        records.querySelectorAll('details').forEach((detail) => { detail.open = state.expanded; });
        expandButton.setAttribute('aria-pressed', String(state.expanded));
        expandButton.textContent = state.expanded ? '收起全部' : '展开全部';
      });

      records.addEventListener('click', (event) => {
        const button = event.target.closest('[data-copy-row]');
        if (!button) return;
        const rowNumber = Number(button.dataset.copyRow);
        const row = activeSheet().rows.find((item) => item.excel_row === rowNumber);
        if (row) copyText(row.values[5] || row.values.join('\n'));
      });

      const topObserver = new IntersectionObserver(([entry]) => {
        backTop.classList.toggle('visible', !entry.isIntersecting);
      });
      topObserver.observe(document.getElementById('topSentinel'));
      backTop.addEventListener('click', () => {
        const reduced = window.matchMedia('(prefers-reduced-motion: reduce)').matches;
        window.scrollTo({ top: 0, behavior: reduced ? 'auto' : 'smooth' });
      });

      setExcelLink();
      renderAll();
    })();
  </script>
  <noscript>请启用浏览器 JavaScript，以使用搜索、筛选和展开按钮。</noscript>
</body>
</html>
'''


def create_html_report_from_workbook(
    workbook_path: str | Path,
    output_path: str | Path,
) -> Path:
    payload = _workbook_payload(workbook_path)
    embedded = DATA_OPEN + _json_for_script(payload) + DATA_CLOSE
    html = HTML_TEMPLATE.replace("__AUDIT_DATA__", embedded)
    target = Path(output_path)
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_text(html, encoding="utf-8", newline="\n")
    return target


def _embedded_payload(html: str) -> dict[str, Any]:
    start = html.find(DATA_OPEN)
    if start < 0:
        raise AuditError("HTML缺少内嵌核销数据")
    start += len(DATA_OPEN)
    end = html.find(DATA_CLOSE, start)
    if end < 0:
        raise AuditError("HTML内嵌核销数据没有正确结束")
    try:
        payload = json.loads(html[start:end])
    except json.JSONDecodeError as exc:
        raise AuditError("HTML内嵌核销数据不是有效JSON") from exc
    if not isinstance(payload, dict):
        raise AuditError("HTML内嵌核销数据结构错误")
    return payload


def verify_html_report(
    html_path: str | Path,
    scenarios: list[str],
    *,
    workbook_path: str | Path | None = None,
) -> dict[str, Any]:
    source = Path(html_path)
    if not source.is_file():
        raise AuditError(f"找不到HTML输出：{source}")
    html = source.read_text(encoding="utf-8")
    payload = _embedded_payload(html)
    expected_sheets = [SCENARIO_SHEETS[scenario] for scenario in scenarios]
    actual_sheets = [str(sheet.get("name") or "") for sheet in payload.get("sheets", [])]
    if actual_sheets != expected_sheets:
        raise AuditError(f"HTML场景名称或顺序错误：{actual_sheets}，期望{expected_sheets}")

    required_controls = (
        'id="scenarioTabs"',
        'id="searchInput"',
        'id="statusFilters"',
        'id="densityButton"',
        'id="expandButton"',
        'id="excelLink"',
        'id="backTop"',
    )
    missing_controls = [control for control in required_controls if control not in html]
    if missing_controls:
        raise AuditError("HTML缺少交互按钮：" + "、".join(missing_controls))
    if re.search(r"<(?:script|link)\b[^>]*(?:src|href)=[\"']https?://", html, re.I):
        raise AuditError("HTML不得依赖在线脚本或样式")
    for forbidden in FORBIDDEN_VISIBLE_TERMS:
        if forbidden in html:
            raise AuditError(f"HTML出现不应交给客户的工程表达：{forbidden}")

    record_counts = {
        str(sheet["name"]): len(sheet.get("rows") or [])
        for sheet in payload.get("sheets", [])
    }
    if workbook_path is not None:
        workbook_payload = _workbook_payload(workbook_path)
        for html_sheet, workbook_sheet in zip(
            payload.get("sheets", []), workbook_payload["sheets"], strict=True
        ):
            if html_sheet.get("headers") != workbook_sheet["headers"]:
                raise AuditError(f"{html_sheet.get('name')} 的HTML表头与Excel不一致")
            if html_sheet.get("rows") != workbook_sheet["rows"]:
                raise AuditError(f"{html_sheet.get('name')} 的HTML记录与Excel不一致")

    return {
        "path": str(source.resolve()),
        "sheet_names": actual_sheets,
        "record_counts": record_counts,
        "button_count": len(required_controls),
        "external_dependency_count": 0,
        "workbook_content_equal": workbook_path is not None,
    }
