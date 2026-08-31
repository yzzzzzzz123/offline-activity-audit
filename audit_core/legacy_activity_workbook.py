from __future__ import annotations

import re
import shutil
import subprocess
import zipfile
from pathlib import Path, PurePosixPath
from typing import Any
from xml.etree import ElementTree

from openpyxl import load_workbook
from PIL import Image

from .common import AuditError, sha256_file


PROJECT_ROOT = Path(__file__).resolve().parents[1]
CONVERTER = (
    PROJECT_ROOT
    / "skills"
    / "audit-self-procured-gift-material"
    / "scripts"
    / "convert-xls-to-xlsx.ps1"
)
DISPIMG_PATTERN = re.compile(
    r'^=(?:_xlfn\.)?DISPIMG\("(?P<image_id>ID_[A-F0-9]+)",\s*1\)$',
    re.I,
)
MAX_IMAGE_BYTES = 50 * 1024 * 1024


def _convert_legacy_workbook(source: Path, target: Path) -> None:
    powershell = shutil.which("powershell") or shutil.which("powershell.exe")
    if powershell is None:
        raise AuditError("读取活动返图.xls需要 Windows PowerShell 和 Microsoft Excel")
    if not CONVERTER.is_file():
        raise AuditError(f"旧式活动返图转换脚本缺失：{CONVERTER}")
    result = subprocess.run(
        [
            powershell,
            "-NoProfile",
            "-ExecutionPolicy",
            "Bypass",
            "-File",
            str(CONVERTER),
            "-InputPath",
            str(source),
            "-OutputPath",
            str(target),
        ],
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
        text=True,
        encoding="utf-8",
        errors="replace",
        check=False,
    )
    if result.returncode != 0 or not target.is_file():
        detail = (result.stderr or result.stdout).strip()[-2000:]
        raise AuditError(f"活动返图.xls无法只读转换为可核验格式：{detail}")


def _cell_image_targets(archive: zipfile.ZipFile) -> dict[str, str]:
    required = {"xl/cellimages.xml", "xl/_rels/cellimages.xml.rels"}
    if not required.issubset(set(archive.namelist())):
        raise AuditError("活动返图工作簿缺少WPS单元格图片关系")
    rel_root = ElementTree.fromstring(archive.read("xl/_rels/cellimages.xml.rels"))
    relationships = {
        str(item.attrib["Id"]): str(item.attrib["Target"])
        for item in rel_root
        if item.attrib.get("Id") and item.attrib.get("Target")
    }
    image_root = ElementTree.fromstring(archive.read("xl/cellimages.xml"))
    name_key = "{http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing}cNvPr"
    blip_key = "{http://schemas.openxmlformats.org/drawingml/2006/main}blip"
    embed_key = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"
    targets: dict[str, str] = {}
    for cell_image in list(image_root):
        name_node = cell_image.find(f".//{name_key}")
        blip_node = cell_image.find(f".//{blip_key}")
        if name_node is None or blip_node is None:
            raise AuditError("活动返图工作簿存在无法绑定的单元格图片")
        image_id = str(name_node.attrib.get("name") or "")
        relationship_id = str(blip_node.attrib.get(embed_key) or "")
        target = relationships.get(relationship_id)
        if not image_id or not target:
            raise AuditError("活动返图工作簿图片ID或关系缺失")
        pure = PurePosixPath("xl") / PurePosixPath(target)
        if pure.is_absolute() or ".." in pure.parts or not str(pure).startswith("xl/media/"):
            raise AuditError(f"活动返图工作簿图片路径不安全：{target}")
        if image_id.casefold() in {key.casefold() for key in targets}:
            raise AuditError(f"活动返图工作簿图片ID重复：{image_id}")
        targets[image_id] = pure.as_posix()
    return targets


def extract_activity_return_workbook(
    source_path: str | Path,
    destination: str | Path,
) -> dict[str, Any]:
    source = Path(source_path).resolve()
    target_root = Path(destination).resolve()
    if source.suffix.lower() != ".xls":
        raise AuditError(f"活动返图必须是.xls：{source.name}")
    if target_root.exists() or target_root.is_symlink():
        raise AuditError(f"活动返图临时目录已存在：{target_root}")
    target_root.mkdir(parents=True)
    converted = target_root / f"{source.stem}.xlsx"
    _convert_legacy_workbook(source, converted)

    formula_book = load_workbook(converted, read_only=True, data_only=False)
    try:
        candidates: list[tuple[Any, dict[str, int]]] = []
        expected_headers = {
            "活动周期": "period",
            "陈列图片": "photo",
            "客户编码": "customer_code",
            "客户名称": "store_name",
            "费用代垫经销商": "dealer_name",
            "费用代垫经销商编码": "dealer_code",
        }
        for worksheet in formula_book.worksheets:
            columns = {
                expected_headers[str(worksheet.cell(1, column).value).strip()]: column
                for column in range(1, worksheet.max_column + 1)
                if str(worksheet.cell(1, column).value or "").strip() in expected_headers
            }
            if set(columns) == set(expected_headers.values()):
                candidates.append((worksheet, columns))
        if len(candidates) != 1:
            raise AuditError(f"活动返图工作簿明细表应唯一，实际{len(candidates)}个")
        worksheet, columns = candidates[0]
        raw_records: list[dict[str, Any]] = []
        seen_ids: set[str] = set()
        for row in range(2, worksheet.max_row + 1):
            formula = str(worksheet.cell(row, columns["photo"]).value or "").strip()
            match = DISPIMG_PATTERN.fullmatch(formula)
            if match is None:
                raise AuditError(f"活动返图工作簿第{row}行图片公式不可核验：{formula}")
            image_id = match.group("image_id")
            if image_id.casefold() in {value.casefold() for value in seen_ids}:
                raise AuditError(f"活动返图工作簿重复使用图片ID：{image_id}")
            seen_ids.add(image_id)
            raw_records.append(
                {
                    "excel_row": row,
                    "period_text": str(worksheet.cell(row, columns["period"]).value or "").strip(),
                    "customer_code": str(worksheet.cell(row, columns["customer_code"]).value or "").strip(),
                    "store_name": str(worksheet.cell(row, columns["store_name"]).value or "").strip(),
                    "dealer_name": str(worksheet.cell(row, columns["dealer_name"]).value or "").strip(),
                    "dealer_code": str(worksheet.cell(row, columns["dealer_code"]).value or "").strip(),
                    "image_id": image_id,
                }
            )
        sheet_name = worksheet.title
    finally:
        formula_book.close()

    with zipfile.ZipFile(converted) as archive:
        targets = _cell_image_targets(archive)
        records: list[dict[str, Any]] = []
        for record in raw_records:
            image_id = str(record["image_id"])
            target = targets.get(image_id)
            if target is None:
                raise AuditError(
                    f"活动返图工作簿第{record['excel_row']}行图片ID没有媒体文件：{image_id}"
                )
            info = archive.getinfo(target)
            if info.file_size <= 0 or info.file_size > MAX_IMAGE_BYTES:
                raise AuditError(f"活动返图图片尺寸异常：{target}={info.file_size}")
            suffix = PurePosixPath(target).suffix.lower()
            if suffix not in {".png", ".jpg", ".jpeg", ".webp"}:
                raise AuditError(f"活动返图图片格式不支持：{target}")
            photo_name = f"活动返图-row-{int(record['excel_row']):02d}{suffix}"
            photo_path = target_root / photo_name
            with archive.open(info) as input_stream, photo_path.open("xb") as output_stream:
                shutil.copyfileobj(input_stream, output_stream, length=1024 * 1024)
            try:
                with Image.open(photo_path) as image:
                    image.verify()
            except Exception as exc:
                raise AuditError(f"活动返图图片无法读取：{photo_name}：{exc}") from exc
            records.append(
                {
                    **record,
                    "photo_file": photo_path.resolve(),
                    "photo_sha256": sha256_file(photo_path),
                }
            )
    if len(records) != len(targets):
        raise AuditError(
            f"活动返图工作簿图片未逐行完整绑定：明细{len(records)}行，媒体{len(targets)}份"
        )
    return {
        "source_file": source.name,
        "source_sha256": sha256_file(source),
        "sheet": sheet_name,
        "record_count": len(records),
        "records": records,
        "converted_file": converted.resolve(),
    }
