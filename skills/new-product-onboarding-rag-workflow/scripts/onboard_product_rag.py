from __future__ import annotations

import argparse
import hashlib
import json
import re
import sys
import unicodedata
from datetime import datetime, timezone
from difflib import SequenceMatcher
from pathlib import Path
from typing import Any

from PIL import Image


SCRIPT_DIR = Path(__file__).resolve().parent
SKILL_ROOT = SCRIPT_DIR.parent
PROJECT_ROOT = SCRIPT_DIR.parents[2]
if str(PROJECT_ROOT / 'skills/orchestrate-offline-audit/scripts') not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT / 'skills/orchestrate-offline-audit/scripts'))

from audit_core.common import AuditError, load_json, sha256_file, validate_json
from audit_core.product_rag import canonical_product_name, catalog_product_name_values, ean13_is_valid
from audit_core.product_database import load_product_catalog
from audit_core.product_oss import manifest_key


OBSERVATION_SCHEMA_PATH = SKILL_ROOT / "references" / "observation-manifest.schema.json"
IMAGE_SUFFIXES = {".jpg", ".jpeg", ".png", ".webp"}
PRODUCT_CODE_PATTERN = re.compile(r"^[A-Za-z0-9][A-Za-z0-9._-]{0,63}$")
CATEGORY_PATTERNS: tuple[tuple[str, re.Pattern[str]], ...] = (
    ("facial_cleanser", re.compile(r"洁面乳|洗面奶")),
    ("shampoo", re.compile(r"洗发水|洗发露")),
    ("conditioner", re.compile(r"护发素|发膜")),
    ("body_wash", re.compile(r"沐浴露|沐浴乳")),
    ("mouthwash", re.compile(r"漱口水")),
    ("oral_spray", re.compile(r"口腔喷雾|口喷")),
    ("toothpaste", re.compile(r"牙膏")),
    ("toothbrush", re.compile(r"牙刷")),
    ("floss", re.compile(r"牙线棒|牙线")),
    ("repellent", re.compile(r"花露水|驱蚊")),
)
UNOBSERVED_MODIFIERS: tuple[tuple[str, re.Pattern[str]], ...] = (
    ("套盒专用", re.compile(r"套盒专用|套装专用")),
    ("箱规", re.compile(r"(?:12|24|36|48)箱规|带中箱|卡箱|外箱")),
    ("升级", re.compile(r"升级版|升级款|升级配方|全新升级")),
    ("代言版本", re.compile(r"代言人|华晨宇")),
    ("客户专供", re.compile(r"专供|定制|思家客|赛维泰|倍加洁|永辉|大润发|克劳丽")),
    ("国际标签", re.compile(r"国际|标签|越南|菲律宾|马来西亚|泰国|新加坡|香港|韩国|日本|英国")),
    ("线上", re.compile(r"线上")),
)
CORE_REMOVALS = (
    "aboutfocus",
    "oralshark",
    "小箭头",
    "参半",
    "线下ka",
    "线下",
    "线上",
    "套盒专用",
    "套装专用",
    "带中箱",
    "全新配方",
    "升级配方",
    "升级版",
    "升级款",
    "成人",
)


def _parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="按视觉名称和精确69码接入商品参考图，并与库存清单对账。"
    )
    subparsers = parser.add_subparsers(dest="command", required=True)

    match = subparsers.add_parser("match", help="查询数据库并生成无凭据图片接入计划")
    match.add_argument("--source-root", required=True)
    match.add_argument("--observations", required=True)
    match.add_argument("--output", required=True)

    validate = subparsers.add_parser("validate", help="重新查库并核验计划与原始图片，不上传")
    validate.add_argument("--plan-file", required=True)

    reconcile = subparsers.add_parser("reconcile", help="对账权威库存清单与正式catalog")
    reconcile.add_argument("--workbook", required=True)
    reconcile.add_argument("--output")
    reconcile.add_argument("--out-of-stock-text", default="无库存，未发出")
    return parser.parse_args()


def _write_json(path: Path, value: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(
        json.dumps(value, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )




def _normalize_text(value: Any) -> str:
    text = unicodedata.normalize("NFKC", str(value or "")).casefold()
    text = text.replace("參半", "参半")
    text = text.replace("双支", "2支").replace("单支", "1支").replace("只", "支")
    return re.sub(r"[^0-9a-z\u4e00-\u9fff+%]+", "", text)


def _core_text(value: Any) -> str:
    text = _normalize_text(value)
    for token in CORE_REMOVALS:
        text = text.replace(_normalize_text(token), "")
    text = re.sub(r"(?:12|24|36|48)箱规", "", text)
    return text


def _grams(value: Any) -> set[str]:
    text = _core_text(value)
    if len(text) < 2:
        return {text} if text else set()
    return {text[index : index + 2] for index in range(len(text) - 1)}


def _measurement_tokens(value: Any) -> set[str]:
    text = unicodedata.normalize("NFKC", str(value or "")).casefold()
    text = text.replace("双支", "2支").replace("单支", "1支").replace("只", "支")
    result: set[str] = set()
    for number_text, unit in re.findall(
        r"(\d+(?:\.\d+)?)\s*(ml|毫升|l|升|g|克|支|条|片|个)",
        text,
    ):
        number = float(number_text)
        canonical_unit = {"毫升": "ml", "升": "l", "克": "g"}.get(unit, unit)
        if canonical_unit == "l":
            number *= 1000
            canonical_unit = "ml"
        number_value = str(int(number)) if number.is_integer() else str(number)
        result.add(f"{number_value}{canonical_unit}")
    return result


def _category(value: Any) -> str | None:
    text = str(value or "")
    for label, pattern in CATEGORY_PATTERNS:
        if pattern.search(text):
            return label
    return None


def _candidate_metrics(
    observation: dict[str, Any],
    candidate: dict[str, str],
) -> dict[str, Any]:
    observed_context = " ".join(
        str(value or "")
        for value in (
            observation.get("observed_product_name"),
            observation.get("observed_specification"),
            observation.get("observed_variant"),
        )
    )
    candidate_name = str(candidate["product_name"])
    left = _core_text(observed_context)
    right = _core_text(candidate_name)
    sequence = SequenceMatcher(None, left, right).ratio() if left and right else 0.0
    left_grams = _grams(observed_context)
    right_grams = _grams(candidate_name)
    union = left_grams | right_grams
    jaccard = len(left_grams & right_grams) / len(union) if union else 0.0
    containment = bool(left and right and (left in right or right in left))

    observed_measurements = _measurement_tokens(observed_context)
    candidate_measurements = _measurement_tokens(candidate_name)
    measurement_conflict = bool(
        observed_measurements
        and candidate_measurements
        and not observed_measurements.intersection(candidate_measurements)
    )
    observed_category = _category(observed_context)
    candidate_category = _category(candidate_name)
    category_conflict = bool(
        observed_category
        and candidate_category
        and observed_category != candidate_category
    )

    variant = _normalize_text(observation.get("observed_variant"))
    variant_match = bool(variant and variant in _normalize_text(candidate_name))
    unobserved_modifiers = [
        label
        for label, pattern in UNOBSERVED_MODIFIERS
        if pattern.search(candidate_name) and not pattern.search(observed_context)
    ]
    score = 0.55 * sequence + 0.45 * jaccard
    if containment:
        score += 0.12
    if observed_measurements.intersection(candidate_measurements):
        score += 0.08
    if variant_match:
        score += 0.45
    compatible = (
        not category_conflict
        and not measurement_conflict
        and score >= 0.43
    )
    return {
        "score": round(score, 4),
        "sequence": round(sequence, 4),
        "bigram_jaccard": round(jaccard, 4),
        "variant_match": variant_match,
        "unobserved_modifiers": unobserved_modifiers,
        "modifier_penalty": len(unobserved_modifiers),
        "measurement_conflict": measurement_conflict,
        "category_conflict": category_conflict,
        "compatible": compatible,
    }


def _valid_exact_candidates(
    database_result: dict[str, Any],
    barcode: str,
) -> list[dict[str, str]]:
    result: list[dict[str, str]] = []
    seen_codes: set[str] = set()
    for raw in database_result.get("list") or []:
        if not isinstance(raw, dict):
            continue
        product_code = str(raw.get("product_code") or "").strip()
        product_name = str(raw.get("product_name") or "").strip()
        returned_barcode = str(raw.get("barcode") or "").strip()
        if (
            returned_barcode != barcode
            or not product_name
            or PRODUCT_CODE_PATTERN.fullmatch(product_code) is None
            or product_code in seen_codes
        ):
            continue
        seen_codes.add(product_code)
        result.append(
            {
                "product_code": product_code,
                "product_name": product_name,
                "barcode": returned_barcode,
            }
        )
    return result


def _select_candidate(
    observation: dict[str, Any],
    candidates: list[dict[str, str]],
) -> tuple[dict[str, str] | None, str, str, list[dict[str, Any]]]:
    scored = [
        {**candidate, **_candidate_metrics(observation, candidate)}
        for candidate in candidates
    ]
    scored.sort(
        key=lambda item: (
            not bool(item["compatible"]),
            int(item["modifier_penalty"]),
            -float(item["score"]),
            str(item["product_code"]),
        )
    )
    if not scored:
        return None, "not_found", "数据库没有同码且三字段有效的候选", scored

    approved_code = str(observation.get("approved_product_code") or "").strip()
    if approved_code:
        approved = next(
            (item for item in scored if item["product_code"] == approved_code),
            None,
        )
        if approved is None:
            return None, "approval_invalid", "人工批准编码不在精确同码候选中", scored
        if not approved["compatible"]:
            return None, "approval_conflict", "人工批准候选与可见名称、品类或规格冲突", scored
        return approved, "explicit_approved", "人工确认了精确同码且名称相容的数据库候选", scored

    compatible = [item for item in scored if item["compatible"]]
    if not compatible:
        return None, "name_conflict", "精确同码候选均与可见产品名、品类或规格不相容", scored
    if len(scored) == 1:
        return compatible[0], "database_unique", "数据库仅有一个精确同码且名称相容的候选", scored
    if len(compatible) == 1:
        return compatible[0], "name_unique", "同码候选中仅一个与可见名称、品类和规格相容", scored

    variant_matches = [item for item in compatible if item["variant_match"]]
    if len(variant_matches) == 1:
        return variant_matches[0], "variant_unique", "仅一个同码候选包含包装可见香型或款式", scored
    narrowed = variant_matches or compatible

    best_penalty = min(int(item["modifier_penalty"]) for item in narrowed)
    standard = [item for item in narrowed if int(item["modifier_penalty"]) == best_penalty]
    if len(standard) == 1:
        return (
            standard[0],
            "unmodified_unique",
            "仅一个名称相容候选不含包装未显示的专用、箱规、升级或渠道修饰",
            scored,
        )

    ranked = sorted(standard, key=lambda item: float(item["score"]), reverse=True)
    if len(ranked) == 1 or float(ranked[0]["score"]) - float(ranked[1]["score"]) >= 0.12:
        return ranked[0], "fuzzy_unique", "产品名称相容度相对次选形成安全唯一差值", scored
    return None, "ambiguous", "精确同码后仍有多个名称相容候选，禁止自动接入", scored


def _validate_source_root(path: Path) -> Path:
    if path.is_symlink() or (hasattr(path, "is_junction") and path.is_junction()) or not path.is_dir():
        raise AuditError("source-root 必须是操作者明确指定的普通素材目录")
    root = path.resolve()
    if root == Path(root.anchor) or root == PROJECT_ROOT or root in PROJECT_ROOT.parents:
        raise AuditError("source-root 不能是磁盘根目录或项目根目录")
    return root


def _inspect_source_folder(
    source_root: Path,
    observation: dict[str, Any],
) -> list[dict[str, Any]]:
    supplied = source_root / str(observation["source_folder"])
    if supplied.is_symlink() or (hasattr(supplied, "is_junction") and supplied.is_junction()):
        raise AuditError("单商品来源不得是链接或联接")
    folder = supplied.resolve()
    if (
        folder.is_symlink()
        or not folder.is_dir()
        or folder.parent != source_root.resolve()
    ):
        raise AuditError(f"观察项必须指向 source-root 的一个普通直接子目录：{folder}")
    children = sorted(folder.iterdir(), key=lambda item: item.name.casefold())
    if any(item.is_symlink() or (hasattr(item, "is_junction") and item.is_junction()) for item in children):
        raise AuditError("来源图片不得是链接或联接")
    nested = [item for item in children if item.is_dir()]
    if nested:
        raise AuditError(f"单商品文件夹不得嵌套目录：{nested[0]}")
    unsupported = [
        item
        for item in children
        if item.is_file() and item.suffix.lower() not in IMAGE_SUFFIXES
    ]
    if unsupported:
        raise AuditError(f"单商品文件夹含非图片文件：{unsupported[0]}")
    images = [item for item in children if item.is_file()]
    if not images:
        raise AuditError(f"单商品文件夹没有图片：{folder}")
    names = {item.name for item in images}
    for evidence_field in ("name_evidence_files", "barcode_evidence_files"):
        unknown = sorted(set(observation[evidence_field]) - names)
        if unknown:
            raise AuditError(
                f"{observation['source_folder']} 的 {evidence_field} 引用了不存在的文件："
                f"{unknown[0]}"
            )
    files: list[dict[str, Any]] = []
    for image_path in images:
        with Image.open(image_path) as image:
            width, height = image.size
        files.append(
            {
                "name": image_path.name,
                "sha256": sha256_file(image_path),
                "width": width,
                "height": height,
            }
        )
    return files


def _match(args: argparse.Namespace) -> int:
    source_root = _validate_source_root(Path(args.source_root))
    observations_path = Path(args.observations).resolve()
    manifest = load_json(observations_path)
    validate_json(manifest, OBSERVATION_SCHEMA_PATH)
    observations = list(manifest["observations"])
    source_ids = [str(item["source_id"]).casefold() for item in observations]
    source_folders = [str(item["source_folder"]).casefold() for item in observations]
    if len(source_ids) != len(set(source_ids)):
        raise AuditError("观察清单 source_id 重复")
    if len(source_folders) != len(set(source_folders)):
        raise AuditError("观察清单 source_folder 重复")
    for observation in observations:
        barcode = str(observation["barcode_69"])
        if not ean13_is_valid(barcode):
            raise AuditError(
                f"观察清单69码未通过EAN-13校验：{observation['source_folder']}={barcode}"
            )

    database_catalog = load_product_catalog()
    database_results: dict[str, dict[str, Any]] = {}
    barcodes = list(dict.fromkeys(str(item["barcode_69"]) for item in observations))
    for barcode in barcodes:
        database_results[barcode] = {"list": [
            {"barcode": p["barcode_69"], "product_name": p["product_name"], "product_code": p["product_code"]}
            for p in database_catalog["products"] if p["barcode_69"] == barcode
        ]}

    plan_products: list[dict[str, Any]] = []
    for observation in observations:
        barcode = str(observation["barcode_69"])
        database_result = database_results[barcode]
        candidates = _valid_exact_candidates(database_result, barcode)
        selected, status, basis, scored = _select_candidate(observation, candidates)
        source_files = _inspect_source_folder(source_root, observation)
        selected_identity = None
        target_directory = None
        if selected is not None:
            selected_identity = {
                "product_code": str(selected["product_code"]),
                "product_name": str(selected["product_name"]),
                "barcode_69": barcode,
            }
            target_directory = manifest_key(selected_identity["product_code"])
        plan_products.append(
            {
                "source_folder": str(observation["source_folder"]),
                "source_id": str(observation["source_id"]),
                "collection": source_root.name,
                "brand": str(observation["brand"]).strip(),
                "observed_product_name": str(observation["observed_product_name"]).strip(),
                "observed_specification": str(observation["observed_specification"]).strip(),
                "observed_variant": observation.get("observed_variant"),
                "barcode_69": barcode,
                "name_evidence_files": list(observation["name_evidence_files"]),
                "barcode_evidence_files": list(observation["barcode_evidence_files"]),
                "source_files": source_files,
                "data_source": "product_catalog.products",
                "selection_status": status,
                "selection_basis": basis,
                "selected": selected_identity,
                "image_manifest_key": target_directory,
                "candidates": scored,
                "approved_product_code": observation.get("approved_product_code"),
            }
        )

    plan = {
        "schema_version": "2.0",
        "kind": "product-oss-onboarding-plan",
        "created_at_utc": datetime.now(timezone.utc).isoformat(),
        "selection_policy": (
            "包装先视觉提取；数据库69码精确过滤；产品名、品类、规格和可见款式唯一模糊匹配；"
            "未显示的套盒、箱规、升级、代言、专供或国际修饰不得抢占；不能唯一收敛即失败关闭。"
        ),
        "database": database_catalog["data_source"],
        "observations_sha256": sha256_file(observations_path),
        "source_root": str(source_root),
        "products": plan_products,
    }
    output = Path(args.output).resolve()
    _write_json(output, plan)
    unresolved = [item for item in plan_products if item["selected"] is None]
    result = {
        "status": "planned" if not unresolved else "unresolved",
        "products": len(plan_products),
        "resolved": len(plan_products) - len(unresolved),
        "unresolved": len(unresolved),
        "database_rows": len(database_catalog["products"]),
        "output": str(output),
    }
    print(json.dumps(result, ensure_ascii=False, separators=(",", ":")))
    return 0 if not unresolved else 2


def _prepare_publish(plan: dict[str, Any]) -> list[dict[str, Any]]:
    """验证计划仍来自当前数据库及未变更原始素材；不执行上传或身份写入。"""
    if plan.get("schema_version") != "2.0" or plan.get("kind") != "product-oss-onboarding-plan":
        raise AuditError("plan-file 不是受支持的数据库到 OSS 接入计划，请重新 match")
    database = load_product_catalog()
    if any((plan.get("database") or {}).get(key) != database["data_source"].get(key)
           for key in ("type", "sha256", "image_manifest_keys_sha256")):
        raise AuditError("商品数据库自计划生成后已变化或计划不是数据库来源，必须重新 match")
    by_code = {p["product_code"]: p for p in database["products"]}
    items = plan.get("products")
    if not isinstance(items, list) or not items:
        raise AuditError("接入计划缺少商品")
    source_path = Path(str(plan.get("source_root") or ""))
    if not source_path.is_absolute():
        raise AuditError("接入计划必须记录显式来源绝对路径")
    source_root = _validate_source_root(source_path)
    operations, seen = [], set()
    for item in items:
        selected = item.get("selected") or {}
        current = by_code.get(selected.get("product_code"))
        if current is None or any(selected.get(key) != current[key] for key in ("product_code", "product_name", "barcode_69")):
            raise AuditError("图片接入计划中的商品身份未通过当前数据库核验")
        if current["product_code"] in seen:
            raise AuditError("多个来源不能同时覆盖同一商品图片集")
        seen.add(current["product_code"])
        expected_key = manifest_key(current["product_code"])
        if item.get("image_manifest_key") != expected_key or item.get("barcode_69") != current["barcode_69"]:
            raise AuditError("计划商品图片集或观察条码与数据库不一致")
        candidates = [{"product_code": p["product_code"], "product_name": p["product_name"], "barcode": p["barcode_69"]}
                      for p in database["products"] if p["barcode_69"] == current["barcode_69"]]
        selected_again, _, _, _ = _select_candidate(item, candidates)
        if selected_again is None or selected_again["product_code"] != current["product_code"]:
            raise AuditError("视觉观察已不能唯一收敛到所选商品，请重新 match")
        files = _inspect_source_folder(source_root, item)
        if files != item.get("source_files"):
            raise AuditError("来源图片自计划生成后已变化，请重新 match")
        operations.append({"product_code": current["product_code"], "source_dir": str(source_root / item["source_folder"]),
                           "image_manifest_key": expected_key, "image_count": len(files)})
    return operations


def _validate_plan(args: argparse.Namespace) -> int:
    operations = _prepare_publish(load_json(Path(args.plan_file)))
    print(json.dumps({"status": "validated", "uploaded": False, "products": operations,
                      "next_step": "审查后使用 shared/product-database/publish_product_images.py 逐商品预演；写入获准才加 --apply"},
                     ensure_ascii=False))
    return 0


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _name_compatibility(
    workbook_name: str,
    product: dict[str, Any],
) -> dict[str, Any]:
    exact = canonical_product_name(workbook_name) == canonical_product_name(product["product_name"])
    ranked: list[dict[str, Any]] = []
    observation = {
        "observed_product_name": workbook_name,
        "observed_specification": "",
        "observed_variant": None,
    }
    for candidate_name in catalog_product_name_values(product):
        metrics = _candidate_metrics(
            observation,
            {
                "product_code": str(product["product_code"]),
                "product_name": candidate_name,
                "barcode": str(product["barcode_69"]),
            },
        )
        ranked.append({"catalog_value": candidate_name, **metrics})
    ranked.sort(key=lambda item: -float(item["score"]))
    best = ranked[0] if ranked else {"compatible": False, "score": 0.0}
    return {
        "status": "exact" if exact else ("fuzzy" if best["compatible"] else "mismatch"),
        "compatible": bool(exact or best["compatible"]),
        "score": 1.0 if exact else float(best["score"]),
        "matched_catalog_value": str(best.get("catalog_value") or product["product_name"]),
    }


def _reconcile(args: argparse.Namespace) -> int:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise AuditError("读取权威xlsx需要项目依赖 openpyxl") from exc

    workbook_path = Path(args.workbook).resolve()
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    if len(workbook.worksheets) != 1:
        raise AuditError("权威清单必须且只能有一个工作表，避免隐式选表")
    sheet = workbook.worksheets[0]
    rows = sheet.iter_rows(values_only=True)
    try:
        header_values = next(rows)
    except StopIteration as exc:
        raise AuditError("权威清单为空") from exc
    headers = {_cell_text(value): index for index, value in enumerate(header_values)}
    required = {
        "*商品名称",
        "*商品编码/组合编码",
        "*申请数量",
        "快递单号",
        "69码",
    }
    missing_headers = sorted(required - set(headers))
    if missing_headers:
        raise AuditError(f"权威清单缺少列：{missing_headers[0]}")

    authority_rows: list[dict[str, Any]] = []
    for excel_row, values in enumerate(rows, start=2):
        product_code = _cell_text(values[headers["*商品编码/组合编码"]])
        if not product_code:
            continue
        authority_rows.append(
            {
                "excel_row": excel_row,
                "product_name": _cell_text(values[headers["*商品名称"]]),
                "product_code": product_code,
                "quantity": _cell_text(values[headers["*申请数量"]]),
                "tracking": _cell_text(values[headers["快递单号"]]),
                "barcode_69": _cell_text(values[headers["69码"]]),
            }
        )
    sheet_title = sheet.title
    workbook.close()
    stock_rows = [
        row
        for row in authority_rows
        if row["tracking"] != str(args.out_of_stock_text).strip()
    ]
    out_of_stock_rows = [row for row in authority_rows if row not in stock_rows]

    def duplicate_values(rows_value: list[dict[str, Any]], key: str) -> dict[str, int]:
        counts: dict[str, int] = {}
        for row in rows_value:
            value = str(row[key])
            counts[value] = counts.get(value, 0) + 1
        return {value: count for value, count in counts.items() if count > 1}

    catalog = load_product_catalog()
    products = list(catalog["products"])
    by_code = {str(product["product_code"]): product for product in products}
    stock_codes = {str(row["product_code"]) for row in stock_rows}
    matches: list[dict[str, Any]] = []
    missing_products: list[dict[str, Any]] = []
    field_mismatches: list[dict[str, Any]] = []
    for row in stock_rows:
        product = by_code.get(str(row["product_code"]))
        if product is None:
            missing_products.append(row)
            continue
        name_result = _name_compatibility(str(row["product_name"]), product)
        barcode_match = str(row["barcode_69"]) == str(product["barcode_69"])
        match = {
            "excel_row": row["excel_row"],
            "product_code": row["product_code"],
            "excel_product_name": row["product_name"],
            "catalog_product_name": str(product["product_name"]),
            "excel_barcode_69": row["barcode_69"],
            "catalog_barcode_69": str(product["barcode_69"]),
            "product_code_match": True,
            "barcode_match": barcode_match,
            "name_match_status": name_result["status"],
            "name_score": name_result["score"],
            "name_compatible": name_result["compatible"],
        }
        matches.append(match)
        if not barcode_match or not name_result["compatible"]:
            field_mismatches.append(match)

    catalog_extra = [
        {
            "product_code": str(product["product_code"]),
            "product_name": str(product["product_name"]),
            "barcode_69": str(product["barcode_69"]),
        }
        for product in products
        if str(product["product_code"]) not in stock_codes
    ]
    duplicate_stock_codes = duplicate_values(stock_rows, "product_code")
    duplicate_all_codes = duplicate_values(authority_rows, "product_code")
    passed = not any(
        (
            missing_products,
            catalog_extra,
            field_mismatches,
            duplicate_stock_codes,
            duplicate_all_codes,
        )
    ) and len(stock_rows) == len(products)
    result = {
        "passed": passed,
        "workbook": str(workbook_path),
        "sheet": sheet_title,
        "out_of_stock_marker": str(args.out_of_stock_text).strip(),
        "workbook_rows": len(authority_rows),
        "out_of_stock_rows": len(out_of_stock_rows),
        "expected_stock_products": len(stock_rows),
        "catalog_products": len(products),
        "products_with_image_manifest": sum(bool(product.get("image_manifest_key")) for product in products),
        "database": catalog["data_source"],
        "matched_products": len(matches),
        "exact_barcode_matches": sum(item["barcode_match"] for item in matches),
        "exact_name_matches": sum(item["name_match_status"] == "exact" for item in matches),
        "fuzzy_name_matches": sum(item["name_match_status"] == "fuzzy" for item in matches),
        "missing_products": missing_products,
        "catalog_extra_products": catalog_extra,
        "field_mismatches": field_mismatches,
        "duplicate_stock_product_codes": duplicate_stock_codes,
        "duplicate_all_product_codes": duplicate_all_codes,
        "out_of_stock_products": out_of_stock_rows,
        "matches": matches,
    }
    if args.output:
        _write_json(Path(args.output).resolve(), result)
    print(
        json.dumps(
            {
                key: value
                for key, value in result.items()
                if key not in {"matches", "out_of_stock_products"}
            },
            ensure_ascii=False,
            separators=(",", ":"),
        )
    )
    return 0 if passed else 2


def main() -> int:
    args = _parse_args()
    if args.command == "match":
        return _match(args)
    if args.command == "validate":
        return _validate_plan(args)
    return _reconcile(args)


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except AuditError as exc:
        print(json.dumps({"error": str(exc)}, ensure_ascii=False), file=sys.stderr)
        raise SystemExit(2) from exc
