"""Deterministic, blocking material diagnostics for safely received archives.

This result describes submitted evidence, never an amount approval or a new
business scenario. Model observations are limited to the submitted visual files.
"""
from __future__ import annotations

from .paths import PROJECT_ROOT

from collections import defaultdict
from copy import deepcopy
from pathlib import Path
import math
import re
from typing import Any
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile, ZipFile

from .common import AuditError, normalize_text, now_utc, validate_json


RESULT_SCHEMA = PROJECT_ROOT / "skills/orchestrate-offline-audit/references/contracts" / "material-diagnostic-result.schema.json"
CLASSIFICATION_REJECTION_SCHEMA = PROJECT_ROOT / "skills/orchestrate-offline-audit/references/contracts" / "classification-rejection-result.schema.json"
EVIDENCE_SCHEMA = PROJECT_ROOT / "skills" / "orchestrate-offline-audit" / "references" / "material-diagnostic-evidence.schema.json"
SCENARIO_LABELS = {
    "personnel_incentive": "人员激励", "promotional_display": "堆头/陈列",
    "poster_material": "海报/展示道具", "other_expense": "其他费用",
    "maintenance_fee": "维护费用", "giveaway_promotion": "额外搭赠",
    "price_difference_support": "价格补差", "pos_target_incentive": "POS达标激励",
    "entry_fee": "条码费", "self_procured_gift_material": "客户自采赠品物料",
}
ROLE_LABELS = {
    "signed_promotional_contract": "已签促销合同", "settlement": "结算单",
    "stamped_pos_data": "盖章 POS 视觉材料", "activity_photo": "活动现场照片",
    "invoice_receipt": "发票或收据", "payment_record": "付款或转账凭证",
    "supporting_document": "费用专项支持材料", "sales_outbound_record": "系统销售或出库单",
    "store_receipt": "门店收货凭证", "special_approval": "特殊审批材料",
    "system_deduction_proof": "系统扣费证明", "entry_contract": "已签进场合同",
    "pos_spreadsheet": "POS 销售电子表（.xlsx 或 .xlsm）",
    "activity_return_workbook": "活动返图工作簿（.xls）",
    "other": "其他材料", "unknown": "类型未确认材料",
}
# An alternative group means one material role can establish presence. Detailed
# content, validity, coverage and amount checks still belong to the scenario.
REQUIREMENTS = {
    "personnel_incentive": [("settlement",), ("payment_record",), ("pos_spreadsheet",)],
    "promotional_display": [("signed_promotional_contract",), ("activity_photo",), ("pos_spreadsheet",)],
    "poster_material": [("signed_promotional_contract",), ("invoice_receipt",), ("settlement",), ("activity_photo",)],
    "other_expense": [("signed_promotional_contract",), ("settlement",), ("supporting_document", "invoice_receipt")],
    "maintenance_fee": [("signed_promotional_contract",), ("settlement",), ("stamped_pos_data",), ("pos_spreadsheet",), ("supporting_document", "activity_photo")],
    "giveaway_promotion": [("signed_promotional_contract",), ("settlement",), ("sales_outbound_record",), ("store_receipt",), ("activity_photo",)],
    "price_difference_support": [("signed_promotional_contract",), ("settlement",), ("stamped_pos_data",), ("pos_spreadsheet",), ("activity_photo",)],
    "pos_target_incentive": [("signed_promotional_contract",), ("settlement",), ("stamped_pos_data",), ("pos_spreadsheet",), ("activity_photo", "store_receipt", "supporting_document")],
    "entry_fee": [("entry_contract",), ("activity_photo",), ("system_deduction_proof",)],
    "self_procured_gift_material": [("signed_promotional_contract",), ("settlement",), ("invoice_receipt",), ("stamped_pos_data",), ("pos_spreadsheet",), ("activity_photo",)],
}
SINGLETON_ROLES = {
    "personnel_incentive": {"settlement", "pos_spreadsheet"},
    "promotional_display": {"signed_promotional_contract", "pos_spreadsheet"},
    "poster_material": {"signed_promotional_contract", "settlement", "invoice_receipt"},
    "other_expense": {"signed_promotional_contract", "settlement"},
    "maintenance_fee": {"signed_promotional_contract", "settlement", "pos_spreadsheet"},
    "giveaway_promotion": {"signed_promotional_contract", "settlement", "sales_outbound_record"},
    "price_difference_support": {"signed_promotional_contract", "settlement", "pos_spreadsheet"},
    "pos_target_incentive": {"signed_promotional_contract", "settlement", "pos_spreadsheet"},
    "entry_fee": {"entry_contract"},
    "self_procured_gift_material": {"signed_promotional_contract", "settlement", "invoice_receipt", "pos_spreadsheet"},
}


def _strings(values: Any) -> list[str]:
    return list(dict.fromkeys(str(value).strip() for value in values or [] if str(value).strip()))


def _readability_limitations(values: Any) -> list[str]:
    """Business uncertainty is not evidence that a submitted file is unreadable."""
    visual_content = r"文字|字迹|印章|名称|数字|图片|图像|照片|扫描件|页面|分辨率|水印"
    reading_problem = re.compile(
        r"缺页|损坏|破损|乱码|不可读|读不清|看不清|辨认不清|"
        r"(?:无法|未能|不能|难以)(?:打开|读取|识读|看清|辨认|辨清)|"
        rf"(?:{visual_content})[^，。；;\n]{{0,24}}(?:模糊|不清晰|不够清晰|难辨|遮挡|重叠)|"
        rf"(?:模糊|遮挡|重叠)[^，。；;\n]{{0,24}}(?:{visual_content})|"
        r"(?:页面|图片|照片|扫描件)[^，。；;\n]{0,12}(?:缺失|不完整|截断|裁切|裁剪)|"
        r"(?:低分辨率|分辨率过低|像素过低)"
    )
    # Keep the specific reading limitation in the issue; all other original
    # caveats remain in the saved document observations and evidence cards.
    return _strings(part.strip() for value in _strings(values)
                    for part in re.split(r"[；;。\n]", value)
                    if reading_problem.search(part))


def _issue(code: str, title: str, observed: str, files: list[str], *, expected: str, action: str, confidence: str = "high") -> dict[str, Any]:
    return {"code": code, "title": title, "observed": observed, "source_files": _strings(files),
            "expected": expected, "impact": "材料尚未形成可确认的核销依据，暂不能自动核销。",
            "action": action, "confidence": confidence}


def _document_groups(items: list[dict[str, Any]]) -> list[list[dict[str, Any]]]:
    """Merge only affirmative distinct pages of the same numbered document."""
    grouped: list[list[dict[str, Any]]] = []
    for item in items:
        number = str(item.get("document_number") or "").strip()
        page = item.get("page_number")
        match = None
        if number and isinstance(page, int) and not isinstance(page, bool) and page > 0:
            for group in grouped:
                if all(str(old.get("document_number") or "").strip() == number
                       and isinstance(old.get("page_number"), int)
                       and old["page_number"] > 0 and old["page_number"] != page
                       and (not old.get("total_pages") or not item.get("total_pages")
                            or old["total_pages"] == item["total_pages"])
                       for old in group):
                    match = group
                    break
        if match is None:
            grouped.append([item])
        else:
            match.append(item)
    return grouped


def _validate_spreadsheet_archive(path: Path) -> None:
    """Check XLSX metadata before a parser loads shared strings or styles."""
    from .archive_input import (
        ArchiveInputError, MAX_ARCHIVE_FILES, MAX_MEMBER_BYTES,
        MAX_TOTAL_BYTES, MAX_COMPRESSION_RATIO, _zip_member_parts,
    )

    with ZipFile(path) as archive:
        members = archive.infolist()
        if len(members) > MAX_ARCHIVE_FILES:
            raise ArchiveInputError("电子表 ZIP 成员数超出安全限制")
        total_bytes = 0
        seen = set()
        for member in members:
            parts = _zip_member_parts(member)
            key = "/".join(parts).casefold()
            if key in seen:
                raise ArchiveInputError("电子表 ZIP 成员路径重复或大小写冲突")
            seen.add(key)
            if member.is_dir():
                continue
            if member.file_size > MAX_MEMBER_BYTES:
                raise ArchiveInputError("电子表 ZIP 单个成员展开量超出安全限制")
            total_bytes += member.file_size
            if total_bytes > MAX_TOTAL_BYTES:
                raise ArchiveInputError("电子表 ZIP 总展开量超出安全限制")
            if member.file_size > 10 * 1024 * 1024 and (
                member.compress_size <= 0
                or member.file_size / member.compress_size > MAX_COMPRESSION_RATIO
            ):
                raise ArchiveInputError("电子表 ZIP 压缩比超出安全限制")


def _pos_spreadsheet_structure(source: dict[str, Any], scenario: str) -> tuple[bool, str]:
    """Inspect bounded local cells without retaining values or doing audit math.

    The ordinary readers also reconcile and aggregate business values. Reuse
    their header vocabulary here, but always close even an invalid workbook.
    Permission, missing-source and dependency failures remain execution errors.
    """
    from openpyxl import LXML, load_workbook
    from openpyxl.utils.exceptions import InvalidFileException
    from .excel_sources import PERSONNEL_HEADERS, DISPLAY_HEADERS, MAINTENANCE_POS_HEADERS
    if LXML:
        from lxml.etree import XMLSyntaxError
    else:
        XMLSyntaxError = ParseError

    target_aliases = {
        "period": {"日期", "期间", "活动时间", "销售期间"},
        "store": {"系统门店名称", "门店名称", "门店", "客户名称"},
        "amount": {"销售金额", "销售额", "销售合计", "金额"},
    }
    self_aliases = {
        "period": {"日期", "期间", "活动时间", "活动周期", "销售期间"},
        "store": {"系统门店名称", "门店名称", "门店", "客户名称"},
        "quantity": {"销售数量", "销量", "数量"},
        "amount": {"销售金额", "销售收入", "销售额", "销售合计", "金额"},
    }
    specifications = {
        "personnel_incentive": (PERSONNEL_HEADERS, {"store", "barcode", "product", "quantity"}),
        "promotional_display": (DISPLAY_HEADERS, {"customer", "product_code", "barcode", "product", "quantity"}),
        "maintenance_fee": (MAINTENANCE_POS_HEADERS, {"product", "quantity", "amount"}),
        "price_difference_support": (MAINTENANCE_POS_HEADERS, {"product", "quantity", "amount"}),
        "pos_target_incentive": (target_aliases, set(target_aliases)),
        "self_procured_gift_material": (self_aliases, set(self_aliases)),
    }
    if scenario not in specifications:
        return False, "本包尚无可适用的 POS 电子表结构要求。"
    aliases, required = specifications[scenario]
    normalized = {field: {normalize_text(label) for label in labels} for field, labels in aliases.items()}
    path = Path(source["path"])
    # Keep security/permission exceptions outside the recoverable parser block.
    try:
        _validate_spreadsheet_archive(path)
    except BadZipFile:
        return False, "电子表文件损坏，或内部结构无法按适用销售表读取。"
    workbook = None
    try:
        workbook = load_workbook(path, read_only=True, data_only=True, keep_links=False)
        if len(workbook.worksheets) > 64:
            return False, "工作表数量超过 64 份检查范围，无法确认完整范围内的唯一主表。"
        candidates = []
        for worksheet in workbook.worksheets:
            rows = worksheet.iter_rows(min_row=1, max_row=1050, max_col=64, values_only=True)
            columns = None
            valid_detail = False
            for row_number, values in enumerate(rows, 1):
                if columns is None and row_number <= 50:
                    mapping = {}
                    for column, value in enumerate(values):
                        label = normalize_text(value)
                        for field, options in normalized.items():
                            if label in options:
                                mapping.setdefault(field, column)
                    if required.issubset(mapping):
                        columns = mapping
                    continue
                if columns is None:
                    break
                if not all(values[columns[field]] not in (None, "") for field in required):
                    continue
                if any(normalize_text(values[columns[field]]) in {"合计", "总计"}
                       for field in required - {"quantity", "amount"}):
                    continue
                numbers_valid = True
                for field in required & {"quantity", "amount"}:
                    value = values[columns[field]]
                    try:
                        if isinstance(value, bool):
                            raise ValueError("boolean")
                        if not math.isfinite(float(str(value).replace(",", ""))):
                            raise ValueError("nonfinite")
                    except (TypeError, ValueError):
                        numbers_valid = False
                if numbers_valid:
                    valid_detail = True
                    break
            if columns is not None:
                candidates.append(valid_detail)
        if len(candidates) == 1 and candidates[0]:
            return True, "程序已识别适用的销售表头及可读取明细；尚未执行金额核验。"
        if len(candidates) > 1:
            return False, "存在多份适用销售表格，主表尚未唯一确认。"
        return False, "未在检查范围内识别到适用的完整销售表头和可读取明细。"
    except (AuditError, BadZipFile, InvalidFileException, ParseError, XMLSyntaxError, KeyError, IndexError, ValueError):
        return False, "电子表文件损坏，或内部结构无法按适用销售表读取。"
    finally:
        if workbook is not None:
            workbook.close()


def build_material_diagnostic_result(case: dict[str, Any], evidence: dict[str, Any]) -> dict[str, Any]:
    validate_json(evidence, EVIDENCE_SCHEMA)
    archives = list(case.get("archives") or [])
    observed_archives = list(evidence.get("archives") or [])
    by_id = {str(item.get("archive_id")): item for item in observed_archives}
    expected_ids = [str(item.get("archive_id")) for item in archives]
    if len(set(expected_ids)) != len(expected_ids) or len(by_id) != len(observed_archives) or set(by_id) != set(expected_ids):
        raise AuditError("材料诊断证据与归档来源不一致")
    archive_results = []
    for archive in archives:
        archive_id = str(archive["archive_id"])
        observed = by_id[archive_id]
        inventory = []
        local_sources = {str(source["file_id"]): source for source in archive.get("inventory") or []}
        for source in archive.get("inventory") or []:
            inventory.append({key: deepcopy(source.get(key)) for key in (
                "file_id", "source_file", "sha256", "suffix", "kind", "parent_source", "limitations"
            ) if key in source})
        index = {str(item.get("file_id")): item for item in inventory}
        if len(index) != len(inventory):
            raise AuditError("材料诊断包含重复来源编号")
        documents = deepcopy(observed.get("documents") or [])
        document_ids = [str(doc.get("file_id")) for doc in documents]
        visual_ids = {str(item["file_id"]) for item in inventory if item.get("kind") == "visual"}
        if len(set(document_ids)) != len(document_ids) or set(document_ids) != visual_ids:
            raise AuditError("材料诊断必须逐份覆盖全部视觉来源，且不能读取电子表格")
        for document in documents:
            document["source_file"] = str(index[str(document["file_id"])]["source_file"])
            if document.get("document_type") not in ROLE_LABELS or document.get("document_type") in {"pos_spreadsheet", "activity_return_workbook"}:
                raise AuditError("材料诊断返回不支持的视觉材料角色")
        issues = []
        candidates = deepcopy(observed.get("scenario_candidates") or [])
        if any(item.get("scenario") not in SCENARIO_LABELS for item in candidates):
            raise AuditError("材料诊断返回未登记的核销类型")
        hint = archive.get("scenario_hint")
        if hint not in SCENARIO_LABELS:
            hint = None
        # Intake derives this unique marker from the submitted ZIP name. Model
        # candidates remain technical observations and never select a checklist
        # or create a cross-scenario business finding.
        scenario = hint
        archive_name = str(archive.get("source_archive") or archive_id).replace("\\", "/").rsplit("/", 1)[-1]
        if not scenario:
            issues.append(_issue("scenario_unconfirmed", "ZIP 名称分类标记待确认",
                f"{archive_name}：ZIP 名称缺少唯一分类标记（未标明或包含多个分类标记）。",
                [archive_name], expected="ZIP 名称须包含唯一的已登记核销类型分类标记。",
                action="确认核销类型，并使 ZIP 文件名只包含该类型的分类标记。"))
        unknown_files = []
        for document in documents:
            readability_limits = _readability_limitations(document.get("limitations"))
            if document.get("document_type") in {"unknown", "other"} or document.get("confidence") == "low":
                unknown_files.append(document["source_file"])
                facts = "；".join(_strings(document.get("visible_facts")))
                limits = "；".join(_strings(document.get("limitations")))
                issues.append(_issue("material_unconfirmed", "材料用途未确认", f"{document['source_file']}：{limits or facts or '可见内容不足以可靠确认材料用途'}。", [document["source_file"]],
                    expected="原始材料内容清晰且能够确认用途。", action="确认该文件的用途；内容不清时提交完整清晰版本。", confidence="medium"))
            elif readability_limits:
                unknown_files.append(document["source_file"])
                issues.append(_issue("material_unreadable", "文件部分内容无法确认", f"{document['source_file']}：{'；'.join(readability_limits)}。", [document["source_file"]],
                    expected="提交文件的全部页面及关键内容清晰完整。", action="补充缺失页或关键内容清晰可读的版本。", confidence="medium"))
        for item in inventory:
            if item.get("kind") in {"unsupported", "container"} and (item.get("kind") == "unsupported" or item.get("limitations")):
                unknown_files.append(str(item["source_file"]))
                issues.append(_issue("material_unreadable", "文件内容无法读取", f"{item['source_file']}：{'；'.join(_strings(item.get('limitations'))) or '当前无法读取此格式的内容'}。", [str(item["source_file"])],
                    expected="材料应能安全打开并读取内容。", action="提供可读取的原始文件，或将相关文档转为清晰图片或 PDF。"))
        hashes: dict[str, list[dict[str, Any]]] = defaultdict(list)
        for item in inventory:
            digest = str(item.get("sha256") or "")
            if digest:
                hashes[digest].append(item)
        duplicate_ids = set()
        for group in hashes.values():
            if len(group) > 1:
                files = [str(item["source_file"]) for item in group]
                duplicate_ids.update(str(item["file_id"]) for item in group[1:])
                issues.append(_issue("duplicate_submission", "文件重复提交", f"以下 {len(files)} 份文件内容完全相同：{'、'.join(files)}。", files,
                    expected="同一份材料只需提交一次；重复副本不能增加材料覆盖。", action="保留其中一份，移除内容完全相同的重复副本。"))
        roles: dict[str, list[dict[str, Any]]] = defaultdict(list)
        unconfirmed_pos_files = []
        for document in documents:
            if str(document["file_id"]) not in duplicate_ids and document.get("confidence") in {"high", "medium"} and document.get("document_type") not in {"other", "unknown"}:
                roles[str(document["document_type"])].append(document)
        for item in inventory:
            if item.get("kind") == "spreadsheet" and str(item["file_id"]) not in duplicate_ids:
                suffix = str(item.get("suffix") or "").lower()
                source_limitations = _strings(item.get("limitations"))
                if source_limitations:
                    unknown_files.append(str(item["source_file"]))
                    issues.append(_issue("material_unconfirmed", "工作簿内容或用途未确认", f"{item['source_file']}：{'；'.join(source_limitations)}。", [str(item["source_file"])],
                        expected="完整读取工作簿内容及其中承载的活动图片后，才能确认材料覆盖。", action="确认工作簿用途，并补充可读取的工作簿或单独提供其中的原始图片。", confidence="medium"))
                if suffix in {".xlsx", ".xlsm"}:
                    recognized, basis = _pos_spreadsheet_structure(local_sources[str(item["file_id"])], scenario)
                    item["spreadsheet_check"] = {"status": "structure_recognized" if recognized else "structure_unconfirmed", "basis": basis}
                    if recognized:
                        roles["pos_spreadsheet"].append(item)
                    else:
                        unconfirmed_pos_files.append(str(item["source_file"]))
                        if scenario not in {"personnel_incentive", "promotional_display", "maintenance_fee", "price_difference_support", "pos_target_incentive", "self_procured_gift_material"}:
                            issues.append(_issue("material_unconfirmed", "电子表用途未确认", f"{item['source_file']}：{basis}", [str(item["source_file"])],
                                expected="依据核销方式确认该电子表的实际用途。", action="确认该电子表用途；核销类型未确定时先确认核销方式。", confidence="medium"))
                elif suffix == ".xls" and scenario == "self_procured_gift_material" and not source_limitations:
                    roles["activity_return_workbook"].append(item)
                else:
                    unknown_files.append(str(item["source_file"]))
                    if not source_limitations:
                        issues.append(_issue("material_unconfirmed", "电子表用途未确认", f"{item['source_file']} 为 {suffix or '未确认格式'} 电子表；不能将旧式活动返图工作簿直接视为 POS 销售电子表。", [str(item["source_file"])],
                            expected="POS 明细与活动返图分别按各自用途核验。", action="确认该工作簿用途；如用于 POS 明细，提供可读取的 .xlsx 或 .xlsm 版本。", confidence="medium"))
        if scenario:
            for alternatives in REQUIREMENTS[scenario]:
                if any(roles[role] for role in alternatives):
                    continue
                label = "或".join(ROLE_LABELS[role] for role in alternatives)
                # Unknown or unreadable inputs might contain this role. We can
                # establish an absent supported spreadsheet from the inventory.
                potential_files = list(unknown_files)
                if alternatives == ("entry_contract",):
                    potential_files.extend(str(doc["source_file"]) for doc in roles["signed_promotional_contract"])
                if alternatives == ("pos_spreadsheet",):
                    potential_files = list(unconfirmed_pos_files)
                uncertain = bool(potential_files)
                title = f"{label}{'未确认' if uncertain else '未提交'}"
                facts = (f"本包未能确认{label}；以下文件用途或内容仍无法确认：{'、'.join(_strings(potential_files))}。"
                         if uncertain else f"已盘点本包全部材料，未发现{label}。")
                if alternatives == ("pos_spreadsheet",) and uncertain:
                    facts += " " + "；".join(f"{item['source_file']}：{item['spreadsheet_check']['basis']}" for item in inventory if item.get("spreadsheet_check", {}).get("status") == "structure_unconfirmed")
                if alternatives == ("entry_contract",) and roles["signed_promotional_contract"]:
                    facts += " 普通促销合同不能直接替代已签进场合同，需要确认其是否包含本次进场条款。"
                issues.append(_issue("required_material_unconfirmed" if uncertain else "missing_material", title, facts, _strings(potential_files) if uncertain else [],
                    expected=f"本核销方式需要{label}。", action=f"{'确认现有文件是否包含，并补充可读取的' if uncertain else '补交'}{label}。", confidence="medium" if uncertain else "high"))
            required_roles = {role for alternatives in REQUIREMENTS[scenario] for role in alternatives}
            for role in sorted(required_roles & SINGLETON_ROLES[scenario]):
                groups = _document_groups(roles[role])
                if len(groups) > 1:
                    files = [str(item["source_file"]) for group in groups for item in group]
                    issues.append(_issue("singleton_role_ambiguous", f"{ROLE_LABELS[role]}主件待确认", f"识别到 {len(groups)} 份{ROLE_LABELS[role]}材料：{'、'.join(files)}；尚不能确定主件与补充件关系，不能据此认定实际重复。", files,
                        expected="主件唯一，或各份材料之间的页码、附件、补充关系清楚。", action="确认各文件用途和主件；属于同一文档的不同页或附件时注明关联。", confidence="medium"))
        # Even a complete inventory cannot substitute for the failed specialized
        # role binding or the remaining deterministic business checks.
        if not issues:
            issues.append(_issue("business_checks_pending", "材料组织需确认", "已完成可读取材料的 AI 用途识别及材料盘点；现有材料尚未完成对应核销方式的完整业务和金额核验。", [archive_name],
                expected="材料角色确认后，需继续完成对应核销规则和金额核验。", action="确认材料主件和补充件关系后继续核销；不能将本次材料诊断视为核销批准。", confidence="high"))
        archive_results.append({"archive_id": archive_id, "source_archive": archive_name, "scenario": scenario,
            "scenario_label": SCENARIO_LABELS.get(scenario, "核销类型待确认"), "scenario_hint": hint,
            "scenario_candidates": candidates, "inventory": inventory, "documents": documents, "issues": issues})

    # Compare only this invocation's diagnosed archives. Cross-archive reuse is
    # reported after role checks, so each package keeps its submitted evidence.
    by_archive_id = {item["archive_id"]: item for item in archive_results}
    archive_hashes: dict[str, list[str]] = defaultdict(list)
    source_hashes: dict[str, list[tuple[str, str]]] = defaultdict(list)
    for archive in archives:
        archive_id = str(archive["archive_id"])
        digest = str(archive.get("archive_sha256") or "")
        if digest:
            archive_hashes[digest].append(archive_id)
        for source in archive.get("inventory") or []:
            digest = str(source.get("sha256") or "")
            if digest:
                source_hashes[digest].append((archive_id, str(source["source_file"])))

    def append_shared_issue(archive_ids: list[str], issue: dict[str, Any]) -> None:
        for archive_id in dict.fromkeys(archive_ids):
            target = by_archive_id[archive_id]
            target["issues"] = [item for item in target["issues"] if item["code"] != "business_checks_pending"]
            target["issues"].append(deepcopy(issue))

    for archive_ids in archive_hashes.values():
        if len(archive_ids) > 1:
            names = [by_archive_id[archive_id]["source_archive"] for archive_id in archive_ids]
            append_shared_issue(archive_ids, _issue("duplicate_submission", "压缩包重复提交",
                f"本次提交的 {len(names)} 个压缩包字节内容完全相同：{'、'.join(names)}。", names,
                expected="同一资料包不应作为多份独立提交重复核销。", action="确认本次应保留的资料包，移除误交的重复包。"))
    for sources in source_hashes.values():
        archive_ids = list(dict.fromkeys(archive_id for archive_id, _ in sources))
        if len(archive_ids) <= 1:
            continue
        names = [f"{by_archive_id[archive_id]['source_archive']} / {source_file}" for archive_id, source_file in sources]
        append_shared_issue(archive_ids, _issue("duplicate_submission", "跨资料包文件内容重复",
            f"本次 {len(archive_ids)} 个资料包中的以下文件字节内容完全相同：{'、'.join(names)}；已保留各包的材料归属，需确认是否为误交或有明确的共同适用关系。", names,
            expected="不同资料包使用相同文件时，应明确其各自用途，不能把重复副本当作独立执行证据。",
            action="确认所列文件对各资料包的适用关系；误交时移除重复副本，属于共同材料时说明关联。"))
    result = {"schema_version": "1.0", "kind": "material_diagnostic", "diagnostic_only": True,
        "generated_at": now_utc(), "summary": {"conclusion": "human_review"},
        "original_intake_error": str(case.get("original_intake_error") or ""), "archives": archive_results}
    validate_json(result, RESULT_SCHEMA)
    return result


def _material_card_evidence(result: dict[str, Any], archive: dict[str, Any], issue: dict[str, Any]) -> dict[str, Any]:
    """Project exact diagnosed sources; never borrow unrelated business facts."""
    selected: list[tuple[dict[str, Any], dict[str, Any], bool]] = []
    requested = set(issue["source_files"])
    whole_archive = issue["code"] in {"scenario_unconfirmed", "business_checks_pending"}
    for owner in result["archives"]:
        for item in owner["inventory"]:
            qualified = f"{owner['source_archive']} / {item['source_file']}"
            local = owner["archive_id"] == archive["archive_id"]
            if qualified in requested or (local and (item["source_file"] in requested or whole_archive)):
                selected.append((owner, item, qualified in requested))
    sources = []
    for owner, item, qualified in selected:
        document = next((doc for doc in owner["documents"] if doc["file_id"] == item["file_id"]), {})
        parent = str(item.get("parent_source") or item["source_file"])
        prefix = f"{owner['source_archive']} / " if qualified else ""
        facts = _strings(document.get("visible_facts"))
        if document.get("title"):
            facts.insert(0, f"可见标题：{document['title']}")
        if document.get("document_number"):
            facts.append(f"可见文档编号：{document['document_number']}")
        limitations = _strings([*(item.get("limitations") or []), *(document.get("limitations") or [])])
        facts.extend(f"读取限制：{limitation}" for limitation in limitations)
        if item.get("spreadsheet_check"):
            facts.append(item["spreadsheet_check"]["basis"])
        if not facts:
            facts.append("已盘点到该提交文件；其可见内容或具体用途仍待确认。")
        locator = f"{owner['source_archive']} 内原始路径：{item['source_file']}"
        if document.get("page_number"):
            locator += f"；可见第 {document['page_number']} 页"
        role = ROLE_LABELS.get(document.get("document_type"), "电子表材料" if item["kind"] == "spreadsheet" else "已提交材料")
        sources.append({"file": prefix + item["source_file"], "original_file": prefix + parent,
                        "kind": "derived" if item.get("parent_source") else "submitted",
                        "role": role, "locator": locator, "facts": facts})
    if issue["title"] == "压缩包重复提交":
        sources = [{"file": name, "original_file": name, "kind": "submitted", "role": "本次提交的原始压缩包",
                    "locator": "本次资料包盘点", "facts": ["与本项列出的其他压缩包字节内容完全相同。"]} for name in issue["source_files"]]
    originals = _strings(source["original_file"] for source in sources)
    limitations = ["此为材料诊断来源，不代表已完成业务规则或金额核验。"]
    if not sources:
        limitations.insert(0, "本项针对未提交或未确认的材料角色；该缺口没有可列出的已确认来源文件，不能据此认为整个资料包未被分析。")
    return {"schema_version": "1.0", "source_files": originals, "source_file_count": len(originals),
            "sources": sources, "derived_file_count": sum(source["kind"] == "derived" for source in sources),
            "references": [], "comparisons": [issue["observed"], issue["expected"]],
            "limitations": limitations, "file_count_complete": True}


def build_material_diagnostic_view(result: dict[str, Any]) -> dict[str, Any]:
    validate_json(result, RESULT_SCHEMA)
    sheets = []
    for archive in result["archives"]:
        rows = []
        for number, issue in enumerate(archive["issues"], 4):
            files = "、".join(issue["source_files"])
            rows.append({"excel_row": number, "kind": "record", "section": "detail", "status": "issue",
                "confidence": issue["confidence"], "heading": f"问题：{issue['title']}",
                "card_evidence": _material_card_evidence(result, archive, issue),
                "values": [f"问题：{issue['title']}\n文件：{files or '本包未确认到该材料'}",
                    f"识别结果：{issue['observed']}", f"材料要求：{issue['expected']}", "审核结论：待人工确认",
                    f"核销影响：{issue['impact']}", f"处理方式：{issue['action']}"]})
        label = archive["scenario_label"]
        sheets.append({"name": f"{label}材料诊断", "scenario": archive["scenario"],
            "projection_kind": "material_diagnostic", "audit_type_label": label,
            "title": f"{label}材料诊断｜{archive['source_archive']}",
            "note": "材料分析已完成，业务核销待确认。", "source_archive": archive["source_archive"],
            "headers": ["问题来源", "已识别内容", "材料要求", "审核结论", "核销影响", "需要补交"],
            "rows": rows, "diagnostic_issues": deepcopy(archive["issues"]),
            "audit_counts": {"source_row_count": len(rows), "error_count": len(rows), "detail_error_count": len(rows), "context_error_count": 0}})
    return {"schema_version": "1.1", "title": "核销材料诊断结果", "sheets": sheets}


def build_classification_rejection_result(archives: list[dict[str, Any]]) -> dict[str, Any]:
    """Record only the ZIP-name rule; this result contains no AI observations."""
    rejected = []
    for source in archives:
        archive = deepcopy(source)
        name = archive["source_archive"]
        matches = archive["matched_scenarios"]
        if matches:
            labels = "、".join(SCENARIO_LABELS[value] for value in matches)
            archive["reason"] = f"{name}：文件名同时包含{labels}，无法确定核销方式。"
        else:
            archive["reason"] = f"{name}：文件名未注明核销方式。"
        archive["action"] = "打回申请；在 ZIP 文件名中注明唯一的核销方式后重新提交。"
        rejected.append(archive)
    result = {"schema_version": "1.0", "kind": "classification_rejection",
              "decision_source": "zip_filename", "generated_at": now_utc(),
              "summary": {"conclusion": "rejected"}, "archives": rejected}
    validate_json(result, CLASSIFICATION_REJECTION_SCHEMA)
    return result


def build_classification_rejection_view(result: dict[str, Any]) -> dict[str, Any]:
    validate_json(result, CLASSIFICATION_REJECTION_SCHEMA)
    sheets = []
    for archive in result["archives"]:
        name = archive["source_archive"]
        reason = archive["reason"]
        sheets.append({
            "name": "核销方式无法确认", "title": f"核销方式无法确认｜{name}",
            "scenario": None, "confirmed_scenario": False,
            "projection_kind": "classification_rejection", "business_decision": "rejected",
            "audit_type_label": "核销方式无法确认", "source_archive": name,
            "note": "已打回。", "headers": ["业务文件", "错误原因", "处理方式"],
            "rows": [{"excel_row": 4, "kind": "record", "section": "detail",
                "status": "issue", "confidence": "high", "heading": "核销方式无法确认",
                "error_reason": reason, "error_reasons": [reason],
                "values": [name, reason, f"处理方式：{archive['action']}"],
                "card_evidence": {"schema_version": "1.0", "source_files": [name],
                    "source_file_count": 1, "sources": [{"file": name, "original_file": name,
                        "kind": "submitted", "role": "原始资料包", "locator": "ZIP 文件名",
                        "facts": [reason]}], "derived_file_count": 0, "references": [],
                    "comparisons": [], "limitations": [], "file_count_complete": True}}],
            "audit_counts": {"source_row_count": 1, "error_count": 1,
                "detail_error_count": 1, "context_error_count": 0},
        })
    return {"schema_version": "1.1", "title": "核销申请打回结果", "sheets": sheets}
