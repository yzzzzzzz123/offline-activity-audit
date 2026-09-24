from __future__ import annotations

from .paths import PROJECT_ROOT

import hashlib
import json
import re
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from .common import AuditError


SHEET_SCENARIOS = {
    "人员激励核销": "personnel_incentive",
    "堆头核销": "promotional_display",
    "海报物料核销": "poster_material",
    "其他费用核销": "other_expense",
    "维护费用核销": "maintenance_fee",
    "额外搭赠核销": "giveaway_promotion",
    "价格补差核销": "price_difference_support",
    "POS达标激励核销": "pos_target_incentive",
    "条码费核销": "entry_fee",
    "自采赠品物料核销": "self_procured_gift_material",
}
SCENARIO_SHEETS = {value: key for key, value in SHEET_SCENARIOS.items()}
SHEET_SCENARIOS["进场费核销"] = "entry_fee"  # Preserve reads of existing archives.
SUMMARY_PREFIXES = (
    "合计",
    "总计",
    "实际申报",
    "实际申请",
    "收款人与日期",
    "收款人/日期",
)
OVERVIEW_PREFIXES = (
    "活动概况",
    "合同销售附件",
)
FORBIDDEN_VISIBLE_TERMS = (
    "详见审计JSON",
    "见内部结果",
    "候选命中",
    "商品ID",
    "视觉RAG",
    "唯一收敛",
)
FORBIDDEN_DENSITY_TOGGLE_MARKERS = (
    'id="densityButton"',
    "densityButton",
    "state.compact",
    "classList.toggle('compact'",
    ">紧凑显示<",
    ">舒展显示<",
)
FORBIDDEN_UI_MARKERS = (
    "页面可离线打开",
    "内容与同名 Excel 保持一致",
    "本页用于",
    "点击记录",
    "合同核心商品 → 知识库：不适用",
    "打开同版 Excel",
    'id="excelLink"',
    "setExcelLink",
    'id="fileName"',
    "复制本条结论",
    "data-copy-row",
    'data-status="summary"',
    "只看汇总",
    "navigator.clipboard",
    "execCommand('copy')",
    'id="toast"',
    'id="printButton"',
    "window.print()",
)
DATA_OPEN = '<script id="audit-data" type="application/json">'
DATA_CLOSE = "</script>"
TEMPLATE_ASSET_PATH = (
    PROJECT_ROOT
    / "skills"
    / "orchestrate-offline-audit"
    / "assets"
    / "legacy"
    / "canban-audit-shell.html"
)
ERROR_ONLY_STYLE_ASSET_PATH = TEMPLATE_ASSET_PATH.with_name("error-only.css")
ERROR_ONLY_SCRIPT_ASSET_PATH = TEMPLATE_ASSET_PATH.with_name("error-only.js")
AUDIT_DATA_MARKER = "__AUDIT_DATA__"
STYLE_SHA256_MARKER = "__CANBAN_STYLE_SHA256__"
SHELL_SHA256_MARKER = "__CANBAN_SHELL_SHA256__"
ERROR_ONLY_STYLE_MARKER = "__ERROR_ONLY_CSS__"
ERROR_ONLY_SCRIPT_MARKER = "__ERROR_ONLY_SCRIPT__"
STYLE_PATTERN = re.compile(
    r'<style id="canban-audit-style">(?P<style>[\s\S]*?)</style>'
)
ERROR_ONLY_STYLE_PATTERN = re.compile(
    r'<style id="error-only-preview-style">(?P<style>[\s\S]*?)</style>'
)
TEMPLATE_VERSION_PATTERN = re.compile(
    r'<meta name="canban-template-version" content="(?P<version>[^"]+)">'
)




def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).replace("\r\n", "\n").replace("\r", "\n").strip()


def _row_kind(values: list[str]) -> str:
    first = values[0].lstrip() if values else ""
    if first.startswith(SUMMARY_PREFIXES + OVERVIEW_PREFIXES):
        return "summary"
    return "record"


def _row_section(
    kind: str,
    scenario: str,
    values: list[str] | None = None,
) -> str:
    if scenario not in {
        "personnel_incentive",
        "promotional_display",
        "poster_material",
        "other_expense",
        "maintenance_fee",
        "giveaway_promotion",
        "price_difference_support",
        "pos_target_incentive",
        "entry_fee",
        "self_procured_gift_material",
    }:
        raise AuditError(f"HTML不支持的核销场景：{scenario}")
    first = values[0].lstrip() if values else ""
    if (
        kind == "summary"
        and scenario == "promotional_display"
        and first.startswith(OVERVIEW_PREFIXES)
    ):
        return "overview"
    return "settlement" if kind == "summary" else "detail"


def _row_status(values: list[str], kind: str) -> str:
    conclusion = values[5] if len(values) > 5 else ""
    full_text = "\n".join(values)
    if "待人工核定" in full_text:
        return "issue"
    no_resubmission = (
        "无需重新提交" in conclusion
        or "无需补交" in conclusion
        or "要重新提交：不用" in conclusion
        or "要重新提交什么：不用" in conclusion
    )
    if kind == "summary":
        if (
            any(token in conclusion for token in ("暂不能核销", "完全不匹配", "需补"))
            or ("要重新提交：" in conclusion and not no_resubmission)
        ):
            return "issue"
        return "summary"
    if no_resubmission and not any(
        token in conclusion for token in ("暂不能核销", "完全不匹配", "需补")
    ):
        return "pass"
    return "issue"


def _row_confidence(values: list[str]) -> str:
    text = "\n".join(values)
    if "置信度：低" in text:
        return "low"
    if "置信度：中" in text:
        return "medium"
    if "置信度：高" in text:
        return "high"
    return "not_applicable"


def _display_store_heading(values: list[str]) -> str | None:
    contract_text = values[0] if values else ""
    store_label = re.search(r"(?:^|\n)门店[：:]\s*([^\n；]+)", contract_text)
    if store_label:
        return store_label.group(1).strip()
    numbered_store = re.search(r"第\s*\d+\s*家[：:]\s*([^\n；]+)", contract_text)
    if numbered_store:
        return numbered_store.group(1).strip()
    current_store = re.search(
        r"本店(?!堆头)(?:[：:]\s*)?([^\n；]+)",
        contract_text,
    )
    if current_store:
        return current_store.group(1).strip()
    photo_text = values[1] if len(values) > 1 else ""
    visible_location = re.search(r"识别地点[：:]\s*([^\n]+)", photo_text)
    if visible_location:
        return visible_location.group(1).strip()
    return None


def _row_heading(
    values: list[str],
    row_number: int,
    kind: str,
    scenario: str,
) -> str:
    if kind == "summary":
        first = values[0] if values else ""
        for line in first.splitlines():
            label = line.strip()
            if label:
                return label[:90]
        return "汇总"
    if scenario == "promotional_display":
        store = _display_store_heading(values)
        if store:
            return store[:90]
    first = values[0] if values else ""
    if scenario == "personnel_incentive":
        title_lines = [line.strip() for line in first.splitlines() if line.strip()]
        if len(title_lines) >= 2 and title_lines[0].startswith("结算第"):
            return title_lines[1][:90]
    for line in first.splitlines():
        label = line.strip()
        if label:
            return label[:90]
    conclusion = values[5] if len(values) > 5 else ""
    for line in conclusion.splitlines():
        label = line.strip()
        if label:
            return label[:90]
    return f"核销记录 {row_number - 3}"


def _workbook_payload(workbook_path: str | Path) -> dict[str, Any]:
    source = Path(workbook_path)
    if not source.is_file():
        raise AuditError(f"找不到待生成HTML的工作簿：{source}")

    workbook = load_workbook(source, read_only=True, data_only=False)
    try:
        unknown = set(workbook.sheetnames) - set(SHEET_SCENARIOS)
        if unknown:
            raise AuditError("HTML不支持的工作表：" + "、".join(sorted(unknown)))
        if not workbook.sheetnames:
            raise AuditError("工作簿没有可展示的工作表")

        sheets: list[dict[str, Any]] = []
        for worksheet in workbook.worksheets:
            scenario = SHEET_SCENARIOS[worksheet.title]
            headers = [_cell_text(worksheet.cell(3, column).value) for column in range(1, 7)]
            if len(headers) != 6 or any(not header for header in headers):
                raise AuditError(f"{worksheet.title} 缺少完整六列表头，不能生成HTML")

            all_rows: list[dict[str, Any]] = []
            for row_number in range(4, worksheet.max_row + 1):
                values = [
                    _cell_text(worksheet.cell(row_number, column).value)
                    for column in range(1, 7)
                ]
                if not any(values):
                    continue
                kind = _row_kind(values)
                all_rows.append(
                    {
                        "excel_row": row_number,
                        "kind": kind,
                        "section": _row_section(kind, scenario, values),
                        "status": _row_status(values, kind),
                        "confidence": _row_confidence(values),
                        "heading": _row_heading(
                            values,
                            row_number,
                            kind,
                            scenario,
                        ),
                        "values": values,
                    }
                )

            issue_rows = [row for row in all_rows if row["status"] == "issue"]
            visible_rows = (
                issue_rows
                if scenario in {"poster_material", "other_expense", "maintenance_fee", "giveaway_promotion", "price_difference_support", "pos_target_incentive", "entry_fee", "self_procured_gift_material"}
                else all_rows
            )
            sheets.append(
                {
                    "name": worksheet.title,
                    "scenario": scenario,
                    "title": _cell_text(worksheet["A1"].value) or worksheet.title,
                    "note": _cell_text(worksheet["A2"].value),
                    "headers": headers,
                    "rows": visible_rows,
                    "audit_counts": {
                        "source_row_count": len(all_rows),
                        "error_count": len(issue_rows),
                        "detail_error_count": sum(
                            row["section"] == "detail" for row in issue_rows
                        ),
                        "context_error_count": sum(
                            row["section"] != "detail" for row in issue_rows
                        ),
                    },
                }
            )
    finally:
        workbook.close()

    return {
        "schema_version": "1.1",
        "title": "线下活动核销结果",
        "sheets": sheets,
    }


def _json_for_script(payload: dict[str, Any]) -> str:
    return (
        json.dumps(payload, ensure_ascii=False, separators=(",", ":"))
        .replace("<", "\\u003c")
        .replace(">", "\\u003e")
        .replace("&", "\\u0026")
        .replace("\u2028", "\\u2028")
        .replace("\u2029", "\\u2029")
    )


def _read_template_asset() -> str:
    if not TEMPLATE_ASSET_PATH.is_file():
        raise AuditError(f"找不到HTML模板资产：{TEMPLATE_ASSET_PATH}")
    template = TEMPLATE_ASSET_PATH.read_text(encoding="utf-8")
    component_paths = {
        ERROR_ONLY_STYLE_MARKER: ERROR_ONLY_STYLE_ASSET_PATH,
        ERROR_ONLY_SCRIPT_MARKER: ERROR_ONLY_SCRIPT_ASSET_PATH,
    }
    for marker, path in component_paths.items():
        if template.count(marker) != 1:
            raise AuditError(f"HTML模板组件占位符缺失或重复：{marker}")
        if not path.is_file():
            raise AuditError(f"找不到HTML模板组件：{path}")
        component = path.read_text(encoding="utf-8").rstrip("\n")
        if marker == ERROR_ONLY_STYLE_MARKER and "</style>" in component.lower():
            raise AuditError("错误清单样式组件不得提前结束style标签")
        if marker == ERROR_ONLY_SCRIPT_MARKER and "</script>" in component.lower():
            raise AuditError("错误清单脚本组件不得提前结束script标签")
        template = template.replace(marker, component)
    required_markers = (
        AUDIT_DATA_MARKER,
        STYLE_SHA256_MARKER,
        SHELL_SHA256_MARKER,
    )
    invalid = [marker for marker in required_markers if template.count(marker) != 1]
    if invalid:
        raise AuditError("HTML模板占位符缺失或重复：" + "、".join(invalid))
    if template.count('id="audit-data"') != 0:
        raise AuditError("HTML模板不得预置核销数据")
    if STYLE_PATTERN.search(template) is None:
        raise AuditError("HTML模板缺少唯一主样式块")
    if ERROR_ONLY_STYLE_PATTERN.search(template) is None:
        raise AuditError("HTML模板缺少错误清单样式块")
    if template.count('id="error-only-preview-script"') != 1:
        raise AuditError("HTML模板缺少唯一错误清单脚本")
    version_matches = list(TEMPLATE_VERSION_PATTERN.finditer(template))
    if len(version_matches) != 1:
        raise AuditError("HTML模板缺少唯一版本标记")
    return template


def _template_metadata(template: str) -> dict[str, str]:
    style_match = STYLE_PATTERN.search(template)
    error_style_match = ERROR_ONLY_STYLE_PATTERN.search(template)
    version_match = TEMPLATE_VERSION_PATTERN.search(template)
    if style_match is None or error_style_match is None or version_match is None:
        raise AuditError("HTML模板版本或样式标记无效")
    combined_style = (
        style_match.group("style")
        + "\n"
        + error_style_match.group("style")
    )
    return {
        "template_version": version_match.group("version"),
        "style_sha256": hashlib.sha256(
            combined_style.encode("utf-8")
        ).hexdigest(),
        "shell_sha256": hashlib.sha256(template.encode("utf-8")).hexdigest(),
    }


def _render_template(payload: dict[str, Any]) -> str:
    if not payload.get("sheets"):
        raise AuditError("HTML没有可展示的核销场景")
    template = _read_template_asset()
    metadata = _template_metadata(template)
    embedded = DATA_OPEN + _json_for_script(payload) + DATA_CLOSE
    return (
        template.replace(STYLE_SHA256_MARKER, metadata["style_sha256"])
        .replace(SHELL_SHA256_MARKER, metadata["shell_sha256"])
        .replace(AUDIT_DATA_MARKER, embedded)
    )


def create_html_report_from_workbook(
    workbook_path: str | Path,
    output_path: str | Path,
) -> Path:
    payload = _workbook_payload(workbook_path)
    html = _render_template(payload)
    target = Path(output_path)
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_text(html, encoding="utf-8", newline="\n")
    return target


def _embedded_payload(html: str) -> dict[str, Any]:
    start = html.find(DATA_OPEN)
    if start < 0:
        raise AuditError("HTML缺少内嵌核销数据")
    start += len(DATA_OPEN)
    end = html.find(DATA_CLOSE, start)
    if end < 0:
        raise AuditError("HTML内嵌核销数据没有正确结束")
    try:
        payload = json.loads(html[start:end])
    except json.JSONDecodeError as exc:
        raise AuditError("HTML内嵌核销数据不是有效JSON") from exc
    if not isinstance(payload, dict):
        raise AuditError("HTML内嵌核销数据结构错误")
    return payload


def verify_html_report(
    html_path: str | Path,
    scenarios: list[str],
    *,
    workbook_path: str | Path | None = None,
) -> dict[str, Any]:
    source = Path(html_path)
    if not source.is_file():
        raise AuditError(f"找不到HTML输出：{source}")
    html = source.read_text(encoding="utf-8")
    payload = _embedded_payload(html)
    expected_sheets = [SCENARIO_SHEETS[scenario] for scenario in scenarios]
    actual_sheets = [str(sheet.get("name") or "") for sheet in payload.get("sheets", [])]
    if actual_sheets != expected_sheets:
        raise AuditError(f"HTML场景名称或顺序错误：{actual_sheets}，期望{expected_sheets}")

    required_controls = (
        'id="error-only-preview-style"',
        'id="error-only-preview-script"',
        'id="eoBackTop"',
        'id="eoErrorList"',
        'id="eoErrorType"',
        'id="eoErrorCategory"',
        'id="eoErrorFilterEmpty"',
        'id="eoPassList"',
        'id="eoPassType"',
        'id="eoPassCategory"',
        'id="eoPassFacetSummary"',
        'id="eoPassFilterEmpty"',
        'class="eo-tab"',
        'class="eo-error-card"',
        'data-error-categories=',
        'data-error-confidence-score=',
        'data-pass-confidence-score=',
        'data-eo-view="passed"',
        "核销类型 ·",
        "筛选正确检查项",
    )
    missing_controls = [control for control in required_controls if control not in html]
    if missing_controls:
        raise AuditError("HTML缺少交互按钮：" + "、".join(missing_controls))
    forbidden_scenario_navigation = (
        'class="eo-scenario-card"',
        'class="eo-enter"',
        'data-open="scenario-',
        "进入错误清单",
        "错误场景队列",
    )
    invalid_navigation = [
        marker for marker in forbidden_scenario_navigation if marker in html
    ]
    if invalid_navigation:
        raise AuditError(
            "HTML不得按错误场景分类或要求再次进入："
            + "、".join(invalid_navigation)
        )
    if any(marker in html for marker in FORBIDDEN_DENSITY_TOGGLE_MARKERS):
        raise AuditError("HTML不得提供紧凑/宽松显示切换")
    forbidden_ui = [marker for marker in FORBIDDEN_UI_MARKERS if marker in html]
    if forbidden_ui:
        raise AuditError("HTML出现无效页面文案或功能：" + "、".join(forbidden_ui))
    if re.search(r"<(?:script|link)\b[^>]*(?:src|href)=[\"']https?://", html, re.I):
        raise AuditError("HTML不得依赖在线脚本或样式")
    for forbidden in FORBIDDEN_VISIBLE_TERMS:
        if forbidden in html:
            raise AuditError(f"HTML出现不应交给客户的工程表达：{forbidden}")

    record_counts = {
        str(sheet["name"]): len(sheet.get("rows") or [])
        for sheet in payload.get("sheets", [])
    }
    for sheet in payload.get("sheets", []):
        for row in sheet.get("rows") or []:
            if (
                sheet.get("scenario") in {"poster_material", "other_expense", "maintenance_fee", "giveaway_promotion", "price_difference_support", "pos_target_incentive", "entry_fee", "self_procured_gift_material"}
                and row.get("status") != "issue"
            ):
                raise AuditError(
                    f"{sheet.get('name')} 包含通过项；该场景HTML只能输出错误或人工处理项"
                )
            expected_section = _row_section(
                str(row.get("kind") or ""),
                str(sheet.get("scenario") or ""),
                [str(value or "") for value in row.get("values") or []],
            )
            if row.get("section") != expected_section:
                raise AuditError(
                    f"{sheet.get('name')} 第{row.get('excel_row')}行分组错误："
                    f"{row.get('section')}，期望{expected_section}"
                )
    if 'class="eo-table"' not in html:
        raise AuditError("HTML缺少合同商品错误明细表格组件")
    expected_html = _render_template(payload)
    if html != expected_html:
        raise AuditError("HTML静态模板与主Skill模板资产不一致")
    template_metadata = _template_metadata(_read_template_asset())
    if workbook_path is not None:
        workbook_payload = _workbook_payload(workbook_path)
        if payload != workbook_payload:
            raise AuditError("HTML内嵌核销内容与Excel不一致")

    return {
        "path": str(source.resolve()),
        "sheet_names": actual_sheets,
        "record_counts": record_counts,
        "button_count": len(required_controls),
        "external_dependency_count": 0,
        "workbook_content_equal": workbook_path is not None,
        "template_version": template_metadata["template_version"],
        "style_sha256": template_metadata["style_sha256"],
        "shell_sha256": template_metadata["shell_sha256"],
        "static_template_equal": True,
    }
