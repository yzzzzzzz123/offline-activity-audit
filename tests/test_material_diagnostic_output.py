from __future__ import annotations

from copy import deepcopy
import hashlib
import io
import json
from pathlib import Path
import tempfile
import unittest
import zipfile
from unittest import mock

from openpyxl import Workbook

from audit_core.common import AuditError, validate_json
from audit_core.archive_input import ArchiveInputError
from audit_core.material_diagnostic_output import (
    CLASSIFICATION_REJECTION_SCHEMA, RESULT_SCHEMA, SCENARIO_LABELS,
    build_classification_rejection_result, build_classification_rejection_view,
    build_material_diagnostic_result,
    build_material_diagnostic_view,
    _pos_spreadsheet_structure,
)
from audit_core.pass_check_log import attach_pass_check_log
from audit_core.workbench_store import WorkbenchRunStore, render_analysis_summary_markdown


def _fixture(scenario="maintenance_fee", roles=()):
    inventory = []
    documents = []
    for number, role in enumerate(roles, 1):
        file_id = f"f{number}"
        name = f"原始资料/文件{number}.jpg"
        item = {"file_id": file_id, "source_file": name, "path": f"/private/input/{name}",
                "kind": "visual", "suffix": ".jpg", "sha256": hashlib.sha256(file_id.encode()).hexdigest(),
                "parent_source": None, "limitations": []}
        if role == "pos_spreadsheet":
            item.update(kind="spreadsheet", suffix=".xlsx", source_file=f"销售{number}.xlsx")
            temporary = tempfile.TemporaryDirectory()
            unittest.addModuleCleanup(temporary.cleanup)
            path = Path(temporary.name) / item["source_file"]
            _sales_workbook(path, scenario)
            item["path"] = str(path)
            item["sha256"] = hashlib.sha256(path.read_bytes()).hexdigest()
        elif role == "activity_return_workbook":
            item.update(kind="spreadsheet", suffix=".xls", source_file=f"返图{number}.xls")
        else:
            documents.append({"file_id": file_id, "document_type": role, "confidence": "high", "title": "可见标题",
                              "document_number": None, "page_number": None, "total_pages": None,
                              "visible_facts": [f"独立观察{number}"], "limitations": []})
        inventory.append(item)
    case = {"kind": "material_diagnostic", "original_intake_error": "结算单应最多1份", "source_archives": ["ai-pack-9.zip"],
            "archives": [{"archive_id": "a1", "source_archive": "ai-pack-9.zip", "scenario_hint": scenario, "inventory": inventory}]}
    evidence = {"schema_version": "1.0", "archives": [{"archive_id": "a1", "scenario_candidates": [], "documents": documents}]}
    return case, evidence


def _sales_workbook(path, scenario):
    fixtures = {
        "personnel_incentive": (["日期", "门店名称", "条码", "名称", "数量"], ["2026-09", "门店一", "6900000000001", "商品甲", 2]),
        "promotional_display": (["客户名称", "业务日期", "商品编码", "商品名称", "条形码", "单位", "数量", "零售价", "合计金额"], ["客户一", "2026-09", "P001", "商品甲", "6900000000001", "件", 2, 10, 20]),
        "maintenance_fee": (["商品名称", "数量", "销售金额"], ["商品甲", 2, 20]),
        "price_difference_support": (["商品名称", "数量", "销售金额"], ["商品甲", 2, 20]),
        "pos_target_incentive": (["日期", "系统门店名称", "销售金额"], ["2026-09", "门店一", 20]),
        "self_procured_gift_material": (["日期", "门店名称", "销售数量", "销售金额"], ["2026-09", "门店一", 2, 20]),
    }
    headers, values = fixtures.get(scenario, fixtures["maintenance_fee"])
    workbook = Workbook()
    workbook.active.append(headers)
    workbook.active.append(values)
    workbook.save(path)
    workbook.close()


def _issues(result, code=None):
    issues = result["archives"][0]["issues"]
    return [issue for issue in issues if code is None or issue["code"] == code]


def _classification_rejections():
    return [
        {"archive_id": "a1", "source_archive": "ai-pack-8.zip", "archive_sha256": "a" * 64,
         "scenario": None, "reason_code": "missing_scenario_marker", "matched_scenarios": []},
        {"archive_id": "a2", "source_archive": "维护费用-进场费.zip", "archive_sha256": "b" * 64,
         "scenario": None, "reason_code": "ambiguous_scenario_markers", "matched_scenarios": ["maintenance_fee", "entry_fee"]},
    ]


class MaterialDiagnosticOutputTests(unittest.TestCase):
    def test_classification_rejection_has_name_reasons_and_no_ai_or_material_findings(self):
        archives = _classification_rejections()
        original = deepcopy(archives)
        result = build_classification_rejection_result(archives)
        self.assertEqual(archives, original)
        self.assertEqual(result["decision_source"], "zip_filename")
        self.assertEqual(result["summary"], {"conclusion": "rejected"})
        unknown, ambiguous = result["archives"]
        self.assertIn("未注明核销方式", unknown["reason"])
        self.assertIn("同时包含维护费用、进场费", ambiguous["reason"])
        for archive in result["archives"]:
            self.assertIn(archive["source_archive"], archive["reason"])
            self.assertNotIn("重新提交", archive["reason"])
            self.assertIn("重新提交", archive["action"])
            for fabricated in ("documents", "inventory", "scenario_candidates", "issues", "approved_amount"):
                self.assertNotIn(fabricated, archive)
        view = attach_pass_check_log(build_classification_rejection_view(result), {})
        self.assertEqual(view["pass_check_log"]["total"], 0)
        for sheet, archive in zip(view["sheets"], result["archives"]):
            self.assertIsNone(sheet["scenario"])
            self.assertFalse(sheet["confirmed_scenario"])
            self.assertEqual(sheet["business_decision"], "rejected")
            self.assertEqual(sheet["projection_kind"], "classification_rejection")
            self.assertEqual(sheet["rows"][0]["card_evidence"]["source_files"], [archive["source_archive"]])
        summary = render_analysis_summary_markdown(view, {})
        self.assertEqual(summary, "核销方式无法确认")
        for fabricated in ("AI", "已签促销合同", "POS", "建议", "重新提交", "打回申请"):
            self.assertNotIn(fabricated, summary)

    def test_classification_rejection_schema_forbids_ai_evidence_approval_and_material_checklists(self):
        original = build_classification_rejection_result(_classification_rejections())
        for invalid in ("ai_source", "approval", "amount", "known_scenario", "documents", "material_findings", "single_marker_ambiguity"):
            with self.subTest(invalid=invalid):
                result = deepcopy(original)
                if invalid == "ai_source":
                    result["decision_source"] = "ai"
                elif invalid == "approval":
                    result["summary"]["conclusion"] = "pass"
                elif invalid == "amount":
                    result["summary"]["suggested_approved_amount"] = 0
                elif invalid == "known_scenario":
                    result["archives"][0]["scenario"] = "maintenance_fee"
                elif invalid == "documents":
                    result["archives"][0]["documents"] = []
                elif invalid == "material_findings":
                    result["archives"][0]["issues"] = [{"code": "missing_material"}]
                else:
                    result["archives"][1]["matched_scenarios"] = ["maintenance_fee"]
                with self.assertRaises(AuditError):
                    validate_json(result, CLASSIFICATION_REJECTION_SCHEMA)

    def test_diagnostic_cards_use_saved_visible_facts_and_exact_source_scope(self):
        result = build_material_diagnostic_result(*_fixture(roles=("settlement", "settlement")))
        view = build_material_diagnostic_view(result)
        ambiguity = next(row for row in view["sheets"][0]["rows"] if "主件待确认" in row["heading"])
        detail = ambiguity["card_evidence"]
        self.assertTrue(detail["file_count_complete"])
        self.assertEqual(detail["source_file_count"], 2)
        self.assertEqual([source["role"] for source in detail["sources"]], ["结算单", "结算单"])
        self.assertIn("独立观察1", detail["sources"][0]["facts"])
        missing = next(row for row in view["sheets"][0]["rows"] if "已签促销合同未提交" in row["heading"])
        self.assertEqual(missing["card_evidence"]["sources"], [])
        self.assertTrue(missing["card_evidence"]["file_count_complete"])
        self.assertNotIn("旧记录", json.dumps(view, ensure_ascii=False))

    def test_unconfirmed_scenario_card_lists_actual_files_used_for_classification(self):
        view = build_material_diagnostic_view(build_material_diagnostic_result(*_fixture(None, ("activity_photo",))))
        detail = view["sheets"][0]["rows"][0]["card_evidence"]
        self.assertEqual(detail["source_file_count"], 1)
        self.assertEqual(detail["sources"][0]["file"], "原始资料/文件1.jpg")
        self.assertIn("独立观察1", detail["sources"][0]["facts"])

    def test_same_archive_bytes_are_reported_with_both_original_zip_names(self):
        case, evidence = _fixture(roles=("settlement",))
        second = deepcopy(case["archives"][0])
        second.update(archive_id="a2", source_archive="另一个原始资料包.zip", archive_sha256="b" * 64)
        case["archives"][0]["archive_sha256"] = "b" * 64
        case["archives"].append(second)
        observations = deepcopy(evidence["archives"][0])
        observations["archive_id"] = "a2"
        evidence["archives"].append(observations)
        result = build_material_diagnostic_result(case, evidence)
        for archive in result["archives"]:
            issue = next(issue for issue in archive["issues"] if issue["title"] == "压缩包重复提交")
            self.assertEqual(issue["code"], "duplicate_submission")
            self.assertEqual(issue["source_files"], ["ai-pack-9.zip", "另一个原始资料包.zip"])

    def test_same_file_across_different_archives_keeps_each_archives_role_presence(self):
        case, evidence = _fixture(roles=("settlement",))
        second = deepcopy(case["archives"][0])
        second.update(archive_id="a2", source_archive="第二个资料包.zip", archive_sha256="c" * 64)
        second["inventory"][0]["source_file"] = "另一目录/独立命名的结算单.jpg"
        case["archives"][0]["archive_sha256"] = "b" * 64
        case["archives"].append(second)
        observations = deepcopy(evidence["archives"][0])
        observations["archive_id"] = "a2"
        evidence["archives"].append(observations)
        result = build_material_diagnostic_result(case, evidence)
        for archive in result["archives"]:
            self.assertFalse(any(issue["title"] == "压缩包重复提交" for issue in archive["issues"]))
            issue = next(issue for issue in archive["issues"] if issue["title"] == "跨资料包文件内容重复")
            self.assertIn("ai-pack-9.zip / 原始资料/文件1.jpg", issue["source_files"])
            self.assertIn("第二个资料包.zip / 另一目录/独立命名的结算单.jpg", issue["source_files"])
            self.assertFalse(any(issue["title"] == "结算单未提交" for issue in archive["issues"]))

    def test_spreadsheet_zip_limits_fail_before_workbook_parser(self):
        case, _ = _fixture(roles=("pos_spreadsheet",))
        source = case["archives"][0]["inventory"][0]
        for limit in ("MAX_ARCHIVE_FILES", "MAX_MEMBER_BYTES", "MAX_TOTAL_BYTES"):
            with self.subTest(limit=limit), mock.patch("audit_core.archive_input." + limit, 1), mock.patch("openpyxl.load_workbook") as parser:
                with self.assertRaises(ArchiveInputError):
                    _pos_spreadsheet_structure(source, "maintenance_fee")
                parser.assert_not_called()
        with zipfile.ZipFile(source["path"], "w", compression=zipfile.ZIP_DEFLATED) as archive:
            archive.writestr("xl/sharedStrings.xml", b"0" * (11 * 1024 * 1024))
        with mock.patch("openpyxl.load_workbook") as parser:
            with self.assertRaisesRegex(ArchiveInputError, "压缩比"):
                _pos_spreadsheet_structure(source, "maintenance_fee")
            parser.assert_not_called()

    def test_more_than_sixty_four_sheets_cannot_hide_an_additional_main_table(self):
        case, evidence = _fixture(roles=("pos_spreadsheet",))
        path = Path(case["archives"][0]["inventory"][0]["path"])
        workbook = Workbook()
        workbook.active.append(["商品名称", "数量", "销售金额"])
        workbook.active.append(["商品甲", 2, 20])
        for index in range(64):
            worksheet = workbook.create_sheet(f"附表{index}")
            if index == 63:
                worksheet.append(["商品名称", "数量", "销售金额"])
                worksheet.append(["商品乙", 3, 30])
        workbook.save(path)
        workbook.close()
        result = build_material_diagnostic_result(case, evidence)
        issue = next(issue for issue in _issues(result) if "POS 销售电子表" in issue["title"])
        self.assertEqual(issue["code"], "required_material_unconfirmed")
        self.assertIn("64 份检查范围", issue["observed"])

    def test_six_sales_formats_are_identified_without_retaining_cells_or_amounts(self):
        for scenario in ("personnel_incentive", "promotional_display", "maintenance_fee", "price_difference_support", "pos_target_incentive", "self_procured_gift_material"):
            with self.subTest(scenario=scenario):
                result = build_material_diagnostic_result(*_fixture(scenario, ("pos_spreadsheet",)))
                item = result["archives"][0]["inventory"][0]
                self.assertEqual(item["spreadsheet_check"]["status"], "structure_recognized")
                text = json.dumps(result, ensure_ascii=False)
                self.assertNotIn("6900000000001", text)
                self.assertNotIn("商品甲", text)
                self.assertFalse(any("POS 销售电子表" in issue["title"] for issue in _issues(result)))

    def test_unrelated_empty_corrupt_or_ambiguous_excel_is_unconfirmed_not_missing(self):
        for kind in ("unrelated", "empty", "corrupt", "corrupt_xml", "multiple"):
            with self.subTest(kind=kind):
                case, evidence = _fixture(roles=("pos_spreadsheet",))
                path = Path(case["archives"][0]["inventory"][0]["path"])
                if kind == "corrupt":
                    path.write_bytes(b"not an Excel workbook")
                elif kind == "corrupt_xml":
                    original = path.read_bytes()
                    with zipfile.ZipFile(io.BytesIO(original)) as source, zipfile.ZipFile(path, "w") as target:
                        for member in source.infolist():
                            target.writestr(member, b"<invalid" if member.filename == "xl/workbook.xml" else source.read(member))
                else:
                    workbook = Workbook()
                    if kind == "unrelated":
                        workbook.active.append(["姓名", "电话"])
                        workbook.active.append(["联系人", "12345678"])
                    elif kind == "empty":
                        workbook.active.append(["商品名称", "数量", "销售金额"])
                    else:
                        for sheet in (workbook.active, workbook.create_sheet("第二份")):
                            sheet.append(["商品名称", "数量", "销售金额"])
                            sheet.append(["商品甲", 2, 20])
                    workbook.save(path)
                    workbook.close()
                result = build_material_diagnostic_result(case, evidence)
                issue = next(issue for issue in _issues(result) if "POS 销售电子表" in issue["title"])
                self.assertEqual(issue["code"], "required_material_unconfirmed")
                self.assertNotIn("未提交", issue["title"])
                self.assertIn("销售1.xlsx", issue["source_files"])

    def test_spreadsheet_permission_and_missing_source_are_execution_failures(self):
        case, _ = _fixture(roles=("pos_spreadsheet",))
        source = case["archives"][0]["inventory"][0]
        with mock.patch("openpyxl.load_workbook", side_effect=PermissionError("blocked")):
            with self.assertRaises(PermissionError):
                _pos_spreadsheet_structure(source, "maintenance_fee")
        source["path"] = str(Path(source["path"]).with_name("missing.xlsx"))
        with self.assertRaises(FileNotFoundError):
            _pos_spreadsheet_structure(source, "maintenance_fee")

    def test_ordinary_promotion_contract_does_not_establish_entry_contract(self):
        result = build_material_diagnostic_result(*_fixture("entry_fee", ("signed_promotional_contract", "activity_photo", "system_deduction_proof")))
        issue = next(issue for issue in _issues(result) if issue["title"].startswith("已签进场合同"))
        self.assertEqual(issue["code"], "required_material_unconfirmed")
        self.assertIn("原始资料/文件1.jpg", issue["source_files"])
        self.assertIn("不能直接替代", issue["observed"])

    def test_unreadable_activity_workbook_does_not_prove_photos_were_not_submitted(self):
        case, evidence = _fixture("self_procured_gift_material", ("settlement", "activity_return_workbook"))
        case["archives"][0]["inventory"][1]["limitations"] = ["活动工作簿的图片绑定无法完整读取"]
        result = build_material_diagnostic_result(case, evidence)
        issue = next(issue for issue in _issues(result) if issue["title"].startswith("活动现场照片"))
        self.assertEqual(issue["code"], "required_material_unconfirmed")
        self.assertNotIn("未提交", issue["title"])
        self.assertIn("返图2.xls", issue["source_files"])
        self.assertTrue(any("图片绑定无法完整读取" in issue["observed"] for issue in _issues(result)))

    def test_missing_materials_have_specific_names_and_no_business_approval(self):
        case, evidence = _fixture(roles=("settlement",))
        result = build_material_diagnostic_result(case, evidence)
        self.assertEqual(result["summary"], {"conclusion": "human_review"})
        self.assertTrue(result["diagnostic_only"])
        missing = " ".join(issue["title"] for issue in _issues(result, "missing_material"))
        self.assertIn("已签促销合同未提交", missing)
        self.assertIn("POS 销售电子表", missing)
        self.assertNotIn("claimed_amount", json.dumps(result))
        self.assertNotIn("/private/input", json.dumps(result))
        self.assertEqual(result["archives"][0]["documents"][0]["visible_facts"], ["独立观察1"])

    def test_all_ten_registered_scenarios_can_report_empty_submission(self):
        for scenario in SCENARIO_LABELS:
            with self.subTest(scenario=scenario):
                result = build_material_diagnostic_result(*_fixture(scenario))
                self.assertEqual(result["archives"][0]["scenario"], scenario)
                self.assertTrue(_issues(result, "missing_material"))
                self.assertEqual(result["summary"]["conclusion"], "human_review")

    def test_two_different_settlements_are_ambiguity_not_duplicate(self):
        result = build_material_diagnostic_result(*_fixture(roles=("settlement", "settlement")))
        self.assertFalse(_issues(result, "duplicate_submission"))
        issue = _issues(result, "singleton_role_ambiguous")[0]
        self.assertEqual(len(issue["source_files"]), 2)
        self.assertIn("不能据此认定实际重复", issue["observed"])

    def test_exact_duplicate_retains_both_names_and_counts_role_once(self):
        case, evidence = _fixture(roles=("settlement", "settlement"))
        case["archives"][0]["inventory"][1]["sha256"] = case["archives"][0]["inventory"][0]["sha256"]
        result = build_material_diagnostic_result(case, evidence)
        self.assertEqual(len(_issues(result, "duplicate_submission")[0]["source_files"]), 2)
        self.assertFalse(_issues(result, "singleton_role_ambiguous"))

    def test_numbered_different_pages_are_one_document(self):
        case, evidence = _fixture(roles=("settlement", "settlement"))
        for number, doc in enumerate(evidence["archives"][0]["documents"], 1):
            doc.update(document_number="JS20260910", page_number=number, total_pages=2)
        result = build_material_diagnostic_result(case, evidence)
        self.assertFalse(_issues(result, "singleton_role_ambiguous"))
        evidence["archives"][0]["documents"][1]["page_number"] = 1
        result = build_material_diagnostic_result(case, evidence)
        self.assertTrue(_issues(result, "singleton_role_ambiguous"))

    def test_service_support_and_multiple_stamped_pos_are_not_conflicts(self):
        result = build_material_diagnostic_result(*_fixture(roles=(
            "settlement", "supporting_document", "stamped_pos_data", "stamped_pos_data",
            "signed_promotional_contract", "pos_spreadsheet")))
        self.assertFalse(_issues(result, "singleton_role_ambiguous"))
        self.assertFalse(_issues(result, "missing_material"))
        self.assertEqual(_issues(result)[0]["code"], "business_checks_pending")

    def test_maintenance_activity_photo_satisfies_special_support_presence(self):
        result = build_material_diagnostic_result(*_fixture(roles=(
            "settlement", "activity_photo", "stamped_pos_data", "signed_promotional_contract", "pos_spreadsheet")))
        self.assertFalse(_issues(result, "missing_material"))

    def test_self_procured_requires_receipt_and_pos_but_not_legacy_container(self):
        result = build_material_diagnostic_result(*_fixture("self_procured_gift_material", (
            "settlement", "payment_record", "activity_photo", "signed_promotional_contract", "activity_return_workbook")))
        missing = " ".join(issue["title"] for issue in _issues(result, "missing_material"))
        self.assertIn("发票或收据未提交", missing)
        self.assertIn("POS 销售电子表", missing)
        self.assertIn("盖章 POS", missing)
        self.assertNotIn("活动现场照片", missing)
        self.assertNotIn("活动返图工作簿", missing)

    def test_unreadable_document_does_not_prove_contract_was_not_submitted(self):
        case, evidence = _fixture(roles=("unknown",))
        evidence["archives"][0]["documents"][0].update(confidence="low", limitations=["照片严重模糊，无法辨认"])
        result = build_material_diagnostic_result(case, evidence)
        contract = next(issue for issue in _issues(result) if issue["title"].startswith("已签促销合同"))
        self.assertEqual(contract["code"], "required_material_unconfirmed")
        self.assertNotIn("未提交", contract["title"])
        self.assertIn("原始资料/文件1.jpg", contract["observed"])

    def test_unknown_scenario_is_never_presented_as_other_expense(self):
        result = build_material_diagnostic_result(*_fixture(None, ("activity_photo",)))
        self.assertIsNone(result["archives"][0]["scenario"])
        self.assertFalse(_issues(result, "missing_material"))
        view = build_material_diagnostic_view(result)
        self.assertEqual(view["sheets"][0]["audit_type_label"], "核销类型待确认")
        self.assertNotIn("其他费用", render_analysis_summary_markdown(view, {}))

    def test_unknown_zip_name_reports_name_problem_and_keeps_ai_candidates_technical(self):
        case, evidence = _fixture(None, tuple("activity_photo" for _ in range(10)))
        candidates = [
            {"scenario": "price_difference_support", "confidence": "high", "basis": "f1明确约定每支补差；f10记载相同机制。"},
            {"scenario": "personnel_incentive", "confidence": "medium", "basis": "f2列示人员激励费用。"},
            {"scenario": "poster_material", "confidence": "low", "basis": "f3仅见货架价格牌。"},
        ]
        evidence["archives"][0]["scenario_candidates"] = candidates
        result = build_material_diagnostic_result(case, evidence)
        observed = _issues(result, "scenario_unconfirmed")[0]["observed"]
        self.assertIsNone(result["archives"][0]["scenario"])
        self.assertFalse(_issues(result, "missing_material"))
        self.assertIn("ai-pack-9.zip", observed)
        self.assertIn("ZIP 名称缺少唯一分类标记", observed)
        self.assertNotIn("f1", observed)
        self.assertNotIn("价格补差", observed)
        self.assertNotIn("人员激励", observed)
        self.assertNotIn("海报/展示道具", observed)
        self.assertNotIn("仅见货架价格牌", observed)
        self.assertEqual(result["archives"][0]["scenario_candidates"], candidates)
        summary = render_analysis_summary_markdown(build_material_diagnostic_view(result), {})
        self.assertNotIn("f10", summary)
        self.assertNotIn("海报/展示道具", summary)
        self.assertNotIn("价格补差", summary)
        self.assertNotIn("人员激励", summary)

    def test_business_verification_limits_do_not_request_clearer_photos(self):
        limits = [
            "水印时间和地点未经独立核验；照片仅呈现局部陈列，无法确认活动期间、双方身份及费用约定。",
            "未见双方身份、活动期间或费用约定，无法确认陈列是否满足合同要求。",
            "照片为局部近景，无法判断整体陈列范围；水印未经独立核验，无法确认活动期间和费用用途。",
            "未显示销售期间、门店范围及POS系统来源，无法仅凭该页确认具体活动归属。",
            "未见明确单据编号、页码及总页数；费用约定不清晰，未能确认费用承担方。",
        ]
        case, evidence = _fixture(roles=("activity_photo",))
        evidence["archives"][0]["documents"][0]["limitations"] = limits
        result = build_material_diagnostic_result(case, evidence)
        self.assertFalse(_issues(result, "material_unreadable"))
        self.assertEqual(result["archives"][0]["documents"][0]["limitations"], limits)
        contract = next(issue for issue in _issues(result) if issue["title"].startswith("已签促销合同"))
        self.assertEqual(contract["code"], "missing_material")

    def test_explicit_reading_limits_still_report_only_the_reading_problem(self):
        for limitation in ("部分文字与印章重叠", "商品规格、部分门店全名和印章主体文字不够清晰",
                           "扫描件缺页", "文件损坏，无法打开", "照片严重模糊，无法辨认"):
            with self.subTest(limitation=limitation):
                case, evidence = _fixture(roles=("stamped_pos_data",))
                evidence["archives"][0]["documents"][0]["limitations"] = [
                    "未显示销售期间，无法确认具体活动归属；" + limitation + "。"]
                result = build_material_diagnostic_result(case, evidence)
                unreadable = _issues(result, "material_unreadable")
                self.assertEqual(len(unreadable), 1)
                self.assertIn(limitation, unreadable[0]["observed"])
                self.assertIn("原始资料/文件1.jpg", unreadable[0]["observed"])
                self.assertNotIn("无法确认具体活动归属", unreadable[0]["observed"])

    def test_each_zip_name_hint_locks_scenario_and_checklist_despite_ai_candidates(self):
        for scenario in SCENARIO_LABELS:
            with self.subTest(scenario=scenario):
                case, evidence = _fixture(scenario, ("settlement",))
                baseline = build_material_diagnostic_result(case, evidence)
                candidates = [{"scenario": other, "confidence": "high", "basis": "其他费用用途的可见描述"}
                              for other in SCENARIO_LABELS if other != scenario]
                evidence["archives"][0]["scenario_candidates"] = candidates
                result = build_material_diagnostic_result(case, evidence)
                self.assertEqual(result["archives"][0]["scenario"], scenario)
                self.assertFalse(_issues(result, "scenario_unconfirmed"))
                self.assertEqual(_issues(result), _issues(baseline))
                self.assertEqual(result["archives"][0]["scenario_candidates"], candidates)
                view = build_material_diagnostic_view(result)
                self.assertEqual(view["sheets"][0]["audit_type_label"], SCENARIO_LABELS[scenario])
                self.assertNotIn("核销类型待确认", render_analysis_summary_markdown(view, {}))

    def test_one_high_ai_candidate_cannot_classify_an_unmarked_or_ambiguous_zip(self):
        for archive_name in ("ai-pack-9.zip", "维护费用-进场费.zip"):
            with self.subTest(archive_name=archive_name):
                case, evidence = _fixture(None, ("activity_photo",))
                case["archives"][0]["source_archive"] = archive_name
                evidence["archives"][0]["scenario_candidates"] = [{"scenario": "entry_fee", "confidence": "high", "basis": "可见进场条款"}]
                result = build_material_diagnostic_result(case, evidence)
                self.assertIsNone(result["archives"][0]["scenario"])
                self.assertFalse(_issues(result, "missing_material"))
                self.assertEqual(len(_issues(result, "scenario_unconfirmed")), 1)
                self.assertIn("ZIP 名称缺少唯一分类标记", _issues(result)[0]["observed"])
                self.assertIn(archive_name, _issues(result)[0]["observed"])
                self.assertEqual(len(result["archives"][0]["documents"]), 1)

    def test_nanxiong_maintenance_hint_keeps_missing_pos_and_ambiguous_settlements(self):
        case, evidence = _fixture("maintenance_fee", ("signed_promotional_contract", "settlement", "settlement", "stamped_pos_data", "activity_photo"))
        case["archives"][0]["source_archive"] = "HX202601160053-广州南雄维护费用申请-9.zip"
        evidence["archives"][0]["scenario_candidates"] = [
            {"scenario": "price_difference_support", "confidence": "high", "basis": "f1明确约定按支补差"},
            {"scenario": "personnel_incentive", "confidence": "high", "basis": "f2列示人员激励"},
            {"scenario": "promotional_display", "confidence": "medium", "basis": "f5可见促销陈列"},
        ]
        result = build_material_diagnostic_result(case, evidence)
        self.assertEqual(result["archives"][0]["scenario"], "maintenance_fee")
        self.assertFalse(_issues(result, "scenario_unconfirmed"))
        self.assertEqual([issue["title"] for issue in _issues(result, "missing_material")], ["POS 销售电子表（.xlsx 或 .xlsm）未提交"])
        self.assertEqual(len(_issues(result, "singleton_role_ambiguous")), 1)
        summary = render_analysis_summary_markdown(build_material_diagnostic_view(result), {})
        self.assertIn("维护费用", summary)
        self.assertNotIn("核销类型待确认", summary)
        self.assertNotIn("价格补差", summary)
        self.assertNotIn("人员激励", summary)

    def test_visual_coverage_rejects_omitted_invented_or_spreadsheet_documents(self):
        for mode in ("omitted", "invented", "spreadsheet"):
            with self.subTest(mode=mode):
                case, evidence = _fixture(roles=("settlement", "pos_spreadsheet"))
                if mode == "omitted":
                    evidence["archives"][0]["documents"] = []
                else:
                    evidence["archives"][0]["documents"][0]["file_id"] = "f2" if mode == "spreadsheet" else "invented"
                with self.assertRaises(AuditError):
                    build_material_diagnostic_result(case, evidence)

    def test_long_duplicate_list_is_complete_in_summary_and_no_pass_checks(self):
        case, evidence = _fixture(roles=tuple("activity_photo" for _ in range(35)))
        for item in case["archives"][0]["inventory"]:
            item["sha256"] = "a" * 64
            item["source_file"] = "客户提交的很长的原始活动照片名称/" + item["source_file"]
        result = build_material_diagnostic_result(case, evidence)
        view = attach_pass_check_log(build_material_diagnostic_view(result), {})
        summary = render_analysis_summary_markdown(view, {})
        for item in case["archives"][0]["inventory"]:
            self.assertIn(item["source_file"].rsplit("/", 1)[-1], summary)
        self.assertNotIn("原始资料/", summary)
        self.assertNotIn("具体明细见完整", summary)
        self.assertEqual(view["pass_check_log"]["total"], 0)
        self.assertTrue(all(row["status"] == "issue" for row in view["sheets"][0]["rows"]))

    def test_schema_rejects_false_approval_amounts_and_empty_issue_result(self):
        original = build_material_diagnostic_result(*_fixture())
        for change in ("approval", "amount", "empty"):
            result = deepcopy(original)
            if change == "approval":
                result["summary"]["conclusion"] = "pass"
            elif change == "amount":
                result["summary"]["suggested_approved_amount"] = 0
            else:
                result["archives"][0]["issues"] = []
            with self.assertRaises(AuditError):
                validate_json(result, RESULT_SCHEMA)

    def test_store_persists_diagnostic_evidence_result_summary_and_stage_order(self):
        case, evidence = _fixture(roles=("settlement", "settlement"))
        result = build_material_diagnostic_result(case, evidence)
        with tempfile.TemporaryDirectory() as temporary:
            workspace = Path(temporary) / "diagnostic"
            workspace.mkdir()
            store = WorkbenchRunStore(workspace, run_id="20260910-diagnostic", producer_model="codex",
                root_html=Path(__file__).resolve().parents[1] / 'skills/orchestrate-offline-audit/assets/offline-activity-audit.html', workbench_url="http://127.0.0.1:8080/",
                source_archives=["ai-pack-9.zip"])
            store.observe("cases.prepared", {"cases": {}, "scenarios": []})
            store.observe("material_diagnostic.started", {"case": case})
            store.observe("material_diagnostic.evidence_validated", {"evidence": evidence})
            store.observe("material_diagnostic.result_validated", {"result": result})
            store.complete(view_payload=build_material_diagnostic_view(result), scenarios=["maintenance_fee"], verification={})
            self.assertEqual(store.manifest["status"], "completed")
            self.assertEqual(store.manifest["scenarios"], ["maintenance_fee"])
            self.assertEqual(store.manifest["source_archives"], ["ai-pack-9.zip"])
            self.assertTrue(store.manifest["analysis_started_at"])
            self.assertTrue((workspace / "analysis/material-diagnostic/evidence.json").is_file())
            self.assertTrue((workspace / "analysis/material-diagnostic/result.json").is_file())
            self.assertIn("主件与补充件关系未明确", store.analysis_summary_path.read_text(encoding="utf-8"))
            digest = hashlib.sha256(store.analysis_summary_path.read_bytes()).hexdigest()
            self.assertEqual(store.manifest["analysis_summary"]["sha256"], digest)


if __name__ == "__main__":
    unittest.main()
