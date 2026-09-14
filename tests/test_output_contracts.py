from __future__ import annotations

from copy import deepcopy
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from audit_core.common import AuditError
from audit_core.common import validate_json
from audit_core.display import (
    _contract_attachment_sales_reconciliation,
    _contract_sales_reconciliation,
    _display_control,
)
from audit_core.html_report import (
    TEMPLATE_ASSET_PATH,
    _row_heading,
    _row_kind,
    _row_section,
    _row_status,
    _render_template,
    _workbook_payload,
    create_html_report_from_workbook,
    verify_html_report,
)
from audit_core.orchestrator import (
    _create_temporary_root,
    _publish_pair_without_overwrite,
    next_output_path,
    normalize_producer_model,
    output_date_from_run_id,
)
from audit_core.report import (
    _contract_requirement_lines,
    create_combined_report,
    verify_workbook,
)


def _personnel_result() -> dict:
    return {
        "scenario": "personnel_incentive",
        "summary": {
            "calculated_line_reward": 30,
            "settlement_line_reward": 30,
            "transfer_total": 30,
            "excel_mapped_quantity": 10,
            "settlement_line_quantity": 10,
            "claimed_amount": 30,
            "store_count": 1,
            "sales_knowledge_matched_count": 0,
            "sales_knowledge_fuzzy_count": 1,
            "sales_knowledge_problem_count": 0,
        },
        "sales": {
            "source_file": "人员销售.xlsx",
            "knowledge_status": "pass",
            "knowledge_matched_count": 0,
            "knowledge_fuzzy_count": 1,
            "knowledge_problem_count": 0,
        },
        "settlement": {
            "source_file": "人员结算单.jpg",
            "activity_start": "2026-08-01",
            "activity_end": "2026-08-31",
        },
        "transfer_evidence": [
            {
                "source_files": ["转账截图.jpg"],
                "identity_visible": False,
                "recipient_name": None,
                "store_name": None,
                "event_date_visible": False,
                "transfer_date": None,
            }
        ],
        "sku_reconciliation": [
            {
                "line_no": 1,
                "settlement_product_name": "示例商品",
                "mapped_barcode": "6970356166832",
                "excel_product_name": "参半示例商品100g",
                "excel_quantity": 10,
                "calculated_settlement_reward": 30,
                "settlement_quantity": 10,
                "settlement_reward_amount": 30,
                "quantity_difference": 0,
                "settlement_line_amount_difference": 0,
                "mapping_status": "matched",
                "mapping_confidence": "medium",
                "knowledge_status": "fuzzy_matched",
                "knowledge_product_id": "canban-example",
                "knowledge_product_code": "CP-KQ-YG-0001",
                "knowledge_product_name": "参半示例商品100g标准名称",
                "knowledge_barcode_69": "6970356166832",
                "knowledge_barcode_match": "exact",
                "knowledge_name_match_type": "fuzzy",
                "knowledge_name_similarity": 0.875,
                "knowledge_basis": "69码精确匹配，产品名模糊匹配。",
                "knowledge_resubmission": None,
            }
        ],
    }


def _display_result() -> dict:
    result = {
        "schema_version": "2.1",
        "generated_at": "2026-08-25T00:00:00+00:00",
        "scenario": "promotional_display",
        "case_id": "display-example",
        "case_name": "示例堆头核销",
        "summary": {
            "conclusion": "pass",
            "activity_budget": 1000,
            "contract_store_count": 1,
            "fee_per_store": 1000,
            "fee_basis": "per_store",
            "contract_stack_count": 1,
            "expected_contract_amount": 1000,
            "claimed_amount": 1000,
            "sales_sku_count": 1,
            "sales_quantity": 20,
            "sales_retail_amount": 398,
            "sales_knowledge_matched_count": 1,
            "sales_knowledge_fuzzy_count": 0,
            "sales_knowledge_problem_count": 0,
            "contract_attachment_present": True,
            "contract_attachment_row_count": 1,
            "contract_attachment_quantity": 20,
            "contract_attachment_amount": 398,
            "contract_attachment_knowledge_problem_count": 0,
            "contract_attachment_sales_problem_count": 0,
            "photo_count": 1,
            "passed_store_count": 1,
            "supplement_store_count": 0,
            "suggested_approved_amount": 1000,
            "supported_amount": 1000,
            "temporarily_held_amount": 0,
            "high_exception_count": 0,
            "medium_exception_count": 0,
            "contract_core_status": "pass",
            "contract_core_problem_count": 0,
            "sales_internal_status": "pass",
        },
        "contract": {
            "source_file": "堆头合同.pdf",
            "customer_name": "示例客户",
            "contract_parties": ["参半公司", "示例客户"],
            "activity_start": "2026-08-01",
            "activity_end": "2026-08-31",
            "activity_budget": 1000,
            "activity_content": "示例门店开展堆头陈列活动",
            "settlement_method": "合同明确按店核销，每店1000元",
            "display_standard": "1平米堆头或4纵陈列",
            "fee_per_store": 1000,
            "fee_basis": "per_store",
            "claimed_amount": 1000,
            "contract_stack_count": 1,
            "watermark_visible": True,
            "seal_visible": True,
            "requires_specific_products": False,
            "required_products": [],
            "required_product_identities": [],
            "requires_promotion": False,
            "required_promotion": None,
            "sales_attachment": {
                "present": True,
                "source_pages": [3],
                "records": [
                    {
                        "line_no": 1,
                        "source_page": 3,
                        "customer_name": "示例客户",
                        "business_date": "2026-08",
                        "product_code": "CP-KQ-YG-0001",
                        "product_name": "示例商品原名",
                        "barcode_69": "6970356167341",
                        "unit": "支",
                        "quantity": 20,
                        "retail_price": 19.9,
                        "total_amount": 398,
                    }
                ],
                "total_quantity": 20,
                "total_amount": 398,
            },
            "stores": [
                {
                    "line_no": 1,
                    "store_name": "示例门店",
                    "address": "示例地址",
                    "stack_count": 1,
                }
            ],
            "notes": [],
        },
        "contract_product_knowledge": {
            "status": "not_applicable",
            "records": [],
            "knowledge_product_ids": [],
            "basis": "合同核心条款未限定具体商品，知识库商品门禁不适用。",
        },
        "contract_attachment_product_knowledge": {
            "status": "pass",
            "record_count": 1,
            "matched_count": 1,
            "fuzzy_count": 0,
            "problem_count": 0,
            "problem_rows": [],
            "records": [
                {
                    "excel_row": 1,
                    "contract_line_no": 1,
                    "source_page": 3,
                    "source_product_code": "CP-KQ-YG-0001",
                    "source_product_name": "示例商品原名",
                    "source_barcode_69": "6970356167341",
                    "source_quantity": 20,
                    "knowledge_status": "matched",
                    "knowledge_product_id": "canban-example",
                    "knowledge_product_name": "示例商品原名",
                    "knowledge_product_code": "CP-KQ-YG-0001",
                    "knowledge_barcode_69": "6970356167341",
                    "basis": "合同附件商品编码、商品名称和69码全部对应。",
                }
            ],
            "basis": "1/1行通过共享商品知识库核验。",
        },
        "contract_core_reconciliation": {
            "status": "pass",
            "source_file": "堆头合同.pdf",
            "field_count": 6,
            "passed_count": 6,
            "problem_count": 0,
            "problem_fields": [],
            "field_checks": [
                {
                    "field": field,
                    "label": label,
                    "contract_value": value,
                    "comparison_value": value,
                    "status": "pass",
                    "basis": "合同字段清晰可核验。",
                }
                for field, label, value in (
                    ("contracting_party", "签订方", "示例客户"),
                    ("activity_budget", "活动预算", "1000元"),
                    ("execution_period", "执行周期", "2026-08-01至2026-08-31"),
                    ("activity_content", "活动内容", "示例门店开展堆头陈列活动"),
                    ("settlement_method", "核销方式", "按店核销"),
                    ("seal", "盖章", "可见"),
                )
            ],
            "basis": "合同六项核心字段全部通过。",
        },
        "contract_pdf": {},
        "contract_sales_reconciliation": {
            "status": "pass",
            "customer_status": "exact",
            "customer_basis": "销售客户与合同签订方名称一致。",
            "customer_matches": [
                {
                    "sales_customer": "示例客户",
                    "contract_party": "示例客户",
                    "status": "exact",
                    "similarity": 1.0,
                }
            ],
            "sales_customers": ["示例客户"],
            "contract_parties": ["示例客户", "参半公司"],
            "period_status": "covered",
            "period_basis": "销售业务日期全部落在合同执行周期内。",
            "sales_period_values": ["2026-08"],
            "period_checks": [
                {
                    "sales_period": "2026-08",
                    "start": "2026-08-01",
                    "end": "2026-08-31",
                    "status": "covered",
                }
            ],
            "watermark_status": "present",
            "seal_status": "present",
            "integrity_status": "pass",
            "integrity_basis": "合同盖章清晰可见；水印状态仅记录，不参与核销。",
        },
        "contract_attachment_sales_reconciliation": {
            "status": "not_applicable",
            "record_count": 0,
            "matched_count": 0,
            "problem_count": 0,
            "records": [],
            "unmatched_contract_rows": [],
            "unmatched_sales_rows": [],
            "quantity_status": "not_applicable",
            "amount_status": "not_applicable",
            "basis": "合同包未附商品销售明细，本项不核验。",
        },
        "sales": {
            "source_file": "堆头销售.xlsx",
            "total_quantity": 20,
            "retail_amount": 398,
            "customers": ["示例客户"],
            "period_values": ["2026-08"],
            "store_level_available": False,
            "internal_status": "pass",
            "internal_problem_rows": [],
            "internal_basis": "销售Excel行内金额及打印合计全部一致。",
            "knowledge_status": "pass",
            "knowledge_matched_count": 1,
            "knowledge_fuzzy_count": 0,
            "knowledge_problem_count": 0,
            "knowledge_problem_rows": [],
            "knowledge_basis": "1/1行唯一对应商品知识库；全部通过",
            "records": [
                {
                    "excel_row": 2,
                    "customer_name": "示例客户",
                    "period_text": "2026-08",
                    "product_code": "CP-KQ-YG-0001",
                    "product_name": "示例商品原名",
                    "barcode": "6970356167341",
                    "unit": "支",
                    "quantity": 20,
                    "retail_price": 19.9,
                    "total_amount": 398,
                }
            ],
            "knowledge_reconciliation": [
                {
                    "excel_row": 2,
                    "source_product_code": "CP-KQ-YG-0001",
                    "source_product_name": "示例商品原名",
                    "source_barcode_69": "6970356167341",
                    "source_quantity": 20,
                    "knowledge_status": "matched",
                    "knowledge_product_id": "canban-example",
                    "knowledge_product_name": "示例商品原名",
                    "knowledge_product_code": "CP-KQ-YG-0001",
                    "knowledge_barcode_69": "6970356167341",
                    "name_match_type": "exact",
                    "name_similarity": None,
                    "matched_fields": ["product_code", "product_name", "barcode_69"],
                    "unmatched_fields": [],
                    "conflicting_fields": [],
                    "missing_fields": [],
                    "candidate_product_ids": ["canban-example"],
                    "field_comparisons": [
                        {
                            "field": "product_code",
                            "source_value": "CP-KQ-YG-0001",
                            "comparison": "matched",
                            "selected_knowledge_value": "CP-KQ-YG-0001",
                            "matching_products": [
                                {
                                    "product_id": "canban-example",
                                    "product_code": "CP-KQ-YG-0001",
                                    "product_name": "示例商品原名",
                                    "barcode_69": "6970356167341",
                                }
                            ],
                        },
                        {
                            "field": "product_name",
                            "source_value": "示例商品原名",
                            "comparison": "matched",
                            "selected_knowledge_value": "示例商品原名",
                            "matching_products": [
                                {
                                    "product_id": "canban-example",
                                    "product_code": "CP-KQ-YG-0001",
                                    "product_name": "示例商品原名",
                                    "barcode_69": "6970356167341",
                                }
                            ],
                        },
                        {
                            "field": "barcode_69",
                            "source_value": "6970356167341",
                            "comparison": "matched",
                            "selected_knowledge_value": "6970356167341",
                            "matching_products": [
                                {
                                    "product_id": "canban-example",
                                    "product_code": "CP-KQ-YG-0001",
                                    "product_name": "示例商品原名",
                                    "barcode_69": "6970356167341",
                                }
                            ],
                        },
                    ],
                    "basis": "销售Excel第2行唯一对应知识库商品CP-KQ-YG-0001",
                }
            ],
        },
        "store_reconciliation": [
            {
                "store_line_no": 1,
                "contract_store_name": "示例门店",
                "photo_files": ["示例门店.jpg"],
                "visible_date": "2026-08-18",
                "visible_location": "示例门店",
                "period_match": "match",
                "store_match": "exact",
                "store_match_basis": "现场可见门店名称与合同门店一致",
                "location_resolution": {
                    "status": "not_needed",
                    "confidence": "not_applicable",
                    "provider": "deterministic_name",
                    "transport": "local",
                    "mcp_attempted": False,
                    "mcp_status": "not_needed",
                    "contract_query": "示例门店",
                    "watermark_query": "示例门店",
                    "city_hint": None,
                    "coordinate_system": None,
                    "contract_poi": None,
                    "watermark_poi": None,
                    "distance_meters": None,
                    "nearby_pass_meters": 100.0,
                    "unrelated_min_meters": 300.0,
                    "evidence": [],
                    "basis": "现场可见门店名称与合同门店一致，无需调用地图MCP",
                },
                "display_match": "pass",
                "display_standard_basis": "four_vertical",
                "display_description": "画面中可清楚数出4列纵向陈列。",
                "display_vertical_facing_count": 4,
                "display_vertical_facing_basis": [
                    "左一绿色产品列",
                    "左二红色产品列",
                    "右一白紫色产品列",
                    "右二窄白色产品列",
                ],
                "display_stack_1sqm_basis": None,
                "duplicate_check": "none",
                "visible_text": ["示例商品原名", "6970356167341"],
                "visible_recognized_products": ["示例商品原名"],
                "recognized_products": [
                    "示例商品原名（产品编码：CP-KQ-YG-0001；69码：6970356167341；视觉RAG：精确命中）"
                ],
                "product_reference_hits": [
                    {
                        "reference_product_id": "canban-example",
                        "product_name": "示例商品原名",
                        "product_code": "CP-KQ-YG-0001",
                        "product_code_aliases": [],
                        "barcode_69": "6970356167341",
                        "confidence": "exact",
                        "matched_view_ids": ["front"],
                        "visible_basis": ["包装正面名称和69码可见"],
                    }
                ],
                "photo_knowledge_match": "exact",
                "photo_knowledge_match_basis": "现场照片商品精确对应正式知识库",
                "promotion_summary": "未识别到明确促销词或组合装",
                "sales_product_names": ["示例商品原名"],
                "sales_product_knowledge_ids": ["canban-example"],
                "sales_product_match": "exact",
                "sales_product_match_basis": "现场照片与销售Excel通过知识库后商品ID完全一致",
                "sales_product_checks": [
                    {
                        "knowledge_product_id": "canban-example",
                        "knowledge_product_code": "CP-KQ-YG-0001",
                        "knowledge_product_code_aliases": [],
                        "knowledge_product_name": "示例商品原名",
                        "knowledge_barcode_69": "6970356167341",
                        "photo_match": "exact",
                        "excel_row": 2,
                        "source_product_code": "CP-KQ-YG-0001",
                        "source_product_name": "示例商品原名",
                        "source_barcode_69": "6970356167341",
                        "source_quantity": 20,
                        "product_code_match": "exact",
                        "name_match": "exact",
                        "name_similarity": 1.0,
                        "barcode_match": "exact",
                        "status": "exact",
                        "confidence": "high",
                        "barcode_comparison_basis": (
                            "现场照片先确定知识库商品，再将该商品登记的69码与销售Excel比较；"
                            "两边一致。"
                        ),
                        "basis": "商品编码和69码严格一致，商品名称可以对应。",
                        "resubmission": None,
                    }
                ],
                "contract_sales_status": "pass",
                "contract_attachment_sales_status": "pass",
                "contract_core_status": "pass",
                "sales_internal_status": "pass",
                "photo_contract_product_status": "exact",
                "photo_contract_product_basis": "现场商品属于合同附件已确认商品范围。",
                "photo_contract_product_checks": [
                    {
                        "reference_product_id": "canban-example",
                        "product_code": "CP-KQ-YG-0001",
                        "product_name": "示例商品原名",
                        "barcode_69": "6970356167341",
                        "visual_confidence": "exact",
                        "contract_sources": ["合同附件第1行（PDF第3页）"],
                        "status": "exact",
                        "basis": "现场包装与合同附件已确认商品一致。",
                    }
                ],
                "contract_attachment_product_names": [],
                "contract_attachment_product_knowledge_ids": [],
                "contract_attachment_product_match": "not_applicable",
                "contract_attachment_product_match_basis": "合同包未附商品销售明细，本项不核验。",
                "contract_attachment_product_checks": [],
                "contract_stack_count": 1,
                "claim_units": 1,
                "amount_rule_status": "pass",
                "amount_rule_basis": "合同明确按店核销：1店×1000元。",
                "promotion_present": False,
                "supported_amount": 1000,
                "status": "pass",
                "supplement_advice": [],
            }
        ],
        "photo_inventory": {},
        "exceptions": [],
    }
    result["contract_attachment_sales_reconciliation"] = (
        _contract_attachment_sales_reconciliation(
            result["contract"],
            result["sales"],
        )
    )
    return result


def _test_ean13(seed: int) -> str:
    body = f"69{seed:010d}"
    weighted = sum(
        int(value) * (1 if index % 2 == 0 else 3)
        for index, value in enumerate(body)
    )
    return body + str((10 - weighted % 10) % 10)


def _attachment_contract_and_sales(row_count: int = 36) -> tuple[dict, dict]:
    attachment_rows: list[dict] = []
    sales_rows: list[dict] = []
    for line_no in range(1, row_count + 1):
        quantity = line_no + 10
        retail_price = line_no + 20
        total_amount = quantity * retail_price
        code = f"CP-TEST-{line_no:04d}"
        barcode = _test_ean13(line_no)
        name = f"参半测试商品{line_no}"
        attachment_rows.append(
            {
                "line_no": line_no,
                "source_page": 3 if line_no <= 23 else 4,
                "customer_name": "东莞市诚成行供应链管理有限公司",
                "business_date": "2026-04",
                "product_code": code,
                "product_name": name,
                "barcode_69": barcode,
                "unit": "支",
                "quantity": quantity,
                "retail_price": retail_price,
                "total_amount": total_amount,
            }
        )
        sales_rows.append(
            {
                "excel_row": line_no + 1,
                "customer_name": "东莞市诚成行供应链管理有限公司",
                "period_text": "2026-04",
                "product_code": code,
                "product_name": name,
                "barcode": barcode,
                "unit": "支",
                "quantity": quantity,
                "retail_price": retail_price,
                "total_amount": total_amount,
            }
        )
    total_quantity = sum(item["quantity"] for item in attachment_rows)
    total_amount = sum(item["total_amount"] for item in attachment_rows)
    contract = {
        "sales_attachment": {
            "present": True,
            "source_pages": [3, 4],
            "records": attachment_rows,
            "total_quantity": total_quantity,
            "total_amount": total_amount,
        }
    }
    sales = {
        "records": sales_rows,
        "total_quantity": total_quantity,
        "retail_amount": total_amount,
    }
    return contract, sales


def _display_result_with_attachment() -> dict:
    return deepcopy(_display_result())


def _row_with_first_cell_prefix(worksheet, prefix: str) -> int:
    for row_index in range(4, worksheet.max_row + 1):
        value = str(worksheet.cell(row_index, 1).value or "").strip()
        if value.startswith(prefix):
            return row_index
    raise AssertionError(f"{worksheet.title} 未找到首列以 {prefix!r} 开头的业务行")


def _display_detail_row(worksheet) -> int:
    return _row_with_first_cell_prefix(worksheet, "门店：")


class ProducerFilenameTests(unittest.TestCase):
    def test_temporary_run_directory_is_short_and_does_not_embed_run_id(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            run_root = _create_temporary_root(root)
            self.assertRegex(run_root.name, r"^\.oa-[0-9a-f]{32}$")
            self.assertLessEqual(len(run_root.name), 36)

    def test_codex_and_repeat_use_date_model_revision(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            first = next_output_path("20260818-simple-007", "codex", root)
            self.assertEqual(first.name, "20260818-codex.html")
            first.touch()
            second = next_output_path("20260818-simple-007", "codex", root)
            self.assertEqual(second.name, "20260818-codex-1.1.html")
            second.touch()
            third = next_output_path("20260818-simple-007", "codex", root)
            self.assertEqual(third.name, "20260818-codex-1.2.html")

    def test_each_model_has_an_independent_sequence(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            (root / "20260818-codex.xlsx").touch()
            self.assertEqual(
                next_output_path("20260818-simple-008", "qwen3.7", root).name,
                "20260818-qwen3.7.html",
            )

    def test_missing_historical_files_do_not_reuse_a_revision_label(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            (root / "20260818-codex-1.3.xlsx").touch()
            self.assertEqual(
                next_output_path("20260818-simple-009", "codex", root).name,
                "20260818-codex-1.4.html",
            )
            (root / "20260818-codex.xlsx").touch()
            (root / "20260818-codex-1.1.xlsx").touch()
            self.assertEqual(
                next_output_path("20260818-simple-010", "codex", root).name,
                "20260818-codex-1.4.html",
            )

    def test_html_only_file_also_reserves_its_revision_label(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            (root / "20260818-codex.html").touch()
            self.assertEqual(
                next_output_path("20260818-simple-011", "codex", root).name,
                "20260818-codex-1.1.html",
            )
            (root / "20260818-codex-1.4.html").touch()
            self.assertEqual(
                next_output_path("20260818-simple-012", "codex", root).name,
                "20260818-codex-1.5.html",
            )

    def test_only_verified_html_is_published(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            temporary_workbook = root / ".result.xlsx"
            temporary_html = root / ".result.html"
            temporary_workbook.write_bytes(b"xlsx-content")
            temporary_html.write_text("<html>result</html>", encoding="utf-8")

            html_path = _publish_pair_without_overwrite(
                temporary_workbook,
                temporary_html,
                "20260818-simple-013",
                "codex",
                root,
            )

            self.assertEqual(html_path.name, "20260818-codex.html")
            self.assertEqual(html_path.read_text(encoding="utf-8"), "<html>result</html>")
            self.assertFalse((root / "20260818-codex.xlsx").exists())
            self.assertFalse(temporary_workbook.exists())
            self.assertFalse(temporary_html.exists())

    def test_run_id_must_start_with_a_valid_business_date(self) -> None:
        self.assertEqual(output_date_from_run_id("20260818-simple-007"), "20260818")
        for value in ("simple-007", "20261301-simple-007", "20260230"):
            with self.subTest(value=value), self.assertRaises(AuditError):
                output_date_from_run_id(value)

    def test_other_model_label_is_normalized(self) -> None:
        self.assertEqual(normalize_producer_model(" QWEN3.7 "), "qwen3.7")

    def test_unsafe_model_label_is_rejected(self) -> None:
        for value in ("../qwen", "qwen/3.7", "", "."):
            with self.subTest(value=value), self.assertRaises(AuditError):
                normalize_producer_model(value)


class HtmlReportTests(unittest.TestCase):
    def test_canonical_template_is_the_only_html_asset_and_empty_payload_fails(self) -> None:
        self.assertTrue(TEMPLATE_ASSET_PATH.is_file())
        self.assertEqual(
            list(TEMPLATE_ASSET_PATH.parent.glob("*.html")),
            [TEMPLATE_ASSET_PATH],
        )
        with self.assertRaisesRegex(AuditError, "没有可展示的核销场景"):
            _render_template({"sheets": []})

    def test_row_sections_and_actual_application_are_deterministic(self) -> None:
        self.assertEqual(_row_kind(["实际申请金额", "", "", "", "", ""]), "summary")
        self.assertEqual(_row_kind(["活动概况", "", "", "", "", ""]), "summary")
        self.assertEqual(_row_kind(["合同销售附件", "", "", "", "", ""]), "summary")
        for scenario in ("personnel_incentive", "promotional_display"):
            self.assertEqual(_row_section("record", scenario), "detail")
            self.assertEqual(_row_section("summary", scenario), "settlement")
        self.assertEqual(
            _row_section(
                "summary",
                "promotional_display",
                ["活动概况", "", "", "", "", ""],
            ),
            "overview",
        )
        self.assertEqual(
            _row_section(
                "summary",
                "promotional_display",
                ["合同销售附件", "", "", "", "", ""],
            ),
            "overview",
        )

    def test_display_attachment_rows_keep_overview_detail_and_settlement_separate(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            workbook_path = Path(temporary) / "result.xlsx"
            create_combined_report([_display_result_with_attachment()], workbook_path)
            payload = _workbook_payload(workbook_path)
            self.assertEqual(len(payload["sheets"]), 1)
            sheet = payload["sheets"][0]
            rows = sheet["rows"]
            self.assertEqual(
                [row["section"] for row in rows],
                ["overview", "overview", "overview", "detail", "settlement"],
            )
            self.assertTrue(rows[0]["heading"].startswith("活动概况"))
            self.assertTrue(rows[1]["heading"].startswith("合同销售附件｜合同主基准"))
            self.assertTrue(rows[2]["heading"].startswith("合同销售附件第1行"))
            self.assertEqual(rows[3]["heading"], "示例门店")
            self.assertTrue(rows[4]["heading"].startswith("合计金额"))
            self.assertEqual(
                sheet["audit_counts"],
                {
                    "source_row_count": 5,
                    "error_count": 0,
                    "detail_error_count": 0,
                    "context_error_count": 0,
                },
            )

            result = _display_result_with_attachment()
            result["sales"]["records"][0]["product_code"] = "CP-WRONG-0001"
            result["contract_attachment_sales_reconciliation"] = (
                _contract_attachment_sales_reconciliation(
                    result["contract"], result["sales"]
                )
            )
            result["store_reconciliation"][0].update(
                {
                    "display_match": "fail",
                    "display_standard_basis": "none",
                    "status": "supplement",
                    "supported_amount": 0,
                }
            )
            result["summary"].update(
                {
                    "passed_store_count": 0,
                    "supplement_store_count": 1,
                    "suggested_approved_amount": 0,
                    "supported_amount": 0,
                    "temporarily_held_amount": 1000,
                }
            )
            create_combined_report([result], workbook_path)
            issue_sheet = _workbook_payload(workbook_path)["sheets"][0]
            rows = issue_sheet["rows"]
            self.assertEqual(
                [row["section"] for row in rows],
                ["overview", "overview", "overview", "detail", "settlement"],
            )
            self.assertEqual(
                [row["heading"] for row in rows if row["status"] == "issue"],
                [
                    "合同销售附件｜合同主基准",
                    "合同销售附件第1行｜PDF第3页",
                    "示例门店",
                    "合计金额",
                ],
            )
            self.assertEqual(issue_sheet["audit_counts"]["error_count"], 4)

    def test_summary_row_with_a_concrete_resubmission_is_still_an_issue(self) -> None:
        values = ["收款人与日期", "", "", "", "", "置信度：低\n要重新提交：补拍转账截图"]
        self.assertEqual(_row_status(values, "summary"), "issue")
        values[5] = "暂不能核销"
        self.assertEqual(_row_status(values, "summary"), "issue")
        values[5] = "完全不匹配，需补材料"
        self.assertEqual(_row_status(values, "summary"), "issue")
        values[5] = "置信度：高\n身份与日期均可确认，无需重新提交"
        self.assertEqual(_row_status(values, "summary"), "summary")

    def test_personnel_card_heading_includes_authoritative_product_name(self) -> None:
        values = [
            "结算第1行\n参半oralshark羟基磷灰石牙膏满陇桂雨味(100g)-线下10.1版",
            "",
            "",
            "",
            "",
            "",
        ]
        self.assertEqual(
            _row_heading(values, 4, "record", "personnel_incentive"),
            "参半oralshark羟基磷灰石牙膏满陇桂雨味(100g)-线下10.1版",
        )

    def test_display_card_heading_prefers_real_store_over_contract_note(self) -> None:
        values = [
            "合同总体见首行\n第2家：百佳华松岗店\n地址：合同未列地址\n本店堆头：1",
            "识别地点：深圳市·百佳华超市（松岗百佳华商场店）",
            "",
            "",
            "",
            "",
        ]
        self.assertEqual(
            _row_heading(values, 5, "record", "promotional_display"),
            "百佳华松岗店",
        )

    def test_html_is_self_contained_interactive_and_equal_to_workbook(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            workbook_path = root / "20260818-codex.xlsx"
            html_path = root / "20260818-codex.html"
            create_combined_report([_personnel_result(), _display_result()], workbook_path)
            create_html_report_from_workbook(workbook_path, html_path)

            verification = verify_html_report(
                html_path,
                ["personnel_incentive", "promotional_display"],
                workbook_path=workbook_path,
            )
            self.assertEqual(
                verification["sheet_names"],
                ["人员激励核销", "堆头核销"],
            )
            self.assertEqual(verification["external_dependency_count"], 0)
            self.assertEqual(verification["button_count"], 20)
            self.assertTrue(verification["workbook_content_equal"])
            self.assertTrue(verification["static_template_equal"])
            self.assertEqual(verification["template_version"], "3.9.0")
            self.assertRegex(verification["style_sha256"], r"^[0-9a-f]{64}$")
            self.assertRegex(verification["shell_sha256"], r"^[0-9a-f]{64}$")

            payload = _workbook_payload(workbook_path)
            self.assertEqual(payload["schema_version"], "1.1")
            self.assertEqual(
                [sheet["scenario"] for sheet in payload["sheets"]],
                ["personnel_incentive", "promotional_display"],
            )
            for sheet in payload["sheets"]:
                self.assertEqual(
                    len(sheet["rows"]),
                    sheet["audit_counts"]["source_row_count"],
                )
                for row in sheet["rows"]:
                    expected = _row_section(
                        row["kind"],
                        sheet["scenario"],
                        row["values"],
                    )
                    self.assertEqual(row["section"], expected)

            html = html_path.read_text(encoding="utf-8")
            for required in (
                'id="error-only-preview-style"',
                'id="error-only-preview-script"',
                'class="eo-home-cockpit"',
                'class="eo-tab"',
                'class="eo-error-card"',
                'aria-label="${escapeHtml(accessibleName)}"',
                'class="eo-error-details"',
                'class="eo-table"',
                'id="eoErrorList"',
                'id="eoErrorType"',
                'id="eoErrorCategory"',
                'id="eoErrorFilterEmpty"',
                'id="eoPassList"',
                'id="eoPassType"',
                'id="eoPassCategory"',
                'id="eoPassFacetSummary"',
                'id="eoPassFilterEmpty"',
                'id="eoBackTop"',
                'role="tablist" aria-label="核销结果视图"',
                'id="eo-tab-home" role="tab"',
                'id="eo-tab-passed" role="tab"',
                'id="home" role="tabpanel" aria-labelledby="eo-tab-home"',
                'id="passed" role="tabpanel" aria-labelledby="eo-tab-passed"',
                'data-eo-view="passed"',
                "筛选错误检查项",
                "筛选正确检查项",
                "种核销方式",
                "核销类型",
                "核销类型 ·",
                "错误原因分类",
                'data-error-categories=',
                'data-error-confidence-score=',
                'data-pass-confidence-score=',
                'class="eo-chip eo-error-reason"',
                "@media (prefers-reduced-motion: reduce)",
                "content-visibility: auto;",
                "contain-intrinsic-size: auto 168px;",
                "@media print",
                'aria-label="${escapeHtml(accessibleLabel)}"',
                "const passFilterRecordByCard = new WeakMap",
                "const passFilterGroups =",
                "const passTypeFacetCount = renderPassFacet",
                "['ArrowUp', 'ArrowDown', 'ArrowLeft', 'ArrowRight', 'Home', 'End']",
                "tab.setAttribute('tabindex', selected ? '0' : '-1')",
            ):
                self.assertIn(required, html)
            self.assertNotIn("单项通过不等于整单核销通过", html)
            self.assertNotIn("全部错误直接向下排列", html)
            self.assertNotIn("每项标明对应核销类型", html)
            self.assertNotIn("当前视图直接列出全部需要处理的错误", html)
            self.assertNotIn('id="eoErrorFacetSummary"', html)
            self.assertNotIn('id="eoErrorConfidence"', html)
            self.assertNotIn('id="eoPassConfidence"', html)
            self.assertNotIn('class="eo-pass-warning"', html)
            self.assertNotIn("核销方式始终展示全部；置信度随核销方式和关键词联动", html)
            self.assertNotIn("核销方式始终展示全部；检查分类和置信度只保留当前存在项", html)
            self.assertNotIn(".eo-chip.eo-error-reason {", html)
            for removed_summary in (
                'id="eoErrorKeyword"',
                'id="eoPassKeyword"',
                "FILTER REJECTED CHECKS",
                "FILTER VERIFIED CHECKS",
                "eo-pass-filter-head",
                'id="eoErrorVisible"',
                'id="eoPassVisible"',
                "重置筛选",
                "REJECTED ERROR RESULTS",
                "核销错误结果",
                "VERIFIED PASS RESULTS",
                "按核销类型归档",
                "eo-queue-head",
                "AUDIT ERROR COMMAND",
                "核销错误处置总览",
                "VERIFIED CHECK LEDGER",
                "<h1>正确检查项日志</h1>",
                "eo-cockpit-hero",
                "eo-total-gauge",
                "待处理总数",
                'class="eo-metric"',
                'class="eo-composition',
                "材料 / 结算错误",
                "商品 / 门店错误",
                "材料 / 商品",
                "门店 / 金额",
                "ERROR COMPOSITION",
                "PASS COMPOSITION",
            ):
                self.assertNotIn(removed_summary, html)

            self.assertIn(
                '<div class="brand-mark" aria-label="参半 CANBAN"><strong>参半</strong><small>CANBAN</small></div>',
                html,
            )
            self.assertIn("const personnelView = (sheet) =>", html)
            self.assertIn("const displayView = (sheet) =>", html)
            self.assertIn("const posterView = (sheet) =>", html)
            self.assertIn("const groupedIssueView = (sheet, auditType) =>", html)
            self.assertIn("合同销售附件第1行", html)
            self.assertIn("problem: groupedRowProblem(sheet, knowledgeRows, 'knowledge')", html)
            self.assertIn("problem: groupedRowProblem(sheet, salesRows, 'sales')", html)
            self.assertNotIn("extra: detailsTable(", html)
            self.assertIn(".eo-table td::before", html)
            self.assertIn("consumed !== content.length", html)
            self.assertNotIn("7350元，转账凭证识别金额为7353元", html)
            self.assertNotIn("销售Excel读取到7家门店", html)
            self.assertNotIn("photoKnowledgeAlreadySupported", html)
            self.assertNotIn(
                "normalizeFilterText(card.textContent).includes(state.keyword)", html
            )

            for forbidden in (
                'id="homeIssueTotal"',
                'id="homeIssueBreakdown"',
                'class="home-issue-anchor"',
                'class="home-head"',
                "scenario-overview-card",
                'class="eo-scenario-card"',
                'class="eo-enter"',
                'data-open="scenario-',
                "进入错误清单",
                "错误场景队列",
                "const openView = (id) =>",
                "__CANBAN_STYLE_SHA256__",
                "__CANBAN_SHELL_SHA256__",
                "__AUDIT_DATA__",
                'class="compact"',
                "densityButton",
                "紧凑显示",
                "舒展显示",
                "state.compact",
                "打印 / 存 PDF",
                'id="printButton"',
                "window.print()",
                "打开同版 Excel",
                'id="excelLink"',
                "setExcelLink",
                'id="fileName"',
                "页面可离线打开",
                "内容与同名 Excel 保持一致",
                "复制本条结论",
                "data-copy-row",
                "navigator.clipboard",
                'id="toast"',
                "只看汇总",
                "Offline activity verification",
                "STKaiti",
                "background-image: url",
                "window.addEventListener('scroll'",
            ):
                self.assertNotIn(forbidden, html)
            self.assertNotRegex(
                html,
                r'<(?:script|link)\b[^>]*(?:src|href)=["\']https?://',
            )

    def test_html_rejects_density_toggle(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            workbook_path = root / "20260818-codex.xlsx"
            html_path = root / "20260818-codex.html"
            create_combined_report([_personnel_result(), _display_result()], workbook_path)
            create_html_report_from_workbook(workbook_path, html_path)
            html = html_path.read_text(encoding="utf-8").replace(
                "</body>",
                '<button id="densityButton">紧凑显示</button></body>',
            )
            html_path.write_text(html, encoding="utf-8")
            with self.assertRaisesRegex(AuditError, "紧凑/宽松"):
                verify_html_report(
                    html_path,
                    ["personnel_incentive", "promotional_display"],
                    workbook_path=workbook_path,
                )

    def test_html_rejects_removed_meta_and_copy_controls(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            workbook_path = root / "result.xlsx"
            html_path = root / "result.html"
            create_combined_report([_personnel_result(), _display_result()], workbook_path)
            create_html_report_from_workbook(workbook_path, html_path)
            html = html_path.read_text(encoding="utf-8").replace(
                "</body>",
                '<button data-copy-row="4">复制本条结论</button></body>',
            )
            html_path.write_text(html, encoding="utf-8")
            with self.assertRaisesRegex(AuditError, "无效页面文案或功能"):
                verify_html_report(
                    html_path,
                    ["personnel_incentive", "promotional_display"],
                    workbook_path=workbook_path,
                )

    def test_html_rejects_static_style_drift_from_canonical_asset(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            workbook_path = root / "result.xlsx"
            html_path = root / "result.html"
            create_combined_report([_personnel_result(), _display_result()], workbook_path)
            create_html_report_from_workbook(workbook_path, html_path)
            html = html_path.read_text(encoding="utf-8").replace(
                "background: var(--page);",
                "background: #ffffff;",
                1,
            )
            html_path.write_text(html, encoding="utf-8")
            with self.assertRaisesRegex(AuditError, "静态模板"):
                verify_html_report(
                    html_path,
                    ["personnel_incentive", "promotional_display"],
                    workbook_path=workbook_path,
                )

    def test_html_keeps_every_error_row_without_top_n_truncation(self) -> None:
        result = deepcopy(_display_result())
        contract_source, sales_source = _attachment_contract_and_sales(row_count=10)
        result["contract"]["sales_attachment"] = contract_source["sales_attachment"]
        result["sales"].update(sales_source)
        for sales_record in result["sales"]["records"]:
            sales_record["product_code"] = (
                f"CP-WRONG-{int(sales_record['excel_row']) - 1:04d}"
            )

        base_knowledge = result["contract_attachment_product_knowledge"]["records"][0]
        knowledge_records: list[dict] = []
        for contract_record in result["contract"]["sales_attachment"]["records"]:
            line_no = int(contract_record["line_no"])
            knowledge_record = deepcopy(base_knowledge)
            knowledge_record.update(
                {
                    "excel_row": line_no,
                    "contract_line_no": line_no,
                    "source_page": contract_record["source_page"],
                    "source_product_code": contract_record["product_code"],
                    "source_product_name": contract_record["product_name"],
                    "source_barcode_69": contract_record["barcode_69"],
                    "source_quantity": contract_record["quantity"],
                    "knowledge_product_id": f"attachment-{line_no}",
                    "knowledge_product_code": contract_record["product_code"],
                    "knowledge_product_name": contract_record["product_name"],
                    "knowledge_barcode_69": contract_record["barcode_69"],
                }
            )
            knowledge_records.append(knowledge_record)
        result["contract_attachment_product_knowledge"].update(
            {
                "record_count": 10,
                "matched_count": 10,
                "problem_count": 0,
                "problem_rows": [],
                "records": knowledge_records,
                "basis": "10/10行通过共享商品知识库核验。",
            }
        )
        result["summary"].update(
            {
                "sales_sku_count": 10,
                "sales_quantity": sales_source["total_quantity"],
                "sales_retail_amount": sales_source["retail_amount"],
                "contract_attachment_row_count": 10,
                "contract_attachment_quantity": contract_source["sales_attachment"]["total_quantity"],
                "contract_attachment_amount": contract_source["sales_attachment"]["total_amount"],
                "contract_attachment_sales_problem_count": 10,
            }
        )
        result["contract_attachment_sales_reconciliation"] = (
            _contract_attachment_sales_reconciliation(
                result["contract"], result["sales"]
            )
        )
        self.assertEqual(
            result["contract_attachment_sales_reconciliation"]["problem_count"],
            10,
        )

        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            workbook_path = root / "result.xlsx"
            html_path = root / "result.html"
            create_combined_report([result], workbook_path)
            create_html_report_from_workbook(workbook_path, html_path)
            verify_html_report(
                html_path,
                ["promotional_display"],
                workbook_path=workbook_path,
            )
            html = html_path.read_text(encoding="utf-8")
            for excel_row in range(2, 12):
                self.assertIn(f"销售Excel第{excel_row}行", html)
            self.assertNotIn("其余", html)


class DisplayContractTests(unittest.TestCase):
    def test_contract_summary_omits_absent_optional_product_and_promotion_terms(self) -> None:
        contract = {
            "display_standard": "1平方米堆头或4纵陈列",
            "requires_specific_products": False,
            "required_products": [],
            "requires_promotion": False,
            "required_promotion": None,
        }
        self.assertEqual(
            _contract_requirement_lines(contract),
            ["陈列：1平方米堆头或4纵陈列"],
        )

        contract.update(
            {
                "requires_specific_products": True,
                "required_products": ["参半示例商品"],
                "requires_promotion": True,
                "required_promotion": "买一赠一",
            }
        )
        self.assertEqual(
            _contract_requirement_lines(contract),
            [
                "陈列：1平方米堆头或4纵陈列",
                "商品：参半示例商品",
                "促销：买一赠一",
            ],
        )

    def test_display_control_requires_consistent_standard(self) -> None:
        self.assertEqual(
            _display_control(
                {"standard_evidence": "meets", "matched_standard": "four_vertical"}
            ),
            ("pass", "four_vertical"),
        )
        with self.assertRaises(AuditError):
            _display_control(
                {"standard_evidence": "meets", "matched_standard": "unclear"}
            )

    def test_contract_sales_checks_party_period_and_seal(self) -> None:
        contract = {
            "customer_name": "郑州诚成商贸有限公司",
            "contract_parties": ["参半健康科技有限公司", "郑州诚成商贸有限公司"],
            "activity_start": "2026-04-01",
            "activity_end": "2026-04-30",
            "watermark_visible": True,
            "seal_visible": True,
        }
        sales = {
            "customers": ["郑州诚成商贸"],
            "period_values": ["2026-04-18 00:00:00"],
        }
        result = _contract_sales_reconciliation(contract, sales)
        self.assertEqual(result["status"], "pass")
        self.assertEqual(result["customer_status"], "exact")
        self.assertEqual(result["period_status"], "covered")
        self.assertEqual(result["integrity_status"], "pass")

        contract["watermark_visible"] = False
        result = _contract_sales_reconciliation(contract, sales)
        self.assertEqual(result["status"], "pass")
        self.assertEqual(result["watermark_status"], "absent")
        self.assertEqual(result["integrity_status"], "pass")

        contract["seal_visible"] = False
        result = _contract_sales_reconciliation(contract, sales)
        self.assertEqual(result["status"], "fail")
        self.assertEqual(result["seal_status"], "absent")
        contract["seal_visible"] = True

        sales["period_values"] = ["2026-05"]
        result = _contract_sales_reconciliation(contract, sales)
        self.assertEqual(result["status"], "fail")
        self.assertEqual(result["period_status"], "mismatch")

    def test_contract_attachment_36_rows_match_sales_excel(self) -> None:
        contract, sales = _attachment_contract_and_sales()
        result = _contract_attachment_sales_reconciliation(contract, sales)
        self.assertEqual(result["status"], "pass")
        self.assertEqual(result["record_count"], 36)
        self.assertEqual(result["matched_count"], 36)
        self.assertEqual(result["problem_count"], 0)
        self.assertEqual(result["quantity_status"], "exact")
        self.assertEqual(result["amount_status"], "exact")
        self.assertTrue(
            all(
                comparison["status"] in {"exact", "fuzzy"}
                for record in result["records"]
                for comparison in record["field_comparisons"]
            )
        )

    def test_contract_attachment_reports_each_mismatched_field(self) -> None:
        cases = {
            "product_code": ("product_code", "CP-WRONG-0001"),
            "barcode_69": ("barcode", _test_ean13(9999)),
            "quantity": ("quantity", 999),
            "retail_price": ("retail_price", 999),
            "total_amount": ("total_amount", 999),
        }
        for expected_field, (sales_field, wrong_value) in cases.items():
            with self.subTest(field=expected_field):
                contract, sales = _attachment_contract_and_sales(row_count=1)
                sales["records"][0][sales_field] = wrong_value
                result = _contract_attachment_sales_reconciliation(contract, sales)
                self.assertEqual(result["status"], "fail")
                comparison = next(
                    item
                    for item in result["records"][0]["field_comparisons"]
                    if item["field"] == expected_field
                )
                self.assertEqual(comparison["status"], "mismatch")
                self.assertIn(expected_field, result["records"][0]["basis"])

    def test_contract_attachment_unit_is_display_only(self) -> None:
        contract, sales = _attachment_contract_and_sales(row_count=1)
        sales["records"][0]["unit"] = "盒"
        result = _contract_attachment_sales_reconciliation(contract, sales)
        self.assertEqual(result["status"], "pass")
        self.assertNotIn(
            "unit",
            {
                item["field"]
                for item in result["records"][0]["field_comparisons"]
            },
        )
        self.assertEqual(contract["sales_attachment"]["records"][0]["unit"], "支")
        self.assertEqual(sales["records"][0]["unit"], "盒")

    def test_contract_attachment_absence_is_unverifiable(self) -> None:
        result = _contract_attachment_sales_reconciliation(
            {"sales_attachment": {"present": False}},
            {"records": []},
        )
        self.assertEqual(result["status"], "unverifiable")
        self.assertEqual(result["record_count"], 0)
        self.assertEqual(result["problem_count"], 0)

    def test_workbook_separates_display_standard_and_photo_reuse(self) -> None:
        validate_json(
            _display_result(),
            Path(__file__).resolve().parents[1] / "contracts" / "audit-result.schema.json",
        )
        with tempfile.TemporaryDirectory() as temporary:
            target = Path(temporary) / "result.xlsx"
            create_combined_report([_personnel_result(), _display_result()], target)
            verification = verify_workbook(
                target, ["personnel_incentive", "promotional_display"]
            )
            self.assertEqual(verification["formula_count"], 0)

            workbook = load_workbook(target, data_only=False)
            try:
                self.assertEqual(str(workbook["人员激励核销"].freeze_panes), "A4")
                self.assertEqual(str(workbook["堆头核销"].freeze_panes), "A4")
                personnel_sheet = workbook["人员激励核销"]
                self.assertEqual(
                    personnel_sheet["A4"].value,
                    "参半示例商品100g标准名称",
                )
                personnel_sales_text = str(personnel_sheet["B4"].value)
                personnel_compare_text = str(personnel_sheet["E4"].value)
                self.assertIn("销售Excel + 商品知识库（代码核验）", str(personnel_sheet["B3"].value))
                self.assertIn(
                    "知识库商品：CP-KQ-YG-0001 / 参半示例商品100g标准名称 / 6970356166832",
                    personnel_sales_text,
                )
                self.assertIn("69码精确匹配；产品名模糊匹配（严格高于0.5）", personnel_sales_text)
                self.assertNotIn("0.875", personnel_sales_text)
                self.assertIn("Excel商品与知识库：", personnel_compare_text)
                self.assertIn("结算单与销售记录：商品、数量、奖励金额全部对应", personnel_compare_text)
                self.assertNotIn("↔", personnel_compare_text)
                personnel_conclusion = str(personnel_sheet["F4"].value)
                self.assertIn("置信度：中", personnel_conclusion)
                self.assertIn("商品名称模糊匹配严格高于0.5", personnel_conclusion)
                self.assertIn("无需重新提交", personnel_conclusion)
                self.assertEqual(personnel_sheet["D4"].value, "转账仅核对总额")

                display_sheet = workbook["堆头核销"]
                activity_row = _row_with_first_cell_prefix(display_sheet, "活动概况")
                attachment_summary_row = _row_with_first_cell_prefix(
                    display_sheet, "合同销售附件｜合同主基准"
                )
                attachment_detail_row = _row_with_first_cell_prefix(
                    display_sheet, "合同销售附件第1行"
                )
                detail_row = _display_detail_row(display_sheet)
                total_row = _row_with_first_cell_prefix(display_sheet, "合计金额")
                contract_text = str(display_sheet.cell(activity_row, 1).value)
                for label in (
                    "签订方：示例客户｜核销：通过",
                    "活动预算：1000元｜核销：通过",
                    "执行周期：2026-08-01至2026-08-31｜核销：通过",
                    "活动内容：示例门店开展堆头陈列活动｜核销：通过",
                    "核销方式：按店核销｜核销：通过",
                    "盖章：可见｜核销：通过",
                ):
                    self.assertIn(label, contract_text)

                attachment_contract_text = str(
                    display_sheet.cell(attachment_detail_row, 1).value
                )
                attachment_sales_text = str(
                    display_sheet.cell(attachment_detail_row, 3).value
                )
                attachment_compare_text = str(
                    display_sheet.cell(attachment_detail_row, 4).value
                )
                self.assertIn("PDF第3页", attachment_contract_text)
                self.assertIn("合同销售附件第1行", attachment_contract_text)
                self.assertIn("销售Excel第2行", attachment_sales_text)
                for label in (
                    "客户名称：精确匹配",
                    "业务日期：一致",
                    "商品编码：精确匹配",
                    "商品名称：模糊匹配（辅助项）",
                    "条形码：精确匹配",
                    "数量：一致",
                    "零售价：一致",
                    "合计金额：一致",
                    "本项结果：全部对应",
                ):
                    self.assertIn(label, attachment_compare_text)
                self.assertIn("单位：支", attachment_contract_text)
                self.assertIn("单位：支", attachment_sales_text)
                self.assertNotIn("单位：精确匹配", attachment_compare_text)
                self.assertIn(
                    "单位：仅展示，不参与对应或结论",
                    str(display_sheet.cell(attachment_summary_row, 4).value),
                )

                photo_text = str(display_sheet.cell(detail_row, 2).value)
                sales_text = str(display_sheet.cell(detail_row, 3).value)
                comparison_text = str(display_sheet.cell(detail_row, 4).value)
                self.assertIn(
                    "销售Excel不作为本店现场证据，也不与现场照片互相核对",
                    sales_text,
                )
                self.assertIn(
                    "知识库结论：置信度：高",
                    photo_text,
                )
                self.assertIn("可见文字：示例商品原名、6970356167341", photo_text)
                self.assertIn(
                    "地点核验：名称直接一致（名称直接核对；名称一致，无需调用地图MCP）",
                    photo_text,
                )
                self.assertIn("地点核验依据：", photo_text)
                self.assertIn("地点核验数据：名称直接对应，无需坐标检索", photo_text)
                self.assertIn(
                    "现场商品：示例商品原名（CP-KQ-YG-0001 / 6970356167341，"
                    "置信度：高）",
                    photo_text,
                )
                self.assertIn(
                    "合同 → 现场：置信度：高"
                    "（门店一致；日期一致；陈列符合",
                    comparison_text,
                )
                self.assertIn(
                    "已确认现场商品 → 合同商品范围：置信度：高",
                    comparison_text,
                )
                self.assertIn("陈列标准核验：符合（达到4纵陈列）", photo_text)
                self.assertIn("视觉依据：画面中可清楚数出4列纵向陈列。", photo_text)
                self.assertIn("可见纵列数：4", photo_text)
                self.assertIn("照片复用检查：未发现跨门店照片复用", photo_text)
                self.assertIn("计算：1 × 1000元", str(display_sheet.cell(detail_row, 5).value))
                self.assertIn("本店支持：1000元", str(display_sheet.cell(detail_row, 5).value))
                self.assertEqual(
                    str(display_sheet.cell(detail_row, 6).value),
                    "结论：通过\n无需重新提交",
                )
                visible_text = "\n".join(
                    str(cell.value or "")
                    for sheet in workbook.worksheets
                    for row_cells in sheet.iter_rows()
                    for cell in row_cells
                )
                self.assertNotIn("匹配结果：", visible_text)
                self.assertEqual(
                    [display_sheet.cell(3, column).value for column in range(1, 7)],
                    [
                        "合同PDF（视觉AI识别）\n堆头合同.pdf",
                        "现场照片（视觉AI识别）\n逐张显示文件名和识别内容",
                        "销售Excel（代码读取）\n堆头销售.xlsx",
                        "与合同具体对比结果",
                        "核销金额",
                        "结论 / 要重新提交什么",
                    ],
                )
                self.assertEqual(display_sheet.max_row, 8)
                self.assertEqual(display_sheet.auto_filter.ref, "A3:F8")
                self.assertTrue(str(display_sheet.cell(total_row, 1).value).startswith("合计金额"))
                for forbidden in (
                    "候选集", "SHA-256", "pHash", "审计JSON", "商品ID", "RAG召回",
                ):
                    self.assertNotIn(forbidden, visible_text)
            finally:
                workbook.close()

    def test_unrelated_global_sales_problem_does_not_pollute_store(self) -> None:
        result = deepcopy(_display_result())
        extra_sales = deepcopy(result["sales"]["records"][0])
        extra_sales.update(
            {
                "excel_row": 3,
                "product_code": "CP-UNMATCHED-0001",
                "product_name": "合同附件未列商品",
                "barcode": _test_ean13(9999),
                "quantity": 5,
                "retail_price": 20,
                "total_amount": 100,
            }
        )
        result["sales"]["records"].append(extra_sales)
        result["sales"]["total_quantity"] = 25
        result["sales"]["retail_amount"] = 498
        result["summary"]["sales_sku_count"] = 2
        result["summary"]["sales_quantity"] = 25
        result["summary"]["sales_retail_amount"] = 498
        result["summary"]["contract_attachment_sales_problem_count"] = 1
        result["contract_attachment_sales_reconciliation"] = (
            _contract_attachment_sales_reconciliation(
                result["contract"], result["sales"]
            )
        )
        self.assertEqual(
            result["contract_attachment_sales_reconciliation"]["unmatched_sales_rows"],
            [3],
        )
        validate_json(
            result,
            Path(__file__).resolve().parents[1] / "contracts" / "audit-result.schema.json",
        )

        with tempfile.TemporaryDirectory() as temporary:
            target = Path(temporary) / "result.xlsx"
            create_combined_report([result], target)
            workbook = load_workbook(target, data_only=False)
            try:
                sheet = workbook["堆头核销"]
                attachment_row = _row_with_first_cell_prefix(
                    sheet, "合同销售附件第1行"
                )
                unmatched_row = _row_with_first_cell_prefix(
                    sheet, "合同销售附件未找到Excel第3行"
                )
                store_row = _display_detail_row(sheet)
                self.assertIn("销售Excel第2行", str(sheet.cell(attachment_row, 3).value))
                unmatched_text = "\n".join(
                    str(sheet.cell(unmatched_row, column).value or "")
                    for column in range(1, 7)
                )
                self.assertIn("Excel第3行", unmatched_text)
                self.assertIn("合同附件未列商品", unmatched_text)
                store_sales_text = str(sheet.cell(store_row, 3).value)
                self.assertNotIn("Excel第2行", store_sales_text)
                self.assertNotIn("Excel第3行", store_sales_text)
                self.assertIn("销售Excel不作为本店现场证据", store_sales_text)
                self.assertEqual(
                    str(sheet.cell(store_row, 6).value),
                    "结论：通过\n无需重新提交",
                )
            finally:
                workbook.close()

    def test_attachment_code_mismatch_names_exact_sources_and_field(self) -> None:
        result = deepcopy(_display_result())
        result["sales"]["records"][0]["product_code"] = "CP-WRONG-0001"
        result["contract_attachment_sales_reconciliation"] = (
            _contract_attachment_sales_reconciliation(
                result["contract"], result["sales"]
            )
        )
        reconciliation = result["contract_attachment_sales_reconciliation"]
        result["summary"]["contract_attachment_sales_problem_count"] = 1
        self.assertEqual(reconciliation["status"], "fail")
        self.assertEqual(reconciliation["records"][0]["sales_excel_row"], 2)
        self.assertEqual(
            {
                item["field"]: item["status"]
                for item in reconciliation["records"][0]["field_comparisons"]
            }["product_code"],
            "mismatch",
        )
        validate_json(
            result,
            Path(__file__).resolve().parents[1] / "contracts" / "audit-result.schema.json",
        )

        with tempfile.TemporaryDirectory() as temporary:
            target = Path(temporary) / "result.xlsx"
            create_combined_report([result], target)
            workbook = load_workbook(target, data_only=False)
            try:
                sheet = workbook["堆头核销"]
                detail_row = _row_with_first_cell_prefix(
                    sheet, "合同销售附件第1行"
                )
                sales_text = str(sheet.cell(detail_row, 3).value)
                self.assertIn(
                    "销售Excel第2行\n客户名称：示例客户\n业务日期：2026-08\n"
                    "商品编码：CP-WRONG-0001\n商品名称：示例商品原名",
                    sales_text,
                )
                comparison_text = str(sheet.cell(detail_row, 4).value)
                self.assertIn(
                    "商品编码：不一致",
                    comparison_text,
                )
                conclusion = str(sheet.cell(detail_row, 6).value)
                self.assertIn("PDF第3页附件第1行与Excel第2行", conclusion)
                self.assertIn("合同附件与销售Excel的商品编码", conclusion)
                self.assertNotIn("审计JSON", sales_text + comparison_text + conclusion)
            finally:
                workbook.close()

    def test_missing_valid_sales_row_does_not_rewrite_an_unrelated_row(self) -> None:
        result = deepcopy(_display_result())
        result["sales"]["records"][0].update(
            {
                "excel_row": 6,
                "product_code": "CP-OTHER-0439",
                "product_name": "完全无关商品",
                "barcode": _test_ean13(439),
            }
        )
        result["contract_attachment_sales_reconciliation"] = (
            _contract_attachment_sales_reconciliation(
                result["contract"], result["sales"]
            )
        )
        reconciliation = result["contract_attachment_sales_reconciliation"]
        result["summary"]["contract_attachment_sales_problem_count"] = 2
        self.assertEqual(reconciliation["status"], "fail")
        self.assertEqual(reconciliation["unmatched_contract_rows"], [1])
        self.assertEqual(reconciliation["unmatched_sales_rows"], [6])
        self.assertIsNone(reconciliation["records"][0]["sales_excel_row"])
        validate_json(
            result,
            Path(__file__).resolve().parents[1] / "contracts" / "audit-result.schema.json",
        )

        with tempfile.TemporaryDirectory() as temporary:
            target = Path(temporary) / "result.xlsx"
            create_combined_report([result], target)
            workbook = load_workbook(target, data_only=False)
            try:
                sheet = workbook["堆头核销"]
                missing_row = _row_with_first_cell_prefix(
                    sheet, "合同销售附件第1行"
                )
                missing_text = "\n".join(
                    str(sheet.cell(missing_row, column).value or "")
                    for column in range(1, 7)
                )
                self.assertIn("销售Excel：未找到唯一对应行", missing_text)
                self.assertIn("补齐对应Excel行", missing_text)
                self.assertNotIn("Excel第6行", missing_text)
                unrelated_row = _row_with_first_cell_prefix(
                    sheet, "合同销售附件未找到Excel第6行"
                )
                unrelated_text = "\n".join(
                    str(sheet.cell(unrelated_row, column).value or "")
                    for column in range(1, 7)
                )
                self.assertIn("完全无关商品", unrelated_text)
                self.assertIn("核实Excel第6行对应的合同附件行", unrelated_text)
            finally:
                workbook.close()

    def test_fuzzy_attachment_sales_name_is_auxiliary_and_still_passes(self) -> None:
        result = deepcopy(_display_result())
        result["sales"]["records"][0]["product_name"] = "示例商品"
        result["contract_attachment_sales_reconciliation"] = (
            _contract_attachment_sales_reconciliation(
                result["contract"], result["sales"]
            )
        )
        reconciliation = result["contract_attachment_sales_reconciliation"]
        self.assertEqual(reconciliation["status"], "pass")
        self.assertEqual(reconciliation["records"][0]["confidence"], "high")
        self.assertEqual(
            {
                item["field"]: item["status"]
                for item in reconciliation["records"][0]["field_comparisons"]
            }["product_name"],
            "fuzzy",
        )
        validate_json(
            result,
            Path(__file__).resolve().parents[1] / "contracts" / "audit-result.schema.json",
        )

        with tempfile.TemporaryDirectory() as temporary:
            target = Path(temporary) / "result.xlsx"
            create_combined_report([result], target)
            workbook = load_workbook(target, data_only=False)
            try:
                sheet = workbook["堆头核销"]
                detail_row = _row_with_first_cell_prefix(
                    sheet, "合同销售附件第1行"
                )
                self.assertIn(
                    "商品名称：模糊匹配",
                    str(sheet.cell(detail_row, 4).value),
                )
                self.assertEqual(
                    str(sheet.cell(detail_row, 6).value),
                    "置信度：高\n无需重新提交",
                )
                store_row = _display_detail_row(sheet)
                self.assertIn(
                    "销售Excel不作为本店现场证据",
                    str(sheet.cell(store_row, 3).value),
                )
            finally:
                workbook.close()

    def test_every_attachment_sales_row_is_shown_without_top_n_truncation(self) -> None:
        result = deepcopy(_display_result())
        attachment_records: list[dict] = []
        sales_records: list[dict] = []
        knowledge_records: list[dict] = []
        base_knowledge = result["contract_attachment_product_knowledge"]["records"][0]
        for line_no in range(1, 6):
            excel_row = line_no + 1
            quantity = line_no + 10
            retail_price = line_no + 20
            total_amount = quantity * retail_price
            code = f"CP-MULTI-{line_no:04d}"
            name = f"合同附件测试商品{line_no}"
            barcode = _test_ean13(100 + line_no)
            attachment_records.append(
                {
                    "line_no": line_no,
                    "source_page": 3,
                    "customer_name": "示例客户",
                    "business_date": "2026-08",
                    "product_code": code,
                    "product_name": name,
                    "barcode_69": barcode,
                    "unit": "支",
                    "quantity": quantity,
                    "retail_price": retail_price,
                    "total_amount": total_amount,
                }
            )
            sales_records.append(
                {
                    "excel_row": excel_row,
                    "customer_name": "示例客户",
                    "period_text": "2026-08",
                    "product_code": code,
                    "product_name": name,
                    "barcode": barcode,
                    "unit": "支",
                    "quantity": quantity,
                    "retail_price": retail_price,
                    "total_amount": total_amount,
                }
            )
            knowledge_record = deepcopy(base_knowledge)
            knowledge_record.update(
                {
                    "excel_row": line_no,
                    "contract_line_no": line_no,
                    "source_page": 3,
                    "source_product_code": code,
                    "source_product_name": name,
                    "source_barcode_69": barcode,
                    "source_quantity": quantity,
                    "knowledge_product_id": f"multi-{line_no}",
                    "knowledge_product_code": code,
                    "knowledge_product_name": name,
                    "knowledge_barcode_69": barcode,
                }
            )
            knowledge_records.append(knowledge_record)

        total_quantity = sum(item["quantity"] for item in attachment_records)
        total_amount = sum(item["total_amount"] for item in attachment_records)
        result["contract"]["sales_attachment"].update(
            {
                "source_pages": [3],
                "records": attachment_records,
                "total_quantity": total_quantity,
                "total_amount": total_amount,
            }
        )
        result["sales"].update(
            {
                "records": sales_records,
                "total_quantity": total_quantity,
                "retail_amount": total_amount,
            }
        )
        result["contract_attachment_product_knowledge"].update(
            {
                "record_count": 5,
                "matched_count": 5,
                "problem_count": 0,
                "problem_rows": [],
                "records": knowledge_records,
                "basis": "5/5行通过共享商品知识库核验。",
            }
        )
        result["summary"].update(
            {
                "sales_sku_count": 5,
                "sales_quantity": total_quantity,
                "sales_retail_amount": total_amount,
                "contract_attachment_row_count": 5,
                "contract_attachment_quantity": total_quantity,
                "contract_attachment_amount": total_amount,
            }
        )
        result["contract_attachment_sales_reconciliation"] = (
            _contract_attachment_sales_reconciliation(
                result["contract"], result["sales"]
            )
        )
        self.assertEqual(result["contract_attachment_sales_reconciliation"]["status"], "pass")
        self.assertEqual(
            len(result["contract_attachment_sales_reconciliation"]["records"]),
            5,
        )
        validate_json(
            result,
            Path(__file__).resolve().parents[1] / "contracts" / "audit-result.schema.json",
        )

        with tempfile.TemporaryDirectory() as temporary:
            target = Path(temporary) / "result.xlsx"
            create_combined_report([result], target)
            workbook = load_workbook(target, data_only=False)
            try:
                sheet = workbook["堆头核销"]
                for line_no in range(1, 6):
                    detail_row = _row_with_first_cell_prefix(
                        sheet, f"合同销售附件第{line_no}行"
                    )
                    self.assertIn(
                        f"销售Excel第{line_no + 1}行",
                        str(sheet.cell(detail_row, 3).value),
                    )
                store_row = _display_detail_row(sheet)
                store_sales_text = str(sheet.cell(store_row, 3).value)
                for excel_row in range(2, 7):
                    self.assertNotIn(f"Excel第{excel_row}行", store_sales_text)
                visible_text = "\n".join(
                    str(cell.value or "")
                    for row_cells in sheet.iter_rows()
                    for cell in row_cells
                )
                self.assertNotIn("其余", visible_text)
                self.assertNotIn("审计JSON", visible_text)
                self.assertEqual(sheet.max_row, 12)
            finally:
                workbook.close()


if __name__ == "__main__":
    unittest.main()
