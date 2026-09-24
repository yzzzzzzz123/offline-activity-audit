from copy import deepcopy
import io
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
from zipfile import ZipFile

import openpyxl
import pymupdf
import jsonschema

from audit_core.common import AuditError
from audit_core.material_intake import prepare_material_diagnosis
from audit_core.pdf_evidence import (audit_decision, audit_schema, classification_decision, calculate, document_map,
                                     validate_documents, AUDIT_SCHEMA, CLASSIFICATION_SCHEMA)
from audit_core.pdf_materials import PdfEvidenceProvider, prepare_units
from audit_core.pdf_policy import MATERIALS, MATERIAL_LABELS, requirements, rules
from audit_core.extraction_chain import render_prompt
from tests.pdf_test_support import bundle, classification, flags_for, image_bytes, unknown_checks


class PdfPolicyTests(unittest.TestCase):
    def test_entry_duplicate_cannot_cite_photo_from_another_claim(self):
        flags = flags_for("entry_fee")
        documents = [{"unit_id": "claim-a-photo", "facts": ["水印地址：示例路18号"]}]
        evidence = unknown_checks("entry_fee", flags)
        check = next(c for c in evidence["checks"] if c["rule_id"] == "entry_duplicate")
        check.update(status="pass", reason="本单可读照片中没有重复门店地址。",
                     source_ids=["claim-a-photo"])
        result = audit_decision("entry_fee", flags, evidence, documents)
        self.assertEqual(next(c for c in result["checks"] if c["rule_id"] == "entry_duplicate")["status"], "pass")
        check.update(status="fail", reason="引用其他核销单或历史照片判定地址重复。",
                     source_ids=["claim-a-photo", "claim-b-photo"])
        with self.assertRaises(AuditError):
            audit_decision("entry_fee", flags, evidence, documents)

    def test_optional_background_does_not_change_matching_for_any_type(self):
        for scenario in MATERIALS:
            with self.subTest(scenario=scenario):
                documents = [{"unit_id": "u", "facts": ["本次核销资料"]}]
                value = classification(scenario, ["u"])
                baseline = classification_decision(value, documents)
                value["background_materials"] = [{"description": "活动申请流程介绍",
                    "reason": "仅介绍申请流转，不含本次活动的核销依据", "source_ids": ["background"]}]
                value["source_ids"] = ["u", "background"]
                documents.append({"unit_id": "background", "facts": ["活动前由业务部门管理BI额度和申请顺序"]})
                actual = classification_decision(value, documents)
                self.assertEqual(actual["scenario"], baseline["scenario"])
                self.assertTrue(actual["matched"])
                self.assertEqual(actual["reasons"], [])
                self.assertEqual(actual["classification"]["background_materials"], value["background_materials"])

    def test_optional_background_does_not_waive_missing_or_extra_audit_materials(self):
        for scenario in MATERIALS:
            required_id = next(r.id for r in requirements(scenario) if r.when == "always")
            for change in ({"missing": [required_id]}, {"extra": [{"description": "其他费用的独立核销凭证",
                          "reason": "不是背景，确属清单外业务资料", "source_ids": ["u"]}]}):
                with self.subTest(scenario=scenario, change=change):
                    value = classification(scenario, ["u"], **change)
                    value["background_materials"] = [{"description": "BI管理流程",
                        "reason": "仅解释业务流程", "source_ids": ["background"]}]
                    value["source_ids"] = ["u", "background"]
                    documents = [{"unit_id": "u"}, {"unit_id": "background"}]
                    self.assertFalse(classification_decision(value, documents)["matched"])

    def test_background_sources_must_be_read_and_cannot_hide_unaccounted_files(self):
        for background_ids in ([], ["unread"], ["background"]):
            with self.subTest(background_ids=background_ids):
                value = classification("poster_material", ["u"])
                value["background_materials"] = [{"description": "BI流程", "reason": "仅作背景", "source_ids": background_ids}]
                value["source_ids"] = ["u", "background", "unaccounted"]
                with self.assertRaises(AuditError):
                    classification_decision(value, [{"unit_id": uid} for uid in value["source_ids"]])

    def test_all_eight_types_audit_without_process_history_checks(self):
        for scenario in MATERIALS:
            with self.subTest(scenario=scenario):
                documents = [{"unit_id": "u", "facts": ["本次核销资料，不含BI、提交、一审或补交记录"]}]
                selected = classification_decision(classification(scenario, ["u"]), documents)
                self.assertTrue(selected["matched"])
                flags = flags_for(scenario)
                result = audit_decision(scenario, flags, unknown_checks(scenario, flags), documents)
                checks = {c["rule_id"]: c for c in result["checks"]}
                self.assertTrue({"bi_quota", "submission_deadline", "review_deadline"}.isdisjoint(checks))
                self.assertTrue({"scene_authenticity", "fraud"}.isdisjoint(checks))
                for process_term in ("BI", "一审", "提交时限", "补交期限"):
                    self.assertFalse(any(process_term in c["reason"] for c in checks.values()))
                for process_facts in (["财务一审已超过2个月", "提交日期已晚于活动结束后3个月"],
                                      ["尚未进行财务一审", "本次为首次提交"]):
                    with_history = [*documents, {"unit_id": "history", "facts": process_facts}]
                    self.assertEqual(audit_decision(scenario, flags, unknown_checks(scenario, flags), with_history), result)
                if scenario == "personnel_incentive":
                    self.assertTrue({"application_quantity_price", "reward_unit_price", "claim_ceiling"} <= checks.keys())

    def test_process_history_findings_are_rejected_by_schema_and_decision_for_all_types(self):
        for scenario in MATERIALS:
            for rule_id in ("bi_quota", "submission_deadline", "review_deadline"):
                with self.subTest(scenario=scenario, rule_id=rule_id):
                    flags = flags_for(scenario)
                    evidence = unknown_checks(scenario, flags)
                    evidence["checks"][0].update(rule_id=rule_id, status="fail",
                                                 reason="流程已超期或未提供流程记录", source_ids=[], comparisons=[])
                    with self.assertRaises(jsonschema.ValidationError):
                        jsonschema.validate(evidence, audit_schema(scenario))
                    with self.assertRaises(AuditError):
                        audit_decision(scenario, flags, evidence, [])

    def test_pos_target_routes_with_pos_excel_settlement_and_contract_only(self):
        ids = ["stamped-pos", "excel", "settlement", "contract"]
        value = classification("pos_target_incentive", ids)
        groups = dict(zip(("pos", "pos_excel", "settlement", "promotion_contract"), ids))
        self.assertEqual({r.id for r in requirements("pos_target_incentive")}, set(groups))
        for material in value["materials"]:
            material["source_ids"] = [groups[material["id"]]]
        result = classification_decision(value, [{"unit_id": uid} for uid in ids])
        self.assertTrue(result["matched"])
        self.assertEqual(result["scenario"], "pos_target_incentive")
        self.assertEqual(len(value["material_matches"]), 8)
        value["material_matches"] = [m for m in value["material_matches"] if m["scenario"] != "pos_target_incentive"]
        with self.assertRaises(AuditError):
            classification_decision(value, [{"unit_id": uid} for uid in ids])

    def test_pos_target_requires_excel_and_cannot_use_training_or_giveaway_exceptions(self):
        documents = [{"unit_id": "u"}]
        value = classification("pos_target_incentive", ["u"], missing=["pos_excel"])
        self.assertFalse(classification_decision(value, documents)["matched"])

        for training in (None, False, True):
            value = classification("pos_target_incentive", ["u"], flags=flags_for("pos_target_incentive", training=training))
            if training:
                with self.assertRaisesRegex(AuditError, "只能用于人员激励"):
                    classification_decision(value, documents)
            else:
                material = next(m for m in value["materials"] if m["id"] == "pos_excel")
                material.update(state="not_applicable", source_ids=[])
                with self.assertRaisesRegex(AuditError, "不能自行免交"):
                    classification_decision(value, documents)

        extra = {"description": "赠送规则及搭赠现场照片", "reason": "独立POS达标激励清单未列此项", "source_ids": ["u"]}
        value = classification("pos_target_incentive", ["u"], extra=[extra])
        self.assertFalse(classification_decision(value, documents)["matched"])
        value = classification("giveaway_promotion", ["u"], missing=["giveaway_evidence"])
        self.assertFalse(classification_decision(value, documents)["matched"])

    def test_pos_target_numeric_rules_recalculate_without_giveaway_or_payment_checks(self):
        flags = flags_for("pos_target_incentive")
        for rule_id, left, right, operator, conclusion in (
            ("reward_unit_price", "2", "1", "eq", "issues_found"),
            ("settlement_pos_quantity", "12", "10", "eq", "issues_found"),
            ("claim_ceiling", "101", "100", "le", "issues_found"),
        ):
            with self.subTest(rule=rule_id):
                documents = [{"unit_id": "s", "facts": [f"结算原值={left}"]},
                             {"unit_id": "p", "facts": [f"对照原值={right}"]}]
                evidence = unknown_checks("pos_target_incentive", flags)
                check = next(c for c in evidence["checks"] if c["rule_id"] == rule_id)
                check.update(status="pass", reason="模型误判为一致", source_ids=["s", "p"], comparisons=[{
                    "label": "本次SKU对应数值", "left": {"unit_id": "s", "quote": f"结算原值={left}", "number": left},
                    "value_kind": {"reward_unit_price": "unit_price", "settlement_pos_quantity": "quantity", "claim_ceiling": "amount"}[rule_id],
                    "right": [{"unit_id": "p", "quote": f"对照原值={right}", "number": right}],
                    "operation": "value", "operator": operator}])
                result = audit_decision("pos_target_incentive", flags, evidence, documents)
                checks = {c["rule_id"]: c for c in result["checks"]}
                self.assertEqual(checks[rule_id]["status"], "fail")
                self.assertEqual(result["summary"]["conclusion"], conclusion)
                self.assertNotIn("pos_seal", checks)
                self.assertNotIn("pos_excel_present", checks)
                self.assertTrue({"giveaway_evidence", "gift_rule_present", "payment_evidence"}.isdisjoint(checks))

    def test_type_rejection_requires_comparing_all_eight_material_lists_first(self):
        documents = [{"unit_id": "u"}]
        for candidates in ([], ["promotional_display", "poster_material"]):
            with self.subTest(candidates=candidates):
                value = classification("poster_material", ["u"], candidates=candidates)
                self.assertFalse(classification_decision(value, documents)["matched"])
                value.pop("material_matches")
                with self.assertRaises(AuditError):
                    classification_decision(value, documents)

    def test_all_type_comparisons_must_cover_their_materials_and_real_sources(self):
        for mutation in (
            lambda v: v["material_matches"].pop(),
            lambda v: v["material_matches"].__setitem__(0, deepcopy(v["material_matches"][1])),
            lambda v: v["material_matches"][0]["materials"].pop(),
            lambda v: v["material_matches"][0]["materials"][0].update(source_ids=["unread"]),
        ):
            with self.subTest(mutation=mutation):
                value = classification("poster_material", ["u"])
                mutation(value)
                with self.assertRaises(AuditError):
                    classification_decision(value, [{"unit_id": "u"}])

    def test_expense_description_cannot_add_a_candidate_without_material_support(self):
        value = classification("promotional_display", ["u"])
        value["candidate_scenarios"].append("poster_material")
        value["reason"] = "费用项目名称写了场地费与物料制作费"
        value["materials"] = []
        with self.assertRaisesRegex(AuditError, "逐类比对结果一致"):
            classification_decision(value, [{"unit_id": "u"}])

    def test_each_supported_type_needs_its_own_read_sources(self):
        value = classification("poster_material", ["u"], candidates=["promotional_display", "poster_material"])
        selected = next(m for m in value["material_matches"] if m["scenario"] == "poster_material")
        selected["source_ids"] = []
        with self.assertRaisesRegex(AuditError, "引用本次材料来源"):
            classification_decision(value, [{"unit_id": "u"}])

    def test_selected_materials_cannot_disagree_with_the_compared_checklist(self):
        value = classification("poster_material", ["u"])
        value["materials"] = deepcopy(value["materials"])
        value["materials"][0].update(state="missing", source_ids=[])
        with self.assertRaisesRegex(AuditError, "已保存的清单比对结果"):
            classification_decision(value, [{"unit_id": "u"}])

    def test_missing_required_pos_prevents_type_identification(self):
        flags = flags_for("giveaway_promotion", uses_pos=False)
        value = classification("giveaway_promotion", ["u"], flags=flags, missing=["pos", "pos_excel"])
        route = classification_decision(value, [{"unit_id": "u"}])
        self.assertFalse(route["matched"])
        self.assertIsNone(route["scenario"])
        self.assertIn(MATERIAL_LABELS["pos"],
                      "；".join(route["reasons"]))
        match = next(m for m in value["material_matches"] if m["scenario"] == "giveaway_promotion")
        match["materials"][0].update(state="present", source_ids=["u"])
        with self.assertRaises(AuditError):
            classification_decision(value, [{"unit_id": "u"}])

    def test_training_can_submit_optional_pos_without_invalidating_the_exemption(self):
        flags = flags_for("personnel_incentive", training=True, uses_pos=True)
        value = classification("personnel_incentive", ["u"], flags=flags)
        for material in value["materials"]:
            if material["id"] in {"pos", "pos_excel"}:
                material.update(state="present", source_ids=["u"])
        self.assertTrue(classification_decision(value, [{"unit_id": "u"}])["matched"])

    def test_training_exemption_does_not_spread_to_other_types(self):
        flags = flags_for("personnel_incentive", training=True, uses_pos=False)
        value = classification("personnel_incentive", ["u"], flags=flags)
        result = classification_decision(value, [{"unit_id": "u"}])
        self.assertTrue(result["matched"])
        self.assertEqual([m["id"] for m in value["materials"] if m["state"] == "not_applicable"], ["pos", "pos_excel", "staff_photos"])
        value["flags"]["training"] = False
        with self.assertRaises(AuditError):
            classification_decision(value, [{"unit_id": "u"}])
        value = classification("price_difference_support", ["u"],
                               flags=flags_for("personnel_incentive", training=True, uses_pos=False))
        with self.assertRaisesRegex(AuditError, "只能用于人员激励"):
            classification_decision(value, [{"unit_id": "u"}])

    def test_inapplicable_conditional_item_cannot_be_silently_absorbed(self):
        for scenario, material_id in (("promotional_display", "atrium_agreement"),
                                      ("personnel_incentive", "staff_photos"),
                                      ("entry_fee", "entry_system")):
            with self.subTest(scenario=scenario):
                value = classification(scenario, ["u"])
                material = next(m for m in value["materials"] if m["id"] == material_id)
                material.update(state="present", source_ids=["u"])
                with self.assertRaisesRegex(AuditError, "不适用的条件资料"):
                    classification_decision(value, [{"unit_id": "u"}])

    def test_conditional_materials_must_be_confirmed_and_missing_names_are_complete(self):
        flags = flags_for("promotional_display", atrium=None)
        value = classification("promotional_display", ["u"], flags=flags, missing=["settlement", "promotion_contract"])
        route = classification_decision(value, [{"unit_id": "u"}])
        self.assertFalse(route["matched"])
        reasons = "；".join(route["reasons"])
        for text in ("商场入场协议", "结算单", "促销合同"):
            self.assertIn(text, reasons)
        match = next(m for m in value["material_matches"] if m["scenario"] == "promotional_display")
        next(m for m in match["materials"] if m["id"] == "atrium_agreement")["state"] = "not_applicable"
        with self.assertRaises(AuditError):
            classification_decision(value, [{"unit_id": "u"}])

    def test_price_proof_is_required_even_when_full_reduction_flag_is_unknown_or_false(self):
        for full_reduction in (None, False, True):
            with self.subTest(full_reduction=full_reduction):
                flags = flags_for("price_difference_support", full_reduction=full_reduction)
                value = classification("price_difference_support", ["u"], flags=flags,
                                       missing=["full_reduction_evidence"])
                self.assertFalse(classification_decision(value, [{"unit_id": "u"}])["matched"])
                materials = next(m for m in value["material_matches"] if m["scenario"] == "price_difference_support")["materials"]
                evidence = unknown_checks("price_difference_support", flags)
                check = next(c for c in evidence["checks"] if c["rule_id"] == "full_reduction_evidence")
                check.update(status="pass", reason="误判为已有证明", source_ids=["u"])
                result = audit_decision("price_difference_support", flags, evidence,
                                        [{"unit_id": "u", "facts": []}], materials)
                actual = next(c for c in result["checks"] if c["rule_id"] == "full_reduction_evidence")
                self.assertEqual(actual["status"], "fail")
                self.assertIn("满减返还活动存在证明", actual["reason"])
                self.assertEqual(len(result["checks"]), len(rules("price_difference_support")))
                check.update(status="fail", reason="未提供对应证明", source_ids=[])
                audit_decision("price_difference_support", flags, evidence,
                               [{"unit_id": "u", "facts": []}], materials)

    def test_each_required_item_must_be_present_for_all_eight_types(self):
        for scenario in MATERIALS:
            flags = flags_for(scenario, **{"atrium": True, "cvs_otc": True})
            flags["uses_pos"] = True
            if scenario == "personnel_incentive":
                flags["temporary_staff"] = True
            for requirement in requirements(scenario):
                with self.subTest(scenario=scenario, material=requirement.id):
                    value = classification(scenario, ["u"], flags=flags, missing=[requirement.id])
                    route = classification_decision(value, [{"unit_id": "u"}])
                    self.assertFalse(route["matched"])
                    self.assertIsNone(route["scenario"])
                    self.assertIn(MATERIAL_LABELS[requirement.id], "；".join(route["reasons"]))

    def test_extra_material_on_an_already_covered_page_blocks_every_type(self):
        extra = {"description": "独立运输单", "reason": "该类PDF资料清单未列此项", "source_ids": ["u"]}
        for scenario in MATERIALS:
            with self.subTest(scenario=scenario):
                value = classification(scenario, ["u"], extra=[extra])
                route = classification_decision(value, [{"unit_id": "u"}])
                self.assertFalse(route["matched"])
                self.assertIn("多出的资料：独立运输单", "；".join(route["reasons"]))
                match = next(m for m in value["material_matches"] if m["scenario"] == scenario)
                match["supported"] = True
                with self.assertRaisesRegex(AuditError, "不多不少"):
                    classification_decision(value, [{"unit_id": "u"}])

    def test_missing_item_cannot_be_overridden_by_model_type_claim(self):
        value = classification("poster_material", ["u"], missing=["invoice"])
        match = next(m for m in value["material_matches"] if m["scenario"] == "poster_material")
        match["supported"] = True
        with self.assertRaisesRegex(AuditError, "不多不少"):
            classification_decision(value, [{"unit_id": "u"}])

    def test_all_read_sources_must_be_accounted_for_in_every_comparison(self):
        value = classification("poster_material", ["u"])
        with self.assertRaisesRegex(AuditError, "全部已读来源"):
            classification_decision(value, [{"unit_id": "u"}, {"unit_id": "extra-page"}])
        value["material_matches"][0]["extra_materials"][0]["source_ids"] = ["unread"]
        with self.assertRaisesRegex(AuditError, "未读取"):
            classification_decision(value, [{"unit_id": "u"}])

    def test_many_pages_and_photos_do_not_add_material_items(self):
        ids = ["contract-page-1", "contract-page-2", "photo-1", "photo-2", "invoice", "settlement"]
        value = classification("poster_material", ids)
        groups = {"promotion_contract": ids[:2], "finished_photos": ids[2:4], "invoice": [ids[4]], "settlement": [ids[5]]}
        for material in value["materials"]:
            material["source_ids"] = groups.get(material["id"], [])
        self.assertTrue(classification_decision(value, [{"unit_id": uid} for uid in ids])["matched"])
        self.assertEqual(sum(m["state"] == "present" for m in value["materials"]), 4)
        self.assertTrue(all(m["state"] == "not_applicable" for m in value["materials"] if m["id"] in {"pos", "pos_excel"}))

    def test_unknown_material_condition_prevents_match_even_with_no_missing_items(self):
        value = classification("entry_fee", ["u"], flags=flags_for("entry_fee", cvs_otc=None))
        route = classification_decision(value, [{"unit_id": "u"}])
        self.assertFalse(route["matched"])
        self.assertIn("尚不能确认的资料或适用情况", "；".join(route["reasons"]))

    def test_giveaway_and_online_price_evidence_allow_one_source_for_multiple_roles(self):
        for scenario, overrides in (("giveaway_promotion", {"photo_evidence": False}),
                                    ("price_difference_support", {"online": True, "full_reduction": True})):
            with self.subTest(scenario=scenario):
                value = classification(scenario, ["u"], flags=flags_for(scenario, **overrides))
                self.assertTrue(classification_decision(value, [{"unit_id": "u"}])["matched"])
                self.assertIn("pos_excel", [r.id for r in requirements(scenario)])

    def test_giveaway_mixed_store_proofs_share_one_item_and_keep_content_findings(self):
        scenario = "giveaway_promotion"
        flags = flags_for(scenario, photo_evidence=True)
        ids = [r.id for r in requirements(scenario) if r.id != "giveaway_evidence"]
        ids += ["photo-a", "receipt-b"]
        documents = [{"unit_id": uid, "roles": [uid], "facts": ["本次活动资料"], "limitations": []}
                     for uid in ids]
        by_id = {d["unit_id"]: d for d in documents}
        by_id["promotion_contract"]["facts"] = ["本次搭赠活动门店为甲店、乙店。"]
        by_id["photo-a"]["facts"] = ["甲店搭赠现场照片，水印显示拍摄时间和甲店地址。"]
        value = classification(scenario, ids, flags=flags)
        selected = next(m for m in value["material_matches"] if m["scenario"] == scenario)
        value["material_matches"] = [selected]
        for material in value["materials"]:
            material["source_ids"] = (["photo-a", "receipt-b"] if material["id"] == "giveaway_evidence"
                                      else [material["id"]])
        for status, receipt, reason in (
            ("pass", "乙店小票显示购买产品、搭赠产品、购买时间和门店。", "甲店照片与乙店小票分别符合要求，共同覆盖本次两家门店。"),
            ("fail", "乙店小票没有显示搭赠产品。", "乙店小票没有显示赠送的商品。"),
            ("unknown", "乙店小票的搭赠产品文字模糊。", "乙店小票看不清赠送的商品名称。"),
        ):
            with self.subTest(status=status):
                by_id["receipt-b"]["facts"] = [receipt]
                gate = classification_decision(value, documents, selected_scenario=scenario)
                self.assertTrue(gate["matched"])
                self.assertEqual(selected["extra_materials"], [])
                self.assertEqual(sum(m["id"] == "giveaway_evidence" for m in value["materials"]), 1)
                evidence = unknown_checks(scenario, flags)
                observed = next(c for c in evidence["checks"] if c["rule_id"] == "giveaway_evidence")
                observed.update(status=status, reason=reason,
                                source_ids=["promotion_contract", "photo-a", "receipt-b"])
                result = audit_decision(scenario, flags, evidence, documents, value["materials"])
                actual = next(c for c in result["checks"] if c["rule_id"] == "giveaway_evidence")
                self.assertEqual(actual["status"], status)
                self.assertEqual(actual["reason"], reason)
                self.assertEqual(actual["source_ids"], ["promotion_contract", "photo-a", "receipt-b"])

    def test_poster_does_not_require_pos_and_entry_has_no_activity_deadline(self):
        flags = flags_for("poster_material")
        self.assertTrue(classification_decision(classification("poster_material", ["u"], flags=flags), [{"unit_id": "u"}])["matched"])
        self.assertEqual(next(r.when for r in requirements("poster_material") if r.id == "pos"), "uses_pos")
        flags = flags_for("entry_fee")
        evidence = unknown_checks("entry_fee", flags)
        result = audit_decision("entry_fee", flags, evidence, [])
        by_id = {c["rule_id"]: c for c in result["checks"]}
        self.assertNotIn("submission_deadline", by_id)
        self.assertNotIn("review_deadline", by_id)
        self.assertEqual(by_id["entry_watermark"]["status"], "unknown")
        self.assertNotIn("entry_system", by_id)

    def test_numeric_mismatch_is_recalculated_and_zero_requires_confirmed_failure(self):
        flags = flags_for("price_difference_support")
        documents = [{"unit_id": "s", "facts": ["SKU甲结算数量=12"]}, {"unit_id": "p", "facts": ["SKU甲POS数量=10"]}]
        evidence = unknown_checks("price_difference_support", flags)
        check = next(c for c in evidence["checks"] if c["rule_id"] == "settlement_pos_quantity")
        check.update(status="pass", source_ids=["s", "p"], reason="模型认为相符", comparisons=[{
            "label": "SKU甲销售数量", "value_kind": "quantity", "left": {"unit_id": "s", "quote": "SKU甲结算数量=12", "number": "12"},
            "right": [{"unit_id": "p", "quote": "SKU甲POS数量=10", "number": "10"}], "operation": "value", "operator": "eq"}])
        result = audit_decision("price_difference_support", flags, evidence, documents)
        self.assertEqual(result["summary"]["conclusion"], "issues_found")
        changed = next(c for c in result["checks"] if c["rule_id"] == "settlement_pos_quantity")
        self.assertEqual(changed["status"], "fail")
        self.assertIn("实际 12，对照 10", changed["reason"])
        self.assertEqual(check["status"], "pass", "原始模型证据保持不变")
        check["comparisons"][0]["operator"] = "ge"
        with self.assertRaisesRegex(AuditError, "比较方向"):
            audit_decision("price_difference_support", flags, evidence, documents)
        check.update(status="unknown", comparisons=[])
        self.assertEqual(audit_decision("price_difference_support", flags, evidence, documents)["summary"]["conclusion"], "issues_found")

    def test_comparison_cannot_fabricate_values_sources_or_fraud(self):
        documents = [{"unit_id": "u", "facts": ["R2: C1=10 | C2=20"]}]
        comparison = {"label": "比较", "value_kind": "quantity", "left": {"unit_id": "u", "quote": "C1=10", "number": "10"},
                      "right": [{"unit_id": "u", "quote": "C2=20", "number": "20"}], "operation": "value", "operator": "eq"}
        for mutation in ({"number": "1"}, {"number": "999"}, {"quote": "伪造金额=10"}, {"unit_id": "unknown"}):
            value = deepcopy(comparison)
            value["left"].update(mutation)
            with self.assertRaises(AuditError):
                calculate(value, document_map(documents))
        flags = flags_for("poster_material")
        evidence = unknown_checks("poster_material", flags)
        check = next(c for c in evidence["checks"] if c["rule_id"] == "finished_photo")
        check.update(status="unknown", source_ids=["u"], comparisons=[comparison])
        with self.assertRaisesRegex(AuditError, "非数值"):
            audit_decision("poster_material", flags, evidence, documents)

    def test_extra_missing_duplicate_rules_and_unread_sources_are_execution_errors(self):
        flags = flags_for("entry_fee")
        for action in (lambda e: e["checks"].pop(), lambda e: e["checks"].append(deepcopy(e["checks"][0])),
                       lambda e: e["checks"][0].update(rule_id="database_ean"),
                       lambda e: e["checks"][1].update(status="pass", source_ids=["unknown"])):
            evidence = unknown_checks("entry_fee", flags)
            action(evidence)
            with self.assertRaises(AuditError):
                audit_decision("entry_fee", flags, evidence, [])

    def test_all_generation_enums_declare_types(self):
        def visit(node):
            if isinstance(node, dict):
                if "enum" in node:
                    self.assertIn("type", node)
                for value in node.values(): visit(value)
            elif isinstance(node, list):
                for value in node: visit(value)
        visit(AUDIT_SCHEMA)
        visit(CLASSIFICATION_SCHEMA)
        for scenario in MATERIALS:
            visit(audit_schema(scenario))

    def test_every_type_rejects_added_omitted_duplicate_and_cross_type_audit_points(self):
        all_ids = {rule.id for scenario in MATERIALS for rule in rules(scenario)}
        for scenario in MATERIALS:
            flags = flags_for(scenario)
            expected = unknown_checks(scenario, flags)
            ids = {check["rule_id"] for check in expected["checks"]}
            for action in (
                lambda value: value["checks"].pop(),
                lambda value: value["checks"].append(deepcopy(value["checks"][0])),
                lambda value: value["checks"].__setitem__(1, deepcopy(value["checks"][0])),
                lambda value: value["checks"][0].update(rule_id="database_ean"),
                lambda value: value["checks"][0].update(rule_id=sorted(all_ids - ids)[0]),
            ):
                with self.subTest(scenario=scenario, mutation=action):
                    value = deepcopy(expected)
                    action(value)
                    with self.assertRaisesRegex(AuditError, "全部且仅覆盖"):
                        audit_decision(scenario, flags, value, [])
            # A valid full checklist still runs every item, regardless of order.
            expected["checks"].reverse()
            result = audit_decision(scenario, flags, expected, [])
            self.assertEqual({check["rule_id"] for check in result["checks"]}, ids)

    def test_model_output_contract_is_closed_for_each_audit_type(self):
        for scenario in MATERIALS:
            schema = audit_schema(scenario)
            validator = jsonschema.Draft202012Validator(schema)
            value = unknown_checks(scenario, flags_for(scenario))
            validator.validate(value)
            value["checks"][0]["rule_id"] = "map_distance"
            with self.assertRaises(jsonschema.ValidationError):
                validator.validate(value)
            value = unknown_checks(scenario, flags_for(scenario))
            value["checks"].pop()
            with self.assertRaises(jsonschema.ValidationError):
                validator.validate(value)

    def test_optional_pos_types_use_actual_pos_presence_and_cannot_hide_missing_pair(self):
        for scenario in ("poster_material", "promotional_display"):
            for uses_pos in (None, True):
                with self.subTest(scenario=scenario, uses_pos=uses_pos):
                    value = classification(scenario, ["u"], flags=flags_for(scenario, uses_pos=uses_pos))
                    route = classification_decision(value, [{"unit_id": "u"}])
                    self.assertEqual(route["matched"], uses_pos is True)
                    if uses_pos:
                        missing = classification(scenario, ["u"], flags=flags_for(scenario, uses_pos=True), missing=["pos_excel"])
                        self.assertFalse(classification_decision(missing, [{"unit_id": "u"}])["matched"])
class PdfReadingTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)

    def prepare(self, files):
        bundle(self.root / "input", files=files)
        intake = prepare_material_diagnosis(self.root / "input", self.root / "sources", original_error="", classify_by_content=True)
        return prepare_units(intake["archives"][0], self.root / "units")

    def test_classification_prompt_is_unchanged_by_archive_and_source_names(self):
        prompts = []
        for index, name in enumerate(("其他.zip", "进场费.zip")):
            temporary = self.root / str(index)
            temporary.mkdir()
            provider = PdfEvidenceProvider("gpt-6-astra", "medium")
            facts = ["单元格记录：同一活动包含场地和搭建明细；照片与票据按内容分别核对"]
            units = [{"unit_id": "u", "file_id": "f", "source_file": name + "/人员激励.xlsx",
                      "locator": name, "image": None, "native_facts": facts, "limitations": []}]

            def fake_call(root, skill, schema, prompt, **kwargs):
                prompts.append(render_prompt(prompt))
                value = classification("poster_material", ["u"])
                kwargs["validator"](value)
                return value

            with patch.object(provider, "_call", side_effect=fake_call):
                result = provider.classify({"archive_id": "a001", "source_archive": name, "units": units}, temporary)
            self.assertEqual({m["scenario"] for m in result["classification"]["material_matches"]}, set(MATERIALS))
        self.assertEqual(prompts[0], prompts[1])
        self.assertNotIn("人员激励.xlsx", prompts[0])
        self.assertNotIn('"audit_points"', prompts[0])

    def test_vector_pdf_has_all_complete_pages_and_nested_photos_keep_original_names(self):
        doc = pymupdf.open()
        for text in ("Contract first page", "Contract second page"):
            page = doc.new_page()
            page.insert_text((50, 60), text)
        payload = doc.tobytes()
        doc.close()
        nested = io.BytesIO()
        with ZipFile(nested, "w") as archive:
            archive.writestr("门店照片.png", image_bytes())
        units = self.prepare({"合同.pdf": payload, "返图.zip": nested.getvalue()})
        pages = [u for u in units if u["source_file"] == "合同.pdf"]
        self.assertEqual(len(pages), 2)
        self.assertTrue(all(u["image"].is_file() for u in units))
        self.assertIn("second page", pages[1]["native_facts"][0])
        self.assertTrue(any(u["source_file"] == "返图.zip/门店照片.png" for u in units))

    def test_spreadsheet_keeps_hidden_sheets_raw_formulas_dates_and_embedded_images(self):
        book = openpyxl.Workbook()
        sheet = book.active
        sheet.append(["SKU", "数量", "单价", "金额"])
        sheet.append(["069001", 3, 10, "=B2*C2"])
        hidden = book.create_sheet("隐藏数据")
        hidden.sheet_state = "hidden"
        hidden["A1"] = 123
        from openpyxl.drawing.image import Image
        sheet.add_image(Image(io.BytesIO(image_bytes())), "F1")
        out = io.BytesIO()
        book.save(out)
        book.close()
        units = self.prepare({"资料.xlsx": out.getvalue()})
        facts = "\n".join(f for u in units for f in u["native_facts"])
        for value in ("069001", "=B2*C2", "无缓存结果", "隐藏数据!R1: C1=123"):
            self.assertIn(value, facts)
        self.assertEqual(sum(bool(u["image"]) for u in units), 1)
        data = next(u for u in units if u["native_facts"])
        document = {"unit_id": data["unit_id"], "roles": ["spreadsheet"], "facts": ["改写数值"], "limitations": []}
        with self.assertRaisesRegex(AuditError, "不得由模型改写"):
            validate_documents({"documents": [document]}, [data])

    def test_permission_and_renderer_failures_are_execution_errors(self):
        book = openpyxl.Workbook()
        out = io.BytesIO()
        book.save(out)
        book.close()
        with patch("audit_core.pdf_materials.openpyxl.load_workbook", side_effect=PermissionError("denied")):
            with self.assertRaises(PermissionError):
                self.prepare({"表格.xlsx": out.getvalue()})


if __name__ == "__main__":
    unittest.main()
