from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .common import AuditError


FONT_NAME = "Microsoft YaHei"
NAVY = "17365D"
BLUE = "D9EAF7"
LIGHT_BLUE = "EAF3F8"
GREEN = "E2F0D9"
YELLOW = "FFF2CC"
RED = "FCE4D6"
GRAY = "E7E6E6"
WHITE = "FFFFFF"
INPUT_BLUE = "0000FF"
FORMULA_BLACK = "000000"
LINK_GREEN = "008000"
THIN = Side(style="thin", color="B7C9D6")


def _font(*, bold: bool = False, color: str = "000000", size: int = 10) -> Font:
    return Font(name=FONT_NAME, size=size, bold=bold, color=color)


def _style_title(ws: Any, title: str, last_col: int, note: str | None = None) -> int:
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col)
    cell = ws.cell(1, 1, title)
    cell.font = _font(bold=True, color=WHITE, size=16)
    cell.fill = PatternFill("solid", fgColor=NAVY)
    cell.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 28
    if note:
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last_col)
        note_cell = ws.cell(2, 1, note)
        note_cell.font = _font(color="666666", size=9)
        note_cell.fill = PatternFill("solid", fgColor=LIGHT_BLUE)
        note_cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.row_dimensions[2].height = 32
        return 3
    return 2


def _header(ws: Any, row: int, headers: list[str]) -> None:
    for col, value in enumerate(headers, 1):
        cell = ws.cell(row, col, value)
        cell.font = _font(bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
    ws.row_dimensions[row].height = 30


def _body_cell(cell: Any, *, formula: bool = False, linked: bool = False) -> None:
    color = LINK_GREEN if linked else FORMULA_BLACK if formula else INPUT_BLUE
    cell.font = _font(color=color)
    cell.border = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
    cell.alignment = Alignment(vertical="top", wrap_text=True)


def _status_fill(cell: Any, value: str) -> None:
    normalized = value.lower()
    if any(word in normalized for word in ("pass", "通过", "匹配")) and not any(
        word in normalized for word in ("partial", "conditional", "补证", "未验证", "异常")
    ):
        cell.fill = PatternFill("solid", fgColor=GREEN)
    elif any(word in normalized for word in ("partial", "conditional", "补证", "未验证", "暂挂")):
        cell.fill = PatternFill("solid", fgColor=YELLOW)
    elif any(word in normalized for word in ("review", "异常", "不一致", "mismatch", "失败")):
        cell.fill = PatternFill("solid", fgColor=RED)


def _set_widths(ws: Any, widths: dict[int, float]) -> None:
    for col, width in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = width


def _freeze_filter(ws: Any, header_row: int, last_row: int, last_col: int) -> None:
    ws.freeze_panes = ws.cell(header_row + 1, 1)
    ws.auto_filter.ref = f"A{header_row}:{get_column_letter(last_col)}{last_row}"
    ws.sheet_view.showGridLines = False


def _write_exceptions(wb: Workbook, title: str, result: dict[str, Any]) -> str:
    ws = wb.create_sheet(title)
    header_row = _style_title(
        ws,
        f"{result['case_name']}：异常与补证清单",
        6,
        "高风险影响数量/金额或核心证据；中风险表示口径、身份或真实性证据仍需说明。",
    )
    headers = ["序号", "等级", "异常代码", "发现", "影响", "补证/处理建议"]
    _header(ws, header_row, headers)
    for index, item in enumerate(result["exceptions"], 1):
        row = header_row + index
        values = [
            index,
            item["severity"],
            item["code"],
            item["message"] + (f"\n来源：{item['source']}" if item.get("source") else ""),
            item["impact"],
            item["suggestion"],
        ]
        for col, value in enumerate(values, 1):
            cell = ws.cell(row, col, value)
            _body_cell(cell)
            if item["severity"] == "high":
                cell.fill = PatternFill("solid", fgColor=RED)
            elif item["severity"] == "medium":
                cell.fill = PatternFill("solid", fgColor=YELLOW)
    last_row = max(header_row + len(result["exceptions"]), header_row)
    _freeze_filter(ws, header_row, last_row, len(headers))
    _set_widths(ws, {1: 8, 2: 10, 3: 34, 4: 60, 5: 42, 6: 52})
    return title


def _add_personnel_sheets(wb: Workbook, result: dict[str, Any]) -> dict[str, str]:
    sku_title = "人员激励-SKU逐项"
    ws = wb.create_sheet(sku_title)
    header_row = _style_title(
        ws,
        "人员激励：Excel销售数量与结算单逐SKU 1:1勾稽",
        18,
        "先按结算行映射唯一条码，再汇总7家门店Excel来源行；数量、单价、行金额和支持金额分别计算。",
    )
    headers = [
        "结算行",
        "结算单奖励单品",
        "对应条码",
        "Excel商品名称",
        "Excel来源行",
        "Excel销售数量",
        "结算销售数量",
        "数量差异",
        "数量状态",
        "奖励单价",
        "Excel应付金额",
        "结算应付金额",
        "结算行奖励金额",
        "行金额差异",
        "金额状态",
        "支持数量",
        "支持金额",
        "映射依据/置信度",
    ]
    _header(ws, header_row, headers)
    start = header_row + 1
    for offset, item in enumerate(result["sku_reconciliation"]):
        row = start + offset
        values = [
            item["line_no"],
            item["settlement_product_name"],
            item["mapped_barcode"],
            item["excel_product_name"],
            ",".join(str(value) for value in item["excel_rows"]),
            item["excel_quantity"],
            item["settlement_quantity"],
            f'=IF(OR(F{row}="",G{row}=""),"",F{row}-G{row})',
            f'=IF(H{row}="","未匹配",IF(H{row}=0,"匹配","异常"))',
            item["unit_reward"],
            f'=IF(F{row}="","",F{row}*J{row})',
            f'=G{row}*J{row}',
            item["settlement_reward_amount"],
            f'=M{row}-L{row}',
            f'=IF(N{row}=0,"匹配","异常")',
            f'=IF(OR(F{row}="",G{row}=""),0,MAX(0,MIN(F{row},G{row})))',
            f'=P{row}*J{row}',
            f"{item['mapping_basis']} / {item['mapping_confidence']}",
        ]
        for col, value in enumerate(values, 1):
            cell = ws.cell(row, col, value)
            formula = isinstance(value, str) and value.startswith("=")
            _body_cell(cell, formula=formula)
            if col in {9, 15}:
                _status_fill(cell, str(item["quantity_status"] if col == 9 else item["amount_status"]))
    total_row = start + len(result["sku_reconciliation"])
    ws.cell(total_row, 1, "合计")
    for col in (6, 7, 8, 10, 11, 12, 13, 14, 16, 17):
        letter = get_column_letter(col)
        ws.cell(total_row, col, f"=SUM({letter}{start}:{letter}{total_row - 1})")
    for col in range(1, 19):
        cell = ws.cell(total_row, col)
        _body_cell(cell, formula=col in {6, 7, 8, 10, 11, 12, 13, 14, 16, 17})
        cell.font = _font(bold=True, color=FORMULA_BLACK if col != 1 else INPUT_BLUE)
        cell.fill = PatternFill("solid", fgColor=BLUE)
    _freeze_filter(ws, header_row, total_row, len(headers))
    _set_widths(
        ws,
        {1: 9, 2: 30, 3: 17, 4: 42, 5: 18, 6: 14, 7: 14, 8: 12, 9: 12,
         10: 11, 11: 15, 12: 15, 13: 17, 14: 13, 15: 12, 16: 12, 17: 13, 18: 30},
    )

    trace_title = "人员激励-门店SKU明细"
    trace_ws = wb.create_sheet(trace_title)
    trace_header = _style_title(
        trace_ws,
        "人员激励：70条门店×SKU销售来源明细",
        11,
        "每一行均保留原Excel行号，并回指结算单行；这是SKU汇总数量的底层证据链。",
    )
    trace_headers = [
        "Excel行",
        "期间",
        "门店",
        "条码",
        "Excel商品名称",
        "销售数量",
        "销售单价",
        "结算行",
        "结算单品",
        "奖励单价",
        "该行应付奖励",
    ]
    _header(trace_ws, trace_header, trace_headers)
    for offset, item in enumerate(result["sales_trace"]):
        row = trace_header + 1 + offset
        values = [
            item["excel_row"], item["period_text"], item["store_name"], item["barcode"],
            item["product_name"], item["quantity"], item["unit_price"], item["settlement_line_no"],
            item["settlement_product_name"], item["unit_reward"],
            f'=IF(OR(F{row}="",J{row}=""),"",F{row}*J{row})',
        ]
        for col, value in enumerate(values, 1):
            _body_cell(trace_ws.cell(row, col, value), formula=isinstance(value, str) and value.startswith("="))
    trace_total = trace_header + len(result["sales_trace"]) + 1
    trace_ws.cell(trace_total, 1, "合计")
    for col in (6, 11):
        letter = get_column_letter(col)
        trace_ws.cell(trace_total, col, f"=SUM({letter}{trace_header + 1}:{letter}{trace_total - 1})")
    for col in range(1, 12):
        _body_cell(trace_ws.cell(trace_total, col), formula=col in {6, 11})
        trace_ws.cell(trace_total, col).fill = PatternFill("solid", fgColor=BLUE)
    _freeze_filter(trace_ws, trace_header, trace_total, len(trace_headers))
    _set_widths(trace_ws, {1: 10, 2: 14, 3: 18, 4: 17, 5: 45, 6: 12, 7: 12, 8: 10, 9: 30, 10: 12, 11: 15})

    transfer_title = "人员激励-门店转账"
    transfer_ws = wb.create_sheet(transfer_title)
    transfer_header = _style_title(
        transfer_ws,
        "人员激励：门店应付与去重转账勾稽",
        9,
        "金额集合一致不等于收款身份一致；截图未显示门店/收款人时单独标记为身份未验证。",
    )
    transfer_headers = [
        "门店", "Excel销售数量", "应付奖励", "转账ID", "去重转账金额", "金额差异", "匹配依据", "身份状态", "结论"
    ]
    _header(transfer_ws, transfer_header, transfer_headers)
    transfer_start = transfer_header + 1
    for offset, item in enumerate(result["store_transfer_reconciliation"]):
        row = transfer_start + offset
        values = [
            item["store_name"], item["excel_quantity"], item["expected_reward_amount"],
            item["transfer_id"], item["transfer_amount"],
            f'=IF(OR(C{row}="",E{row}=""),"",E{row}-C{row})',
            item["match_basis"], item["identity_status"], item["status"],
        ]
        for col, value in enumerate(values, 1):
            cell = transfer_ws.cell(row, col, value)
            _body_cell(cell, formula=isinstance(value, str) and value.startswith("="))
            if col == 9:
                _status_fill(cell, str(value))
    transfer_total = transfer_start + len(result["store_transfer_reconciliation"])
    transfer_ws.cell(transfer_total, 1, "合计")
    for col in (2, 3, 5, 6):
        letter = get_column_letter(col)
        transfer_ws.cell(transfer_total, col, f"=SUM({letter}{transfer_start}:{letter}{transfer_total - 1})")
    for col in range(1, 10):
        _body_cell(transfer_ws.cell(transfer_total, col), formula=col in {2, 3, 5, 6})
        transfer_ws.cell(transfer_total, col).fill = PatternFill("solid", fgColor=BLUE)
    _freeze_filter(transfer_ws, transfer_header, transfer_total, len(transfer_headers))
    _set_widths(transfer_ws, {1: 18, 2: 15, 3: 14, 4: 13, 5: 16, 6: 12, 7: 22, 8: 25, 9: 30})

    voucher_title = "人员激励-转账凭证"
    voucher_ws = wb.create_sheet(voucher_title)
    voucher_header = _style_title(
        voucher_ws,
        "人员激励：转账截图去重明细",
        9,
        "同一笔转账的发送方与接收方气泡合并为一个transfer_id，同时保留画面出现次数。",
    )
    voucher_headers = ["转账ID", "金额", "画面出现次数", "身份可见", "日期可见", "门店/收款人", "来源文件", "去重依据", "备注"]
    _header(voucher_ws, voucher_header, voucher_headers)
    for offset, item in enumerate(result["transfer_evidence"]):
        row = voucher_header + 1 + offset
        identity = item.get("store_name") or item.get("recipient_name")
        values = [
            item["transfer_id"], item["amount"], item["occurrence_count"],
            "是" if item["identity_visible"] else "否", "是" if item["event_date_visible"] else "否",
            identity, "、".join(item["source_files"]), item["dedup_basis"], "；".join(item.get("notes") or []),
        ]
        for col, value in enumerate(values, 1):
            _body_cell(voucher_ws.cell(row, col, value))
    voucher_last = voucher_header + len(result["transfer_evidence"])
    _freeze_filter(voucher_ws, voucher_header, max(voucher_last, voucher_header), len(voucher_headers))
    _set_widths(voucher_ws, {1: 13, 2: 13, 3: 15, 4: 12, 5: 12, 6: 20, 7: 46, 8: 55, 9: 35})

    exception_title = _write_exceptions(wb, "人员激励-异常补证", result)
    return {
        "detail_sheet": sku_title,
        "detail_total_row": str(total_row),
        "supported_cell": f"'{sku_title}'!Q{total_row}",
        "transfer_total_cell": f"'{transfer_title}'!E{transfer_total}",
        "exception_sheet": exception_title,
    }


def _add_display_sheets(wb: Workbook, result: dict[str, Any]) -> dict[str, str]:
    store_title = "堆头-逐店核验"
    ws = wb.create_sheet(store_title)
    header_row = _style_title(
        ws,
        "堆头核销：合同门店与照片逐店证据核验",
        14,
        "单店费用仅在照片、期间、门店、陈列和重复检查全部通过时计入支持金额。",
    )
    headers = [
        "序号",
        "合同门店",
        "照片文件",
        "照片数",
        "可见日期",
        "可见地点/水印",
        "期间",
        "门店匹配",
        "陈列标准",
        "重复检查",
        "结论",
        "支持金额",
        "风险说明",
        "补证建议",
    ]
    _header(ws, header_row, headers)
    start = header_row + 1
    fee = result["summary"]["fee_per_store"]
    status_labels = {"pass": "通过", "supplement": "补证"}
    period_labels = {"match": "是", "mismatch": "否", "unverifiable": "无法判断"}
    store_labels = {"exact": "是", "compatible": "基本匹配", "mismatch": "不一致", "filename_only": "弱（仅文件名）"}
    display_labels = {"pass": "是", "fail": "否", "uncertain": "无法判断"}
    duplicate_labels = {"none": "未发现", "exact": "完全重复", "possible": "疑似重复", "unverifiable": "无法判断"}
    for offset, item in enumerate(result["store_reconciliation"]):
        row = start + offset
        status = status_labels[item["status"]]
        values = [
            item["store_line_no"],
            item["contract_store_name"],
            "、".join(item["photo_files"]),
            item["photo_count"],
            item["visible_date"],
            item["visible_location"],
            period_labels[item["period_match"]],
            store_labels[item["store_match"]],
            display_labels[item["display_match"]],
            duplicate_labels[item["duplicate_check"]],
            status,
            f'=IF(K{row}="通过",{fee},0)',
            "；".join(item["risk_notes"]),
            "；".join(item["supplement_advice"]),
        ]
        for col, value in enumerate(values, 1):
            cell = ws.cell(row, col, value)
            _body_cell(cell, formula=isinstance(value, str) and value.startswith("="))
            if col == 11:
                _status_fill(cell, status)
    total_row = start + len(result["store_reconciliation"])
    ws.cell(total_row, 1, "合计")
    ws.cell(total_row, 4, f"=SUM(D{start}:D{total_row - 1})")
    ws.cell(total_row, 12, f"=SUM(L{start}:L{total_row - 1})")
    for col in range(1, 15):
        _body_cell(ws.cell(total_row, col), formula=col in {4, 12})
        ws.cell(total_row, col).fill = PatternFill("solid", fgColor=BLUE)
    _freeze_filter(ws, header_row, total_row, len(headers))
    _set_widths(ws, {1: 8, 2: 31, 3: 46, 4: 10, 5: 13, 6: 36, 7: 13, 8: 15, 9: 15, 10: 14, 11: 12, 12: 14, 13: 50, 14: 52})

    sales_title = "堆头-销售检查"
    sales_ws = wb.create_sheet(sales_title)
    sales_header = _style_title(
        sales_ws,
        "堆头核销：销售Excel完整性与明细检查",
        8,
        "经销商汇总销售可证明活动期总体销售，但不能替代20家合同门店的逐店执行证据。",
    )
    metrics = [
        ("客户", "、".join(result["sales"]["customers"]), "销售Excel"),
        ("业务期间", "、".join(result["sales"]["period_values"]), "销售Excel"),
        ("SKU数量", result["sales"]["sku_count"], "有效明细行"),
        ("销售数量", result["sales"]["total_quantity"], "数量列求和"),
        ("零售金额", result["sales"]["retail_amount"], "缓存值可用时求和"),
        ("门店级销售", "否", "当前文件为经销商汇总"),
        ("外部链接公式", len(result["sales"]["external_formula_cells"]), "含[工作簿]引用的公式单元格"),
    ]
    _header(sales_ws, sales_header, ["检查项", "结果", "来源/说明"])
    for offset, (name, value, note) in enumerate(metrics):
        row = sales_header + 1 + offset
        for col, cell_value in enumerate((name, value, note), 1):
            _body_cell(sales_ws.cell(row, col, cell_value))
    detail_header = sales_header + len(metrics) + 2
    detail_headers = ["Excel行", "客户", "期间", "商品编码", "条码", "商品名称", "数量", "零售金额"]
    _header(sales_ws, detail_header, detail_headers)
    for offset, item in enumerate(result["sales"]["records"]):
        row = detail_header + 1 + offset
        values = [
            item["excel_row"], item["customer_name"], item["period_text"], item["product_code"],
            item["barcode"], item["product_name"], item["quantity"], item["total_amount"],
        ]
        for col, value in enumerate(values, 1):
            _body_cell(sales_ws.cell(row, col, value))
    detail_total = detail_header + len(result["sales"]["records"]) + 1
    sales_ws.cell(detail_total, 1, "合计")
    for col in (7, 8):
        letter = get_column_letter(col)
        sales_ws.cell(detail_total, col, f"=SUM({letter}{detail_header + 1}:{letter}{detail_total - 1})")
    for col in range(1, 9):
        _body_cell(sales_ws.cell(detail_total, col), formula=col in {7, 8})
        sales_ws.cell(detail_total, col).fill = PatternFill("solid", fgColor=BLUE)
    sales_ws.freeze_panes = f"A{detail_header + 1}"
    sales_ws.auto_filter.ref = f"A{detail_header}:H{detail_total}"
    sales_ws.sheet_view.showGridLines = False
    _set_widths(sales_ws, {1: 14, 2: 34, 3: 25, 4: 16, 5: 17, 6: 55, 7: 12, 8: 15})

    photo_title = "堆头-照片文件检查"
    photo_ws = wb.create_sheet(photo_title)
    photo_header = _style_title(
        photo_ws,
        "堆头核销：照片文件哈希与原始元数据",
        9,
        "SHA-256用于完全重复检测，dHash用于快速识别视觉相同候选；EXIF/GPS缺失需在风险中保留。",
    )
    photo_headers = ["文件名", "尺寸", "文件大小", "SHA-256", "dHash", "EXIF时间", "GPS", "绑定门店", "状态"]
    _header(photo_ws, photo_header, photo_headers)
    bound_store: dict[str, str] = {}
    for item in result["store_reconciliation"]:
        for file_name in item["photo_files"]:
            bound_store[Path(file_name).name] = item["contract_store_name"]
    for offset, item in enumerate(result["photo_inventory"]["files"]):
        row = photo_header + 1 + offset
        values = [
            item["file_name"], f"{item['width']}×{item['height']}", item["size"], item["sha256"],
            item["dhash"], item["exif_datetime"], "有" if item["gps_present"] else "无",
            bound_store.get(item["file_name"]), "已绑定" if item["file_name"] in bound_store else "未绑定",
        ]
        for col, value in enumerate(values, 1):
            _body_cell(photo_ws.cell(row, col, value))
    photo_last = photo_header + len(result["photo_inventory"]["files"])
    _freeze_filter(photo_ws, photo_header, max(photo_last, photo_header), len(photo_headers))
    _set_widths(photo_ws, {1: 48, 2: 15, 3: 14, 4: 68, 5: 20, 6: 23, 7: 10, 8: 35, 9: 12})

    exception_title = _write_exceptions(wb, "堆头-异常补证", result)
    return {
        "detail_sheet": store_title,
        "supported_cell": f"'{store_title}'!L{total_row}",
        "exception_sheet": exception_title,
    }


def create_combined_report(results: list[dict[str, Any]], output_path: str | Path) -> Path:
    if not results:
        raise AuditError("Cannot create a report without audit results")
    workbook = Workbook()
    workbook.remove(workbook.active)
    summary_ws = workbook.create_sheet("核销总览")
    personnel_refs: dict[str, str] | None = None
    display_refs: dict[str, str] | None = None
    for result in results:
        if result["scenario"] == "personnel_incentive":
            if personnel_refs is not None:
                raise AuditError("Combined report currently supports one personnel incentive case")
            personnel_refs = _add_personnel_sheets(workbook, result)
        elif result["scenario"] == "promotional_display":
            if display_refs is not None:
                raise AuditError("Combined report currently supports one promotional display case")
            display_refs = _add_display_sheets(workbook, result)
        else:
            raise AuditError(f"Unsupported report scenario: {result['scenario']}")

    header_row = _style_title(
        summary_ws,
        "AI 核销试审结果（详细证据链版）",
        10,
        "数量与门店证据先逐项核验，再汇总金额；AI结论用于辅助复核，不替代最终审批。",
    )
    headers = [
        "项目", "场景", "结算/申请金额", "证据支持金额", "建议通过金额", "暂挂金额",
        "AI结论", "逐项核验覆盖", "关键异常", "明细入口",
    ]
    _header(summary_ws, header_row, headers)
    start = header_row + 1
    result_by_scenario = {item["scenario"]: item for item in results}
    row = start
    if "personnel_incentive" in result_by_scenario and personnel_refs:
        result = result_by_scenario["personnel_incentive"]
        summary = result["summary"]
        values = [
            result["case_name"], "人员激励", summary["claimed_amount"],
            f"=MIN({personnel_refs['supported_cell']},{personnel_refs['transfer_total_cell']})",
            f"=MIN(C{row},D{row})", f"=C{row}-E{row}", summary["conclusion"],
            f"{summary['quantity_match_count']}/{summary['settlement_line_count']}个SKU数量匹配；{summary['excel_detail_row_count']}条门店SKU来源行",
            "；".join(item["message"] for item in result["exceptions"][:3]),
            "查看“人员激励-SKU逐项”",
        ]
        for col, value in enumerate(values, 1):
            cell = summary_ws.cell(row, col, value)
            formula = isinstance(value, str) and value.startswith("=")
            _body_cell(cell, formula=formula, linked=formula and col == 4)
            if col == 7:
                _status_fill(cell, str(value))
        summary_ws.cell(row, 10).hyperlink = "#'人员激励-SKU逐项'!A1"
        summary_ws.cell(row, 10).style = "Hyperlink"
        row += 1
    if "promotional_display" in result_by_scenario and display_refs:
        result = result_by_scenario["promotional_display"]
        summary = result["summary"]
        values = [
            result["case_name"], "堆头陈列", summary["claimed_amount"],
            f"={display_refs['supported_cell']}", f"=MIN(C{row},D{row})", f"=C{row}-E{row}",
            summary["conclusion"],
            f"{summary['passed_store_count']}/{summary['contract_store_count']}家通过；{summary['supplement_store_count']}家补证",
            "；".join(item["message"] for item in result["exceptions"][:3]),
            "查看“堆头-逐店核验”",
        ]
        for col, value in enumerate(values, 1):
            cell = summary_ws.cell(row, col, value)
            formula = isinstance(value, str) and value.startswith("=")
            _body_cell(cell, formula=formula, linked=formula and col == 4)
            if col == 7:
                _status_fill(cell, str(value))
        summary_ws.cell(row, 10).hyperlink = "#'堆头-逐店核验'!A1"
        summary_ws.cell(row, 10).style = "Hyperlink"
        row += 1
    total_row = row
    summary_ws.cell(total_row, 1, "合计")
    for col in (3, 4, 5, 6):
        letter = get_column_letter(col)
        summary_ws.cell(total_row, col, f"=SUM({letter}{start}:{letter}{total_row - 1})")
    summary_ws.cell(total_row, 7, f'=IF(F{total_row}=0,"全部通过","存在暂挂/补证")')
    for col in range(1, 11):
        _body_cell(summary_ws.cell(total_row, col), formula=col in {3, 4, 5, 6, 7})
        summary_ws.cell(total_row, col).fill = PatternFill("solid", fgColor=BLUE)
        summary_ws.cell(total_row, col).font = _font(bold=True, color=FORMULA_BLACK)
    _freeze_filter(summary_ws, header_row, total_row, len(headers))
    _set_widths(summary_ws, {1: 28, 2: 14, 3: 17, 4: 17, 5: 17, 6: 15, 7: 20, 8: 42, 9: 72, 10: 28})

    summary_ws.sheet_properties.tabColor = NAVY
    workbook._sheets.remove(summary_ws)
    workbook._sheets.insert(0, summary_ws)
    workbook.calculation.calcMode = "auto"
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    target = Path(output_path).resolve()
    target.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(target)
    workbook.close()
    return target


def workbook_formula_inventory(path: str | Path) -> dict[str, Any]:
    workbook = load_workbook(path, read_only=True, data_only=False)
    formulas: list[str] = []
    formula_errors: list[str] = []
    known_sheets = set(workbook.sheetnames)
    for ws in workbook.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    location = f"{ws.title}!{cell.coordinate}"
                    formulas.append(location)
                    if any(token in cell.value for token in ("#REF!", "#DIV/0!", "#VALUE!", "#NAME?")):
                        formula_errors.append(location)
                    for quoted in __import__("re").findall(r"'([^']+)'!", cell.value):
                        if quoted not in known_sheets:
                            formula_errors.append(location)
    cached_errors: list[str] = []
    workbook.close()
    cached_book = load_workbook(path, read_only=True, data_only=True)
    for ws in cached_book.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and any(
                    token in cell.value
                    for token in ("#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A")
                ):
                    cached_errors.append(f"{ws.title}!{cell.coordinate}")
    cached_book.close()
    return {
        "formula_count": len(formulas),
        "locations": formulas,
        "static_formula_errors": sorted(set(formula_errors)),
        "cached_formula_errors": sorted(set(cached_errors)),
        "status": "success" if not formula_errors and not cached_errors else "errors_found",
    }
