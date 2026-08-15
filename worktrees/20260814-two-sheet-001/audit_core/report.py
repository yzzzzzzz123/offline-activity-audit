from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .analysis_dimensions import build_analysis_dimension_rows
from .common import AuditError


FONT_NAME = "Microsoft YaHei"
NAVY = "17365D"
BLUE = "D9EAF7"
LIGHT_BLUE = "EAF3F8"
GREEN = "E2F0D9"
YELLOW = "FFF2CC"
RED = "FCE4D6"
GRAY = "E7E6E6"
WHITE = "FFFFFF"
INPUT_BLUE = "0000FF"
FORMULA_BLACK = "000000"
LINK_GREEN = "008000"
THIN = Side(style="thin", color="B7C9D6")
AI_REVIEW_SCOPE = [
    (
        "文件分类与图像预处理",
        2,
        3000,
        "完成收件、分析、待复核、驳回、完成通知及群状态调整",
    ),
    (
        "照片真实性风险辅助检测",
        3,
        4500,
        "读取EXIF/GPS/水印并进行pHash重复和跨活动复用检测",
    ),
    (
        "现场/陈列照片多模态识别",
        7,
        10500,
        "按固定活动规则识别门店、商品、排面、物料、位置和照片质量",
    ),
    (
        "发票及票据OCR结构化",
        4,
        6000,
        "提取发票代码、号码、金额、日期、购销方等关键字段",
    ),
    (
        "合同、费用明细及规则抽取",
        4,
        6000,
        "抽取合同金额、活动周期、Excel费用明细和核销政策字段",
    ),
    (
        "跨材料一致性与规则判断",
        6,
        9000,
        "对照片、票据、费用、预算、合同和门店信息进行字段及计算比对",
    ),
    (
        "结果汇总、异常重试与用量日志",
        2,
        3000,
        "汇总通过/异常/待补件结论，处理失败重试并记录模型调用量",
    ),
]


def _font(*, bold: bool = False, color: str = "000000", size: int = 10) -> Font:
    return Font(name=FONT_NAME, size=size, bold=bold, color=color)


def _style_title(ws: Any, title: str, last_col: int, note: str | None = None) -> int:
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col)
    cell = ws.cell(1, 1, title)
    cell.font = _font(bold=True, color=WHITE, size=16)
    cell.fill = PatternFill("solid", fgColor=NAVY)
    cell.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 28
    if note:
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last_col)
        note_cell = ws.cell(2, 1, note)
        note_cell.font = _font(color="666666", size=9)
        note_cell.fill = PatternFill("solid", fgColor=LIGHT_BLUE)
        note_cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.row_dimensions[2].height = 42
        return 3
    return 2


def _header(ws: Any, row: int, headers: list[str]) -> None:
    for col, value in enumerate(headers, 1):
        cell = ws.cell(row, col, value)
        cell.font = _font(bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
    ws.row_dimensions[row].height = 30


def _body_cell(cell: Any, *, formula: bool = False, linked: bool = False) -> None:
    color = LINK_GREEN if linked else FORMULA_BLACK if formula else INPUT_BLUE
    cell.font = _font(color=color)
    cell.border = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
    cell.alignment = Alignment(vertical="top", wrap_text=True)


def _status_fill(cell: Any, value: str) -> None:
    normalized = value.lower()
    if "本案未触发" in normalized or "不适用" in normalized:
        cell.fill = PatternFill("solid", fgColor=GRAY)
        return
    if any(word in normalized for word in ("pass", "通过", "匹配")) and not any(
        word in normalized
        for word in ("partial", "conditional", "部分", "补证", "未验证", "异常", "证据不足")
    ):
        cell.fill = PatternFill("solid", fgColor=GREEN)
    elif any(
        word in normalized
        for word in ("partial", "conditional", "部分", "补证", "未验证", "暂挂", "证据不足")
    ):
        cell.fill = PatternFill("solid", fgColor=YELLOW)
    elif any(word in normalized for word in ("review", "异常", "不一致", "mismatch", "失败")):
        cell.fill = PatternFill("solid", fgColor=RED)


def _set_widths(ws: Any, widths: dict[int, float]) -> None:
    for col, width in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = width


def _freeze_filter(ws: Any, header_row: int, last_row: int, last_col: int) -> None:
    ws.freeze_panes = ws.cell(header_row + 1, 1)
    ws.auto_filter.ref = f"A{header_row}:{get_column_letter(last_col)}{last_row}"
    ws.sheet_view.showGridLines = False


def _sheet_link(cell: Any, sheet_name: str) -> None:
    cell.hyperlink = f"#'{sheet_name}'!A1"
    cell.style = "Hyperlink"


def _scenario_result(
    results: list[dict[str, Any]],
    scenario: str,
) -> dict[str, Any] | None:
    return next((item for item in results if item.get("scenario") == scenario), None)


def _manifest_scenario(
    input_manifest: dict[str, Any] | None,
    scenario: str,
) -> dict[str, Any]:
    if not isinstance(input_manifest, dict):
        return {}
    extractions = input_manifest.get("extractions")
    if not isinstance(extractions, dict):
        return {}
    value = extractions.get(scenario)
    return value if isinstance(value, dict) else {}


def _manifest_archive_name(value: dict[str, Any]) -> str:
    return Path(str(value.get("archive") or "")).name


def _file_reading_description(
    scenario: str,
    file_path: str,
    result: dict[str, Any] | None,
) -> tuple[str, str, str]:
    name = Path(file_path).name
    suffix = Path(file_path).suffix.lower()
    if scenario == "personnel_incentive" and result:
        summary = result["summary"]
        if suffix in {".xlsx", ".xlsm"}:
            sales = result["sales"]
            return (
                "销售Excel",
                f"工作表{sales['sheet']}；表头第{sales['header_row']}行；读取期间、门店、条码、商品、数量、单价和金额公式；保留原Excel行号。",
                f"有效明细{sales['detail_row_count']}行、{sales['store_count']}家门店、{sales['sku_count']}个SKU；按条码汇总后与最终结算单{len(result['settlement']['lines'])}行逐项比较。",
            )
        if "结算单" in name:
            settlement = result["settlement"]
            return (
                "最终结算单",
                "按原图读取客户、活动期间、每个奖励单品、销售数量、奖励单价、行奖励金额、总数量、总奖励、实际申报、盖章和签署日期。",
                f"取得{len(settlement['lines'])}条结算行；逐行映射销售Excel唯一条码，再比较数量和奖励金额。",
            )
        return (
            "转账截图",
            "读取每个金额气泡、发送/接收方向、可见收款身份、可见日期和同一转账的镜像画面关系。",
            f"两张截图合并为{len(result['transfer_evidence'])}笔去重转账；与{summary['store_count']}家门店应付金额集合比较，身份和日期另行判定。",
        )
    if scenario == "promotional_display" and result:
        summary = result["summary"]
        if suffix == ".pdf":
            contract_pdf = result["contract_pdf"]
            return (
                "促销合同",
                "逐页读取客户、活动起止日、陈列标准、单店费用、申报金额、盖章及完整编号门店清单。",
                f"合同{contract_pdf['page_count']}页、{summary['contract_store_count']}家门店、单店{summary['fee_per_store']}元；作为逐店照片核验基准。",
            )
        if suffix in {".xlsx", ".xlsm"}:
            sales = result["sales"]
            rows = sales.get("records") or []
            return (
                "销售Excel",
                f"工作表{sales['sheet']}；表头第{sales['header_row']}行；读取客户、业务期间、商品编码、条码、商品、数量、零售价和合计金额，并检查外部链接公式。",
                f"有效明细{len(rows)}行、数量{sales['total_quantity']}、零售金额{sales['retail_amount']}；文件为经销商汇总，不能替代{summary['contract_store_count']}家门店逐店执行证据。",
            )
        return (
            "现场陈列照片",
            "按原始分辨率读取画面日期、地点/门店水印、陈列范围和清晰度；同时记录SHA-256、pHash、dHash、尺寸、EXIF时间和GPS。",
            f"与合同门店逐店绑定；共{summary['photo_count']}张照片支持{summary['passed_store_count']}家通过、{summary['supplement_store_count']}家补证。",
        )
    return ("源文件", "读取文件可见内容和技术元数据。", "纳入对应场景核验。")


def _add_file_inventory_sheet(
    wb: Workbook,
    results: list[dict[str, Any]],
    input_manifest: dict[str, Any] | None,
) -> Any:
    ws = wb.create_sheet("文件读取清单")
    header_row = _style_title(
        ws,
        "输入文件读取清单：从两个ZIP到每个原始文件",
        10,
        "每行说明文件来自哪个ZIP、读取了哪些位置/字段、取得什么数据以及后续与谁比较；SHA-256用于固定收到后的文件身份，不等于真伪鉴定。",
    )
    headers = [
        "场景",
        "来源ZIP",
        "ZIP SHA-256",
        "压缩包内文件",
        "文件角色",
        "文件大小(字节)",
        "文件SHA-256",
        "具体读取位置/字段",
        "读取结果/数据量",
        "后续比对工作",
    ]
    _header(ws, header_row, headers)
    scenario_labels = {"personnel_incentive": "人员激励", "promotional_display": "堆头"}
    rows: list[list[Any]] = []
    for scenario in ("personnel_incentive", "promotional_display"):
        extraction = _manifest_scenario(input_manifest, scenario)
        result = _scenario_result(results, scenario)
        archive_name = _manifest_archive_name(extraction)
        archive_sha = extraction.get("archive_sha256")
        if extraction:
            rows.append(
                [
                    scenario_labels[scenario],
                    archive_name,
                    archive_sha,
                    "（ZIP输入包本身）",
                    "输入压缩包",
                    extraction.get("archive_bytes"),
                    archive_sha,
                    "检查ZIP路径穿越、符号链接、加密、重名、文件数量、单文件大小、总展开大小和压缩比；根据文件组合与名称路由。",
                    f"安全解压{extraction.get('file_count')}个文件，共{extraction.get('expanded_bytes')}字节；路由到{scenario_labels[scenario]}子Skill。",
                    "冻结ZIP及所有解压文件哈希，再开始业务核验。",
                ]
            )
        files = extraction.get("files") if isinstance(extraction.get("files"), list) else []
        for item in files:
            if not isinstance(item, dict):
                continue
            file_path = str(item.get("path") or "")
            role, reading, result_text = _file_reading_description(scenario, file_path, result)
            rows.append(
                [
                    scenario_labels[scenario],
                    archive_name,
                    archive_sha,
                    file_path,
                    role,
                    item.get("bytes"),
                    item.get("sha256"),
                    reading,
                    result_text,
                    (
                        "逐SKU/逐门店/逐转账勾稽，详见人员激励明细页。"
                        if scenario == "personnel_incentive"
                        else "逐合同门店比较日期、地点、陈列和重复情况，详见堆头明细页。"
                    ),
                ]
            )
    if not rows:
        for result in results:
            scenario = str(result["scenario"])
            label = scenario_labels.get(scenario, scenario)
            sources: list[str] = []
            for key in ("settlement", "sales", "contract"):
                value = result.get(key)
                if isinstance(value, dict) and value.get("source_file"):
                    sources.append(str(value["source_file"]))
            for source in sources:
                role, reading, result_text = _file_reading_description(scenario, source, result)
                rows.append([label, None, None, source, role, None, None, reading, result_text, "详见对应明细页。"])
    for offset, values in enumerate(rows):
        row = header_row + 1 + offset
        for col, value in enumerate(values, 1):
            _body_cell(ws.cell(row, col, value))
        ws.row_dimensions[row].height = 78
    last_row = max(header_row + len(rows), header_row)
    _freeze_filter(ws, header_row, last_row, len(headers))
    _set_widths(ws, {1: 12, 2: 42, 3: 68, 4: 58, 5: 18, 6: 16, 7: 68, 8: 66, 9: 60, 10: 54})
    return ws


def _add_analysis_dimensions_sheet(
    wb: Workbook,
    results: list[dict[str, Any]],
    input_manifest: dict[str, Any] | None,
) -> Any:
    ws = wb.create_sheet("1.4.1分析维度")
    header_row = _style_title(
        ws,
        "1.4.1 分析维度：真实性验证 + 合规性审查",
        11,
        "逐场景覆盖方案规定的7项控制。状态说明：通过=关键证据完整；部分通过=部分对象未通过；部分验证=技术能力或证据不完整；证据不足=适用但缺证；本案未触发=本场景未提交该类材料且未规定为必交。",
    )
    headers = [
        "序号",
        "分析维度",
        "检测内容",
        "1.4.1判定标准",
        "场景",
        "本次读取文件/证据",
        "实际执行的检查",
        "状态",
        "本次判定依据",
        "能力边界/补证要求",
        "关联明细页",
    ]
    _header(ws, header_row, headers)
    rows = build_analysis_dimension_rows(results, input_manifest)
    for offset, item in enumerate(rows):
        row = header_row + 1 + offset
        values = [
            item["dimension_no"],
            item["analysis_dimension"],
            item["check_content"],
            item["proposal_standard"],
            item["scenario"],
            item["files"],
            item["executed_check"],
            item["status"],
            item["basis"],
            item["limitation"],
            item["detail_sheet"],
        ]
        for col, value in enumerate(values, 1):
            cell = ws.cell(row, col, value)
            _body_cell(cell)
            if col == 8:
                _status_fill(cell, str(value))
        detail_sheet = str(item["detail_sheet"])
        if detail_sheet in wb.sheetnames:
            _sheet_link(ws.cell(row, 11), detail_sheet)
        ws.row_dimensions[row].height = 96
    last_row = max(header_row + len(rows), header_row)
    _freeze_filter(ws, header_row, last_row, len(headers))
    _set_widths(ws, {1: 8, 2: 16, 3: 25, 4: 48, 5: 12, 6: 48, 7: 58, 8: 14, 9: 68, 10: 68, 11: 24})
    return ws


def _add_workflow_sheet(
    wb: Workbook,
    results: list[dict[str, Any]],
    input_manifest: dict[str, Any] | None,
) -> Any:
    ws = wb.create_sheet("核销步骤说明")
    header_row = _style_title(
        ws,
        "核销步骤说明：每一步读什么文件、取什么数据、怎样比对",
        9,
        "建议按本页从上到下阅读，再点击最后一列进入底层明细。最终外部交付只有这一份Excel，不保存运行日志。",
    )
    headers = ["步骤", "场景", "读取文件", "读取位置/字段", "处理动作", "比对对象", "本步结果", "结论/限制", "明细入口"]
    _header(ws, header_row, headers)
    personnel = _scenario_result(results, "personnel_incentive")
    display = _scenario_result(results, "promotional_display")
    p_manifest = _manifest_scenario(input_manifest, "personnel_incentive")
    d_manifest = _manifest_scenario(input_manifest, "promotional_display")
    steps: list[tuple[Any, ...]] = []
    if p_manifest or d_manifest:
        steps.append(
            (
                1,
                "主流程",
                f"{_manifest_archive_name(p_manifest)}；{_manifest_archive_name(d_manifest)}",
                "两个ZIP及全部压缩包成员",
                "安全校验、解压、计算ZIP/文件SHA-256，并按文件组合路由到两个子Skill。",
                "必须恰好1个人员激励包+1个堆头包",
                f"人员激励{p_manifest.get('file_count', 0)}个文件；堆头{d_manifest.get('file_count', 0)}个文件。",
                "哈希固定收到后的文件身份，不等于真伪证明。",
                "文件读取清单",
            )
        )
    step_no = len(steps) + 1
    if personnel:
        sales = personnel["sales"]
        settlement = personnel["settlement"]
        p_rows = sales.get("records") or []
        row_range = (
            f"原Excel第{min(item['excel_row'] for item in p_rows)}至{max(item['excel_row'] for item in p_rows)}行中的有效行"
            if p_rows
            else "有效销售明细行"
        )
        steps.extend(
            [
                (
                    step_no,
                    "人员激励",
                    Path(str(sales["source_file"])).name,
                    f"工作表{sales['sheet']}，表头第{sales['header_row']}行；{row_range}；字段为期间、门店、条码、商品、数量、单价、金额。",
                    "跳过空白/小计/合计，保留每条有效记录的原Excel行号，再按条码和门店汇总。",
                    "销售Excel自身明细与汇总",
                    f"取得{sales['detail_row_count']}条明细、{sales['store_count']}家门店、{sales['sku_count']}个SKU、总数量{sales['total_quantity']}。",
                    "这是结算数量的底层来源，不先看总金额。",
                    "人员激励-门店SKU明细",
                ),
                (
                    step_no + 1,
                    "人员激励",
                    Path(str(settlement["source_file"])).name,
                    "最终结算单原图；逐行读取奖励单品、销售数量、奖励单价、奖励金额，以及总数量/总额/实际申报。",
                    "把结算单每一行作为独立待核对象；不把总额相等当作逐行通过。",
                    "最终结算单10行",
                    f"取得{len(settlement['lines'])}条结算SKU、结算数量{personnel['summary']['settlement_line_quantity']}、行奖励合计{personnel['summary']['settlement_line_reward']}元、申报{personnel['summary']['claimed_amount']}元。",
                    "结算单无条码时，商品名称映射必须保留依据和置信度。",
                    "人员激励-SKU逐项",
                ),
                (
                    step_no + 2,
                    "人员激励",
                    "销售Excel + 最终结算单",
                    "结算单每行 ↔ Excel唯一条码及其全部来源行",
                    "先做商品1:1映射，再把7家门店同一条码的数量相加；逐SKU比较Excel数量与结算数量，并计算数量×奖励单价。",
                    "10条结算SKU与10个Excel条码",
                    f"{personnel['summary']['quantity_match_count']}/{personnel['summary']['settlement_line_count']}个SKU数量匹配；Excel {personnel['summary']['excel_mapped_quantity']}={personnel['summary']['settlement_line_quantity']}结算数量；奖励计算{personnel['summary']['calculated_line_reward']}元。",
                    "4个名称映射为中置信度，仍需商品条码/主数据补证；数量事实已逐行展示。",
                    "人员激励-SKU逐项",
                ),
                (
                    step_no + 3,
                    "人员激励",
                    "全部转账截图",
                    "每个金额气泡、方向、身份、日期和重复画面关系",
                    "把同一笔发送/接收镜像合并为一个transfer_id，再将7家门店应付金额集合与7笔转账金额集合比较。",
                    "门店应付 ↔ 去重转账",
                    f"去重转账{len(personnel['transfer_evidence'])}笔、合计{personnel['summary']['transfer_total']}元，与结算行奖励合计一致。",
                    "金额一致不证明收款门店身份或转账日期；当前两项不可见。",
                    "人员激励-门店转账",
                ),
            ]
        )
        step_no += 4
    if display:
        contract = display["contract"]
        sales = display["sales"]
        steps.extend(
            [
                (
                    step_no,
                    "堆头",
                    Path(str(contract["source_file"])).name,
                    "合同逐页；读取客户、活动期、陈列标准、单店费用、申报金额和编号1—20门店。",
                    "建立20行合同门店基准，任何照片都必须回绑某一合同原始序号。",
                    "合同条款 ↔ 申报/照片",
                    f"合同{display['contract_pdf']['page_count']}页，{display['summary']['contract_store_count']}家×{display['summary']['fee_per_store']}元={display['summary']['expected_contract_amount']}元。",
                    "合同是逐店判定基准；文件名只能帮助归组，不能证明地点。",
                    "堆头-逐店核验",
                ),
                (
                    step_no + 1,
                    "堆头",
                    Path(str(sales["source_file"])).name,
                    f"工作表{sales['sheet']}，表头第{sales['header_row']}行；读取客户、期间、商品编码、条码、商品、数量、零售价和金额。",
                    "直接读取公式和缓存值，汇总数量/金额并判断数据粒度。",
                    "销售期间/客户 ↔ 合同；销售数据粒度 ↔ 逐店执行要求",
                    f"{len(sales['records'])}条SKU，数量{sales['total_quantity']}，零售金额{sales['retail_amount']}；仅经销商汇总，无门店字段。",
                    "只能支持活动总体销售，不能证明20家逐店陈列。",
                    "堆头-销售检查",
                ),
                (
                    step_no + 2,
                    "堆头",
                    f"现场陈列照片{display['summary']['photo_count']}张",
                    "每张原图的画面日期、地点/门店水印、陈列范围、清晰度；SHA-256、pHash、dHash、尺寸、EXIF时间和GPS。",
                    "逐张查重并绑定合同门店；每家分别判断照片、期间、门店、陈列、重复5项控制。",
                    "21张照片 ↔ 合同20家门店",
                    f"{display['summary']['passed_store_count']}家通过，{display['summary']['supplement_store_count']}家补证，支持{display['summary']['supported_amount']}元、暂挂{display['summary']['temporarily_held_amount']}元。",
                    "全部照片无可用EXIF/GPS；6家具体不一致见逐店明细。",
                    "堆头-逐店核验",
                ),
            ]
        )
        step_no += 3
    final_parts: list[str] = []
    if personnel:
        final_parts.append(
            f"人员激励建议通过{personnel['summary']['suggested_approved_amount']}元、暂挂{personnel['summary']['temporarily_held_amount']}元"
        )
    if display:
        final_parts.append(
            f"堆头建议通过{display['summary']['suggested_approved_amount']}元、暂挂{display['summary']['temporarily_held_amount']}元"
        )
    steps.extend(
        [
            (
                step_no,
                "主流程",
                "两个ZIP内全部材料",
                "方案1.4.1规定的真实性3项、合规性4项",
                "对两个场景分别写明实际检查、状态、依据、能力边界和补证；缺证不默认为通过。",
                "七维标准 ↔ 本次可得证据",
                "形成14行（7维×2场景）分析矩阵。",
                "SHA/dHash、EXIF、水印、预算、发票和合同各自有独立边界。",
                "1.4.1分析维度",
            ),
            (
                step_no + 1,
                "主流程",
                "上述全部核验结果",
                "逐SKU、逐门店、逐转账、七维分析及异常补证",
                "先取证据支持金额，再与申报金额比较，合并成唯一Excel交付。",
                "人员激励与堆头结果",
                "；".join(final_parts) + "。",
                "最终金额可追溯到本工作簿各明细页；不输出日志。",
                "核销总览",
            ),
        ]
    )
    for offset, values in enumerate(steps):
        row = header_row + 1 + offset
        for col, value in enumerate(values, 1):
            _body_cell(ws.cell(row, col, value))
        sheet_name = str(values[-1])
        if sheet_name in wb.sheetnames:
            _sheet_link(ws.cell(row, 9), sheet_name)
        ws.row_dimensions[row].height = 96
    last_row = max(header_row + len(steps), header_row)
    _freeze_filter(ws, header_row, last_row, len(headers))
    _set_widths(ws, {1: 8, 2: 12, 3: 48, 4: 62, 5: 66, 6: 48, 7: 60, 8: 60, 9: 26})
    return ws


def _add_model_usage_sheet(wb: Workbook, results: list[dict[str, Any]]) -> Any:
    ws = wb.create_sheet("模型用量")
    header_row = _style_title(
        ws,
        "模型调用量、失败重试与报价功能范围",
        12,
        "只保存结构化用量统计，不保存模型事件流或stdout/stderr日志。报价表中的“工作量（人天）/报价”不是调用次数或Token预算；实际调用量取自模型执行返回的usage字段。",
    )
    usage_headers = [
        "场景",
        "案件",
        "模型",
        "调用次数",
        "成功次数",
        "失败次数",
        "重试次数",
        "输入Token",
        "缓存输入Token",
        "输出Token",
        "推理输出Token",
        "总Token(输入+输出)",
    ]
    _header(ws, header_row, usage_headers)
    scenario_labels = {"personnel_incentive": "人员激励", "promotional_display": "堆头"}
    start = header_row + 1
    usage_rows: list[dict[str, Any]] = []
    for result in results:
        usage = result.get("model_usage")
        if not isinstance(usage, dict):
            usage = {
                "model": "未调用（使用预提取证据）",
                "model_invocation_count": 0,
                "successful_invocation_count": 0,
                "failed_invocation_count": 0,
                "retry_count": 0,
                "input_tokens": 0,
                "cached_input_tokens": 0,
                "output_tokens": 0,
                "reasoning_output_tokens": 0,
                "total_tokens": 0,
                "attempts": [],
            }
        usage_rows.append({"result": result, "usage": usage})
    for offset, item in enumerate(usage_rows):
        row = start + offset
        result = item["result"]
        usage = item["usage"]
        values = [
            scenario_labels.get(str(result.get("scenario")), str(result.get("scenario"))),
            result.get("case_name"),
            usage.get("model"),
            usage.get("model_invocation_count", 0),
            usage.get("successful_invocation_count", 0),
            usage.get("failed_invocation_count", 0),
            usage.get("retry_count", 0),
            usage.get("input_tokens", 0),
            usage.get("cached_input_tokens", 0),
            usage.get("output_tokens", 0),
            usage.get("reasoning_output_tokens", 0),
            usage.get("total_tokens", 0),
        ]
        for col, value in enumerate(values, 1):
            _body_cell(ws.cell(row, col, value))
            if col in {6, 7} and int(value or 0) > 0:
                ws.cell(row, col).fill = PatternFill("solid", fgColor=YELLOW)
        ws.row_dimensions[row].height = 42
    total_row = start + len(usage_rows)
    ws.cell(total_row, 1, "合计")
    for col in range(4, 13):
        letter = get_column_letter(col)
        ws.cell(total_row, col, f"=SUM({letter}{start}:{letter}{total_row - 1})")
    for col in range(1, 13):
        _body_cell(ws.cell(total_row, col), formula=col >= 4)
        ws.cell(total_row, col).fill = PatternFill("solid", fgColor=BLUE)
        ws.cell(total_row, col).font = _font(bold=True, color=FORMULA_BLACK)

    attempt_header = total_row + 2
    _header(
        ws,
        attempt_header,
        [
            "场景",
            "案件",
            "尝试序号",
            "状态",
            "退出码",
            "失败原因摘要",
            "输入Token",
            "缓存输入Token",
            "输出Token",
            "推理输出Token",
            "总Token",
        ],
    )
    attempt_row = attempt_header + 1
    for item in usage_rows:
        result = item["result"]
        usage = item["usage"]
        for attempt in usage.get("attempts") or []:
            attempt_usage = attempt.get("usage") or {}
            values = [
                scenario_labels.get(str(result.get("scenario")), str(result.get("scenario"))),
                result.get("case_name"),
                attempt.get("attempt_no"),
                "成功" if attempt.get("status") == "success" else "失败后重试",
                attempt.get("exit_code"),
                attempt.get("failure_reason"),
                attempt_usage.get("input_tokens", 0),
                attempt_usage.get("cached_input_tokens", 0),
                attempt_usage.get("output_tokens", 0),
                attempt_usage.get("reasoning_output_tokens", 0),
                attempt_usage.get("total_tokens", 0),
            ]
            for col, value in enumerate(values, 1):
                _body_cell(ws.cell(attempt_row, col, value))
            _status_fill(ws.cell(attempt_row, 4), str(values[3]))
            ws.row_dimensions[attempt_row].height = 42
            attempt_row += 1
    if attempt_row == attempt_header + 1:
        ws.cell(attempt_row, 1, "本次使用预提取证据，未发生模型调用。")
        ws.merge_cells(start_row=attempt_row, start_column=1, end_row=attempt_row, end_column=11)
        _body_cell(ws.cell(attempt_row, 1))
        attempt_row += 1

    scope_header = attempt_row + 1
    _header(ws, scope_header, ["报价功能项", "工作量(人天)", "报价(元)", "原报价备注", "本项目落点"])
    implementation = {
        "文件分类与图像预处理": "主Skill安全解压、文件分类、图片/PDF/Excel读取预处理。",
        "照片真实性风险辅助检测": "堆头子Skill读取EXIF/GPS/水印并做SHA-256/dHash同包去重；跨活动历史库未接入时明确标注未验证。",
        "现场/陈列照片多模态识别": "堆头子Skill逐合同门店识别日期、地点、商品陈列、范围和质量。",
        "发票及票据OCR结构化": "当前两个样例包无发票，状态为本案未触发；以后有发票时按字段结构化。",
        "合同、费用明细及规则抽取": "读取堆头合同、两个销售Excel、人员结算单和活动规则字段。",
        "跨材料一致性与规则判断": "人员激励逐SKU/转账，堆头逐门店，并执行1.4.1七维分析。",
        "结果汇总、异常重试与用量日志": "汇总通过/异常/待补件；模型提取失败最多3次；只统计调用数和Token，不保存原始日志。",
    }
    for offset, (name, days, quote, note) in enumerate(AI_REVIEW_SCOPE, 1):
        row = scope_header + offset
        values = [name, days, quote, note, implementation[name]]
        for col, value in enumerate(values, 1):
            _body_cell(ws.cell(row, col, value))
        ws.row_dimensions[row].height = 60
    scope_total = scope_header + len(AI_REVIEW_SCOPE) + 1
    ws.cell(scope_total, 1, "小计")
    ws.cell(scope_total, 2, f"=SUM(B{scope_header + 1}:B{scope_total - 1})")
    ws.cell(scope_total, 3, f"=SUM(C{scope_header + 1}:C{scope_total - 1})")
    for col in range(1, 6):
        _body_cell(ws.cell(scope_total, col), formula=col in {2, 3})
        ws.cell(scope_total, col).fill = PatternFill("solid", fgColor=BLUE)
        ws.cell(scope_total, col).font = _font(bold=True, color=FORMULA_BLACK)

    ws.freeze_panes = f"A{header_row + 1}"
    ws.sheet_view.showGridLines = False
    _set_widths(ws, {1: 28, 2: 34, 3: 24, 4: 13, 5: 13, 6: 15, 7: 13, 8: 16, 9: 18, 10: 16, 11: 18, 12: 20})
    return ws


def _write_exceptions(wb: Workbook, title: str, result: dict[str, Any]) -> str:
    ws = wb.create_sheet(title)
    header_row = _style_title(
        ws,
        f"{result['case_name']}：异常与补证清单",
        6,
        "高风险影响数量/金额或核心证据；中风险表示口径、身份或真实性证据仍需说明。",
    )
    headers = ["序号", "等级", "异常代码", "发现", "影响", "补证/处理建议"]
    _header(ws, header_row, headers)
    for index, item in enumerate(result["exceptions"], 1):
        row = header_row + index
        values = [
            index,
            item["severity"],
            item["code"],
            item["message"] + (f"\n来源：{item['source']}" if item.get("source") else ""),
            item["impact"],
            item["suggestion"],
        ]
        for col, value in enumerate(values, 1):
            cell = ws.cell(row, col, value)
            _body_cell(cell)
            if item["severity"] == "high":
                cell.fill = PatternFill("solid", fgColor=RED)
            elif item["severity"] == "medium":
                cell.fill = PatternFill("solid", fgColor=YELLOW)
    last_row = max(header_row + len(result["exceptions"]), header_row)
    _freeze_filter(ws, header_row, last_row, len(headers))
    _set_widths(ws, {1: 8, 2: 10, 3: 34, 4: 60, 5: 42, 6: 52})
    return title


def _add_personnel_sheets(wb: Workbook, result: dict[str, Any]) -> dict[str, str]:
    sku_title = "人员激励-SKU逐项"
    ws = wb.create_sheet(sku_title)
    header_row = _style_title(
        ws,
        "人员激励：Excel销售数量与结算单逐SKU 1:1勾稽",
        18,
        "先按结算行映射唯一条码，再汇总7家门店Excel来源行；数量、单价、行金额和支持金额分别计算。",
    )
    headers = [
        "结算行",
        "结算单奖励单品",
        "对应条码",
        "Excel商品名称",
        "Excel来源行",
        "Excel销售数量",
        "结算销售数量",
        "数量差异",
        "数量状态",
        "奖励单价",
        "Excel应付金额",
        "结算应付金额",
        "结算行奖励金额",
        "行金额差异",
        "金额状态",
        "支持数量",
        "支持金额",
        "映射依据/置信度",
    ]
    _header(ws, header_row, headers)
    start = header_row + 1
    for offset, item in enumerate(result["sku_reconciliation"]):
        row = start + offset
        values = [
            item["line_no"],
            item["settlement_product_name"],
            item["mapped_barcode"],
            item["excel_product_name"],
            ",".join(str(value) for value in item["excel_rows"]),
            item["excel_quantity"],
            item["settlement_quantity"],
            f'=IF(OR(F{row}="",G{row}=""),"",F{row}-G{row})',
            f'=IF(H{row}="","未匹配",IF(H{row}=0,"匹配","异常"))',
            item["unit_reward"],
            f'=IF(F{row}="","",F{row}*J{row})',
            f'=G{row}*J{row}',
            item["settlement_reward_amount"],
            f'=M{row}-L{row}',
            f'=IF(N{row}=0,"匹配","异常")',
            f'=IF(OR(F{row}="",G{row}=""),0,MAX(0,MIN(F{row},G{row})))',
            f'=P{row}*J{row}',
            f"{item['mapping_basis']} / {item['mapping_confidence']}",
        ]
        for col, value in enumerate(values, 1):
            cell = ws.cell(row, col, value)
            formula = isinstance(value, str) and value.startswith("=")
            _body_cell(cell, formula=formula)
            if col in {9, 15}:
                _status_fill(cell, str(item["quantity_status"] if col == 9 else item["amount_status"]))
    total_row = start + len(result["sku_reconciliation"])
    ws.cell(total_row, 1, "合计")
    for col in (6, 7, 8, 10, 11, 12, 13, 14, 16, 17):
        letter = get_column_letter(col)
        ws.cell(total_row, col, f"=SUM({letter}{start}:{letter}{total_row - 1})")
    for col in range(1, 19):
        cell = ws.cell(total_row, col)
        _body_cell(cell, formula=col in {6, 7, 8, 10, 11, 12, 13, 14, 16, 17})
        cell.font = _font(bold=True, color=FORMULA_BLACK if col != 1 else INPUT_BLUE)
        cell.fill = PatternFill("solid", fgColor=BLUE)
    _freeze_filter(ws, header_row, total_row, len(headers))
    _set_widths(
        ws,
        {1: 9, 2: 30, 3: 17, 4: 42, 5: 18, 6: 14, 7: 14, 8: 12, 9: 12,
         10: 11, 11: 15, 12: 15, 13: 17, 14: 13, 15: 12, 16: 12, 17: 13, 18: 30},
    )

    trace_title = "人员激励-门店SKU明细"
    trace_ws = wb.create_sheet(trace_title)
    trace_header = _style_title(
        trace_ws,
        "人员激励：70条门店×SKU销售来源明细",
        11,
        "每一行均保留原Excel行号，并回指结算单行；这是SKU汇总数量的底层证据链。",
    )
    trace_headers = [
        "Excel行",
        "期间",
        "门店",
        "条码",
        "Excel商品名称",
        "销售数量",
        "销售单价",
        "结算行",
        "结算单品",
        "奖励单价",
        "该行应付奖励",
    ]
    _header(trace_ws, trace_header, trace_headers)
    for offset, item in enumerate(result["sales_trace"]):
        row = trace_header + 1 + offset
        values = [
            item["excel_row"], item["period_text"], item["store_name"], item["barcode"],
            item["product_name"], item["quantity"], item["unit_price"], item["settlement_line_no"],
            item["settlement_product_name"], item["unit_reward"],
            f'=IF(OR(F{row}="",J{row}=""),"",F{row}*J{row})',
        ]
        for col, value in enumerate(values, 1):
            _body_cell(trace_ws.cell(row, col, value), formula=isinstance(value, str) and value.startswith("="))
    trace_total = trace_header + len(result["sales_trace"]) + 1
    trace_ws.cell(trace_total, 1, "合计")
    for col in (6, 11):
        letter = get_column_letter(col)
        trace_ws.cell(trace_total, col, f"=SUM({letter}{trace_header + 1}:{letter}{trace_total - 1})")
    for col in range(1, 12):
        _body_cell(trace_ws.cell(trace_total, col), formula=col in {6, 11})
        trace_ws.cell(trace_total, col).fill = PatternFill("solid", fgColor=BLUE)
    _freeze_filter(trace_ws, trace_header, trace_total, len(trace_headers))
    _set_widths(trace_ws, {1: 10, 2: 14, 3: 18, 4: 17, 5: 45, 6: 12, 7: 12, 8: 10, 9: 30, 10: 12, 11: 15})

    transfer_title = "人员激励-门店转账"
    transfer_ws = wb.create_sheet(transfer_title)
    transfer_header = _style_title(
        transfer_ws,
        "人员激励：门店应付与去重转账勾稽",
        9,
        "金额集合一致不等于收款身份一致；截图未显示门店/收款人时单独标记为身份未验证。",
    )
    transfer_headers = [
        "门店", "Excel销售数量", "应付奖励", "转账ID", "去重转账金额", "金额差异", "匹配依据", "身份状态", "结论"
    ]
    _header(transfer_ws, transfer_header, transfer_headers)
    transfer_start = transfer_header + 1
    for offset, item in enumerate(result["store_transfer_reconciliation"]):
        row = transfer_start + offset
        values = [
            item["store_name"], item["excel_quantity"], item["expected_reward_amount"],
            item["transfer_id"], item["transfer_amount"],
            f'=IF(OR(C{row}="",E{row}=""),"",E{row}-C{row})',
            item["match_basis"], item["identity_status"], item["status"],
        ]
        for col, value in enumerate(values, 1):
            cell = transfer_ws.cell(row, col, value)
            _body_cell(cell, formula=isinstance(value, str) and value.startswith("="))
            if col == 9:
                _status_fill(cell, str(value))
    transfer_total = transfer_start + len(result["store_transfer_reconciliation"])
    transfer_ws.cell(transfer_total, 1, "合计")
    for col in (2, 3, 5, 6):
        letter = get_column_letter(col)
        transfer_ws.cell(transfer_total, col, f"=SUM({letter}{transfer_start}:{letter}{transfer_total - 1})")
    for col in range(1, 10):
        _body_cell(transfer_ws.cell(transfer_total, col), formula=col in {2, 3, 5, 6})
        transfer_ws.cell(transfer_total, col).fill = PatternFill("solid", fgColor=BLUE)
    _freeze_filter(transfer_ws, transfer_header, transfer_total, len(transfer_headers))
    _set_widths(transfer_ws, {1: 18, 2: 15, 3: 14, 4: 13, 5: 16, 6: 12, 7: 22, 8: 25, 9: 30})

    voucher_title = "人员激励-转账凭证"
    voucher_ws = wb.create_sheet(voucher_title)
    voucher_header = _style_title(
        voucher_ws,
        "人员激励：转账截图去重明细",
        9,
        "同一笔转账的发送方与接收方气泡合并为一个transfer_id，同时保留画面出现次数。",
    )
    voucher_headers = ["转账ID", "金额", "画面出现次数", "身份可见", "日期可见", "门店/收款人", "来源文件", "去重依据", "备注"]
    _header(voucher_ws, voucher_header, voucher_headers)
    for offset, item in enumerate(result["transfer_evidence"]):
        row = voucher_header + 1 + offset
        identity = item.get("store_name") or item.get("recipient_name")
        values = [
            item["transfer_id"], item["amount"], item["occurrence_count"],
            "是" if item["identity_visible"] else "否", "是" if item["event_date_visible"] else "否",
            identity, "、".join(item["source_files"]), item["dedup_basis"], "；".join(item.get("notes") or []),
        ]
        for col, value in enumerate(values, 1):
            _body_cell(voucher_ws.cell(row, col, value))
    voucher_last = voucher_header + len(result["transfer_evidence"])
    _freeze_filter(voucher_ws, voucher_header, max(voucher_last, voucher_header), len(voucher_headers))
    _set_widths(voucher_ws, {1: 13, 2: 13, 3: 15, 4: 12, 5: 12, 6: 20, 7: 46, 8: 55, 9: 35})

    image_title = "人员激励-图片文件检查"
    image_ws = wb.create_sheet(image_title)
    image_header = _style_title(
        image_ws,
        "人员激励：结算单与转账截图真实性风险辅助信息",
        10,
        "SHA-256检查完全重复，pHash筛查近似画面，dHash作为辅助；两张转账截图可能因同一聊天内容重叠而成为近似候选，需结合transfer_id去重，不能仅凭pHash判定造假。",
    )
    image_headers = ["文件名", "角色", "尺寸", "大小(字节)", "SHA-256", "pHash", "dHash", "EXIF时间", "GPS", "pHash候选"]
    _header(image_ws, image_header, image_headers)
    image_inventory = result.get("image_inventory") or {}
    candidate_by_file: dict[str, list[str]] = {}
    for pair in image_inventory.get("phash_candidate_pairs") or []:
        names = [str(value) for value in pair.get("files") or []]
        distance = pair.get("distance")
        for name in names:
            others = [value for value in names if value != name]
            candidate_by_file.setdefault(name, []).append(
                f"{'、'.join(others)}（距离{distance}）"
            )
    for offset, item in enumerate(image_inventory.get("files") or []):
        row = image_header + 1 + offset
        file_name = str(item["file_name"])
        values = [
            file_name,
            "最终结算单" if "结算单" in file_name else "转账截图",
            f"{item['width']}×{item['height']}",
            item["size"],
            item["sha256"],
            item.get("phash"),
            item["dhash"],
            item["exif_datetime"],
            "有" if item["gps_present"] else "无",
            "；".join(candidate_by_file.get(file_name, [])) or "无",
        ]
        for col, value in enumerate(values, 1):
            _body_cell(image_ws.cell(row, col, value))
    image_last = image_header + len(image_inventory.get("files") or [])
    _freeze_filter(image_ws, image_header, max(image_last, image_header), len(image_headers))
    _set_widths(image_ws, {1: 42, 2: 16, 3: 14, 4: 16, 5: 68, 6: 20, 7: 20, 8: 22, 9: 10, 10: 44})

    exception_title = _write_exceptions(wb, "人员激励-异常补证", result)
    return {
        "detail_sheet": sku_title,
        "detail_total_row": str(total_row),
        "supported_cell": f"'{sku_title}'!Q{total_row}",
        "transfer_total_cell": f"'{transfer_title}'!E{transfer_total}",
        "exception_sheet": exception_title,
    }


def _add_display_sheets(wb: Workbook, result: dict[str, Any]) -> dict[str, str]:
    store_title = "堆头-逐店核验"
    ws = wb.create_sheet(store_title)
    header_row = _style_title(
        ws,
        "堆头核销：合同门店与照片逐店证据核验",
        14,
        "单店费用仅在照片、期间、门店、陈列和重复检查全部通过时计入支持金额。",
    )
    headers = [
        "序号",
        "合同门店",
        "照片文件",
        "照片数",
        "可见日期",
        "可见地点/水印",
        "期间",
        "门店匹配",
        "陈列标准",
        "重复检查",
        "结论",
        "支持金额",
        "风险说明",
        "补证建议",
    ]
    _header(ws, header_row, headers)
    start = header_row + 1
    fee = result["summary"]["fee_per_store"]
    status_labels = {"pass": "通过", "supplement": "补证"}
    period_labels = {"match": "是", "mismatch": "否", "unverifiable": "无法判断"}
    store_labels = {"exact": "是", "compatible": "基本匹配", "mismatch": "不一致", "filename_only": "弱（仅文件名）"}
    display_labels = {"pass": "是", "fail": "否", "uncertain": "无法判断"}
    duplicate_labels = {"none": "未发现", "exact": "完全重复", "possible": "疑似重复", "unverifiable": "无法判断"}
    for offset, item in enumerate(result["store_reconciliation"]):
        row = start + offset
        status = status_labels[item["status"]]
        values = [
            item["store_line_no"],
            item["contract_store_name"],
            "、".join(item["photo_files"]),
            item["photo_count"],
            item["visible_date"],
            item["visible_location"],
            period_labels[item["period_match"]],
            store_labels[item["store_match"]],
            display_labels[item["display_match"]],
            duplicate_labels[item["duplicate_check"]],
            status,
            f'=IF(K{row}="通过",{fee},0)',
            "；".join(item["risk_notes"]),
            "；".join(item["supplement_advice"]),
        ]
        for col, value in enumerate(values, 1):
            cell = ws.cell(row, col, value)
            _body_cell(cell, formula=isinstance(value, str) and value.startswith("="))
            if col == 11:
                _status_fill(cell, status)
    total_row = start + len(result["store_reconciliation"])
    ws.cell(total_row, 1, "合计")
    ws.cell(total_row, 4, f"=SUM(D{start}:D{total_row - 1})")
    ws.cell(total_row, 12, f"=SUM(L{start}:L{total_row - 1})")
    for col in range(1, 15):
        _body_cell(ws.cell(total_row, col), formula=col in {4, 12})
        ws.cell(total_row, col).fill = PatternFill("solid", fgColor=BLUE)
    _freeze_filter(ws, header_row, total_row, len(headers))
    _set_widths(ws, {1: 8, 2: 31, 3: 46, 4: 10, 5: 13, 6: 36, 7: 13, 8: 15, 9: 15, 10: 14, 11: 12, 12: 14, 13: 50, 14: 52})

    sales_title = "堆头-销售检查"
    sales_ws = wb.create_sheet(sales_title)
    sales_header = _style_title(
        sales_ws,
        "堆头核销：销售Excel完整性与明细检查",
        8,
        "经销商汇总销售可证明活动期总体销售，但不能替代20家合同门店的逐店执行证据。",
    )
    metrics = [
        ("客户", "、".join(result["sales"]["customers"]), "销售Excel"),
        ("业务期间", "、".join(result["sales"]["period_values"]), "销售Excel"),
        ("SKU数量", result["sales"]["sku_count"], "有效明细行"),
        ("销售数量", result["sales"]["total_quantity"], "数量列求和"),
        ("零售金额", result["sales"]["retail_amount"], "缓存值可用时求和"),
        ("门店级销售", "否", "当前文件为经销商汇总"),
        ("外部链接公式", len(result["sales"]["external_formula_cells"]), "含[工作簿]引用的公式单元格"),
    ]
    _header(sales_ws, sales_header, ["检查项", "结果", "来源/说明"])
    for offset, (name, value, note) in enumerate(metrics):
        row = sales_header + 1 + offset
        for col, cell_value in enumerate((name, value, note), 1):
            _body_cell(sales_ws.cell(row, col, cell_value))
    detail_header = sales_header + len(metrics) + 2
    detail_headers = ["Excel行", "客户", "期间", "商品编码", "条码", "商品名称", "数量", "零售金额"]
    _header(sales_ws, detail_header, detail_headers)
    for offset, item in enumerate(result["sales"]["records"]):
        row = detail_header + 1 + offset
        values = [
            item["excel_row"], item["customer_name"], item["period_text"], item["product_code"],
            item["barcode"], item["product_name"], item["quantity"], item["total_amount"],
        ]
        for col, value in enumerate(values, 1):
            _body_cell(sales_ws.cell(row, col, value))
    detail_total = detail_header + len(result["sales"]["records"]) + 1
    sales_ws.cell(detail_total, 1, "合计")
    for col in (7, 8):
        letter = get_column_letter(col)
        sales_ws.cell(detail_total, col, f"=SUM({letter}{detail_header + 1}:{letter}{detail_total - 1})")
    for col in range(1, 9):
        _body_cell(sales_ws.cell(detail_total, col), formula=col in {7, 8})
        sales_ws.cell(detail_total, col).fill = PatternFill("solid", fgColor=BLUE)
    sales_ws.freeze_panes = f"A{detail_header + 1}"
    sales_ws.auto_filter.ref = f"A{detail_header}:H{detail_total}"
    sales_ws.sheet_view.showGridLines = False
    _set_widths(sales_ws, {1: 14, 2: 34, 3: 25, 4: 16, 5: 17, 6: 55, 7: 12, 8: 15})

    photo_title = "堆头-照片文件检查"
    photo_ws = wb.create_sheet(photo_title)
    photo_header = _style_title(
        photo_ws,
        "堆头核销：照片文件哈希与原始元数据",
        10,
        "SHA-256用于完全重复；pHash用于近似重复候选，dHash作为快速辅助指纹；EXIF/GPS缺失和跨活动历史库未接入均需保留风险。",
    )
    photo_headers = ["文件名", "尺寸", "文件大小", "SHA-256", "pHash", "dHash", "EXIF时间", "GPS", "绑定门店", "状态"]
    _header(photo_ws, photo_header, photo_headers)
    bound_store: dict[str, str] = {}
    for item in result["store_reconciliation"]:
        for file_name in item["photo_files"]:
            bound_store[Path(file_name).name] = item["contract_store_name"]
    for offset, item in enumerate(result["photo_inventory"]["files"]):
        row = photo_header + 1 + offset
        values = [
            item["file_name"], f"{item['width']}×{item['height']}", item["size"], item["sha256"],
            item.get("phash"), item["dhash"], item["exif_datetime"], "有" if item["gps_present"] else "无",
            bound_store.get(item["file_name"]), "已绑定" if item["file_name"] in bound_store else "未绑定",
        ]
        for col, value in enumerate(values, 1):
            _body_cell(photo_ws.cell(row, col, value))
    photo_last = photo_header + len(result["photo_inventory"]["files"])
    _freeze_filter(photo_ws, photo_header, max(photo_last, photo_header), len(photo_headers))
    _set_widths(photo_ws, {1: 48, 2: 15, 3: 14, 4: 68, 5: 20, 6: 20, 7: 23, 8: 10, 9: 35, 10: 12})

    exception_title = _write_exceptions(wb, "堆头-异常补证", result)
    return {
        "detail_sheet": store_title,
        "supported_cell": f"'{store_title}'!L{total_row}",
        "exception_sheet": exception_title,
    }


COMPACT_LAST_COL = 13


def _section_title(
    ws: Any,
    row: int,
    title: str,
    note: str | None = None,
    *,
    last_col: int = COMPACT_LAST_COL,
) -> int:
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=last_col)
    cell = ws.cell(row, 1, title)
    cell.font = _font(bold=True, color=WHITE, size=12)
    cell.fill = PatternFill("solid", fgColor=NAVY)
    cell.alignment = Alignment(vertical="center")
    ws.row_dimensions[row].height = 24
    row += 1
    if note:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=last_col)
        cell = ws.cell(row, 1, note)
        cell.font = _font(color="555555", size=9)
        cell.fill = PatternFill("solid", fgColor=LIGHT_BLUE)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.row_dimensions[row].height = 42
        row += 1
    return row


def _write_row(
    ws: Any,
    row: int,
    values: list[Any],
    *,
    status_columns: set[int] | None = None,
    height: float | None = None,
    number_formats: dict[int, str] | None = None,
) -> None:
    status_columns = status_columns or set()
    number_formats = number_formats or {}
    for col, value in enumerate(values, 1):
        cell = ws.cell(row, col, value)
        formula = isinstance(value, str) and value.startswith("=")
        _body_cell(cell, formula=formula)
        if col in status_columns:
            _status_fill(cell, str(value))
        if col in number_formats:
            cell.number_format = number_formats[col]
    if height is not None:
        ws.row_dimensions[row].height = height


def _write_merged_row(
    ws: Any,
    row: int,
    blocks: list[tuple[int, int]],
    values: list[Any],
    *,
    header: bool = False,
    status_indexes: set[int] | None = None,
    height: float | None = None,
    number_formats: dict[int, str] | None = None,
) -> None:
    status_indexes = status_indexes or set()
    number_formats = number_formats or {}
    for index, ((start_col, end_col), value) in enumerate(zip(blocks, values, strict=True), 1):
        if end_col > start_col:
            ws.merge_cells(start_row=row, start_column=start_col, end_row=row, end_column=end_col)
        cell = ws.cell(row, start_col, value)
        if header:
            cell.font = _font(bold=True, color=WHITE)
            cell.fill = PatternFill("solid", fgColor=NAVY)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
        else:
            _body_cell(cell, formula=isinstance(value, str) and value.startswith("="))
            if index in status_indexes:
                _status_fill(cell, str(value))
        if index in number_formats:
            cell.number_format = number_formats[index]
    if height is not None:
        ws.row_dimensions[row].height = height


def _write_summary_cards(ws: Any, cards: list[dict[str, Any]], *, start_row: int = 3) -> None:
    blocks = [(1, 3), (4, 6), (7, 9), (10, COMPACT_LAST_COL)]
    for index, card in enumerate(cards):
        band = index // 4
        block = index % 4
        label_row = start_row + band * 3
        value_row = label_row + 1
        start_col, end_col = blocks[block]
        ws.merge_cells(start_row=label_row, start_column=start_col, end_row=label_row, end_column=end_col)
        ws.merge_cells(start_row=value_row, start_column=start_col, end_row=value_row, end_column=end_col)
        label_cell = ws.cell(label_row, start_col, card["label"])
        label_cell.font = _font(bold=True, color=NAVY, size=9)
        label_cell.fill = PatternFill("solid", fgColor=BLUE)
        label_cell.alignment = Alignment(horizontal="center", vertical="center")
        value_cell = ws.cell(value_row, start_col, card["value"])
        value_cell.font = _font(bold=True, color=FORMULA_BLACK, size=14)
        value_cell.fill = PatternFill("solid", fgColor=card.get("fill", WHITE))
        value_cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        if card.get("number_format"):
            value_cell.number_format = str(card["number_format"])
        if card.get("status"):
            _status_fill(value_cell, str(card["status"]))
        ws.row_dimensions[label_row].height = 20
        ws.row_dimensions[value_row].height = 30


def _comparison_target(scenario: str, role: str) -> str:
    if scenario == "personnel_incentive":
        return {
            "输入压缩包": "固定收到的材料范围并路由到人员激励Skill。",
            "销售Excel": "与最终结算单逐SKU比较数量；按门店汇总应付奖励。",
            "最终结算单": "与销售Excel逐SKU比较数量和奖励；与转账总额比较。",
            "转账截图": "去重后与7家门店应付金额集合比较，并单独检查身份和日期。",
        }.get(role, "纳入人员激励证据链。")
    return {
        "输入压缩包": "固定收到的材料范围并路由到堆头Skill。",
        "促销合同": "与销售Excel比较客户/期间；与照片逐店比较日期、地点和陈列；与申报比较金额。",
        "销售Excel": "与合同比较客户和期间；仅支持总体销售，不能替代逐店照片。",
        "现场陈列照片": "绑定合同原始门店序号，逐店比较期间、地点、陈列和重复情况。",
    }.get(role, "纳入堆头证据链。")


def _scenario_file_rows(
    result: dict[str, Any],
    input_manifest: dict[str, Any] | None,
) -> list[list[Any]]:
    scenario = str(result["scenario"])
    extraction = _manifest_scenario(input_manifest, scenario)
    archive_name = _manifest_archive_name(extraction)
    archive_sha = extraction.get("archive_sha256")
    rows: list[list[Any]] = []
    if extraction:
        rows.append(
            [
                archive_name,
                "（ZIP输入包本身）",
                "输入压缩包",
                "检查路径穿越、符号链接、加密、重名、文件数、大小和压缩比。",
                f"安全解压{extraction.get('file_count')}个文件，共{extraction.get('expanded_bytes')}字节。",
                _comparison_target(scenario, "输入压缩包"),
                extraction.get("archive_bytes"),
                archive_sha,
            ]
        )
    files = extraction.get("files") if isinstance(extraction.get("files"), list) else []
    for item in files:
        if not isinstance(item, dict):
            continue
        file_path = str(item.get("path") or "")
        role, reading, read_result = _file_reading_description(scenario, file_path, result)
        rows.append(
            [
                archive_name,
                file_path,
                role,
                reading,
                read_result,
                _comparison_target(scenario, role),
                item.get("bytes"),
                item.get("sha256"),
            ]
        )
    return rows


def _write_file_section(
    ws: Any,
    row: int,
    result: dict[str, Any],
    input_manifest: dict[str, Any] | None,
) -> int:
    row = _section_title(
        ws,
        row,
        "一、读取了哪些文件、读到了什么",
        "每个原始文件单独列示来源、读取字段、读取结果和后续比对对象；SHA-256只固定收到后的文件身份。",
    )
    headers = [
        "序号", "来源ZIP", "压缩包内文件", "文件角色", "具体读取位置/字段",
        "读取结果/数据量", "后续比对工作", "大小与文件SHA-256",
    ]
    blocks = [(1, 1), (2, 3), (4, 6), (7, 7), (8, 9), (10, 11), (12, 12), (13, 13)]
    _write_merged_row(ws, row, blocks, headers, header=True, height=30)
    header_row = row
    row += 1
    for index, values in enumerate(_scenario_file_rows(result, input_manifest), 1):
        source_zip, file_path, role, reading, read_result, comparison, size, sha256 = values
        size_hash = f"{int(size or 0):,}字节\n{sha256 or ''}"
        _write_merged_row(
            ws,
            row,
            blocks,
            [index, source_zip, file_path, role, reading, read_result, comparison, size_hash],
            height=78,
        )
        row += 1
    ws.row_dimensions.group(header_row + 1, row - 1, outline_level=1, hidden=False)
    return row + 1


def _write_analysis_section(
    ws: Any,
    row: int,
    result: dict[str, Any],
    input_manifest: dict[str, Any] | None,
) -> int:
    row = _section_title(
        ws,
        row,
        "1.4.1 七项分析维度",
        "七个维度分别说明读取证据、实际检查、判定依据和能力边界；没有发票、预算或历史照片库时不会默认判为通过。",
    )
    headers = [
        "序号", "分析维度", "检测内容", "1.4.1判定标准", "本次读取文件/证据",
        "实际执行检查", "状态", "判定依据及能力边界/补证",
    ]
    blocks = [(1, 1), (2, 2), (3, 4), (5, 6), (7, 8), (9, 10), (11, 11), (12, 13)]
    _write_merged_row(ws, row, blocks, headers, header=True, height=30)
    row += 1
    for item in build_analysis_dimension_rows([result], input_manifest):
        values = [
            item["dimension_no"], item["analysis_dimension"], item["check_content"],
            item["proposal_standard"], item["files"], item["executed_check"],
            item["status"], f"判定依据：{item['basis']}\n能力边界/补证：{item['limitation']}",
        ]
        _write_merged_row(ws, row, blocks, values, status_indexes={7}, height=108)
        row += 1
    return row + 1


def _write_exception_section(ws: Any, row: int, result: dict[str, Any]) -> int:
    row = _section_title(
        ws,
        row,
        "异常、影响与补证清单",
        "高风险影响数量、金额或核心证据；中风险表示映射、身份、日期或真实性证据仍需说明。",
    )
    headers = ["序号", "等级", "异常代码", "发现及来源", "金额/结论影响", "补证或处理建议"]
    blocks = [(1, 1), (2, 2), (3, 4), (5, 8), (9, 10), (11, 13)]
    _write_merged_row(ws, row, blocks, headers, header=True, height=30)
    row += 1
    for index, item in enumerate(result.get("exceptions") or [], 1):
        finding = item["message"] + (f"\n来源：{item['source']}" if item.get("source") else "")
        severity_label = {"high": "高", "medium": "中", "low": "低"}.get(item["severity"], item["severity"])
        _write_merged_row(
            ws,
            row,
            blocks,
            [index, severity_label, item["code"], finding, item["impact"], item["suggestion"]],
            height=108,
        )
        fill = RED if item["severity"] == "high" else YELLOW
        for start_col, _ in blocks:
            ws.cell(row, start_col).fill = PatternFill("solid", fgColor=fill)
        row += 1
    return row + 1


def _usage_or_default(result: dict[str, Any]) -> dict[str, Any]:
    usage = result.get("model_usage")
    if isinstance(usage, dict):
        return usage
    return {
        "model": "未调用（使用预提取证据）",
        "model_invocation_count": 0,
        "successful_invocation_count": 0,
        "failed_invocation_count": 0,
        "retry_count": 0,
        "input_tokens": 0,
        "cached_input_tokens": 0,
        "output_tokens": 0,
        "reasoning_output_tokens": 0,
        "total_tokens": 0,
        "attempts": [],
    }


def _write_model_usage_section(ws: Any, row: int, result: dict[str, Any]) -> int:
    row = _section_title(
        ws,
        row,
        "模型调用量与失败重试",
        "这里只保存结构化调用次数、重试和Token统计，不保存stdout/stderr、模型事件流或其他运行日志。",
    )
    headers = [
        "模型", "调用次数", "成功次数", "失败次数", "重试次数", "输入Token",
        "缓存输入Token", "输出Token", "推理输出Token", "总Token",
    ]
    _header(ws, row, headers)
    usage = _usage_or_default(result)
    row += 1
    values = [
        usage.get("model"), usage.get("model_invocation_count", 0),
        usage.get("successful_invocation_count", 0), usage.get("failed_invocation_count", 0),
        usage.get("retry_count", 0), usage.get("input_tokens", 0),
        usage.get("cached_input_tokens", 0), usage.get("output_tokens", 0),
        usage.get("reasoning_output_tokens", 0), usage.get("total_tokens", 0),
    ]
    _write_row(ws, row, values, number_formats={col: "#,##0" for col in range(2, 11)}, height=36)
    if int(usage.get("failed_invocation_count", 0) or 0) or int(usage.get("retry_count", 0) or 0):
        for col in (4, 5):
            ws.cell(row, col).fill = PatternFill("solid", fgColor=YELLOW)
    row += 2
    _header(
        ws,
        row,
        ["尝试序号", "状态", "退出码", "失败原因摘要", "输入Token", "缓存输入Token", "输出Token", "推理Token", "总Token"],
    )
    row += 1
    attempts = usage.get("attempts") or []
    if not attempts:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=9)
        ws.cell(row, 1, "本次使用预提取证据，未发生模型调用。")
        _body_cell(ws.cell(row, 1))
        row += 1
    else:
        for attempt in attempts:
            attempt_usage = attempt.get("usage") or {}
            status = "成功" if attempt.get("status") == "success" else "失败后重试"
            values = [
                attempt.get("attempt_no"), status, attempt.get("exit_code"), attempt.get("failure_reason"),
                attempt_usage.get("input_tokens", 0), attempt_usage.get("cached_input_tokens", 0),
                attempt_usage.get("output_tokens", 0), attempt_usage.get("reasoning_output_tokens", 0),
                attempt_usage.get("total_tokens", 0),
            ]
            _write_row(ws, row, values, status_columns={2}, height=36)
            row += 1
    return row + 1


def _add_personnel_audit_sheet(
    wb: Workbook,
    result: dict[str, Any],
    input_manifest: dict[str, Any] | None,
) -> Any:
    ws = wb.create_sheet("人员激励核销")
    _style_title(
        ws,
        "人员激励核销：文件读取、逐SKU、逐门店转账与最终结论",
        COMPACT_LAST_COL,
        "阅读顺序：销售Excel底层行 → 最终结算单10行 → 逐SKU 1:1数量/金额 → 门店应付与转账截图 → 1.4.1与补证。",
    )
    summary = result["summary"]
    sales = result["sales"]
    settlement = result["settlement"]
    run_id = input_manifest.get("run_id") if isinstance(input_manifest, dict) else None
    ws.merge_cells(start_row=8, start_column=1, end_row=8, end_column=COMPACT_LAST_COL)
    ws.cell(
        8,
        1,
        f"运行ID：{run_id or '未提供'}　案件：{result['case_name']}　活动期：{settlement.get('activity_start')} 至 {settlement.get('activity_end')}　"
        f"核心原则：结算单没有条码时保留商品映射依据和置信度，不能只看总金额。",
    )
    ws.cell(8, 1).font = _font(color="555555", size=9)
    ws.cell(8, 1).fill = PatternFill("solid", fgColor=GRAY)
    ws.cell(8, 1).alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[8].height = 30

    row = 9
    row = _section_title(ws, row, "核销步骤总览", "每一步都列明读什么、怎样计算、与什么比较以及本步限制。")
    step_blocks = [(1, 1), (2, 3), (4, 5), (6, 7), (8, 9), (10, 11), (12, 13)]
    _write_merged_row(
        ws,
        row,
        step_blocks,
        ["步骤", "读取文件", "读取内容", "处理/计算", "对比对象", "本步结果", "结论/限制"],
        header=True,
        height=30,
    )
    row += 1
    screenshot_files = sorted(
        {
            Path(str(file_name)).name
            for item in result.get("transfer_evidence") or []
            for file_name in item.get("source_files") or []
        }
    )
    steps = [
        [1, Path(str(sales["source_file"])).name, "Sheet1有效销售行；期间、门店、条码、商品、数量、单价、金额", "跳过空白/小计/合计并保留原Excel行号；按条码、门店汇总", "销售Excel自身明细与合计", f"{sales['detail_row_count']}行、{sales['store_count']}家、{sales['sku_count']}个SKU、数量{sales['total_quantity']}", "销售数量是结算核验的底层来源"],
        [2, Path(str(settlement["source_file"])).name, "客户、活动期、10个奖励商品、数量、奖励单价、行金额、总额、实际申报", "把结算单每一行作为独立核销对象", "结算单10行", f"数量{summary['settlement_line_quantity']}、行奖励{summary['settlement_line_reward']}元、申报{summary['claimed_amount']}元", "结算单无条码，名称映射须保留依据"],
        [3, "销售Excel + 最终结算单", "结算每行与Excel唯一条码及全部来源行", "7家门店同条码数量求和；计算数量×奖励单价", "10个结算SKU ↔ 10个Excel条码", f"{summary['quantity_match_count']}/{summary['settlement_line_count']}个数量匹配，计算奖励{summary['calculated_line_reward']}元", "中置信度映射仍需条码/主数据补证"],
        [4, "、".join(screenshot_files), "金额气泡、方向、身份、日期和镜像画面", "发送/接收镜像合并为transfer_id；比较金额集合", "7家门店应付 ↔ 7笔去重转账", f"转账合计{summary['transfer_total']}元", "金额相同不证明收款门店身份或日期"],
        [5, "上述全部材料", "逐SKU、逐门店、逐转账、真实性与合规性证据", "先计算证据支持金额，再与申报金额比较", "申报金额 ↔ 证据支持金额", f"建议通过{summary['suggested_approved_amount']}元、暂挂{summary['temporarily_held_amount']}元", "异常和补证要求在本页末尾列示"],
    ]
    for values in steps:
        _write_merged_row(ws, row, step_blocks, values, height=66)
        row += 1
    row += 1
    row = _write_file_section(ws, row, result, input_manifest)

    row = _section_title(
        ws,
        row,
        "二、销售Excel与最终结算单逐SKU 1:1比对",
        f"结算来源：{Path(str(settlement['source_file'])).name}；销售来源：{Path(str(sales['source_file'])).name}。",
    )
    sku_headers = [
        "结算行", "结算单读取商品", "结算数量", "销售Excel读取（条码/商品）", "Excel来源行",
        "Excel汇总数量", "数量差异", "奖励单价", "Excel计算奖励", "结算行奖励",
        "金额差异", "证据支持金额", "结论/映射依据",
    ]
    _header(ws, row, sku_headers)
    sku_start = row + 1
    row = sku_start
    quantity_labels = {"match": "匹配", "mismatch": "异常", "unmatched": "未匹配"}
    amount_labels = {"match": "匹配", "mismatch": "异常", "unmatched": "未匹配"}
    for item in result["sku_reconciliation"]:
        confidence = str(item["mapping_confidence"])
        confidence_label = {"high": "高", "medium": "中", "low": "低"}.get(confidence, confidence)
        settlement_name = str(item["settlement_product_name"])
        if confidence == "high":
            mapping_note = "结算单无条码；按商品核心名称在销售Excel中唯一映射。"
        elif settlement_name.startswith("SP-4"):
            mapping_note = "结算单使用SP-4别名，销售Excel未显示该别名；按一对一未映射SKU及汇总数量佐证，需补条码主数据。"
        else:
            mapping_note = "结算单与销售Excel名称存在简称、前缀或规格差异；按唯一候选及汇总数量佐证，需补条码主数据。"
        conclusion = (
            f"数量{quantity_labels.get(item['quantity_status'], item['quantity_status'])}；"
            f"金额{amount_labels.get(item['amount_status'], item['amount_status'])}；"
            f"映射置信度{confidence_label}\n{mapping_note}"
        )
        values = [
            item["line_no"], item["settlement_product_name"], item["settlement_quantity"],
            f"{item['mapped_barcode']}\n{item['excel_product_name']}",
            ",".join(str(value) for value in item["excel_rows"]), item["excel_quantity"],
            f"=F{row}-C{row}", item["unit_reward"], f"=F{row}*H{row}",
            item["settlement_reward_amount"], f"=J{row}-I{row}",
            f"=MAX(0,MIN(C{row},F{row}))*H{row}", conclusion,
        ]
        _write_row(
            ws,
            row,
            values,
            height=78,
            number_formats={3: "#,##0", 6: "#,##0", 7: "#,##0", 8: "#,##0.00", 9: "#,##0.00", 10: "#,##0.00", 11: "#,##0.00", 12: "#,##0.00"},
        )
        ws.cell(row, 13).fill = PatternFill("solid", fgColor=GREEN if confidence == "high" else YELLOW)
        row += 1
    sku_total = row
    ws.cell(row, 1, "合计")
    for col in (3, 6, 7, 9, 10, 11, 12):
        letter = get_column_letter(col)
        ws.cell(row, col, f"=SUM({letter}{sku_start}:{letter}{row - 1})")
    for col in range(1, COMPACT_LAST_COL + 1):
        _body_cell(ws.cell(row, col), formula=col in {3, 6, 7, 9, 10, 11, 12})
        ws.cell(row, col).fill = PatternFill("solid", fgColor=BLUE)
        ws.cell(row, col).font = _font(bold=True, color=FORMULA_BLACK)
    row += 2

    row = _section_title(
        ws,
        row,
        "三、销售Excel 70条底层来源明细",
        f"来源：{Path(str(sales['source_file'])).name} / {sales['sheet']}。每一行保留原Excel行号并回指结算行。",
    )
    trace_headers = ["Excel行", "期间", "门店", "条码", "Excel商品", "销售数量", "销售单价", "结算行", "结算商品", "奖励单价", "本行应付奖励"]
    _header(ws, row, trace_headers)
    trace_start = row + 1
    row = trace_start
    for item in result["sales_trace"]:
        values = [
            item["excel_row"], item["period_text"], item["store_name"], item["barcode"], item["product_name"],
            item["quantity"], item["unit_price"], item["settlement_line_no"], item["settlement_product_name"],
            item["unit_reward"], f"=F{row}*J{row}",
        ]
        _write_row(ws, row, values, height=30, number_formats={6: "#,##0", 7: "#,##0.00", 10: "#,##0.00", 11: "#,##0.00"})
        row += 1
    trace_total = row
    ws.cell(row, 1, "合计")
    ws.cell(row, 6, f"=SUM(F{trace_start}:F{row - 1})")
    ws.cell(row, 11, f"=SUM(K{trace_start}:K{row - 1})")
    for col in range(1, 12):
        _body_cell(ws.cell(row, col), formula=col in {6, 11})
        ws.cell(row, col).fill = PatternFill("solid", fgColor=BLUE)
    ws.row_dimensions.group(trace_start, trace_total - 1, outline_level=1, hidden=False)
    row += 2

    row = _section_title(
        ws,
        row,
        "四、7家门店应付奖励与转账截图比对",
        "Excel按门店汇总10个SKU奖励；截图按发送/接收镜像去重。金额集合匹配不等于收款身份匹配。",
    )
    transfer_headers = [
        "门店", "Excel销售数量", "奖励单价", "Excel应付奖励", "转账ID", "去重转账金额",
        "金额差异", "截图来源", "可见身份", "可见日期", "匹配依据", "结论", "限制",
    ]
    _header(ws, row, transfer_headers)
    transfer_start = row + 1
    row = transfer_start
    evidence_by_id = {item["transfer_id"]: item for item in result.get("transfer_evidence") or []}
    for item in result["store_transfer_reconciliation"]:
        evidence = evidence_by_id.get(item["transfer_id"], {})
        unit_reward = (
            float(item["expected_reward_amount"]) / float(item["excel_quantity"])
            if item.get("excel_quantity")
            else 0
        )
        visible_identity = evidence.get("store_name") or evidence.get("recipient_name") or "不可见"
        limitation = "金额仅按集合配对；未取得门店与收款人的身份映射。"
        if not evidence.get("event_date_visible"):
            limitation += "未见完整日历日期。"
        values = [
            item["store_name"], item["excel_quantity"], unit_reward, f"=B{row}*C{row}",
            item["transfer_id"], item["transfer_amount"], f"=F{row}-D{row}",
            "、".join(Path(str(value)).name for value in evidence.get("source_files") or []),
            visible_identity, "是" if evidence.get("event_date_visible") else "否",
            "金额集合匹配" if item["match_basis"] == "amount_multiset" else item["match_basis"],
            "金额匹配，身份未验证" if item["status"] == "amount_match_identity_unverified" else item["status"],
            limitation,
        ]
        _write_row(
            ws,
            row,
            values,
            status_columns={12},
            height=60,
            number_formats={2: "#,##0", 3: "#,##0.00", 4: "#,##0.00", 6: "#,##0.00", 7: "#,##0.00"},
        )
        row += 1
    transfer_total = row
    ws.cell(row, 1, "合计")
    for col in (2, 4, 6, 7):
        letter = get_column_letter(col)
        ws.cell(row, col, f"=SUM({letter}{transfer_start}:{letter}{row - 1})")
    for col in range(1, COMPACT_LAST_COL + 1):
        _body_cell(ws.cell(row, col), formula=col in {2, 4, 6, 7})
        ws.cell(row, col).fill = PatternFill("solid", fgColor=BLUE)
        ws.cell(row, col).font = _font(bold=True, color=FORMULA_BLACK)
    row += 2

    row = _section_title(
        ws,
        row,
        "五、转账截图去重明细",
        "同一笔转账的发送方与接收方气泡合并为一个transfer_id，同时保留出现次数和去重依据。",
    )
    voucher_headers = ["转账ID", "金额", "画面次数", "身份可见", "日期可见", "门店/收款人", "来源文件", "去重依据", "备注"]
    _header(ws, row, voucher_headers)
    voucher_start = row + 1
    row = voucher_start
    for item in result.get("transfer_evidence") or []:
        values = [
            item["transfer_id"], item["amount"], item["occurrence_count"],
            "是" if item["identity_visible"] else "否", "是" if item["event_date_visible"] else "否",
            item.get("store_name") or item.get("recipient_name"),
            "、".join(Path(str(value)).name for value in item.get("source_files") or []),
            item["dedup_basis"], "；".join(item.get("notes") or []),
        ]
        _write_row(ws, row, values, height=72, number_formats={2: "#,##0.00", 3: "#,##0"})
        row += 1
    ws.row_dimensions.group(voucher_start, row - 1, outline_level=1, hidden=False)
    row += 1

    row = _section_title(
        ws,
        row,
        "六、结算单与转账截图真实性风险辅助检查",
        "SHA-256检查完全重复，pHash/dHash筛查近似画面；EXIF/GPS缺失不能自动判假，也不能自动判真。",
    )
    image_headers = ["文件名", "角色", "尺寸", "大小(字节)", "SHA-256", "pHash", "dHash", "EXIF时间", "GPS", "pHash候选"]
    _header(ws, row, image_headers)
    row += 1
    image_inventory = result.get("image_inventory") or {}
    candidate_by_file: dict[str, list[str]] = {}
    for pair in image_inventory.get("phash_candidate_pairs") or []:
        names = [str(value) for value in pair.get("files") or []]
        for name in names:
            others = [value for value in names if value != name]
            candidate_by_file.setdefault(name, []).append(f"{'、'.join(others)}（距离{pair.get('distance')}）")
    for item in image_inventory.get("files") or []:
        file_name = str(item["file_name"])
        values = [
            file_name, "最终结算单" if "结算单" in file_name else "转账截图",
            f"{item['width']}×{item['height']}", item["size"], item["sha256"], item.get("phash"),
            item["dhash"], item["exif_datetime"], "有" if item["gps_present"] else "无",
            "；".join(candidate_by_file.get(file_name, [])) or "无",
        ]
        _write_row(ws, row, values, height=78, number_formats={4: "#,##0"})
        row += 1
    row += 1

    row = _write_analysis_section(ws, row, result, input_manifest)
    row = _write_exception_section(ws, row, result)
    row = _write_model_usage_section(ws, row, result)

    exception_count = int(summary["high_exception_count"]) + int(summary["medium_exception_count"])
    cards = [
        {"label": "结算/申请金额（元）", "value": summary["claimed_amount"], "number_format": "#,##0.00"},
        {"label": "证据支持金额（元）", "value": f"=MIN(L{sku_total},F{transfer_total})", "number_format": "#,##0.00", "fill": GREEN},
        {"label": "建议通过金额（元）", "value": "=MIN(A4,D4)", "number_format": "#,##0.00", "fill": GREEN},
        {"label": "暂挂金额（元）", "value": "=A4-G4", "number_format": "#,##0.00", "fill": YELLOW},
        {"label": "结算SKU", "value": f"{summary['settlement_line_count']}个"},
        {"label": "逐SKU数量通过", "value": f"{summary['quantity_match_count']}/{summary['settlement_line_count']}", "fill": GREEN},
        {"label": "异常/待说明", "value": f"{exception_count}项", "fill": YELLOW if exception_count else GREEN},
        {"label": "最终结论", "value": "待补件（金额可通过）" if exception_count else "通过", "status": "待补件" if exception_count else "通过"},
    ]
    _write_summary_cards(ws, cards)
    ws.freeze_panes = "A9"
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 70
    ws.sheet_properties.tabColor = "70AD47"
    ws.sheet_properties.outlinePr.summaryBelow = True
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_title_rows = "1:8"
    _set_widths(
        ws,
        {1: 10, 2: 28, 3: 16, 4: 38, 5: 20, 6: 15, 7: 13, 8: 13, 9: 17, 10: 17, 11: 15, 12: 18, 13: 52},
    )
    return ws


def _add_display_audit_sheet(
    wb: Workbook,
    result: dict[str, Any],
    input_manifest: dict[str, Any] | None,
) -> Any:
    ws = wb.create_sheet("堆头核销")
    _style_title(
        ws,
        "堆头核销：合同、销售、逐店照片、真实性风险与最终结论",
        COMPACT_LAST_COL,
        "阅读顺序：合同20家及陈列规则 → 销售客户/期间 → 21张照片逐店绑定 → 日期/地点/陈列/重复五项控制 → 支持金额。",
    )
    summary = result["summary"]
    contract = result["contract"]
    sales = result["sales"]
    run_id = input_manifest.get("run_id") if isinstance(input_manifest, dict) else None
    ws.merge_cells(start_row=8, start_column=1, end_row=8, end_column=COMPACT_LAST_COL)
    ws.cell(
        8,
        1,
        f"运行ID：{run_id or '未提供'}　案件：{result['case_name']}　活动期：{contract.get('activity_start')} 至 {contract.get('activity_end')}　"
        f"合同规则：{contract.get('display_standard')}　核心原则：照片文件名只用于归组，不作为地点证明。",
    )
    ws.cell(8, 1).font = _font(color="555555", size=9)
    ws.cell(8, 1).fill = PatternFill("solid", fgColor=GRAY)
    ws.cell(8, 1).alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[8].height = 30

    row = 9
    row = _section_title(ws, row, "核销步骤总览", "每家1000元只有在照片、期间、门店、陈列和重复检查全部通过时才计入支持金额。")
    step_blocks = [(1, 1), (2, 3), (4, 5), (6, 7), (8, 9), (10, 11), (12, 13)]
    _write_merged_row(
        ws,
        row,
        step_blocks,
        ["步骤", "读取文件", "读取内容", "处理/计算", "对比对象", "本步结果", "结论/限制"],
        header=True,
        height=30,
    )
    row += 1
    steps = [
        [1, Path(str(contract["source_file"])).name, "客户、活动期、20家门店、陈列标准、单店费用、申报金额、印章", "建立合同原始门店1—20行基准", "合同条款 ↔ 申报/照片", f"{summary['contract_store_count']}家×{summary['fee_per_store']}元={summary['expected_contract_amount']}元", "合同门店是逐店判定基准"],
        [2, Path(str(sales["source_file"])).name, "客户、期间、商品编码、条码、商品、数量、零售价、金额", "汇总36个SKU并判断是否有门店字段", "销售客户/期间 ↔ 合同", f"数量{summary['sales_quantity']}、零售金额{summary['sales_retail_amount']}元", "仅经销商汇总，不能证明20家逐店执行"],
        [3, f"现场陈列照片{summary['photo_count']}张", "画面日期、地点水印、陈列范围、清晰度；SHA/pHash/dHash、尺寸、EXIF、GPS", "按合同序号绑定照片并逐张查重", "21张照片 ↔ 合同20家门店", f"已绑定{summary['photo_count']}张照片", "第10家有2张，其余每家1张"],
        [4, "合同 + 逐店照片", "照片存在、期间、门店地点、陈列标准、重复情况", "五项控制逐店判断，全部通过才支持1000元", "合同每家 ↔ 对应现场照片", f"{summary['passed_store_count']}家通过、{summary['supplement_store_count']}家补证", "无EXIF/GPS时保留真实性边界"],
        [5, "上述全部材料", "逐店结果、销售支持、1.4.1和异常补证", "先计算逐店支持金额，再与申报比较", "申报金额 ↔ 逐店证据支持金额", f"建议通过{summary['suggested_approved_amount']}元、暂挂{summary['temporarily_held_amount']}元", "6家补证要求在本页列示"],
    ]
    for values in steps:
        _write_merged_row(ws, row, step_blocks, values, height=66)
        row += 1
    row += 1
    row = _write_file_section(ws, row, result, input_manifest)

    row = _section_title(
        ws,
        row,
        "二、合同、销售Excel与申报基础对比",
        "销售Excel只验证客户、活动期间和总体销售；因为没有门店字段，不能直接支持任何单店1000元。",
    )
    baseline_headers = ["检查项", "合同PDF读取值", "销售Excel/申报读取值", "对比规则", "差异/覆盖", "结果", "限制"]
    baseline_blocks = [(1, 1), (2, 3), (4, 5), (6, 7), (8, 9), (10, 10), (11, 13)]
    _write_merged_row(ws, row, baseline_blocks, baseline_headers, header=True, height=30)
    baseline_start = row + 1
    row = baseline_start
    contract_period = f"{contract.get('activity_start')} 至 {contract.get('activity_end')}"
    sales_period = "、".join(str(value) for value in sales.get("period_values") or [])
    baseline_rows: list[list[Any]] = [
        ["客户", contract.get("customer_name"), "、".join(sales.get("customers") or []), "名称一致性", "一致", "匹配", "仅证明经销商层级销售"],
        ["活动期间", contract_period, sales_period, "销售期间应在合同活动期内", "完整覆盖4月", "匹配", "不代表逐店均在期间内完成陈列"],
        ["数据粒度", f"合同门店{summary['contract_store_count']}家", "销售Excel无门店字段", "门店级执行需要门店级证据", "0家可由销售Excel单独证明", "证据不足", "逐店必须回到照片核验"],
        ["陈列规则", contract.get("display_standard"), f"现场照片{summary['photo_count']}张", "逐店判断1平方米堆头或至少4纵陈列", f"{summary['passed_store_count']}家综合通过", "部分通过", "照片未提供尺量数据"],
        ["费用金额", summary["expected_contract_amount"], summary["claimed_amount"], "合同门店数×单店费用 ↔ 申报", f"=D{row + 4}-B{row + 4}", f"=IF(H{row + 4}=0,\"匹配\",\"异常\")", "独立活动预算表未提交"],
        ["总体销售", "合同无SKU销售目标", summary["sales_quantity"], "作为活动期总体销售背景", summary["sales_retail_amount"], "总体支持", "不能分摊给20家门店"],
    ]
    for values in baseline_rows:
        _write_merged_row(
            ws,
            row,
            baseline_blocks,
            values,
            status_indexes={6},
            height=54,
            number_formats={2: "#,##0.00", 3: "#,##0.00", 5: "#,##0.00"},
        )
        row += 1
    row += 1

    row = _section_title(
        ws,
        row,
        "三、合同20家门店与21张照片逐店比对",
        "五项控制：有照片、日期在活动期、可见地点与合同门店一致、达到陈列标准、未发现重复；全部通过才支持单店费用。",
    )
    store_headers = [
        "合同序号", "合同门店", "照片文件", "可见日期", "可见地点/水印", "期间",
        "门店匹配", "陈列标准", "重复检查", "结论", "支持金额", "风险/判断依据", "补证建议",
    ]
    _header(ws, row, store_headers)
    store_start = row + 1
    row = store_start
    status_labels = {"pass": "通过", "supplement": "补证"}
    period_labels = {"match": "是", "mismatch": "否", "unverifiable": "无法判断"}
    store_labels = {"exact": "是", "compatible": "基本匹配", "mismatch": "不一致", "filename_only": "弱（仅文件名）"}
    display_labels = {"pass": "是", "fail": "否", "uncertain": "无法判断"}
    duplicate_labels = {"none": "未发现", "exact": "完全重复", "possible": "疑似重复", "unverifiable": "无法判断"}
    for item in result["store_reconciliation"]:
        status = status_labels[item["status"]]
        values = [
            item["store_line_no"], item["contract_store_name"], "、".join(item["photo_files"]),
            item["visible_date"], item["visible_location"], period_labels[item["period_match"]],
            store_labels[item["store_match"]], display_labels[item["display_match"]],
            duplicate_labels[item["duplicate_check"]], status,
            f"=IF(J{row}=\"通过\",{summary['fee_per_store']},0)",
            "；".join(item["risk_notes"]), "；".join(item["supplement_advice"]),
        ]
        _write_row(ws, row, values, status_columns={10}, height=108, number_formats={11: "#,##0.00"})
        row += 1
    store_total = row
    ws.cell(row, 1, "合计")
    ws.cell(row, 3, f"照片{summary['photo_count']}张")
    ws.cell(row, 10, f'=COUNTIF(J{store_start}:J{row - 1},"通过")&"家通过"')
    ws.cell(row, 11, f"=SUM(K{store_start}:K{row - 1})")
    for col in range(1, COMPACT_LAST_COL + 1):
        _body_cell(ws.cell(row, col), formula=col in {10, 11})
        ws.cell(row, col).fill = PatternFill("solid", fgColor=BLUE)
        ws.cell(row, col).font = _font(bold=True, color=FORMULA_BLACK)
    row += 2

    row = _section_title(
        ws,
        row,
        "四、销售Excel 36条SKU明细",
        f"来源：{Path(str(sales['source_file'])).name} / {sales['sheet']}。客户和期间与合同一致，但没有门店字段。",
    )
    sales_headers = ["Excel行", "客户", "期间", "商品编码", "条码", "商品名称", "数量", "零售金额", "在本案中的作用"]
    _header(ws, row, sales_headers)
    sales_start = row + 1
    row = sales_start
    for item in sales.get("records") or []:
        values = [
            item["excel_row"], item["customer_name"], item["period_text"], item["product_code"],
            item["barcode"], item["product_name"], item["quantity"], item["total_amount"],
            "支持合同客户和活动期内的总体销售；不能证明单店陈列。",
        ]
        _write_row(ws, row, values, height=36, number_formats={7: "#,##0", 8: "#,##0.00"})
        row += 1
    sales_total = row
    ws.cell(row, 1, "合计")
    ws.cell(row, 7, f"=SUM(G{sales_start}:G{row - 1})")
    ws.cell(row, 8, f"=SUM(H{sales_start}:H{row - 1})")
    for col in range(1, 10):
        _body_cell(ws.cell(row, col), formula=col in {7, 8})
        ws.cell(row, col).fill = PatternFill("solid", fgColor=BLUE)
    ws.row_dimensions.group(sales_start, sales_total - 1, outline_level=1, hidden=False)
    row += 2

    row = _section_title(
        ws,
        row,
        "五、21张照片文件哈希与原始元数据",
        "技术检查用于发现完全重复、近似重复和元数据缺口；本包内未重复不等于历史活动从未复用。",
    )
    photo_headers = ["文件名", "尺寸", "文件大小", "SHA-256", "pHash", "dHash", "EXIF时间", "GPS", "绑定合同门店", "状态"]
    _header(ws, row, photo_headers)
    photo_start = row + 1
    row = photo_start
    bound_store: dict[str, str] = {}
    for item in result["store_reconciliation"]:
        for file_name in item["photo_files"]:
            bound_store[Path(str(file_name)).name] = item["contract_store_name"]
    for item in result["photo_inventory"].get("files") or []:
        file_name = str(item["file_name"])
        values = [
            file_name, f"{item['width']}×{item['height']}", item["size"], item["sha256"],
            item.get("phash"), item["dhash"], item["exif_datetime"],
            "有" if item["gps_present"] else "无", bound_store.get(file_name),
            "已绑定" if file_name in bound_store else "未绑定",
        ]
        _write_row(ws, row, values, status_columns={10}, height=66, number_formats={3: "#,##0"})
        row += 1
    ws.row_dimensions.group(photo_start, row - 1, outline_level=1, hidden=False)
    row += 1

    row = _write_analysis_section(ws, row, result, input_manifest)
    row = _write_exception_section(ws, row, result)
    row = _write_model_usage_section(ws, row, result)

    cards = [
        {"label": "合同/申请金额（元）", "value": summary["claimed_amount"], "number_format": "#,##0.00"},
        {"label": "逐店证据支持金额（元）", "value": f"=K{store_total}", "number_format": "#,##0.00", "fill": GREEN},
        {"label": "建议通过金额（元）", "value": "=MIN(A4,D4)", "number_format": "#,##0.00", "fill": GREEN},
        {"label": "暂挂金额（元）", "value": "=A4-G4", "number_format": "#,##0.00", "fill": YELLOW},
        {"label": "合同门店", "value": f"{summary['contract_store_count']}家"},
        {"label": "逐店通过", "value": f'=COUNTIF(J{store_start}:J{store_total - 1},"通过")&"/{summary["contract_store_count"]}"', "fill": GREEN},
        {"label": "待补件门店", "value": f'=COUNTIF(J{store_start}:J{store_total - 1},"补证")&"家"', "fill": YELLOW},
        {"label": "最终结论", "value": "待补件" if summary["supplement_store_count"] else "通过", "status": "待补件" if summary["supplement_store_count"] else "通过"},
    ]
    _write_summary_cards(ws, cards)
    ws.freeze_panes = "A9"
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 70
    ws.sheet_properties.tabColor = "ED7D31"
    ws.sheet_properties.outlinePr.summaryBelow = True
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_title_rows = "1:8"
    _set_widths(
        ws,
        {1: 10, 2: 30, 3: 34, 4: 17, 5: 36, 6: 14, 7: 15, 8: 15, 9: 18, 10: 14, 11: 15, 12: 52, 13: 52},
    )
    return ws


def _create_legacy_combined_report(
    results: list[dict[str, Any]],
    output_path: str | Path,
    *,
    input_manifest: dict[str, Any] | None = None,
) -> Path:
    if not results:
        raise AuditError("Cannot create a report without audit results")
    workbook = Workbook()
    workbook.remove(workbook.active)
    summary_ws = workbook.create_sheet("核销总览")
    personnel_refs: dict[str, str] | None = None
    display_refs: dict[str, str] | None = None
    for result in results:
        if result["scenario"] == "personnel_incentive":
            if personnel_refs is not None:
                raise AuditError("Combined report currently supports one personnel incentive case")
            personnel_refs = _add_personnel_sheets(workbook, result)
        elif result["scenario"] == "promotional_display":
            if display_refs is not None:
                raise AuditError("Combined report currently supports one promotional display case")
            display_refs = _add_display_sheets(workbook, result)
        else:
            raise AuditError(f"Unsupported report scenario: {result['scenario']}")

    header_row = _style_title(
        summary_ws,
        "线下活动核销结果：通过 / 异常 / 待补件汇总",
        13,
        "先逐SKU、逐门店、逐转账核验，再汇总通过、异常、待补件和金额结论；最终外部交付仅此一份Excel，不保存运行日志。",
    )
    headers = [
        "项目",
        "场景",
        "核验对象",
        "通过",
        "异常",
        "待补件",
        "结算/申请金额",
        "证据支持金额",
        "建议通过金额",
        "暂挂金额",
        "汇总结论",
        "关键原因",
        "明细入口",
    ]
    _header(summary_ws, header_row, headers)
    start = header_row + 1
    result_by_scenario = {item["scenario"]: item for item in results}
    row = start
    if "personnel_incentive" in result_by_scenario and personnel_refs:
        result = result_by_scenario["personnel_incentive"]
        summary = result["summary"]
        exception_count = summary["high_exception_count"] + summary["medium_exception_count"]
        values = [
            result["case_name"],
            "人员激励",
            f"{summary['settlement_line_count']}个结算SKU、{summary['store_count']}家门店转账",
            f"{summary['quantity_match_count']}个SKU数量/金额通过",
            f"{exception_count}项异常提示（高{summary['high_exception_count']}、中{summary['medium_exception_count']}）",
            f"{exception_count}项待说明/补证",
            summary["claimed_amount"],
            f"=MIN({personnel_refs['supported_cell']},{personnel_refs['transfer_total_cell']})",
            f"=MIN(G{row},H{row})",
            f"=G{row}-I{row}",
            "待补件（金额可通过）" if exception_count else "通过",
            "；".join(item["message"] for item in result["exceptions"][:3]),
            "查看“人员激励-SKU逐项”",
        ]
        for col, value in enumerate(values, 1):
            cell = summary_ws.cell(row, col, value)
            formula = isinstance(value, str) and value.startswith("=")
            _body_cell(cell, formula=formula, linked=formula and col == 4)
            if col == 11:
                _status_fill(cell, str(value))
        _sheet_link(summary_ws.cell(row, 13), "人员激励-SKU逐项")
        summary_ws.row_dimensions[row].height = 72
        row += 1
    if "promotional_display" in result_by_scenario and display_refs:
        result = result_by_scenario["promotional_display"]
        summary = result["summary"]
        exception_count = summary["high_exception_count"] + summary["medium_exception_count"]
        values = [
            result["case_name"],
            "堆头陈列",
            f"{summary['contract_store_count']}家合同门店",
            f"{summary['passed_store_count']}家通过",
            f"{exception_count}项异常提示（高{summary['high_exception_count']}、中{summary['medium_exception_count']}）",
            f"{summary['supplement_store_count']}家待补件",
            summary["claimed_amount"],
            f"={display_refs['supported_cell']}",
            f"=MIN(G{row},H{row})",
            f"=G{row}-I{row}",
            "待补件" if summary["supplement_store_count"] else "通过",
            "；".join(item["message"] for item in result["exceptions"][:3]),
            "查看“堆头-逐店核验”",
        ]
        for col, value in enumerate(values, 1):
            cell = summary_ws.cell(row, col, value)
            formula = isinstance(value, str) and value.startswith("=")
            _body_cell(cell, formula=formula, linked=formula and col == 4)
            if col == 11:
                _status_fill(cell, str(value))
        _sheet_link(summary_ws.cell(row, 13), "堆头-逐店核验")
        summary_ws.row_dimensions[row].height = 72
        row += 1
    total_row = row
    summary_ws.cell(total_row, 1, "合计")
    for col in (7, 8, 9, 10):
        letter = get_column_letter(col)
        summary_ws.cell(total_row, col, f"=SUM({letter}{start}:{letter}{total_row - 1})")
    summary_ws.cell(total_row, 11, f'=IF(J{total_row}=0,"全部通过","存在异常/待补件")')
    for col in range(1, 14):
        _body_cell(summary_ws.cell(total_row, col), formula=col in {7, 8, 9, 10, 11})
        summary_ws.cell(total_row, col).fill = PatternFill("solid", fgColor=BLUE)
        summary_ws.cell(total_row, col).font = _font(bold=True, color=FORMULA_BLACK)
    _freeze_filter(summary_ws, header_row, total_row, len(headers))
    _set_widths(summary_ws, {1: 28, 2: 14, 3: 32, 4: 28, 5: 32, 6: 24, 7: 17, 8: 17, 9: 17, 10: 15, 11: 24, 12: 72, 13: 28})

    summary_ws.sheet_properties.tabColor = NAVY
    model_usage_ws = _add_model_usage_sheet(workbook, results)
    workflow_ws = _add_workflow_sheet(workbook, results, input_manifest)
    analysis_ws = _add_analysis_dimensions_sheet(workbook, results, input_manifest)
    inventory_ws = _add_file_inventory_sheet(workbook, results, input_manifest)
    lead_sheets = [summary_ws, model_usage_ws, workflow_ws, analysis_ws, inventory_ws]
    remaining_sheets = [sheet for sheet in workbook._sheets if sheet not in lead_sheets]
    workbook._sheets = lead_sheets + remaining_sheets
    workbook.calculation.calcMode = "auto"
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    target = Path(output_path).resolve()
    target.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(target)
    workbook.close()
    return target


def create_combined_report(
    results: list[dict[str, Any]],
    output_path: str | Path,
    *,
    input_manifest: dict[str, Any] | None = None,
) -> Path:
    """Create one complete worksheet per supported audit method."""
    if not results:
        raise AuditError("Cannot create a report without audit results")
    by_scenario: dict[str, dict[str, Any]] = {}
    for result in results:
        scenario = str(result.get("scenario"))
        if scenario not in {"personnel_incentive", "promotional_display"}:
            raise AuditError(f"Unsupported report scenario: {scenario}")
        if scenario in by_scenario:
            raise AuditError(f"Combined report supports only one case per scenario: {scenario}")
        by_scenario[scenario] = result

    workbook = Workbook()
    workbook.remove(workbook.active)
    expected_sheets: list[str] = []
    if "personnel_incentive" in by_scenario:
        _add_personnel_audit_sheet(workbook, by_scenario["personnel_incentive"], input_manifest)
        expected_sheets.append("人员激励核销")
    if "promotional_display" in by_scenario:
        _add_display_audit_sheet(workbook, by_scenario["promotional_display"], input_manifest)
        expected_sheets.append("堆头核销")
    if workbook.sheetnames != expected_sheets:
        raise AuditError(
            f"Compact workbook sheet contract failed: expected {expected_sheets}, got {workbook.sheetnames}"
        )

    workbook.calculation.calcMode = "auto"
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    target = Path(output_path).resolve()
    target.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(target)
    workbook.close()
    return target


def workbook_formula_inventory(path: str | Path) -> dict[str, Any]:
    workbook = load_workbook(path, read_only=True, data_only=False)
    formulas: list[str] = []
    formula_errors: list[str] = []
    known_sheets = set(workbook.sheetnames)
    for ws in workbook.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    location = f"{ws.title}!{cell.coordinate}"
                    formulas.append(location)
                    if any(token in cell.value for token in ("#REF!", "#DIV/0!", "#VALUE!", "#NAME?")):
                        formula_errors.append(location)
                    for quoted in __import__("re").findall(r"'([^']+)'!", cell.value):
                        if quoted not in known_sheets:
                            formula_errors.append(location)
    cached_errors: list[str] = []
    workbook.close()
    cached_book = load_workbook(path, read_only=True, data_only=True)
    for ws in cached_book.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and any(
                    token in cell.value
                    for token in ("#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A")
                ):
                    cached_errors.append(f"{ws.title}!{cell.coordinate}")
    cached_book.close()
    return {
        "formula_count": len(formulas),
        "locations": formulas,
        "static_formula_errors": sorted(set(formula_errors)),
        "cached_formula_errors": sorted(set(cached_errors)),
        "status": "success" if not formula_errors and not cached_errors else "errors_found",
    }
